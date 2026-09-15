"""
Module for exporting recommendations to CSV, Excel, and HTML.
"""
import csv
import json
import os
import re
import webbrowser
from datetime import datetime
from html import escape
from pathlib import Path
from .copilot_readiness_import import FLAG_COLUMNS


PREFERRED_STATUS_ORDER = [
    "Critical",
    "Action Required",
    "Attention Required",
    "Warning",
    "PendingInput",
    "PendingActivation",
    "Not Assessed",
    "Success",
    "Unknown",
]

from .new_recommendation import (  # noqa: F401  (re-exported)
    NOT_ASSESSED_STATUS,
    CATEGORY_SCAN_COVERAGE,
)
from .assessment_model import (
    DISPOSITION_ACTION,
    DISPOSITION_ASSURANCE,
    DISPOSITION_COVERAGE,
    DISPOSITION_OPPORTUNITY,
    enrich_assessment_records,
    summarize_readiness,
)

RECOMMENDATION_EXPORT_FIELDS = [
    "RecommendationId",
    "Service",
    "Disposition",
    "ReadinessStage",
    "ImpactArea",
    "AIApplicability",
    "Category",
    "Feature",
    "AlsoLicensedVia",
    "Status",
    "Priority",
    "Observation",
    "Recommendation",
    "LinkText",
    "LinkUrl",
    "EvidenceAvailable",
    "EvidenceSheet",
    "EvidenceSummary",
    "EvidenceBasis",
    "Confidence",
]


def summarize_finding(observation, max_length=110):
    """Condense an observation into a scannable card heading.

    Cards were titled by licence, so a reader scanning the Entra section saw "Microsoft Entra ID
    P1" six times over six unrelated findings (Conditional Access coverage, MFA enrolment,
    passwordless adoption, group-based licensing...). The licence is still shown, as a subtitle;
    the heading now says what was actually found.
    """
    text = str(observation or "").strip()
    if not text:
        return ""

    # Several observations open with a headline followed by a blank line and a breakdown
    # ("Copilot Security Posture: NOT READY\n\n1 Critical Gaps: ..."). Keep the headline.
    first_block = re.split(r"\n\s*\n", text, maxsplit=1)[0]
    if first_block.strip():
        text = first_block

    # Drop markdown emphasis and collapse whitespace so headings stay on one line.
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return ""

    # First sentence, avoiding common abbreviations and decimals that contain a period.
    match = re.search(r"(?<![A-Z0-9])\.(?=\s+[A-Z(])", text)
    if match:
        text = text[:match.start()].strip()

    if len(text) > max_length:
        cut = text[:max_length].rsplit(" ", 1)[0].rstrip(" ,;:-")
        text = f"{cut}…"
    return text


def customer_coverage_label(record):
    """Give decision-limiting evidence gaps a plain-language customer label."""
    service = str(record.get("Service", "") or "")
    feature = str(record.get("Feature", "") or record.get("Service", "") or "Required evidence")
    text = f"{service} {feature}".lower()
    if "sharepoint sharing and oversharing" in text:
        return "SharePoint sharing settings and DAG status"
    if "intune" in text or "managed device" in text:
        return "Endpoint access-control model"
    if "data governance evidence" in text:
        return "Purview DLP and governance configuration"
    if service == "Data Exposure" and "sharepoint" in text:
        return "SharePoint Data Access Governance report"
    if service == "Data Exposure" and ("sensitive-data" in text or "dspm" in text):
        return "Purview DSPM data-risk assessment"
    return feature


def sanitize_filename_component(value):
    """Convert tenant names into filesystem-friendly filename segments."""
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value or "").strip())
    safe = safe.strip("._-").lower()
    return safe[:80]


def build_report_filename(extension, tenant_name=None):
    """Build a timestamped report filename, optionally including tenant name."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    tenant_slug = sanitize_filename_component(tenant_name)
    if tenant_slug:
        return f"m365_recommendations_{tenant_slug}_{timestamp}.{extension}"
    return f"m365_recommendations_{timestamp}.{extension}"


def sort_values_with_preferences(values, preferred_order):
    """Sort values using a preferred order first, then alphabetically."""
    preferred_positions = {value.lower(): index for index, value in enumerate(preferred_order)}
    unique_values = []
    seen = set()

    for value in values:
        normalized = str(value or "").strip() or "Unknown"
        key = normalized.lower()
        if key in seen:
            continue
        seen.add(key)
        unique_values.append(normalized)

    return sorted(
        unique_values,
        key=lambda value: (
            preferred_positions.get(value.lower(), len(preferred_positions)),
            value.lower(),
        ),
    )


def _autosize_columns(ws, headers):
    for index, header in enumerate(headers, start=1):
        max_length = len(str(header or ""))
        column_letter = ws.cell(row=1, column=index).column_letter
        for cell in ws[column_letter]:
            cell_value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(cell_value), 80))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 60)


def _apply_excel_sheet_formatting(ws, headers, header_fill, header_font, wrap_alignment):
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap_alignment

    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_alignment

    _autosize_columns(ws, headers)


def _safe_table_name(name):
    sanitized = re.sub(r"[^A-Za-z0-9_]+", "_", str(name or "Table"))
    sanitized = sanitized.strip("_") or "Table"
    if sanitized[0].isdigit():
        sanitized = f"T_{sanitized}"
    return sanitized[:200]


def _add_excel_table(ws, table_name):
    if ws.max_row <= 1 or ws.max_column <= 0:
        return

    from openpyxl.worksheet.table import Table, TableStyleInfo

    # A table supplies its own AutoFilter. Keeping a worksheet AutoFilter over the
    # same range creates duplicate filter definitions that openpyxl tolerates but
    # some desktop Excel builds report as a workbook repair/corruption condition.
    ws.auto_filter.ref = None
    base_name = _safe_table_name(table_name)
    used_names = {name.casefold() for sheet in ws.parent for name in sheet.tables}
    unique_name = base_name
    suffix = 2
    while unique_name.casefold() in used_names:
        unique_name = f"{base_name}_{suffix}"
        suffix += 1
    table = Table(displayName=unique_name, ref=ws.dimensions)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)

def export_to_csv(recommendations, filename=None, tenant_name=None):
    """
    Export recommendations to CSV file
    
    Args:
        recommendations: List of recommendation dictionaries
        filename: Output filename (optional, generates timestamp-based name if not provided)
    
    Returns:
        str: Path to created CSV file
    """
    recommendations = enrich_assessment_records(recommendations)

    # Create Reports folder if it doesn't exist
    recommendations_dir = Path("Reports")
    recommendations_dir.mkdir(exist_ok=True)
    
    if not filename:
        filename = build_report_filename("csv", tenant_name=tenant_name)
    
    if not filename.endswith('.csv'):
        filename += '.csv'
    
    # Build full path
    filepath = recommendations_dir / filename
    
    if not recommendations:
        print("No recommendations to export.")
        return None
    
    # Define CSV headers
    fieldnames = RECOMMENDATION_EXPORT_FIELDS

    with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
        # Recommendation dicts carry internal fields (EvidenceKey) that are not exported.
        # Without extrasaction='ignore' DictWriter raises on every row.
        writer = csv.DictWriter(csvfile, fieldnames=fieldnames, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(recommendations)
    
    return str(filepath)

def export_to_json(recommendations, filename=None, tenant_name=None):
    """
    Export recommendations to JSON file
    
    Args:
        recommendations: List of recommendation dictionaries
        filename: Output filename (optional, generates timestamp-based name if not provided)
    
    Returns:
        str: Path to created JSON file
    """
    recommendations = enrich_assessment_records(recommendations)

    # Create Reports folder if it doesn't exist
    recommendations_dir = Path("Reports")
    recommendations_dir.mkdir(exist_ok=True)
    
    if not filename:
        filename = build_report_filename("json", tenant_name=tenant_name)
    
    if not filename.endswith('.json'):
        filename += '.json'
    
    # Build full path
    filepath = recommendations_dir / filename
    
    if not recommendations:
        print("No recommendations to export.")
        return None
    
    with open(filepath, 'w', encoding='utf-8') as jsonfile:
        json.dump(recommendations, jsonfile, indent=2, ensure_ascii=False)
    
    print(f"Recommendations exported to JSON: {filepath}")
    return str(filepath)

def _append_dict_rows_to_sheet(ws, rows, header_fill, header_font, wrap_alignment, table_name=None):
    if not rows:
        return

    headers = list(dict.fromkeys(header for row in rows for header in row))
    ws.append(headers)
    for cell in ws[1]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
    for row_number, row in enumerate(rows, start=2):
        values = [json.dumps(row.get(header), ensure_ascii=False, sort_keys=True) if isinstance(row.get(header), (dict, list, tuple)) else row.get(header, "") for header in headers]
        ws.append(values)
        # Evidence can include user-supplied CSV text. Preserve it as text rather
        # than allowing an imported value to become an executable Excel formula.
        for column, value in enumerate(values, start=1):
            if isinstance(value, str) and value.startswith("="):
                ws.cell(row=row_number, column=column).data_type = "s"

    _apply_excel_sheet_formatting(ws, headers, header_fill, header_font, wrap_alignment)
    _add_excel_table(ws, table_name or ws.title)


def _rollout_progress_rows(result):
    return [
        {'Stage': stage.get('label'), 'Stage status': stage.get('status'),
         'Requirement': requirement.get('title'), 'Status': requirement.get('status'),
         'Reason': requirement.get('reason'),
         'Action IDs': '; '.join(str(value) for value in requirement.get('action_ids', [])),
         'Control ID': requirement.get('control_id', ''),
         'Responsible role': requirement.get('owner_role', '')}
        for stage in (result.get('rollout_progress') or {}).get('stages', [])
        for requirement in stage.get('requirements', [])
    ]


def _readiness_review_rows(result):
    review = result.get('readiness_review') or {}
    rows = []
    for name, label in (('pilot_plan', 'Pilot plan'), ('pilot_outcomes', 'Pilot outcomes'),
                        ('expansion_approval', 'Expansion approval')):
        value = review.get(name)
        if isinstance(value, dict) and value:
            rows.append({'Review type': label, **value})
    for name, label in (('control_reviews', 'Control review'), ('pilot_conditions', 'Pilot condition')):
        rows.extend({'Review type': label, **value} for value in review.get(name, []) if isinstance(value, dict))
    if review.get('errors'):
        rows.append({'Review type': 'Validation', 'Status': review.get('status'), 'Errors': review['errors']})
    return rows


def _prior_sheet_title(workbook, original_title):
    """Keep historical tabs identifiable while satisfying Excel's naming rules."""
    clean_title = re.sub(r"[\\/*?:\[\]]", "_", str(original_title or "Evidence")).strip(" '")
    base_title = f"Prior {clean_title or 'Evidence'}"
    used_titles = {title.casefold() for title in workbook.sheetnames}
    title = base_title[:31]
    suffix = 2
    while title.casefold() in used_titles:
        ending = f" ({suffix})"
        title = f"{base_title[:31 - len(ending)]}{ending}"
        suffix += 1
    return title


def _append_prior_report_workbook(workbook, prior, header_fill, header_font, wrap_alignment):
    """Preserve earlier report evidence separately from this build's decisions."""
    if not prior.get("available"):
        return
    source_title = _prior_sheet_title(workbook, "Report Sources")
    source_sheet = workbook.create_sheet(source_title)
    source_rows = []
    prior_sheets = dict(prior.get("sheets") or {})
    existing = {str(title).casefold() for title in prior_sheets}
    if "recommendations" not in existing and prior.get("recommendations"):
        prior_sheets["Recommendations"] = {"rows": prior["recommendations"]}
    if "collection coverage" not in existing and prior.get("collection_coverage"):
        prior_sheets["Collection Coverage"] = {"rows": prior["collection_coverage"]}
    for original_title, original_sheet in prior_sheets.items():
        title = _prior_sheet_title(workbook, original_title)
        sheet = workbook.create_sheet(title)
        _append_dict_rows_to_sheet(
            sheet, original_sheet.get("rows") or [], header_fill, header_font,
            wrap_alignment, table_name=title,
        )
        source_rows.append({
            "Source File": prior.get("source_file", ""),
            "Original Report Date": prior.get("generated_at") or "Not established",
            "Original Sheet": original_title,
            "New Workbook Tab": title,
            "Evidence Context": "Historical report; tenant controls were not recollected.",
        })
    if not source_rows:
        source_rows.append({
            "Source File": prior.get("source_file", ""),
            "Original Report Date": prior.get("generated_at") or "Not established",
            "Original Sheet": "", "New Workbook Tab": "",
            "Evidence Context": "Historical report contains no importable rows.",
        })
    _append_dict_rows_to_sheet(
        source_sheet, source_rows, header_fill, header_font, wrap_alignment,
        table_name=source_title,
    )


def export_to_excel(recommendations, filename=None, tenant_name=None, evidence_bundle=None):
    """
    Export recommendations to Excel file
    Requires openpyxl: pip install openpyxl
    
    Args:
        recommendations: List of recommendation dictionaries
        filename: Output filename (optional, generates timestamp-based name if not provided)
    
    Returns:
        str: Path to created Excel file
    """
    from .assessment_result import build_assessment_result
    evidence_bundle = evidence_bundle if evidence_bundle is not None else {}
    result = evidence_bundle.get('assessment_result')
    if result is None:
        result = build_assessment_result(recommendations, evidence_bundle)
        evidence_bundle['assessment_result'] = result
    recommendations = result['recommendations']

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("Warning: openpyxl not installed. Install it with: pip install openpyxl")
        print("Falling back to CSV export...")
        return export_to_csv(recommendations, filename, tenant_name=tenant_name)
    
    # Create Reports folder if it doesn't exist
    recommendations_dir = Path("Reports")
    recommendations_dir.mkdir(exist_ok=True)
    
    if not filename:
        filename = build_report_filename("xlsx", tenant_name=tenant_name)
    
    if not filename.endswith('.xlsx'):
        filename += '.xlsx'
    
    # Build full path
    filepath = recommendations_dir / filename
    
    if not recommendations:
        print("No recommendations to export.")
        return None
    
    # Create workbook. The first tabs answer "what should we do and why?"; the
    # full legacy register remains available under its established name.
    wb = Workbook()
    ws = wb.active
    ws.title = "Action Plan"

    # Style headers
    header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    wrap_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Priority colors
    priority_colors = {
        "High": "FF6B6B",
        "Medium": "FFD93D",
        "Low": "95E1D3"
    }

    action_rows = [{
        "Action": index,
        "RecommendationId": rec.get("RecommendationId", ""),
        "Domain": rec.get("Domain", ""),
        "Type": rec.get("ActionType", ""),
        "Priority": rec.get("Priority", ""),
        "What We Found": rec.get("Observation", ""),
        "Recommended Action": rec.get("Recommendation", ""),
        "Responsible Role": rec.get("OwnerRole", ""),
        "Rollout Stage": rec.get("ReadinessStage", ""),
        "Target Date": "",
        "Completion Evidence": rec.get("CompletionEvidence", ""),
        "Observed": rec.get("ObservationDate", ""),
        "Qualification": rec.get("Qualification", ""),
        "Evidence": "Evidence Index",
    } for index, rec in enumerate(result['actions'], 1)]
    _append_dict_rows_to_sheet(ws, action_rows or [{
        "Priority": "", "What We Found": "No deployment actions were identified from the evidence collected.",
        "Recommended Action": "Continue monitoring the tenant as conditions and intended AI use cases change.",
        "Owner": "", "Target Date": "", "Completion Evidence": "",
    }], header_fill, header_font, wrap_alignment, table_name="ActionPlan")

    if evidence_bundle:
        # License availability belongs in Service Plan Inventory and the compatibility
        # Recommendations register. It is not source evidence for a control decision.
        evidence_index = [{
            'RecommendationId': row.get('RecommendationId', ''),
            'Domain': row.get('Domain', ''), 'Finding': row.get('Feature', ''),
            'Original Date': row.get('ObservationDate', ''),
            'Evidence Basis': row.get('EvidenceBasis', ''), 'Confidence': row.get('Confidence', ''),
            'Qualification': row.get('Qualification', ''), 'Workbook Tab': row.get('EvidenceSheet', ''),
            'Source File': row.get('SourceFile', ''),
            'Supporting Observation': row.get('Evidence') or row.get('Observation', ''),
        } for row in recommendations]
        evidence_index.extend([
            row for row in evidence_bundle.get('evidence_index', [])
            if str(row.get('Evidence Basis', row.get('EvidenceBasis', ''))).lower() != 'license signal'
            and row.get('RecommendationId') not in {r.get('RecommendationId') for r in recommendations}
        ])
        if evidence_index:
            index_sheet = wb.create_sheet("Evidence Index")
            _append_dict_rows_to_sheet(index_sheet, evidence_index, header_fill, header_font, wrap_alignment, table_name="EvidenceIndex")

        coverage_rows = []
        collection_context = evidence_bundle.get("collection_context", {}) or {}
        migration = collection_context.get('methodology_migration')
        if migration:
            coverage_rows.append({
                'Source': 'Methodology compatibility migration', 'State': 'migrated',
                'Records': '', 'Pages': '', 'Truncated': '',
                'Refresh Date': collection_context.get('collected_at', ''),
                'Freshness': '',
                'Reason': f"{migration['from']} -> {migration['to']}: {migration['reason']}",
                'Source File': collection_context.get('source_file', ''), 'Age (Days)': '',
            })
        if collection_context.get("mode") == "offline":
            coverage_rows.append({
                "Source": "Offline report build", "State": "offline",
                "Records": "", "Pages": "", "Truncated": "",
                "Refresh Date": collection_context.get("collected_at", ""),
                "Freshness": collection_context.get("freshness", "Unknown"),
                "Reason": collection_context.get("scope", "") or "Tenant controls were not collected during this report build.",
                "Source File": collection_context.get("source_file", ""),
                "Age (Days)": collection_context.get("age_days", ""),
            })
        for source, state in (evidence_bundle.get('source_statuses', {}) or {}).items():
            state = state if isinstance(state, dict) else {"availability_status": str(state)}
            coverage_rows.append({
                "Source": source,
                "State": state.get("availability_status", state.get("status", "unknown")),
                "Records": state.get("records_collected", state.get("record_count", "")),
                "Pages": state.get("pages_collected", ""),
                "Truncated": state.get("truncated", False),
                "Refresh Date": state.get("refresh_date", ""),
                "Freshness": state.get("freshness", ""),
                "Reason": state.get("reason", ""),
            })
        if not coverage_rows:
            coverage_rows.append({
                "Source": "Collection metadata", "State": "unavailable", "Records": "",
                "Pages": "", "Truncated": "", "Refresh Date": "", "Freshness": "",
                "Reason": "This export call did not receive collection metadata.",
            })
        coverage_sheet = wb.create_sheet("Collection Coverage")
        _append_dict_rows_to_sheet(coverage_sheet, coverage_rows, header_fill, header_font, wrap_alignment, table_name="CollectionCoverage")

        action_evidence = set()
        for rec in recommendations:
            if rec.get("Disposition") == DISPOSITION_ACTION:
                action_evidence.update(key.strip().lower() for key in str(rec.get("EvidenceKey", "")).split(";") if key.strip())
                evidence_titles = {title.strip().lower() for title in str(rec.get("EvidenceSheet", "")).split(";") if title.strip()}
                for key, sheet in evidence_bundle.get('sheets', {}).items():
                    if str(sheet.get('title', '')).strip().lower() in evidence_titles:
                        action_evidence.add(key)
        for key, sheet in evidence_bundle.get('sheets', {}).items():
            if key not in action_evidence or not sheet.get('rows'):
                continue
            detail_sheet = wb.create_sheet(sheet.get('title', 'Evidence'))
            _append_dict_rows_to_sheet(detail_sheet, sheet['rows'], header_fill, header_font, wrap_alignment, table_name=sheet.get('title', 'Evidence'))

    ws = wb.create_sheet("Recommendations")
    # Add full compatibility register.
    summary_headers = [
        "RecommendationId",
        "Control ID",
        "Methodology Version",
        "Finding Fingerprint",
        "Service",
        "Disposition",
        "Readiness Stage",
        "Impact Area",
        "AI Applicability",
        "Category",
        "Feature",
        "Also Licensed Via",
        "Status",
        "Priority",
        "Observation",
        "Recommendation",
        "Link Text",
        "Link URL",
        "Evidence Available",
        "Evidence Sheet",
        "Evidence Summary",
        "Evidence Basis",
        "Confidence",
    ]
    ws.append(summary_headers)
    for rec in recommendations:
        row = [
            rec.get("RecommendationId", ""),
            rec.get("ControlId", ""),
            rec.get("MethodologyVersion", ""),
            rec.get("FindingFingerprint", ""),
            rec.get("Service", ""),
            rec.get("Disposition", ""),
            rec.get("ReadinessStage", ""),
            rec.get("ImpactArea", ""),
            rec.get("AIApplicability", ""),
            rec.get("Category", ""),
            rec.get("Feature", ""),
            rec.get("AlsoLicensedVia", ""),
            rec.get("Status", ""),
            rec.get("Priority", ""),
            rec.get("Observation", ""),
            rec.get("Recommendation", ""),
            rec.get("LinkText", ""),
            rec.get("LinkUrl", ""),
            rec.get("EvidenceAvailable", ""),
            rec.get("EvidenceSheet", ""),
            rec.get("EvidenceSummary", ""),
            rec.get("EvidenceBasis", ""),
            rec.get("Confidence", ""),
        ]
        ws.append(row)

        # Color code priority
        priority = rec.get("Priority", "")
        if priority in priority_colors:
            row_num = ws.max_row
            priority_cell = ws.cell(row=row_num, column=14)  # Priority column
            priority_cell.fill = PatternFill(start_color=priority_colors[priority], 
                                            end_color=priority_colors[priority], 
                                            fill_type="solid")

    _apply_excel_sheet_formatting(ws, summary_headers, header_fill, header_font, wrap_alignment)
    _add_excel_table(ws, "Recommendations")

    # Create control, decision, reproducibility, and remaining evidence sheets.
    if evidence_bundle:
        integrity_issues = evidence_bundle.get('integrity', {}).get('issues', [])
        integrity_rows = (
            [{"Status": "Failed", "Issue": issue} for issue in integrity_issues]
            if integrity_issues
            else [{"Status": "Passed", "Issue": "No pre-export integrity issues were detected."}]
        )
        for title, rows, table_name in (
            ("Control Results", evidence_bundle.get('control_results', []), "ControlResults"),
            ("Provider Register", evidence_bundle.get('provider_evidence', {}).get('rows', []), "ProviderRegister"),
            ("Use Case Readiness", evidence_bundle.get('conclusions', {}).get('use_cases', []), "UseCaseReadiness"),
            ("Improvement Tracking", evidence_bundle.get('baseline_comparison', {}).get('rows', []), "ImprovementTracking"),
            ("Run Manifest", evidence_bundle.get('run_manifest', {}).get('rows', []), "RunManifest"),
            ("Integrity Checks", integrity_rows, "IntegrityChecks"),
        ):
            if rows:
                extra_sheet = wb.create_sheet(title)
                _append_dict_rows_to_sheet(extra_sheet, rows, header_fill, header_font, wrap_alignment, table_name=table_name)

        existing_titles = set(wb.sheetnames)
        for key, sheet in evidence_bundle.get('sheets', {}).items():
            sheet_rows = sheet.get('rows', [])
            if not sheet_rows or sheet.get('title', 'Evidence') in existing_titles:
                continue
            detail_sheet = wb.create_sheet(sheet.get('title', 'Evidence'))
            _append_dict_rows_to_sheet(
                detail_sheet,
                sheet_rows,
                header_fill,
                header_font,
                wrap_alignment,
                table_name=sheet.get('title', 'Evidence'),
            )

    if evidence_bundle:
        _append_prior_report_workbook(
            wb, evidence_bundle.get("prior_report") or {}, header_fill,
            header_font, wrap_alignment,
        )

    for title, rows in (
        ('Assessment Summary', [{'Item': 'Deployment decision', 'Value': result['decision']},
            {'Item': 'Rationale', 'Value': result['rationale']},
            {'Item': 'Evaluation date', 'Value': result.get('evaluation_date')},
            *[{'Item': key, 'Value': value} for key, value in result['counts'].items()]]),
        ('Domain Results', [{'Domain': row['title'], 'Status': row.get('status'),
            'Summary': row.get('summary'), 'Actions': row.get('action_count'),
            'Responsible Role': row.get('owner_role')} for row in result['domains']]),
        ('Evidence Observations', result.get('observations', result.get('evidence', []))),
        ('Adoption Metrics', result.get('adoption_metrics', [])),
        ('Import Receipt', evidence_bundle.get('import_receipt', [])),
        ('Portal Review', (evidence_bundle.get('portal_review') or {}).get('rows', [])),
        ('PDF Extracted Text', [{'Source File': capture.get('source_name'), 'Page': page['page'],
                                'Method': page['method'], 'Extracted text': page['text'],
                                'Qualification': capture.get('qualification')}
                               for capture in (evidence_bundle.get('portal_review') or {}).get('captures', [])
                               for page in capture.get('extracted_pages', [])]),
        ('Copilot Admin Data', (evidence_bundle.get('copilot_admin_review') or {}).get('rows', [])),
        ('Copilot DLP Detail', ((evidence_bundle.get('copilot_admin_review') or {}).get('dlp') or {}).get('detail', [])),
        ('Portal Review Remaining', (evidence_bundle.get('copilot_admin_review') or {}).get('manual_checks', [])),
        ('Rollout Progress', _rollout_progress_rows(result)),
        ('Readiness Reviews', _readiness_review_rows(result)),
    ):
        if rows:
            sheet = wb.create_sheet(title)
            _append_dict_rows_to_sheet(sheet, rows, header_fill, header_font, wrap_alignment, table_name=title)

    # Every action has a navigable evidence reference, including historical checks.
    if 'Evidence Index' in wb:
        index_sheet = wb['Evidence Index']
        headers = {cell.value: cell.column for cell in index_sheet[1]}
        prior_tabs = {}
        if 'Prior Report Sources' in wb:
            mapping = list(wb['Prior Report Sources'].values)
            for values in mapping[1:]:
                row = dict(zip(mapping[0], values))
                prior_tabs[row.get('Original Sheet')] = row.get('New Workbook Tab')
        historical_ids = {r.get('RecommendationId') for r in recommendations if r.get('SourceType') == 'prior_assessment'}
        for i in range(2, index_sheet.max_row + 1):
            if index_sheet.cell(i, headers['RecommendationId']).value in historical_ids:
                cell = index_sheet.cell(i, headers['Workbook Tab'])
                original_titles = str(cell.value or '').split(';')
                targets = [prior_tabs[t.strip()] for t in original_titles if t.strip() in prior_tabs]
                cell.value = '; '.join(targets) or prior_tabs.get('Recommendations', '')
        index_rows = {index_sheet.cell(i, headers['RecommendationId']).value: i for i in range(2, index_sheet.max_row + 1)}
        action_sheet = wb['Action Plan']
        action_headers = {cell.value: cell.column for cell in action_sheet[1]}
        if 'RecommendationId' in action_headers:
            for i in range(2, action_sheet.max_row + 1):
                target = index_rows.get(action_sheet.cell(i, action_headers['RecommendationId']).value)
                if target:
                    cell = action_sheet.cell(i, action_headers['Evidence'])
                    cell.hyperlink = f"#'Evidence Index'!A{target}"
                    cell.style = 'Hyperlink'
        for i in range(2, index_sheet.max_row + 1):
            cell = index_sheet.cell(i, headers['Workbook Tab'])
            candidates = str(cell.value or '').split(';')
            target = next((item.strip() for item in candidates if item.strip() in wb.sheetnames), None)
            if target:
                cell.hyperlink = "#'" + target.replace("'", "''") + "'!A1"
                cell.style = 'Hyperlink'

    # All imported strings, including register cells and headers, remain inert.
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'

    # Save workbook
    wb.save(filepath)
    return str(filepath)

def export_to_html(recommendations, filename=None, tenant_name=None, evidence_bundle=None, excel_path=None):
    """Render the same assessed result used by the workbook and operator receipt."""
    from .assessment_result import build_assessment_result
    from .customer_report import render_customer_report
    bundle = evidence_bundle if evidence_bundle is not None else {}
    result = bundle.get('assessment_result')
    if result is None:
        result = build_assessment_result(recommendations, bundle)
        bundle['assessment_result'] = result
    folder = Path('Reports')
    folder.mkdir(exist_ok=True)
    name = filename or build_report_filename('html', tenant_name=tenant_name)
    if not name.endswith('.html'):
        name += '.html'
    path = folder / name
    path.write_text(render_customer_report(result, bundle, tenant_name, excel_path), encoding='utf-8')
    return str(path)


def open_html_report(html_path):
    """
    Open the generated HTML report in the default browser.
    
    Args:
        html_path: Filesystem path to HTML report
    
    Returns:
        bool: True if browser launch was requested successfully
    """
    if not html_path:
        return False
    
    try:
        report_uri = Path(html_path).resolve().as_uri()
        return webbrowser.open(report_uri)
    except Exception as e:
        print(f"Warning: Unable to open HTML report automatically: {e}")
        return False

def print_recommendations_summary(
    recommendations,
    csv_path=None,
    excel_path=None,
    html_path=None,
    tenant_name=None,
    evidence_bundle=None,
):
    """
    Print a summary of recommendations grouped by service and priority
    
    Args:
        recommendations: List of recommendation dictionaries
        csv_path: Path to CSV export (optional)
        excel_path: Path to Excel export (optional)
        html_path: Path to HTML export (optional)
    """
    if not recommendations:
        print("\nℹ️  No recommendations generated")
        return
    
    from .console_reporting import detail, display_path, is_verbose, section, status, style, print_source_gaps
    section('ASSESSMENT')
    if tenant_name:
        print(style(tenant_name, 'important'))
    
    recommendations = enrich_assessment_records(recommendations)
    from .assessment_result import build_assessment_result
    readiness = (evidence_bundle or {}).get('assessment_result') or build_assessment_result(recommendations, evidence_bundle or {})
    findings = readiness['actions']
    opportunities = readiness['opportunities']
    assurances = readiness['strengths']
    coverage_items = readiness['coverage']

    # Group by priority
    high_priority = [r for r in findings if r.get("Priority") == "High"]
    medium_priority = [r for r in findings if r.get("Priority") == "Medium"]
    low_priority = [r for r in findings if r.get("Priority") == "Low"]
    tone = 'success' if readiness['decision'] == 'Ready for a controlled pilot' else 'error' if readiness['decision'] == 'Not ready for pilot' else 'warning'
    progress = readiness.get('rollout_progress') or {}
    if progress:
        status(f"Rollout stage: {progress['current_stage']}", tone)
        checks = progress.get('counts') or {}
        status(f"Required controls: {checks.get('established_controls', 0)} established | "
               f"{checks.get('observed_issues', 0)} with issues | {checks.get('unconfirmed_controls', 0)} unconfirmed")
    status(f"Deployment decision: {readiness['decision']}", tone)
    detail(readiness['rationale'])
    counts = readiness['counts']
    print(f"Actions: {len(findings)} — {counts['remediation']} remediation, {counts['confirmation']} confirmation, {counts['evidence_gaps']} evidence checks")
    if counts.get('critical'):
        status(f"Critical priority: {counts['critical']}", 'error')
    print('Priority: ' + style(f'{len(high_priority)} high', 'error' if high_priority else 'muted')
          + ' | ' + style(f'{len(medium_priority)} medium', 'warning' if medium_priority else 'muted')
          + ' | ' + style(f'{len(low_priority)} low', 'info'))
    detail(f"Service-plan/context items (workbook): {len(opportunities)}")
    strength_count = len(assurances)
    status(f"Verified strengths: {strength_count}", 'success' if strength_count else 'muted')
    if coverage_items:
        detail(f"Coverage gaps: {len(coverage_items)} (included in the action total)")

    if is_verbose():
        section('Actions by assessment area')
    for domain in readiness['domains']:
        count = domain.get('action_count', 0)
        if count and is_verbose():
            action_label = "action" if count == 1 else "actions"
            print(f"  • {domain['title']}: {count} {action_label}")

    print_source_gaps((evidence_bundle or {}).get('source_statuses'))
    if html_path or excel_path or csv_path:
        section('OUTPUT FILES')
        for label, path in (('HTML', html_path), ('Workbook', excel_path), ('CSV', csv_path)):
            if path:
                print(f"{label}: {style(display_path(path), 'path')}")
