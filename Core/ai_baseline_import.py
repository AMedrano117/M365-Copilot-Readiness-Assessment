"""Read Compliance Manager AI Baseline action exports as dated reference evidence.

Microsoft's implementation and testing states never become local control results.
Workbook filenames and filesystem timestamps cannot establish assessment identity
or evidence dates. Unfamiliar exports are rejected rather than guessed.
"""

from collections import Counter
from datetime import date, datetime
import hashlib
from itertools import islice
from pathlib import Path
import re
from uuid import UUID
from zipfile import BadZipFile


QUALIFICATION = ('Microsoft-reported compliance actions, retained as a dated reference. '
                 'These statuses do not pass technical controls, close local actions or change the readiness decision. '
                 'Compliance Manager AI Baseline and DSPM policy recommendations are separate sources.')
ALIASES = {
    'Improvement action': ('improvementaction', 'improvementactions', 'improvementactionname', 'actionname', 'actiontitle'),
    'Action ID': ('improvementactionid', 'actionid'),
    'Assessment': ('assessment', 'assessmentname', 'assessments'),
    'Regulation': ('regulation', 'regulationname', 'assessmenttemplate', 'regulations'),
    'Control': ('control', 'controlname', 'controltitle'),
    'Control ID': ('controlid',),
    'Responsibility': ('responsibility', 'managedby'),
    'Action type': ('actiontype',),
    'Solution': ('solution', 'microsoftsolutionname'),
    'Points achieved': ('pointsachieved', 'pointsacheived'),
    'Implementation status': ('implementationstatus', 'implementationstate'),
    'Test result': ('testresult', 'testresults', 'teststatus', 'testingstatus'),
    'Test date': ('testdate', 'testeddate', 'lasttested', 'lasttestdate', 'testingdate'),
    'Owner': ('owner', 'assignedto', 'actionowner'),
    'Tenant ID': ('tenantid', 'tenantidentifier'),
    'Export date': ('exportdate', 'exportedon', 'reportdate', 'generateddate', 'exportedat'),
}


def _key(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())


def _text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def _baseline(value):
    text = _text(value).lower()
    return bool(re.search(r'\bai[\s_-]*baseline\b', text)
                or ('baseline' in text and ('microsoft 365 copilot' in text or 'copilot chat' in text)))


def _columns(values):
    keys = [_key(value) for value in values]
    # A metadata row is not a table header, even when labels repeat there.
    if not any(key in ALIASES['Improvement action'] for key in keys):
        return {}
    if not any(key in ALIASES['Implementation status'] + ALIASES['Test result'] for key in keys):
        return {}
    columns = {}
    for label, aliases in ALIASES.items():
        matches = [index for index, key in enumerate(keys) if key in aliases]
        if len(matches) > 1:
            raise ValueError(f'AI Baseline export has ambiguous {label} columns.')
        if matches:
            columns[label] = matches[0]
    return columns


def _identifier(value):
    if not _text(value):
        return ''
    try:
        return str(UUID(_text(value)))
    except ValueError:
        raise ValueError('AI Baseline export contains an invalid Tenant ID; retain the original tenant GUID.') from None


def _date(value):
    text = _text(value)
    if not text:
        return ''
    try:
        return datetime.fromisoformat(text.replace('Z', '+00:00')).isoformat() if 'T' in text else date.fromisoformat(text).isoformat()
    except ValueError:
        return ''  # Preserve the original value in Original fields; never guess regional dates.


def load_ai_baseline_export(path, expected_tenant_id=None):
    path = Path(path)
    if path.suffix.lower() != '.xlsx' or not path.is_file():
        raise ValueError('AI Baseline input must be an existing Compliance Manager Export actions .xlsx file.')
    if path.stat().st_size > 50 * 1024 * 1024:
        raise ValueError('AI Baseline export exceeds the 50 MiB input limit.')
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError) as exc:
        raise ValueError('AI Baseline input is not a readable Excel workbook.') from exc
    rows, identities, metadata = [], set(), {}
    baseline_name = ''
    tables = []
    try:
        from .ai_baseline_native import read_native_actions
        native = read_native_actions(workbook, path)
        if native is not None:
            identifier = ''  # This native layout does not export a tenant GUID.
            for row in native:
                row.update({'Tenant ID': identifier, 'Export date': '', 'Qualification': QUALIFICATION,
                            'Identity verification': 'Unverified: confirm the export belongs to this customer'})
            return {'available': bool(native), 'rows': native, 'tenant_id': identifier,
                    'source_file': path.name, 'source_hash': hashlib.sha256(path.read_bytes()).hexdigest(),
                    'assessment': native[0]['Assessment'] if native else 'AI Baseline',
                    'export_date': '', 'identity_verified': False, 'qualification': QUALIFICATION,
                    'format': 'Compliance Manager native multi-sheet export'}
        # Read all sheet preambles before actions, so a Summary sheet can identify
        # the assessment for separate Microsoft/customer action sheets.
        for sheet in workbook:
            preamble = list(islice(sheet.iter_rows(values_only=True), 40))
            sheet_metadata = {}
            sheet_baseline = sheet.title if _baseline(sheet.title) else ''
            found_table = False
            for number, values in enumerate(preamble, 1):
                columns = _columns(values)
                if ('Improvement action' in columns and
                        ('Implementation status' in columns or 'Test result' in columns)):
                    tables.append((sheet, number, values, columns, sheet_metadata, sheet_baseline))
                    found_table = True
                    break
                for index, value in enumerate(values):
                    text = _text(value)
                    if _baseline(text):
                        sheet_baseline = text
                    field = next((name for name, aliases in ALIASES.items() if _key(value) in aliases), None)
                    if field and index + 1 < len(values) and _text(values[index + 1]):
                        sheet_metadata[field] = values[index + 1]
                        if field == 'Tenant ID':
                            identities.add(_identifier(values[index + 1]))
            if not found_table:
                metadata.update(sheet_metadata)
                baseline_name = sheet_baseline or baseline_name
        if not tables:
            raise ValueError('AI Baseline export needs an Improvement action column and Implementation status or Test result column within the first 40 rows of an action sheet.')
        for sheet, header_row, headers, columns, sheet_metadata, sheet_baseline in tables:
            effective_metadata = {**metadata, **sheet_metadata}
            for number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                if not any(value is not None and _text(value) for value in values):
                    continue
                def value(label):
                    index = columns.get(label)
                    return values[index] if index is not None and index < len(values) else None
                if not _text(value('Improvement action')):
                    continue
                if len(rows) >= 10000:
                    raise ValueError('AI Baseline export exceeds 10,000 action rows; narrow the export scope.')
                assessment = _text(value('Assessment')) or _text(effective_metadata.get('Assessment')) or sheet_baseline or baseline_name
                regulation = _text(value('Regulation')) or _text(effective_metadata.get('Regulation'))
                if not _baseline(regulation or assessment):
                    raise ValueError('The export does not identify an AI Baseline assessment. Preserve its assessment/regulation metadata; filenames do not establish identity.')
                identifier = _identifier(value('Tenant ID') or effective_metadata.get('Tenant ID'))
                if identifier:
                    identities.add(identifier)
                row = {label: _text(value(label)) for label in ALIASES}
                row.update({'Assessment': assessment, 'Regulation': regulation, 'Tenant ID': identifier,
                            'Test date': _date(value('Test date')),
                            'Export date': _date(value('Export date') or effective_metadata.get('Export date')),
                            'Source file': path.name, 'Source sheet': sheet.title, 'Source row': number,
                            'Original fields': {str(header): _text(item) for header, item in zip(headers, values) if header is not None},
                            'Qualification': QUALIFICATION})
                if not row['Responsibility']:
                    row['Responsibility'] = ('Microsoft' if 'microsoft' in sheet.title.lower() else
                                             'Customer' if 'your' in sheet.title.lower() or 'customer' in sheet.title.lower() else '')
                rows.append(row)
        if not baseline_name and not rows and not any(table[5] for table in tables):
            raise ValueError('The empty export does not identify an AI Baseline assessment.')
        if len(identities) > 1:
            raise ValueError('AI Baseline export identifies multiple tenants. Use one tenant per build.')
        identifier = next(iter(identities), '')
        if identifier and expected_tenant_id and identifier != _identifier(expected_tenant_id):
            raise ValueError('AI Baseline export tenant ID does not match the assessed tenant. Use the tenant GUID for comparison.')
        verified = bool(identifier and expected_tenant_id)
        identity = 'Verified' if verified else 'Unverified: confirm the export belongs to this customer'
        for row in rows:
            row['Identity verification'] = identity
        return {'available': bool(rows), 'rows': rows, 'tenant_id': identifier,
                'source_file': path.name, 'source_hash': hashlib.sha256(path.read_bytes()).hexdigest(),
                'assessment': baseline_name or (rows[0]['Assessment'] if rows else ''),
                'export_date': _date(metadata.get('Export date')),
                'identity_verified': verified, 'qualification': QUALIFICATION}
    finally:
        workbook.close()


def is_ai_baseline_export(path):
    if Path(path).suffix.lower() != '.xlsx':
        return False
    try:
        load_ai_baseline_export(path)
        return True
    except (ValueError, OSError):
        return False


def load_ai_baseline_exports(paths, expected_tenant_id=None):
    rows, receipt, sources, hashes, tenants = [], [], [], set(), set()
    for path in dict.fromkeys(paths or []):
        data = load_ai_baseline_export(path, expected_tenant_id)
        if data['tenant_id']:
            tenants.add(data['tenant_id'])
        if len(tenants) > 1:
            raise ValueError('AI Baseline exports identify multiple tenants. Use one tenant per build.')
        duplicate = data['source_hash'] in hashes
        hashes.add(data['source_hash'])
        receipt.append({'Source': data['source_file'], 'Status': 'duplicate' if duplicate else 'accepted' if data['available'] else 'no_rows',
                        'Original Date': data['export_date'], 'Scope': 'Compliance Manager AI Baseline actions',
                        'Reason': 'Identical export; included once.' if duplicate else QUALIFICATION})
        if not duplicate:
            rows.extend(data['rows'])
            sources.append({key: value for key, value in data.items() if key not in {'rows', 'available', 'qualification'}})
    from .ai_baseline_priorities import prioritize_actions
    queue = prioritize_actions(rows)
    return {'available': bool(rows), 'rows': rows, 'sources': sources, 'receipt': receipt,
            'review_queue': queue,
            'tenant_id': next(iter(tenants), ''), 'qualification': QUALIFICATION,
            'test_results': dict(Counter(row['Test result'] or 'Not recorded' for row in rows)),
            'implementation_statuses': dict(Counter(row['Implementation status'] or 'Not recorded' for row in rows))}
