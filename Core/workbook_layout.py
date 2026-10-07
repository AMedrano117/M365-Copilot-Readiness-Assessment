"""Lay out two workbooks from the in-memory export, without reopening an XLSX.

InvestigationRange, RawEvidenceRanges and EvidenceRecordsRange describe the
shared evidence model and must remain unchanged. Only the new workbook fields
refer to this presentation. The source-row map also retains origins when old
per-finding log sheets are consolidated into a dataset register.
"""

from collections import OrderedDict, defaultdict
from copy import copy
from pathlib import Path
import json
import re
from hashlib import sha256

ASSESSMENT_ORDER = ('Start Here', 'Action Plan', 'Findings', 'Coverage', 'Configuration',
                    'Devices', 'Users & Identity', 'Apps & Consent', 'Sites & Sharing')
COMPATIBILITY_TITLES = ('Recommendations', 'Evidence Index', 'Control Results',
                        'Collection Coverage', 'Run Manifest', 'Integrity Checks',
                        # These optional registers are also read by the existing audit tools.
                        'Findings Register', 'App Access Detail', 'Admin Role Detail', 'AI Adoption Usage')
FINDING_COLUMNS = ('ID', 'Priority', 'Readiness effect', 'Area', 'Finding', 'What we found',
                   'What to do', 'Affected', 'Evidence', 'Observed', 'Owner', 'Status', 'Target date')
EVIDENCE_COLUMNS = {
    'Configuration': ('ID', 'Setting', 'Current', 'Recommended', 'Status', 'Observed', 'Evidence ID'),
    'Devices': ('ID', 'Device', 'Primary user', 'OS', 'Issue', 'Detail', 'Last reported', 'Intune compliance', 'Evidence ID'),
    'Users & Identity': ('ID', 'User', 'UPN', 'User type', 'Admin', 'Issue', 'Detail', 'Last sign-in or observed', 'Evidence ID'),
    'Apps & Consent': ('ID', 'App', 'App ID', 'Publisher / verified', 'Permission or grant', 'Consent type', 'Issue', 'Detail', 'Evidence ID'),
    'Sites & Sharing': ('ID', 'Site', 'URL', 'Owner', 'Exposure type', 'Measure', 'Detail', 'Evidence ID'),
    # Incidents, adoption metrics and custom service records have no truthful
    # entity projection onto the five layouts above. Keep them in this explicit
    # fallback instead of inventing a device, account, application or setting.
    'Other Evidence': ('ID', 'Record', 'Issue', 'Detail', 'Observed', 'Evidence ID'),
}
RANGE = re.compile(r"^'((?:[^']|'')+)'!([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?$")
ERROR_VALUES = {'#REF!', '#VALUE!', '#NAME?', '#DIV/0!'}


class WorkbookLayoutError(ValueError):
    """A workbook reconciliation failure must fail the build, not fall back to CSV."""


def technical_workbook_path(assessment_path):
    """Use a short sibling name while keeping repeated custom exports distinct."""
    from .export_paths import REPORT_STEM
    path = Path(assessment_path)
    stem = path.stem
    suffix = stem[len(REPORT_STEM):] if stem.startswith(REPORT_STEM) else ''
    name = 'Technical Evidence' + suffix if stem.startswith(REPORT_STEM) else stem + ' Technical Evidence'
    return path.with_name(name + '.xlsx')


def _rows(sheet):
    values = sheet.iter_rows(values_only=True)
    headers = next(values, ())
    return [dict(zip(headers, row)) for row in values]


def _write(workbook, title, rows, columns=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from .export_recommendations import _append_dict_rows_to_sheet
    sheet = workbook.create_sheet(title)
    if columns:
        rows = [OrderedDict((name, row.get(name)) for name in columns) for row in rows]
    if title in {'Findings', 'Action Plan', *EVIDENCE_COLUMNS}:
        from datetime import datetime, timezone
        date_columns = {'Observed', 'Last reported', 'Last sign-in or observed', 'Target date', 'Target Date'}
        rows = [dict(row) for row in rows]
        for row in rows:
            for name in date_columns & row.keys():
                value = row[name]
                if isinstance(value, str) and re.fullmatch(r'\d{4}-\d\d-\d\d(?:[T ].*)?', value):
                    try:
                        parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
                        row[name] = (parsed.astimezone(timezone.utc).replace(tzinfo=None) if parsed.tzinfo else parsed) if len(value) > 10 else parsed.date()
                    except ValueError:
                        pass
    _append_dict_rows_to_sheet(sheet, rows, PatternFill('solid', fgColor='102B40'),
                              Font(color='FFFFFF', bold=True), Alignment(vertical='top', wrap_text=True), title)
    if not rows and columns:
        sheet.append(list(columns))
        sheet.freeze_panes = 'A2'
    return sheet


def _clone(source, workbook, title=None):
    """Copy cells and links in memory; style objects belong to the destination."""
    target = workbook.create_sheet(title or source.title)
    for row in source:
        for cell in row:
            dest = target.cell(cell.row, cell.column, cell.value)
            dest.data_type = cell.data_type
            if cell.hyperlink:
                dest.hyperlink = copy(cell.hyperlink)
            if cell.comment:
                dest.comment = copy(cell.comment)
    from .export_recommendations import _add_excel_table
    _add_excel_table(target, target.title)
    target.freeze_panes = 'A2'
    return target


def _ref(title, first, last, width):
    from .investigation_details import _ref as location
    return location(title, first, last, width)


def _cell_link(sheet, number, field, target):
    columns = {cell.value: cell.column for cell in sheet[1]}
    if field not in columns or not target:
        return
    cell = sheet.cell(number, columns[field])
    cell.hyperlink = target
    cell.style = 'Hyperlink'


def _value(fields, *names):
    normalized = {re.sub(r'[^a-z0-9]', '', str(key).lower()): value for key, value in fields.items()}
    for name in names:
        item = normalized.get(re.sub(r'[^a-z0-9]', '', name.lower()))
        if item is not None and item != '':
            return item
    return None


def _text(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    return str(value)


def _label(value):
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    labels = {'noncompliant': 'Noncompliant', 'compliant': 'Compliant', 'unknown': 'Unknown',
              'notapplied': 'Not applied', 'notregistered': 'Not registered',
              'member': 'Member', 'guest': 'Guest', 'granted': 'Granted', 'denied': 'Denied',
              'disabled': 'Disabled', 'enabled': 'Enabled', 'true': 'Yes', 'false': 'No'}
    return labels.get(str(value).lower(), _text(value)) if value is not None else ''


def _detail(fields, names):
    return '; '.join(f'{label}: {_label(value)}' for label, keys in names
                     if (value := _value(fields, *keys)) is not None)


def _category(row, fields, dataset, kind):
    from .customer_report import SHARING_SETTING_LABELS
    from .finding_evidence import concern_for
    if _value(fields, 'setting', 'Setting Key', 'property') in SHARING_SETTING_LABELS:
        return 'Configuration'
    if dataset == 'compliance_setting_states' or _value(fields, 'settingName', 'Setting Name') is not None:
        return 'Devices'
    if kind == 'configuration' or dataset in {'security_defaults', 'authorization_policy', 'consent_policies',
                                             'ca_policies', 'sharepoint_tenant_settings', 'sharepoint_graph_settings'}:
        return 'Configuration'
    concern = concern_for(row)
    if concern == 'devices-endpoint' or dataset in {'windows_protection', 'managed_devices', 'machines', 'antivirus_health'}:
        return 'Devices'
    if concern in {'legacy-authentication', 'mfa-registration', 'conditional-access', 'admin-access', 'identity-risk'}:
        return 'Users & Identity'
    if concern in {'app-consent', 'connectors-agents', 'external-ai'}:
        return 'Apps & Consent'
    if concern == 'content-sharing' or dataset in {'sites', 'content_exposure', 'sensitive_exposure', 'label_coverage'}:
        return 'Sites & Sharing'
    if _value(fields, 'userPrincipalName', 'upn') is not None:
        return 'Users & Identity'
    return 'Other Evidence'


def _projection(row, record, devices):
    from .customer_report import SHARING_SETTING_LABELS, _sharing_setting_value, _sharing_check, _heading
    fields = record['fields']
    dataset = record.get('dataset') or ''
    kind = record.get('evidence_kind')
    title = _category(row, fields, dataset, kind)
    observed = _value(fields, 'lastReportedDateTime', 'lastSeen', 'createdDateTime', 'Created UTC', 'Observed At', 'Report Date', 'lastUpdatedDateTime') or row.get('ObservationDate')
    common = {'ID': row['RecommendationId'], 'Evidence ID': record['evidence_id']}
    issue = _heading(row)
    if title == 'Configuration':
        key = _value(fields, 'setting', 'Setting Key', 'property')
        current = _value(fields, 'value', 'Current Value', 'Current', 'state', 'status')
        recommended = _value(fields, 'recommended', 'Recommended Value') or (row.get('TechnicalGuidance') or {}).get('change') or row.get('Recommendation')
        status = 'Change' if row.get('Disposition') == 'Action' else 'Review'
        if key in SHARING_SETTING_LABELS:
            current = _sharing_setting_value(key, current)
            recommended, state = _sharing_check(key, current, record.get('anyone', True))
            status = {'change': 'Change', 'review': 'Review', 'ok': 'Meets'}[state]
        setting = SHARING_SETTING_LABELS.get(key) or key or _value(fields, 'displayName', 'Name', 'Setting', 'Policy') or issue
        if current is None:
            # A policy object has no single "value"; show its relevant retained
            # enforcement state rather than treating unknown configuration as 0.
            current = _detail(fields, [('State', ('enabled', 'isEnabled', 'state', 'Mode')),
                                      ('Result', ('result', 'review_result'))]) or 'Unknown'
        return title, dict(common, Setting=setting, Current=_label(current), Recommended=recommended,
                           Status=status, Observed=observed)
    if title == 'Devices':
        parent = _value(fields, 'ParentId', 'managedDeviceId', 'deviceId', 'id')
        inventory = devices.get(_text(parent).casefold(), {})
        merged = {**inventory, **fields}
        device = _value(merged, 'deviceName', 'computerDnsName', 'Device Name', 'Device', 'displayName')
        primary_user = _value(merged, 'userPrincipalName', 'Primary User', 'userDisplayName')
        os_name = _value(merged, 'operatingSystem', 'osPlatform', 'OS')
        version = _value(merged, 'osVersion', 'OS Version')
        device_issue = _value(fields, 'settingName', 'Setting Name')
        if _value(fields, 'realTimeProtectionEnabled') is False:
            device_issue = 'Real-time protection off'
        details = _detail(fields, [('Tamper protection', ('isTamperProtected', 'tamperProtectionEnabled')),
                                  ('Signatures overdue', ('signatureUpdateOverdue',)),
                                  ('Signatures current', ('avIsSignatureUpToDate',)),
                                  ('State', ('state', 'Status')), ('Risk', ('riskScore',)),
                                  ('Protection', ('healthStatus',)), ('Policy', ('policyName', 'Policy'))])
        return title, dict(common, Device=device or 'Unknown', **{'Primary user': primary_user,
                           'OS': ' '.join(_text(item) for item in (os_name, version) if item is not None),
                           'Issue': device_issue or issue, 'Detail': details, 'Last reported': observed,
                           'Intune compliance': _label(_value(merged, 'complianceState', 'Compliance State'))})
    if title == 'Users & Identity':
        return title, dict(common, User=_value(fields, 'userDisplayName', 'displayName', 'User', 'Principal', 'Entity'),
                           UPN=_value(fields, 'userPrincipalName', 'upn', 'User Principal Name'),
                           **{'User type': _label(_value(fields, 'userType', 'User Type')),
                              'Admin': _label(_value(fields, 'isAdmin', 'Admin')), 'Issue': issue,
                              'Detail': _detail(fields, [('MFA registered', ('isMfaRegistered',)),
                                                       ('Methods', ('methodsRegistered', 'authenticationMethods')),
                                                       ('Role', ('roleName', 'Role', 'Role Name')),
                                                       ('Risk', ('riskLevel', 'riskState')),
                                                       ('Client', ('clientAppUsed',)), ('Outcome', ('outcome', 'Record Status')),
                                                       ('Application', ('appDisplayName',)), ('Result', ('state',))]),
                              'Last sign-in or observed': observed})
    if title == 'Apps & Consent':
        return title, dict(common, App=_value(fields, 'appDisplayName', 'displayName', 'App Display Name', 'appName', 'Entity'),
                           **{'App ID': _value(fields, 'appId', 'App ID', 'servicePrincipalId', 'id'),
                              'Publisher / verified': _detail(fields, [('Publisher', ('publisherName', 'Publisher')),
                                                                      ('Verified', ('verifiedPublisher', 'Publisher Verified'))]),
                              'Permission or grant': _value(fields, 'scope', 'permission', 'Permission', 'appRoleId', 'Permissions'),
                              'Consent type': _value(fields, 'consentType', 'Consent Type'), 'Issue': issue,
                              'Detail': _detail(fields, [('Reason', ('Flagged Because', 'reason')),
                                                       ('Activity', ('Last Activity', 'lastSignInDateTime')), ('State', ('state', 'status'))])})
    if title == 'Sites & Sharing':
        measure = _value(fields, 'measure', 'Metric', 'metric', 'reportedMeasure')
        count = _value(fields, 'count', 'value', 'Value')
        if measure is None:
            measures = [(key, value) for key, value in fields.items() if isinstance(value, (int, float)) and not isinstance(value, bool)]
            measure = '; '.join(f'Reported {key}: {value}' for key, value in measures)
        elif count is not None:
            measure = f'Reported {measure}: {count}'
        return title, dict(common, Site=_value(fields, 'siteName', 'Site Name', 'Title', 'Name', 'Site'),
                           URL=_value(fields, 'siteUrl', 'Site URL', 'Url', 'URL'),
                           Owner=_value(fields, 'owner', 'Owner', 'Primary Admin', 'Owner Email'),
                           **{'Exposure type': _value(fields, 'exposureType', 'reportType', 'type') or issue,
                              'Measure': measure, 'Detail': _detail(fields, [('Scope', ('scope', 'Report scope')),
                                                                          ('Qualification', ('qualification', 'reason'))])})
    return title, dict(common, Record=_value(fields, 'displayName', 'Name', 'title', 'id', 'Record ID') or record.get('detail_record_id'),
                       Issue=issue, Detail='; '.join(f'{key}: {_label(value)}' for key, value in fields.items()
                                                   if value is not None), Observed=observed)


def _affected(row, records, finding):
    if not records:
        return 'Unknown'
    datasets = {record.get('dataset') for record in records}
    if datasets == {'compliance_setting_states'} or all(_value(record['fields'], 'settingName', 'Setting Name') is not None for record in records):
        return f'{len(records)} settings'
    if all(_projection_kind(record) == 'setting' for record in records):
        return f'{len(records)} settings'
    evidence = (finding or {}).get('evidence') or {}
    if datasets and datasets <= {'windows_protection', 'managed_devices', 'machines', 'antivirus_health', 'directory_devices'}:
        identifiers = [_value(record['fields'], 'ParentId', 'managedDeviceId', 'deviceId', 'id') for record in records]
        if all(value is not None for value in identifiers):
            return f'{len({_text(value) for value in identifiers})} devices'
    if 'compliance_setting_states' not in datasets and evidence.get('affected_entity_count') is not None:
        return f"{evidence['affected_entity_count']} {evidence.get('entity_unit') or 'entities'}"
    return f"{len(records)} {evidence.get('record_unit') or row.get('RecordUnit') or 'records'}"


def _projection_kind(record):
    from .customer_report import SHARING_SETTING_LABELS
    return 'setting' if (_value(record['fields'], 'setting', 'Setting Key', 'property') in SHARING_SETTING_LABELS
                         or record.get('evidence_kind') == 'configuration') else 'record'


def _coverage(result, bundle):
    controls = {str(row.get('Control ID') or ''): row for row in bundle.get('control_results') or result.get('control_results') or []}
    required = list(result.get('domain_coverage') or [])
    present = {str(row.get('Check ID') or row.get('Control ID') or '') for row in required}
    required.extend(row for key, row in controls.items() if key not in present)
    output = []
    for row in required:
        identifier = row.get('Check ID') or row.get('Control ID') or row.get('Control')
        control = controls.get(str(row.get('Control ID') or identifier), {})
        text = _text(row.get('Failures and limitations')).lower()
        source_states = row.get('Source metadata') or []
        row_state = str(row.get('State') or row.get('Status') or '').lower()
        if any(source.get('truncated') is True for source in source_states) or 'truncated' in text:
            state = 'truncated'
        elif 'stale' in text or 'expired' in text:
            state = 'stale'
        elif row.get('Owner review') or row_state in {'reviewed', 'not_applicable'}:
            state = 'owner review'
        elif row_state in {'partial', 'conflict'} or 'incomplete' in text or 'partial' in text:
            state = 'incomplete'
        elif row.get('Missing evidence') or row_state in {'missing', 'not assessed', 'not established', 'unavailable', 'not collected'}:
            state = 'missing'
        else:
            state = 'available'
        output.append({'Check ID': identifier, 'Check': row.get('Check') or row.get('Description') or row.get('Control'),
                       'Area': row.get('Domain') or control.get('Domain'), 'Evidence state': state,
                       'Applicability': row.get('Applicability'),
                       'Readiness effect': row.get('Readiness effect') or row.get('Pilot impact') or row.get('AI readiness effect'),
                       'Configuration': row.get('Configuration') or control.get('Configuration result'),
                       'Operational result': row.get('Operational result') or control.get('Operational result'),
                       'Result': control.get('Result') or row.get('Status'),
                       'Missing evidence': row.get('Missing evidence'), 'Owner': row.get('Owner') or row.get('Responsible role')})
    return output


def split_workbook_layout(source_workbook, result, bundle, assessment_path, technical_path, *, tenant_name=None):
    """Project the compact workbook and retain every source row in a visible companion."""
    from openpyxl import Workbook
    from .dashboard_export import build_dashboard_export
    from .finding_evidence import build_finding_evidence
    from .customer_report import _heading
    from .workbook_navigation import apply_workbook_navigation
    dashboard = build_dashboard_export(result, bundle, tenant_name=tenant_name)
    model = bundle.get('finding_evidence') or build_finding_evidence(dashboard)
    model_findings = {item['finding_id']: item for item in model['findings']}
    finding_rows = {row['RecommendationId']: row for row in result['recommendations']}
    assessment, technical = Workbook(), Workbook()
    assessment.remove(assessment.active)
    technical.remove(technical.active)
    source_map, sheet_map = {}, {}
    # Logs generated by a producer may have arbitrary titles. Keys identify
    # per-finding sheets; titles are never guessed from the finding's wording.
    source_sheets = bundle.get('sheets') or {}
    declared = {sheet['title']: key for key, sheet in source_sheets.items() if key.startswith('declared_investigation.')}
    grouped_logs = OrderedDict()
    log_titles = {}
    private_dataset_names = set()
    for sheet in source_workbook:
        if sheet.title in {'Start Here', 'Action Plan', 'Findings Register'}:
            continue
        if sheet.title in declared:
            identifier = declared[sheet.title].split('declared_investigation.', 1)[1].split('.part', 1)[0]
            declaration = (finding_rows.get(identifier) or {}).get('InvestigationEvidence') or {}
            origin = declaration.get('source') or {}
            finding = model_findings.get(identifier) or {}
            native_datasets = [model['tables'][key].get('dataset') for key in (finding.get('evidence') or {}).get('tables') or []
                               if not str(model['tables'][key].get('dataset')).startswith(('declared.', 'investigation.'))]
            dataset = next(iter(dict.fromkeys(native_datasets)), None) or origin.get('dataset') or origin.get('api') or origin.get('source_api') or origin.get('name') or origin.get('file') or 'Declared records'
            if dataset == origin.get('file'):
                private_dataset_names.add(_text(dataset))
            for number, row in enumerate(_rows(sheet), 2):
                row = {'RecommendationId': row.get('RecommendationId') or identifier,
                       **{key: value for key, value in row.items() if key != 'RecommendationId'}}
                grouped_logs.setdefault(_text(dataset), []).append((sheet.title, number, row))
            continue
        target = _clone(sheet, technical)
        sheet_map[sheet.title] = target.title
        for number in range(2, sheet.max_row + 1):
            source_map[(sheet.title, number)] = (target.title, number)
    for dataset, members in grouped_logs.items():
        # A stable dataset title avoids one worksheet per finding. Excel's name
        # limit applies to titles only; the full dataset stays in each row.
        from .investigation_contract import _title
        label = 'Declared records' if dataset in private_dataset_names else dataset.rsplit('/', 1)[-1].replace('_', ' ')
        title = _title('Logs ' + label, {name.casefold() for name in technical.sheetnames})
        target = _write(technical, title, [row for _, _, row in members])
        log_titles[dataset] = target.title
        for number, (origin, source_number, _) in enumerate(members, 2):
            source_map[(origin, source_number)] = (target.title, number)
            sheet_map.setdefault(origin, target.title)

    # Exact native evidence identifiers index the raw inventory. Derived records
    # also receive a retained raw row, so every assessment Evidence ID has an
    # exact, visible technical destination even without a native collector.
    evidence_targets = {}
    native_fields = {}
    for sheet in technical:
        headers = {cell.value: cell.column for cell in sheet[1]}
        if 'Evidence Record ID' in headers:
            for number in range(2, sheet.max_row + 1):
                identifier = sheet.cell(number, headers['Evidence Record ID']).value
                evidence_targets[identifier] = (sheet.title, number)
    evidence_records = {record['record_id']: record for record in dashboard.get('evidence_records') or []}
    derived = []
    for identifier, record in evidence_records.items():
        raw = record.get('raw')
        native_fields[identifier] = raw if isinstance(raw, dict) else {'value': raw}
        if identifier not in evidence_targets:
            derived.append({'Evidence Record ID': identifier, 'Source dataset': record.get('dataset'),
                            **native_fields[identifier]})
    if derived:
        sheet = _write(technical, 'Raw Derived Records', derived)
        for number, row in enumerate(derived, 2):
            evidence_targets[row['Evidence Record ID']] = (sheet.title, number)

    selected = defaultdict(list)
    for table in model['tables'].values():
        if table['role'] != 'selected':
            continue
        for record in table['rows']:
            # COMMON_COLUMNS can grow; use its actual names rather than assuming
            # every source has a particular schema or particular number of fields.
            from .finding_evidence import COMMON_COLUMNS
            fields = {key.removesuffix(' (record field)'): value for key, value in record.items() if key not in COMMON_COLUMNS}
            identifiers = [part.strip() for part in _text(record.get('evidence_record_ids')).split(';') if part.strip()]
            selected[record['finding_id']].append({'fields': fields, 'dataset': record.get('dataset'),
                                                  'evidence_kind': record.get('evidence_kind'),
                                                  'detail_record_id': record.get('detail_record_id'),
                                                  'evidence_ids': identifiers,
                                                  'evidence_id': next(iter(identifiers), record.get('detail_record_id'))})
    # Older/imported bundles can retain selected worksheet rows without their
    # native assessment_sources object. The dashboard deliberately cannot infer
    # missing native events. Retain those existing selected rows here, with their
    # worksheet provenance, without relabelling them as a recollected raw API.
    from .raw_evidence import KEY_SOURCES
    by_title = {sheet['title']: (key, sheet) for key, sheet in source_sheets.items()}
    touched = set()
    for identifier, row in finding_rows.items():
        evidence = (model_findings.get(identifier) or {}).get('evidence') or {}
        if selected.get(identifier) or not row.get('InvestigationCount') or evidence.get('kind') == 'supporting_context':
            continue
        seen = set()
        def retain(location, depth=0):
            match = RANGE.fullmatch(_text(location))
            if not match or depth > 2:
                return
            title = match[1].replace("''", "'")
            key, source = by_title.get(title, ('', {}))
            data = source.get('rows') or []
            for number in range(int(match[3]), int(match[5] or match[3]) + 1):
                if (title, number) in seen or number < 2 or number > len(data) + 1:
                    continue
                seen.add((title, number))
                record = data[number - 2]
                owners = {part.strip() for part in _text(record.get('RecommendationId')).split(';') if part.strip()}
                if owners and identifier not in owners:
                    continue
                if title == 'Investigation Items' and record.get('Source Detail'):
                    retain(record['Source Detail'], depth + 1)
                    continue
                fields = record.get('Raw record')
                if isinstance(fields, str):
                    chunks = [fields]
                    index = 2
                    while f'Raw record (continued {index})' in record:
                        chunks.append(record[f'Raw record (continued {index})'])
                        index += 1
                    try:
                        fields = json.loads(''.join(chunks))
                    except ValueError:
                        fields = None
                fields = fields if isinstance(fields, dict) else dict(record)
                evidence_id = record.get('Evidence Record ID') or record.get('Evidence ID') or 'WB-' + sha256(
                    json.dumps([title, number, fields], ensure_ascii=False, sort_keys=True, default=str).encode('utf-8')).hexdigest()[:24]
                destination = source_map.get((title, number))
                if not destination:
                    continue
                evidence_targets[evidence_id] = destination
                native_fields[evidence_id] = fields
                target = technical[destination[0]]
                headers = {cell.value: cell.column for cell in target[1]}
                if evidence_id not in [cell.value for cell in target[destination[1]]]:
                    column = headers.get('Workbook Evidence ID') or target.max_column + 1
                    target.cell(1, column, 'Workbook Evidence ID')
                    target.cell(destination[1], column, evidence_id)
                    touched.add(target.title)
                dataset = next(iter(KEY_SOURCES.get(key.split('.part', 1)[0], [])), key) or 'Retained details'
                selected[identifier].append({'fields': fields, 'dataset': dataset, 'evidence_kind': evidence.get('kind'),
                                            'detail_record_id': evidence_id, 'evidence_ids': [evidence_id], 'evidence_id': evidence_id})
        for location in row.get('InvestigationRanges') or [row.get('InvestigationRange')]:
            retain(location)
    if touched:
        from .export_recommendations import _add_excel_table
        for title in touched:
            sheet = technical[title]
            sheet.tables.clear()
            _add_excel_table(sheet, title)
    # Sharing findings often retain one wide tenant response. Project its
    # relevant setting properties separately, retaining the exact source ID on
    # every row. This intentionally changes the presentation unit to settings;
    # it leaves model row counts and all original lineage ranges untouched.
    from .customer_report import SHARING_SETTING_LABELS
    for identifier, records in list(selected.items()):
        row = finding_rows[identifier]
        if not str(row.get('FindingKey') or '').startswith('sharepoint.'):
            continue
        expanded = []
        key = row.get('FindingKey')
        relevant = ({'DefaultSharingLinkType'} if key == 'sharepoint.sharing.organization_default' else
                    {'LegacyAuthProtocolsEnabled'} if key == 'sharepoint.authentication.legacy_permitted' else
                    {'SharingCapability', 'OneDriveSharingCapability', 'DefaultSharingLinkType',
                     'RequireAnonymousLinksExpireInDays', 'FileAnonymousLinkType', 'FolderAnonymousLinkType'}
                    if key in {'sharepoint.sharing.permissive_anonymous_defaults', 'sharepoint.sharing.anyone_enabled'} else
                    set(SHARING_SETTING_LABELS))
        for record in records:
            native = next((native_fields[value] for value in record['evidence_ids'] if value in native_fields), record['fields'])
            values = {name: _value(native, name) for name in SHARING_SETTING_LABELS if _value(native, name) is not None}
            anyone = any(_text(values.get(name)).lower() in {'2', 'externaluserandguestsharing'} for name in ('SharingCapability', 'OneDriveSharingCapability'))
            settings = [name for name in SHARING_SETTING_LABELS if name in values and name in relevant]
            if settings:
                expanded.extend({**record, 'fields': {'setting': name, 'value': values[name]},
                                 'evidence_kind': 'configuration', 'anyone': anyone} for name in settings)
            else:
                expanded.append(record)
        selected[identifier] = expanded
    devices = {}
    for dataset in ('managed_devices', 'machines', 'directory_devices'):
        for entry in (bundle.get('assessment_sources') or {}).get(dataset, []):
            for record in entry.get('records') or []:
                for key in ('id', 'deviceId', 'managedDeviceId', 'azureADDeviceId', 'aadDeviceId'):
                    if record.get(key):
                        devices[_text(record[key]).casefold()] = record

    # Full selected source records, consolidated per dataset, retain all returned
    # fields. A shared source record is repeated once for each finding it supports.
    full = OrderedDict()
    full_members = defaultdict(list)
    for identifier, records in selected.items():
        seen = set()
        for record in records:
            for evidence_id in record['evidence_ids'] or [record['evidence_id']]:
                dataset = record['dataset'] or 'Retained details'
                if str(dataset).startswith(('declared.', 'investigation.')):
                    declaration = (finding_rows.get(identifier) or {}).get('InvestigationEvidence') or {}
                    origin = declaration.get('source') or {}
                    dataset = origin.get('dataset') or origin.get('api') or origin.get('source_api') or origin.get('name') or origin.get('file') or 'Retained details'
                dataset = _text(dataset)
                marker = (dataset, evidence_id)
                if marker in seen:
                    continue
                seen.add(marker)
                fields = native_fields.get(evidence_id) or record['fields']
                data = {'RecommendationId': identifier, 'Evidence ID': evidence_id, 'Detail record ID': record['detail_record_id']}
                data.update(('Record.' + key if key in data else key, value) for key, value in fields.items())
                full.setdefault(dataset, []).append(data)
                full_members[identifier].append((dataset, len(full[dataset]) + 1))
                if evidence_id not in evidence_targets:
                    evidence_targets[evidence_id] = None  # filled by the selected-record register
    record_titles = {}
    full_destinations = {}
    for dataset, rows in full.items():
        from .investigation_contract import _title
        # Derived dataset names include an ID, but their grouping is by source
        # declaration, not by finding, just as for the older Logs registers.
        group_name = 'Retained details' if str(dataset).startswith(('declared.', 'investigation.')) else dataset
        title = log_titles.get(dataset)
        if title:
            # Enrich the consolidated legacy rows in place instead of creating
            # a second selected-record table for the same source dataset. All
            # old log columns/values survive; EVD IDs occupy a distinct column
            # from the older INV Evidence ID bookkeeping.
            sheet = technical[title]
            existing = _rows(sheet)
            used = set()
            headers = {cell.value: cell.column for cell in sheet[1]}
            from .export_recommendations import _excel_safe_rows
            for index, data in enumerate(_excel_safe_rows(rows, title), 2):
                native_id = _value(data, 'id', 'Record ID', 'source_record_id')
                number = next((position for position, old in enumerate(existing, 2)
                               if position not in used and old.get('RecommendationId') == data['RecommendationId']
                               and (native_id is None or _value(old, 'id', 'Record ID', 'source_record_id') == native_id)), None)
                if number is None:
                    number = sheet.max_row + 1
                used.add(number)
                for field, value in data.items():
                    field = 'Evidence Record ID' if field == 'Evidence ID' else field
                    if field not in headers:
                        headers[field] = sheet.max_column + 1
                        sheet.cell(1, headers[field], field)
                    cell = sheet.cell(number, headers[field])
                    if cell.value is None:
                        cell.value = value if not isinstance(value, (dict, list, tuple)) else json.dumps(value, ensure_ascii=False, sort_keys=True)
                        if isinstance(cell.value, str):
                            cell.data_type = 's'
                full_destinations[(dataset, index)] = (title, number)
            from .export_recommendations import _add_excel_table
            sheet.tables.clear()
            _add_excel_table(sheet, title)
        else:
            title = _title('Records ' + group_name.replace('_', ' ').title(), {name.casefold() for name in technical.sheetnames})
            sheet = _write(technical, title, rows)
            full_destinations.update({(dataset, number): (title, number) for number in range(2, len(rows) + 2)})
        record_titles[dataset] = title
        for number, row in enumerate(rows, 2):
            if not evidence_targets.get(row['Evidence ID']):
                evidence_targets[row['Evidence ID']] = full_destinations[(dataset, number)]

    def remap(location, external=False):
        match = RANGE.fullmatch(_text(location))
        if not match:
            return []
        original = match[1].replace("''", "'")
        first, last = int(match[3]), int(match[5] or match[3])
        targets = [source_map.get((original, number)) for number in range(first, last + 1)]
        targets = [item for item in targets if item]
        blocks = []
        for title, number in targets:
            if blocks and blocks[-1][0] == title and blocks[-1][2] + 1 == number:
                blocks[-1][2] = number
            else:
                blocks.append([title, number, number])
        prefix = Path(technical_path).name + '#' if external else ''
        return [prefix + _ref(title, start, end, technical[title].max_column) for title, start, end in blocks]

    # Replace hyperlink destinations (never the shared range fields) when a
    # source sheet moved or was consolidated. Remove links that cannot resolve.
    for sheet in technical:
        for cells in sheet:
            for cell in cells:
                if cell.hyperlink and cell.hyperlink.target and cell.hyperlink.target.startswith('#'):
                    locations = remap(cell.hyperlink.target[1:])
                    if not locations:
                        old = cell.hyperlink.target[1:].split('!')[0].strip("'").replace("''", "'")
                        title = sheet_map.get(old)
                        locations = [_ref(title, 1, 1, 1)] if title else []
                    cell.hyperlink = '#' + locations[0] if locations else None

    assessment_rows = defaultdict(list)
    evidence_links = {}
    reconciliation = []
    issues = []
    for row in result['recommendations']:
        identifier = row['RecommendationId']
        row['AssessmentEvidenceRange'] = ''
        row['AssessmentEvidenceRanges'] = []
        row['TechnicalEvidenceRanges'] = []
        records = selected.get(identifier, [])
        for record in records:
            title, projected = _projection(row, record, devices)
            assessment_rows[title].append(projected)
            destination = evidence_targets.get(record['evidence_id'])
            if destination:
                evidence_links[(title, len(assessment_rows[title]) + 1)] = Path(technical_path).name + '#' + _ref(destination[0], destination[1], destination[1], 1)
            else:
                issues.append(f'{identifier}: unmatched evidence record {record["evidence_id"]}.')
        for dataset, number in full_members[identifier]:
            title, destination_row = full_destinations[(dataset, number)]
            row['TechnicalEvidenceRanges'].append(_ref(title, destination_row, destination_row, technical[title].max_column))
        # Compact adjacent ranges so large populations do not create megabyte
        # lineage cells and HTML links. Raw/lineage references remain available.
        row['TechnicalEvidenceRanges'] = _compact_ranges(row['TechnicalEvidenceRanges'])
        if not row['TechnicalEvidenceRanges']:
            for location in [*(row.get('InvestigationRanges') or [row.get('InvestigationRange')]),
                             *(row.get('RawEvidenceRanges') or [])]:
                row['TechnicalEvidenceRanges'].extend(remap(location))
        row['TechnicalEvidenceRanges'] = list(dict.fromkeys(row['TechnicalEvidenceRanges']))

    for title in EVIDENCE_COLUMNS:
        if not assessment_rows[title]:
            continue
        sheet = _write(assessment, title, assessment_rows[title], EVIDENCE_COLUMNS[title])
        blocks = OrderedDict()
        for number, row in enumerate(assessment_rows[title], 2):
            blocks.setdefault(row['ID'], []).append(number)
            _cell_link(sheet, number, 'Evidence ID', evidence_links.get((title, number)))
        for identifier, numbers in blocks.items():
            if numbers != list(range(numbers[0], numbers[-1] + 1)):
                issues.append(f'{identifier}: assessment evidence rows are not contiguous.')
            finding_rows[identifier]['AssessmentEvidenceRanges'].append(_ref(title, numbers[0], numbers[-1], sheet.max_column))
    for row in result['recommendations']:
        identifier = row['RecommendationId']
        ranges = row['AssessmentEvidenceRanges']
        if ranges:
            row['AssessmentEvidenceRange'] = max(ranges, key=lambda location: int(RANGE.fullmatch(location)[5]) - int(RANGE.fullmatch(location)[3]))
        written = sum(int(RANGE.fullmatch(location)[5]) - int(RANGE.fullmatch(location)[3]) + 1 for location in ranges)
        expected = len(selected.get(identifier, []))
        reconciliation.append({'Status': 'Passed' if expected == written else 'Failed', 'Check': 'Finding evidence rows',
                               'Source': identifier, 'Source rows': expected, 'Written rows': written,
                               'Issue': '' if expected == written else 'Affected records do not reconcile.'})
        if written != expected:
            issues.append(f'{identifier}: affected evidence rows do not reconcile.')

    ordered = sorted(result['recommendations'], key=lambda row: (
        {'Action': 0, 'Coverage': 1, 'Assurance': 2, 'Opportunity': 3, 'Reference': 4}.get(row.get('Disposition'), 5),
        {'Critical': 0, 'High': 1, 'Medium': 2, 'Low': 3}.get(row.get('Priority'), 4), _heading(row)))
    findings = []
    for row in ordered:
        affected = _affected(row, selected.get(row['RecommendationId'], []), model_findings.get(row['RecommendationId']))
        row['AssessmentAffected'] = affected
        row['AssessmentEvidenceCount'] = len(selected.get(row['RecommendationId'], []))
        findings.append(dict(zip(FINDING_COLUMNS, (row['RecommendationId'], row.get('Priority'), row.get('PilotImpact'),
                        row.get('Domain') or row.get('ImpactArea'), _heading(row), row.get('Observation'),
                        row.get('Recommendation'), affected, affected + ' ›' if row['AssessmentEvidenceRange'] else
                        row.get('InvestigationStatus') or 'Evidence unavailable', row.get('ObservationDate'),
                        row.get('OwnerRole'), 'Open', None))))
    register = _write(assessment, 'Findings', findings, FINDING_COLUMNS)
    for number, row in enumerate(ordered, 2):
        _cell_link(register, number, 'Evidence', '#' + row['AssessmentEvidenceRange'] if row['AssessmentEvidenceRange'] else None)
    action_rows = []
    for number, action in enumerate(result.get('actions') or [], 1):
        row = finding_rows.get(action['RecommendationId'], action)
        # Action numbers use result.actions, the exact ordering used by HTML.
        action.update({key: row[key] for key in ('AssessmentEvidenceRange', 'AssessmentEvidenceRanges',
                                                'TechnicalEvidenceRanges', 'AssessmentEvidenceCount', 'AssessmentAffected') if key in row})
        action_rows.append({'Action': number, 'RecommendationId': row['RecommendationId'], 'Priority': row.get('Priority'),
                            'Finding': _heading(row), 'What We Found': row.get('Observation'),
                            'Recommended Action': row.get('Recommendation'), 'Affected': row.get('AssessmentAffected'),
                            'Evidence': row.get('AssessmentEvidenceRange') or row.get('InvestigationStatus'),
                            'Responsible Role': row.get('OwnerRole'), 'Rollout Stage': row.get('ReadinessStage'),
                            'Target Date': None, 'Status': 'Open', 'Completion Evidence': row.get('CompletionEvidence'),
                            'Observed': row.get('ObservationDate'),
                            'Qualification': ' '.join(_text(row.get(key)) for key in ('Qualification', 'InvestigationQualification') if row.get(key)),
                            'Investigation Details': (row.get('InvestigationSummary') or row.get('InvestigationStatus'))
                            if row.get('InvestigationCount') else ': '.join(dict.fromkeys(_text(row.get(key)) for key in
                            ('InvestigationStatus', 'InvestigationSummary') if row.get(key)))})
    actions = _write(assessment, 'Action Plan', action_rows or [{'What We Found': 'No deployment actions were identified from the evidence collected.',
                                                               'Recommended Action': 'Continue monitoring the tenant as conditions and intended AI use cases change.'}])
    for number, action in enumerate(result.get('actions') or [], 2):
        location = action.get('AssessmentEvidenceRange')
        _cell_link(actions, number, 'Evidence', '#' + location if location else None)
        _cell_link(actions, number, 'Investigation Details', '#' + location if location else None)
        from openpyxl.comments import Comment
        if action.get('InvestigationNote'):
            column = next((cell.column for cell in actions[1] if cell.value == 'Investigation Details'), None)
            if column:
                actions.cell(number, column).comment = Comment(action['InvestigationNote'], 'Assessment')
    _write(assessment, 'Coverage', _coverage(result, bundle),
           ('Check ID', 'Check', 'Area', 'Evidence state', 'Applicability', 'Readiness effect', 'Configuration',
            'Operational result', 'Result', 'Missing evidence', 'Owner'))

    lineage_rows = []
    legacy = {row.get('RecommendationId'): row for row in _rows(source_workbook['Findings Register'])} if 'Findings Register' in source_workbook else {}
    for row in ordered:
        lineage = dict(legacy.get(row['RecommendationId']) or row)
        lineage['RecommendationId'] = row['RecommendationId']
        lineage.update({'Assessment evidence': row['AssessmentEvidenceRanges'],
                        'Technical evidence': row['TechnicalEvidenceRanges']})
        for field in ('Supporting evidence', 'Raw source evidence', 'Evidence records'):
            if field in lineage:
                locations = remap(lineage[field])
                lineage[field] = next(iter(locations), lineage[field])
        for field in ('Supporting ranges', 'Raw source ranges', 'Evidence record ranges'):
            if field in lineage:
                lineage[field] = [target for location in lineage[field] or [] for target in remap(location)]
        lineage_rows.append(lineage)
    lineage = _write(technical, 'Findings Lineage', lineage_rows)
    for number, row in enumerate(ordered, 2):
        for field in ('Supporting evidence', 'Raw source evidence', 'Evidence records'):
            locations = remap((legacy.get(row['RecommendationId']) or {}).get(field))
            _cell_link(lineage, number, field, '#' + locations[0] if locations else None)
        _cell_link(lineage, number, 'Technical evidence', '#' + next(iter(row['TechnicalEvidenceRanges'])) if row['TechnicalEvidenceRanges'] else None)
        _cell_link(lineage, number, 'Assessment evidence', Path(assessment_path).name + '#' + row['AssessmentEvidenceRange'] if row['AssessmentEvidenceRange'] else None)
    if 'Technical Recommendations' in technical:
        fixes = technical['Technical Recommendations']
        positions = {row['RecommendationId']: number for number, row in enumerate(ordered, 2)}
        fix_positions = {}
        for number, fix in enumerate(_rows(fixes), 2):
            identifier = fix.get('RecommendationId')
            fix_positions[identifier] = number
            if identifier in positions:
                _cell_link(fixes, number, 'RecommendationId', f"#'Findings Lineage'!A{positions[identifier]}")
        for identifier, number in positions.items():
            if identifier in fix_positions:
                _cell_link(lineage, number, 'Technical fix', f"#'Technical Recommendations'!A{fix_positions[identifier]}")

    # Source-row reconciliation uses the actual destination cells, including
    # consolidated logs. Identical source values and field names are retained.
    for key, source in source_sheets.items():
        title = source['title']
        rows = source.get('rows') or []
        written = sum(bool(destination := source_map.get((title, number))) and
                      destination[1] <= technical[destination[0]].max_row for number in range(2, len(rows) + 2))
        reconciliation.append({'Status': 'Passed' if written == len(rows) else 'Failed', 'Check': 'Source worksheet rows',
                               'Source': title, 'Source rows': len(rows), 'Written rows': written,
                               'Issue': '' if written == len(rows) else 'Source rows do not reconcile.'})
        if written != len(rows):
            issues.append(f'{title}: source rows do not reconcile.')
    reconciliation.append({'Status': 'Passed' if not issues else 'Failed', 'Check': 'Unmatched records',
                           'Source': 'Workbook layout', 'Source rows': 0, 'Written rows': len(issues),
                           'Issue': '; '.join(issues)})
    for title in COMPATIBILITY_TITLES:
        if title == 'Integrity Checks':
            continue
        if title in source_workbook:
            target = _clone(source_workbook[title], assessment)
        elif title in COMPATIBILITY_TITLES[:6]:
            columns = {'Control Results': ('Control ID', 'Status', 'Result'),
                       'Run Manifest': ('Item', 'Value'), 'Collection Coverage': ('Source', 'State', 'Records')}.get(title, ('RecommendationId',))
            target = _write(assessment, title, [], columns)
        else:
            continue
        target.sheet_state = 'hidden'
        for cells in target:
            for cell in cells:
                if cell.hyperlink and cell.hyperlink.target and cell.hyperlink.target.startswith('#'):
                    locations = remap(cell.hyperlink.target[1:], external=True)
                    cell.hyperlink = locations[0] if locations else None
    manifest_rows = _rows(assessment['Run Manifest'])
    present = {row.get('Item') for row in manifest_rows}
    manifest_rows.extend({'Item': item, 'Value': value} for item, value in
                         [('Tenant', tenant_name), ('Tenant ID', result.get('tenant_id')),
                          ('Evaluation date', result.get('evaluation_date')), ('Methodology Version', result.get('methodology_version'))]
                         if item not in present and value is not None)
    manifest_rows = [row for row in manifest_rows if row.get('Item') not in {'Assessment workbook', 'Technical evidence workbook'}]
    manifest_rows.extend([{'Item': 'Assessment workbook', 'Value': Path(assessment_path).name},
                          {'Item': 'Technical evidence workbook', 'Value': Path(technical_path).name}])
    for book in (assessment, technical):
        if 'Run Manifest' in book:
            book.remove(book['Run Manifest'])
        _write(book, 'Run Manifest', manifest_rows, ('Item', 'Value'))
    bundle.setdefault('run_manifest', {})['rows'] = manifest_rows
    previous = _rows(source_workbook['Integrity Checks']) if 'Integrity Checks' in source_workbook else []
    for book in (assessment, technical):
        if 'Integrity Checks' in book:
            book.remove(book['Integrity Checks'])
        _write(book, 'Integrity Checks', previous + reconciliation)
    # PDF excerpts span health, adoption and security without a finding ID or
    # affected-record population. Keep this optional context sheet visible in
    # both books, rather than forcing them into a per-finding evidence layout.
    if 'PDF Highlights' in technical:
        _clone(technical['PDF Highlights'], assessment)
        pages = {(row.get('Capture ID'), row.get('Page')): number
                 for number, row in enumerate(_rows(technical['PDF Extracted Text']), 2)} if 'PDF Extracted Text' in technical else {}
        captures = {row.get('Capture ID'): number for number, row in enumerate(_rows(technical['Portal Review']), 2)} if 'Portal Review' in technical else {}
        for book in (assessment, technical):
            sheet = book['PDF Highlights']
            for number, row in enumerate(_rows(sheet), 2):
                page = pages.get((row.get('Capture ID'), row.get('Page')))
                capture = captures.get(row.get('Capture ID'))
                location = f"'PDF Extracted Text'!A{page}" if page else f"'Portal Review'!A{capture}" if capture else None
                prefix = '#' if book is technical else Path(technical_path).name + '#'
                _cell_link(sheet, number, 'Source detail', prefix + location if location else None)
    bundle['workbook_reconciliation'] = reconciliation
    bundle['technical_excel_path'] = str(technical_path)
    bundle['assessment_excel_path'] = str(assessment_path)
    for book, role, other in ((assessment, 'assessment', technical_path), (technical, 'technical', assessment_path)):
        apply_workbook_navigation(book, result, bundle, role=role, companion_path=other, tenant_name=tenant_name)
    validation = validate_workbook_layout(assessment, technical, result, assessment_path, technical_path)
    issues.extend(validation)
    if issues:
        integrity = bundle.setdefault('integrity', {'valid': True, 'issues': []})
        integrity['valid'] = False
        integrity.setdefault('issues', []).extend(issues)
        integrity['banner'] = 'Validation incomplete—do not use for deployment approval'
        # Preserve failed checks for inspection before the caller fails the build.
        for book in (assessment, technical):
            book['Integrity Checks'].append(['Failed', '; '.join(issues), 'Workbook links and cells'])
        assessment.save(assessment_path)
        technical.save(technical_path)
        raise WorkbookLayoutError('; '.join(issues))
    return assessment, technical


def _compact_ranges(locations):
    blocks = []
    for location in locations:
        match = RANGE.fullmatch(location)
        title, first, last = match[1].replace("''", "'"), int(match[3]), int(match[5] or match[3])
        if blocks and blocks[-1][0] == title and blocks[-1][2] + 1 == first:
            blocks[-1][2] = last
        else:
            blocks.append([title, first, last, match[4] or match[2]])
    return ["'" + title.replace("'", "''") + f"'!A{first}:{column}{last}" for title, first, last, column in blocks]


def validate_workbook_layout(assessment, technical, result, assessment_path, technical_path):
    """Validate cell destinations, per-finding IDs and workbook structural invariants."""
    from openpyxl.utils.cell import range_boundaries
    books = {Path(assessment_path).name: assessment, Path(technical_path).name: technical}
    issues = []
    for book in books.values():
        tables = [name.casefold() for sheet in book for name in sheet.tables]
        if len(tables) != len(set(tables)):
            issues.append('Duplicate Excel table names.')
        for sheet in book:
            for cells in sheet:
                for cell in cells:
                    if isinstance(cell.value, str) and any(error in cell.value for error in ERROR_VALUES):
                        issues.append(f'{sheet.title}: Excel error value at {cell.coordinate}.')
                    target = cell.hyperlink.target if cell.hyperlink else None
                    if not target or '#' not in target or target.startswith(('https:', 'http:')):
                        continue
                    filename, location = target.split('#', 1)
                    match = RANGE.fullmatch(location)
                    destination = books.get(filename) if filename else book
                    if not match or destination is None:
                        issues.append(f'{sheet.title}: unresolved workbook link at {cell.coordinate}.')
                        continue
                    title = match[1].replace("''", "'")
                    if title not in destination or destination[title].sheet_state != 'visible':
                        issues.append(f'{sheet.title}: workbook link targets a missing or hidden sheet at {cell.coordinate}.')
                        continue
                    first_col, first, last_col, last = range_boundaries(location.rsplit('!', 1)[1])
                    if not (1 <= first <= last <= destination[title].max_row and 1 <= first_col <= last_col <= destination[title].max_column):
                        issues.append(f'{sheet.title}: workbook link is outside retained cells at {cell.coordinate}.')
    for row in result['recommendations']:
        for location in row.get('AssessmentEvidenceRanges') or []:
            match = RANGE.fullmatch(location)
            if not match:
                issues.append(f'{row["RecommendationId"]}: invalid assessment evidence range.')
                continue
            title = match[1].replace("''", "'")
            if title not in assessment:
                issues.append(f'{row["RecommendationId"]}: missing assessment evidence sheet.')
                continue
            if any(assessment[title].cell(number, 1).value != row['RecommendationId']
                   for number in range(int(match[3]), int(match[5] or match[3]) + 1)):
                issues.append(f'{row["RecommendationId"]}: assessment evidence link targets another finding.')
    return list(dict.fromkeys(issues))
