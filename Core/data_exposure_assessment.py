"""Evidence-backed SharePoint and Purview data exposure assessment.

Microsoft's authoritative oversharing scans are asynchronous administrative reports rather
than ordinary Microsoft Graph inventory endpoints.  This module intentionally consumes their
exports instead of inferring exposure from site or file counts.  A missing export becomes a
coverage item; it never becomes a clean bill of health.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from .new_recommendation import (
    CATEGORY_SCAN_COVERAGE,
    NOT_ASSESSED_STATUS,
    new_recommendation,
)


SUPPORTED_SUFFIXES = {".csv", ".tsv", ".json", ".xlsx", ".zip"}
MAX_EVIDENCE_ROWS = 25000
DEFAULT_SAM_MAX_AGE_DAYS = 35
DEFAULT_DSPM_MAX_AGE_DAYS = 8
DEFAULT_LIFECYCLE_MAX_AGE_DAYS = 90


def _normalized_key(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _text(value):
    if value is None:
        return ""
    return str(value).strip()


def _number(value):
    if isinstance(value, bool):
        return int(value)
    text = _text(value).replace(",", "")
    if not text:
        return 0
    match = re.fullmatch(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return 0
    try:
        return int(float(match.group(0)))
    except ValueError:
        return 0


def _positive_int_env(name, default):
    try:
        value = int(os.environ.get(name, default))
        return value if value > 0 else default
    except (TypeError, ValueError):
        return default


def _parse_date(value):
    text = _text(value)
    if not text:
        return None
    normalized = text.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(normalized)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc)
    except ValueError:
        pass
    for pattern in ("%m/%d/%Y", "%m/%d/%Y %H:%M:%S", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def _is_true(value):
    return _text(value).lower() in {"1", "true", "yes", "y", "enabled", "active"}


def _value(record, *aliases):
    normalized = {_normalized_key(key): value for key, value in (record or {}).items()}
    for alias in aliases:
        key = _normalized_key(alias)
        if key in normalized:
            return normalized[key]
    return ""


def _has_column(record, *aliases):
    keys = {_normalized_key(key) for key in (record or {}).keys()}
    return any(_normalized_key(alias) in keys for alias in aliases)


EXPOSURE_COLUMNS = {
    "Anyone links": ("Anyone link count", "Anonymous link count"),
    "Everyone permissions": ("Everyone permission count",),
    "EEEU permissions": ("EEEU permission count", "Everyone except external users count"),
    "Organization links": ("Organization link count", "People in your organization link count", "PeopleInYourOrg link count", "Company link count"),
    "External exposure": ("External user count", "External sharing link count", "Externally shared item count", "Shared externally count", "Guest user permissions", "External participant permissions", "External participant permissionst"),
    "Potentially overshared items": ("Potentially overshared items", "Overshared item count"),
    "Sensitive items": ("Sensitive data detected", "Sensitive item count", "Sensitive data count"),
    "Unlabeled sensitive items": ("Unlabeled sensitive item count", "Unlabeled item count", "Is unlabeled sensitive", "Unlabeled sensitive", "Sensitive and unlabeled"),
}
SAM_DOMAINS = set(list(EXPOSURE_COLUMNS)[:5])
DSPM_DOMAINS = {"Potentially overshared items", "Unlabeled sensitive items"}


def detect_sharepoint_report(record_or_headers):
    """Identify supported Microsoft report families from headers, including empty CSVs.

    Names and dates in filenames are deliberately not schema or freshness evidence.
    """
    record = record_or_headers if isinstance(record_or_headers, dict) else dict.fromkeys(record_or_headers or [])
    location = _has_column(record, "Site URL", "URL", "Site ID", "Item URL", "Site Name")
    if location and _has_column(record, "Is inactive") and _has_column(record, "Is ownerless"):
        return "content_management_assessment"
    if location and _has_column(record, "Labeled files") and _has_column(record, "Site sensitivity label ID", "Site Sensitivity"):
        return "label_inventory"
    if _has_column(record, "Recipient", "UserPrincipalName", "Permission Recipient") and _has_column(record, "Item URL", "ItemType", "Role definition"):
        return "special_group_permissions"
    if location and _has_column(record, "Number of users having access", "Number of users with permissions", "Permissioned users") and _has_column(record, "Anyone link count", "EEEU permission count", "Everyone permission count"):
        return "permission_snapshot"
    if any(_has_column(record, *EXPOSURE_COLUMNS[domain]) for domain in SAM_DOMAINS):
        return "exposure_counts"
    if location and _has_column(record, "Link Type", "Sharing Link Type", "Permission Recipient"):
        return "sharing_details"
    return ""


def _matches_authoritative_schema(record, source_kind):
    if source_kind == "sam":
        return detect_sharepoint_report(record) in {"permission_snapshot", "special_group_permissions", "exposure_counts", "sharing_details"}
    # A sensitivity label, site owner or generic Access column alone is not a DSPM
    # assessment. In particular, CMA and label inventories must not close this gap.
    return detect_sharepoint_report(record) not in {"content_management_assessment", "label_inventory", "permission_snapshot", "special_group_permissions"} and any(
        _has_column(record, *EXPOSURE_COLUMNS[domain])
        for domain in (*DSPM_DOMAINS, "Sensitive items")
    )


def _split_env_paths(name):
    raw = os.environ.get(name, "").strip()
    if not raw:
        return []
    # Windows paths commonly contain a colon, so use the platform path separator only.
    return [item.strip() for item in raw.split(os.pathsep) if item.strip()]


def _expand_paths(paths):
    expanded = []
    seen = set()
    for raw_path in paths or []:
        path = Path(raw_path).expanduser()
        candidates = (
            sorted(item for item in path.iterdir() if item.suffix.lower() in SUPPORTED_SUFFIXES)
            if path.is_dir()
            else [path]
        )
        for candidate in candidates:
            identity = str(candidate.resolve()).lower() if candidate.exists() else str(candidate).lower()
            if identity not in seen:
                seen.add(identity)
                expanded.append(candidate)
    return expanded


def _iter_json_records(payload, sheet_name=""):
    if isinstance(payload, list):
        for item in payload:
            if isinstance(item, dict):
                row = dict(item)
                if sheet_name:
                    row.setdefault("_source_sheet", sheet_name)
                yield row
        return
    if not isinstance(payload, dict):
        return
    for key in ("rows", "items", "value", "results", "data"):
        nested = payload.get(key)
        if isinstance(nested, list):
            yield from _iter_json_records(nested, sheet_name=sheet_name or key)
            return
    # A dictionary of named report arrays is also a common export envelope.
    yielded = False
    for key, nested in payload.items():
        if isinstance(nested, list) and any(isinstance(item, dict) for item in nested):
            yielded = True
            yield from _iter_json_records(nested, sheet_name=key)
    if not yielded:
        row = dict(payload)
        if sheet_name:
            row.setdefault("_source_sheet", sheet_name)
        yield row


def _read_records(path):
    suffix = path.suffix.lower()
    if suffix == ".zip":
        # Read members in memory; no customer-controlled archive path is extracted.
        with zipfile.ZipFile(path) as archive:
            for member in archive.infolist():
                member_suffix = Path(member.filename).suffix.lower()
                if member.is_dir() or member_suffix not in {".csv", ".tsv"}:
                    continue
                with archive.open(member) as raw, io.TextIOWrapper(raw, encoding="utf-8-sig", errors="strict", newline="") as handle:
                    reader = csv.DictReader(handle, delimiter="\t" if member_suffix == ".tsv" else ",")
                    if reader.fieldnames:
                        yield {**dict.fromkeys(reader.fieldnames, ""), "_headers_only": True, "_source_member": member.filename}
                    for row in reader:
                        yield {**row, "_source_member": member.filename}
        return
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", errors="strict", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=delimiter)
            if reader.fieldnames:
                yield {**dict.fromkeys(reader.fieldnames, ""), "_headers_only": True}
            yield from reader
        return
    if suffix == ".json":
        with path.open("r", encoding="utf-8-sig") as handle:
            yield from _iter_json_records(json.load(handle))
        return
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                headers = next(rows, None)
                if not headers:
                    continue
                headers = [_text(header) or f"Column {index + 1}" for index, header in enumerate(headers)]
                yield {**dict.fromkeys(headers, ""), "_headers_only": True, "_source_sheet": worksheet.title}
                for values in rows:
                    row = dict(zip(headers, values))
                    row["_source_sheet"] = worksheet.title
                    yield row
        finally:
            workbook.close()
        return
    raise ValueError(f"Unsupported report type: {suffix or 'no extension'}")


def _record_identity(record):
    site_url = _text(_value(record, "Site URL", "SiteUrl", "Site Address", "Location URL", "URL"))
    item_url = _text(_value(record, "Item URL", "File URL", "Object URL", "Item Path", "Path"))
    site_name = _text(_value(record, "Site Name", "SiteName", "Location Name", "Name"))
    return site_url, item_url, site_name


def _report_date(record):
    return _parse_date(_value(record, "Report Date", "Scan Date", "Assessment Date", "Completed Date",
        "Completion Date", "Data As Of", "Generated Date", "CreatedDateTime",
        "Report End Time", "ReportEndTime", "TriggeredDateTime"))


def _workload(record):
    explicit = _text(_value(record, "Workload", "Data Source", "Service"))
    if explicit:
        return "OneDrive" if "onedrive" in explicit.lower() else explicit
    site_url, item_url, _ = _record_identity(record)
    address = (site_url or item_url).lower()
    # The tenant's -my.sharepoint root can appear in a SharePoint snapshot.
    # Only personal sites (or an explicit workload) establish OneDrive scope.
    if "/personal/" in address or "spspers" in _text(_value(record, "Site Template", "Template")).lower():
        return "OneDrive"
    if address or detect_sharepoint_report(record) == "content_management_assessment":
        return "SharePoint"
    return "Unknown"


def _classify_record(record, source_kind):
    """Return normalized risk signals backed by explicit report fields."""
    site_url, item_url, site_name = _record_identity(record)
    label = _text(_value(record, "Sensitivity Label", "Site Sensitivity", "Label Name", "Label"))
    owner = _text(_value(record, "Primary admin email")) or _text(_value(record, "Primary Admin", "Site Owner", "Owner", "Owner Email"))
    workload = _workload(record)
    counts = {name: max(0, _number(_value(record, *aliases))) for name, aliases in EXPOSURE_COLUMNS.items()}
    # These are distinct permission categories, not alternative spellings of a
    # single count. The portal currently includes a trailing 't' in one header.
    explicit_external = sum(max(0, _number(_value(record, *aliases))) for aliases in (
        ("Guest user permissions",), ("External participant permissions", "External participant permissionst"),
    ))
    counts["External exposure"] = max(counts["External exposure"], explicit_external)
    if _is_true(_value(record, "Sensitive data detected")):
        counts["Sensitive items"] = max(counts["Sensitive items"], 1)

    link_type = _text(_value(record, "Link Type", "Sharing Link Type", "SharingLinkType", "Access Type"))
    permission_recipient = _text(_value(record, "Permission Recipient", "Recipient", "UserPrincipalName", "Principal", "Group Name", "Shared With"))
    access_text = " ".join([link_type, permission_recipient, _text(_value(record, "Access", "Exposure", "Sharing Type"))]).lower()

    if any(marker in access_text for marker in ("anonymous", "anyone")):
        counts["Anyone links"] = max(counts["Anyone links"], 1)
    if "everyone except external" in access_text or "eeeu" in access_text:
        counts["EEEU permissions"] = max(counts["EEEU permissions"], 1)
    elif re.search(r"\beveryone\b", access_text):
        counts["Everyone permissions"] = max(counts["Everyone permissions"], 1)
    external_access_text = re.sub(r"everyone\s+except\s+external\s+users", "", access_text)
    if any(marker in external_access_text for marker in ("external", "guest")):
        counts["External exposure"] = max(counts["External exposure"], 1)
    if any(marker in access_text for marker in ("organization", "peopleinyourorg", "company-wide", "company wide")):
        counts["Organization links"] = max(counts["Organization links"], 1)

    explicit_sensitive = counts["Sensitive items"] > 0 or bool(label and label.lower() not in {"none", "unlabeled", "not labeled", "n/a", "public", "general", "non-business"})
    unlabeled_sensitive = counts["Unlabeled sensitive items"] > 0 or _is_true(
        _value(record, "Is unlabeled sensitive", "Unlabeled sensitive", "Sensitive and unlabeled")
    )
    if unlabeled_sensitive:
        counts["Unlabeled sensitive items"] = max(counts["Unlabeled sensitive items"], 1)

    # Sensitive data is context that raises the severity of an exposure; its mere existence is
    # not a failure. Keep it out of the risk totals unless it is explicitly unlabeled.
    signals = {
        name: count for name, count in counts.items()
        if count > 0 and name != "Sensitive items"
    }
    if not signals:
        return None

    broad = sum(signals.get(key, 0) for key in (
        "Anyone links", "Everyone permissions", "EEEU permissions", "Organization links",
        "External exposure", "Potentially overshared items",
    ))
    severity = "High" if explicit_sensitive and (broad or unlabeled_sensitive) else "Medium"
    if set(signals) == {"Ownerless site"}:
        severity = "Medium"

    return {
        "Source": "SharePoint Advanced Management" if source_kind == "sam" else "Microsoft Purview DSPM",
        "Source File": _text(record.get("_source_file")),
        "Source Sheet": _text(record.get("_source_sheet")),
        "Report Type": record.get("_report_type", detect_sharepoint_report(record) if source_kind == "sam" else "dspm_assessment"),
        "Report Date": (_report_date(record).isoformat() if _report_date(record) else ""),
        "Tenant ID": _text(_value(record, "Tenant ID")),
        "Site ID": _text(_value(record, "Site ID")),
        "Workload": workload,
        "Site Name": site_name,
        "Site URL": site_url,
        "Item URL": item_url,
        "Owner": owner,
        "Sensitivity Label": label,
        "Risk Signals": "; ".join(signals.keys()),
        "Signal Count": sum(signals.values()),
        "Severity": severity,
        "RecommendationId": "",
        "Flagged By": "",
        "_signals": signals,
    }


def _flag(value):
    text = _text(value).lower()
    if text in {"true", "yes", "1", "y"}:
        return "Yes"
    if text in {"false", "no", "0", "n"}:
        return "No"
    return "Unknown"


def _known_domain_value(record, domain):
    value = _value(record, *EXPOSURE_COLUMNS[domain])
    if domain in {"Sensitive items", "Unlabeled sensitive items"} and _flag(value) != "Unknown":
        return True
    def numeric(value):
        return bool(re.fullmatch(r"\d+(?:\.\d+)?", _text(value).replace(",", "")))
    if domain == "External exposure" and _has_column(record, "Guest user permissions", "External participant permissions", "External participant permissionst"):
        return all(numeric(_value(record, *aliases)) for aliases in (
            ("Guest user permissions",), ("External participant permissions", "External participant permissionst"),
        ) if _has_column(record, *aliases))
    return numeric(value)


def _lifecycle_row(record, report_date=None, date_basis="", source_hash=""):
    site_url, _, site_name = _record_identity(record)
    date = report_date or _report_date(record)
    return {
        "Source": "SharePoint Content Management Assessment",
        "Source File": record.get("_source_file", ""),
        "Source Sheet": record.get("_source_sheet", ""),
        "Report Date": date.isoformat() if date else "",
        "Report Date Basis": date_basis or ("Report date field" if date else "Unknown"),
        "Source Hash": source_hash,
        "Workload": _workload(record), "Site Name": site_name, "Site URL": site_url,
        "Owner": _text(_value(record, "Email address of site owners")),
        "Is Ownerless": _flag(_value(record, "Is ownerless")),
        "Is Inactive": _flag(_value(record, "Is inactive")),
        "Last Activity Date": _text(_value(record, "Last activity date (UTC)")),
        "Site Creation Date": _text(_value(record, "Site creation date (UTC)")),
        "Sensitivity Label": _text(_value(record, "Sensitivity label")),
        "Retention Policy": _text(_value(record, "Retention policy")),
        "Site Lock State": _text(_value(record, "Site lock state")),
        "Storage Used (GB)": _text(_value(record, "Storage used (GB)")),
        "Template": _text(_value(record, "Template")),
        "Connected to Teams": _flag(_value(record, "Connected to Teams")),
    }


def _scan_reports(paths, source_kind, evaluation_date=None, max_age_days=None, *,
                  lifecycle_max_age_days=None, lifecycle_report_dates=None):
    evaluated_at = _parse_date(evaluation_date) if evaluation_date else datetime.now(timezone.utc)
    max_age = max_age_days or _positive_int_env(
        "SAM_REPORT_MAX_AGE_DAYS" if source_kind == "sam" else "DSPM_REPORT_MAX_AGE_DAYS",
        DEFAULT_SAM_MAX_AGE_DAYS if source_kind == "sam" else DEFAULT_DSPM_MAX_AGE_DAYS,
    )
    lifecycle_max_age = lifecycle_max_age_days or _positive_int_env(
        "LIFECYCLE_REPORT_MAX_AGE_DAYS", DEFAULT_LIFECYCLE_MAX_AGE_DAYS)
    lifecycle_report_dates = lifecycle_report_dates or {}
    result = {
        "kind": source_kind, "files_requested": len(paths), "files_parsed": 0,
        "files_loaded": 0, "files_recognized": 0, "records_read": 0,
        "records_selected": 0, "errors": [], "risk_rows": [], "reports": [],
        "lifecycle_rows": [], "evidence_truncated": False, "signals": {},
        "affected_sites": [], "max_age_days": max_age, "observations": [],
    }
    blocks = []
    snapshot_url_ids = defaultdict(set)
    for path in paths:
        if not path.exists() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            result["errors"].append(f"{path}: " + ("file not found" if not path.exists() else "unsupported file type"))
            continue
        first_observation = len(result['observations'])
        try:
            documents = {}
            file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
            for raw in _read_records(path):
                record = dict(raw or {})
                headers_only = bool(record.pop("_headers_only", False))
                if not headers_only and not any(_text(v) for k, v in record.items() if not str(k).startswith("_")):
                    continue
                if not headers_only:
                    result["records_read"] += 1
                report_type = detect_sharepoint_report(record) if source_kind == "sam" else (
                    "dspm_assessment" if _matches_authoritative_schema(record, source_kind) else ""
                )
                member = _text(record.get("_source_member"))
                sheet = _text(record.get("_source_sheet"))
                source_file = path.name + ("!" + member if member else "")
                doc = documents.setdefault((member, sheet, report_type), {"header": record, "type": report_type, "blocks": {}})
                if headers_only or not report_type:
                    continue
                record["_source_file"] = source_file
                record["_report_type"] = report_type
                date = _report_date(record)
                date_basis = "Report date field" if date else "Unknown"
                if report_type == "content_management_assessment" and not date:
                    # Only this exact file's explicitly confirmed generation date
                    # may fill a missing report date. Filenames and site activity
                    # timestamps do not establish when a report was generated.
                    confirmed_date = _parse_date(lifecycle_report_dates.get(file_hash))
                    if confirmed_date:
                        date, date_basis = confirmed_date, "Operator-confirmed report date"
                workload = _workload(record)
                tenant = _text(_value(record, "Tenant ID"))
                if not tenant:
                    address = _record_identity(record)[0]
                    tenant = address.split("/")[2].lower() if "://" in address else ""
                signature = tuple(sorted(_normalized_key(k) for k in record if not str(k).startswith("_")))
                block_key = (workload, tenant, date, signature)
                block = doc["blocks"].setdefault(block_key, {
                    "source_file": source_file, "source_sheet": sheet, "report_type": report_type,
                    "workload": workload, "tenant_id": tenant, "report_date": date.isoformat() if date else "",
                    "date_basis": date_basis, "source_hash": file_hash,
                    "records_read": 0, "status": "selected", "domains": set(), "domains_unknown": set(),
                    "_date": date, "_signature": signature, "_path": str(path),
                    "_risks": {}, "_maxima": {}, "_lifecycle": [], "_baseline_signals": set(), "_detail_sites": {},
                    "_population": set(),
                })
                block["records_read"] += 1
                site_url = _record_identity(record)[0]
                site_id = _text(_value(record, 'Site ID')).lower()
                object_identity = (site_url or site_id).lower()
                if report_type == 'permission_snapshot':
                    normalized_url = site_url.lower().rstrip('/')
                    object_identity = 'site:' + site_id if site_id else ('url:' + normalized_url if normalized_url else '')
                if object_identity:
                    block['_population'].add(object_identity)
                # Retain each measured field (including zero/unknown) with its
                # object, unit and source. These facts do not assert control pass.
                for metric, aliases in EXPOSURE_COLUMNS.items():
                    if not _has_column(record, *aliases):
                        continue
                    known = _known_domain_value(record, metric)
                    value = max(0, _number(_value(record, *aliases))) if known else None
                    if known and metric == 'External exposure':
                        value = max(value, sum(max(0, _number(_value(record, *group))) for group in (
                            ('Guest user permissions',), ('External participant permissions', 'External participant permissionst'))))
                    unit = 'links' if 'links' in metric.lower() else 'permissions' if 'permission' in metric.lower() or metric == 'External exposure' else 'items'
                    result['observations'].append({
                        'metric_id': source_kind + '.' + _normalized_key(metric), 'label': metric,
                        'domain_id': 'content' if source_kind == 'sam' else 'data_protection',
                        'control_id': 'CONTENT.PERMISSIONS' if source_kind == 'sam' else 'DATA.EXPOSURE',
                        'value': value, 'unit': unit, 'availability': 'available' if known else 'unknown',
                        'tenant_id': _text(_value(record, 'Tenant ID')),
                        'population': workload, 'scope': object_identity or 'Exported row',
                        'affected_objects': [object_identity] if object_identity else [],
                        'observed_at': date.isoformat() if date else '', 'complete': False,
                        'source_type': 'portal_export', 'source_file': source_file,
                        'source_schema': report_type, 'source_hash': file_hash, 'max_age_days': max_age,
                        **({'site_id': site_id, 'site_url': site_url, '_identity_tenant': tenant.lower()}
                           if report_type == 'permission_snapshot' else {}),
                    })
                block["domains"].update(name for name, aliases in EXPOSURE_COLUMNS.items() if _has_column(record, *aliases))
                block["domains_unknown"].update(name for name, aliases in EXPOSURE_COLUMNS.items() if _has_column(record, *aliases) and not _known_domain_value(record, name))
                if report_type == "content_management_assessment":
                    block["_lifecycle"].append(_lifecycle_row(record, date, date_basis, file_hash))
                    continue
                if report_type == "label_inventory":
                    continue
                if report_type == "permission_snapshot" and site_id:
                    block["_baseline_signals"].update((name, tenant.lower(), workload.lower(), site_id) for name in SAM_DOMAINS if _has_column(record, *EXPOSURE_COLUMNS[name]) and _known_domain_value(record, name))
                risk = _classify_record(record, source_kind)
                if not risk:
                    continue
                signals = risk.pop("_signals")
                identity = (object_identity if report_type == 'permission_snapshot' else
                            risk["Item URL"] or risk["Site URL"] or risk["Site ID"] or risk["Site Name"]) or f"{source_file}:{block['records_read']}"
                identity = (tenant.lower(), workload.lower(), identity.lower())
                for name, count in signals.items():
                    key = (name, *identity)
                    block["_maxima"][key] = max(block["_maxima"].get(key, 0), count)
                    if report_type == "special_group_permissions" and site_id:
                        block["_detail_sites"][key] = (name, tenant.lower(), workload.lower(), site_id)
                evidence_key = (*identity, risk["Risk Signals"])
                if evidence_key in block["_risks"] or len(block["_risks"]) < MAX_EVIDENCE_ROWS:
                    previous = block["_risks"].get(evidence_key)
                    if not previous or risk["Signal Count"] > previous["Signal Count"]:
                        block["_risks"][evidence_key] = risk
                else:
                    result["evidence_truncated"] = True
            recognized = False
            for (member, sheet, report_type), doc in documents.items():
                if not report_type:
                    continue
                recognized = True
                if doc["blocks"]:
                    blocks.extend(doc["blocks"].values())
                else:
                    # Headers establish the report family only. Zero rows reveal
                    # neither the report date nor the absence of tenant exposure.
                    blocks.append({
                        "source_file": path.name + ("!" + member if member else ""),
                        "source_sheet": sheet, "report_type": report_type, "workload": "Unknown",
                        "tenant_id": "", "report_date": "", "records_read": 0, "status": "empty",
                        "domains": set(), "domains_unknown": set(), "_date": None, "_signature": (), "_path": str(path),
                        "_risks": {}, "_maxima": {}, "_lifecycle": [], "_baseline_signals": set(), "_detail_sites": {},
                        "_population": set(),
                    })
            result["files_parsed"] += 1
            if recognized:
                result["files_recognized"] += 1
            else:
                result["errors"].append(f"{path}: no recognized {source_kind.upper()} report schema")
        except Exception as exc:
            # Do not use a partially read workbook/archive as complete evidence.
            del result['observations'][first_observation:]
            result["errors"].append(f"{path}: {type(exc).__name__}: {exc}")

    # Build aliases only from successfully read files. A corrupt archive or
    # workbook cannot contribute identities after its observations are discarded.
    for observation in result['observations']:
        if observation.get('source_schema') == 'permission_snapshot' and observation.get('site_id') and observation.get('site_url'):
            snapshot_url_ids[(observation['_identity_tenant'], observation['population'].lower(),
                              observation['site_url'].lower().rstrip('/'))].add(observation['site_id'])

    def snapshot_identity(tenant, workload, identity):
        # A missing ID may use a URL only when that URL identifies a single
        # explicit site ID within this tenant/workload. Reused URLs remain
        # ambiguous, and tenants or workloads never share an identity map.
        if identity.startswith('url:'):
            candidates = snapshot_url_ids.get((tenant.lower(), workload.lower(), identity[4:]), set())
            if len(candidates) == 1:
                return 'site:' + next(iter(candidates))
        return identity

    for block in blocks:
        if block['report_type'] != 'permission_snapshot':
            continue
        tenant, workload = block['tenant_id'], block['workload']
        block['_population'] = {snapshot_identity(tenant, workload, identity) for identity in block['_population']}
        maxima = {}
        for key, count in block['_maxima'].items():
            canonical = (*key[:3], snapshot_identity(tenant, workload, key[3]))
            maxima[canonical] = max(maxima.get(canonical, 0), count)
        block['_maxima'] = maxima
        risks = {}
        for key, risk in block['_risks'].items():
            canonical = (*key[:2], snapshot_identity(tenant, workload, key[2]), *key[3:])
            if canonical not in risks or risk['Signal Count'] > risks[canonical]['Signal Count']:
                risks[canonical] = risk
        block['_risks'] = risks
    for observation in result['observations']:
        if observation.get('source_schema') == 'permission_snapshot':
            identity = snapshot_identity(observation.pop('_identity_tenant', ''), observation['population'], observation['scope'])
            observation['scope'] = identity
            observation['affected_objects'] = [identity] if identity else []

    # A CSV does not establish that omitted sites were removed from the tenant.
    # Replace a snapshot only for a comparable population. Otherwise retain the
    # older objects with their original dates and flag the scope for confirmation.
    newest = {}
    newest_objects = {}
    for block in blocks:
        if (block["report_type"] == "permission_snapshot" and block["_date"]
                and block["_date"].date() <= evaluated_at.date() and block["tenant_id"]):
            key = (block["report_type"], block["tenant_id"].lower(), block["workload"], block["_signature"], tuple(sorted(block['_population'])))
            newest[key] = max(newest.get(key, block["_date"]), block["_date"])
            for identity in block['_population']:
                object_key = (block['tenant_id'].lower(), block['workload'].lower(), identity)
                newest_objects[object_key] = max(newest_objects.get(object_key, block['_date']), block['_date'])
    maxima, evidence, lifecycle, sites, dates, domains, workloads, loaded_paths = {}, {}, {}, set(), [], set(), set(), set()
    unknown_domains, baseline_dates, detail_sites, detail_dates = set(), {}, {}, {}
    selected_reports = []
    for block in blocks:
        key = (block["report_type"], block["tenant_id"].lower(), block["workload"], block["_signature"], tuple(sorted(block['_population'])))
        if block["report_type"] == "permission_snapshot" and block["_date"] and block["_date"] < newest.get(key, block["_date"]):
            block["status"] = "superseded"
        age = (evaluated_at.date() - block["_date"].date()).days if block["_date"] else None
        if block['report_type'] == 'permission_snapshot' and age is not None and age < 0:
            block['status'] = 'future_date'
        block["age_days"] = age
        block["max_age_days"] = lifecycle_max_age if block["report_type"] == "content_management_assessment" else max_age
        block["freshness"] = "unknown" if age is None or age < 0 else ("fresh" if age <= block["max_age_days"] else "stale")
        block['population_count'] = len(block['_population'])
        block['scope_note'] = 'Exported objects only; tenant-wide completeness is not established.'
        if block['status'] == 'future_date':
            block['scope_note'] += ' Report date is after the evaluation date; retained for review and excluded from exposure totals.'
        if block['report_type'] == 'permission_snapshot' and block['status'] == 'selected':
            comparable = [other for other in blocks if other is not block and other['report_type'] == block['report_type']
                and other['tenant_id'] == block['tenant_id'] and other['workload'] == block['workload']]
            if any(other['_population'] != block['_population'] for other in comparable):
                block['scope_note'] += ' Snapshot populations differ. Older objects absent from the newer file remain dated observations requiring confirmation.'
        block["domains"] = sorted(block["domains"])
        block["domains_unknown"] = sorted(block["domains_unknown"])
        result["reports"].append({k: v for k, v in block.items() if not k.startswith("_")})
        if block["status"] != "selected":
            continue
        for row in block["_lifecycle"]:
            identity = row["Site URL"].lower() or row["Site Name"].lower() or f"{row['Source File']}:{len(lifecycle)}"
            previous = lifecycle.get(identity)
            if not previous or (row["Report Date"] and row["Report Date"] > previous["Report Date"]):
                lifecycle[identity] = row
        if block["report_type"] in {"content_management_assessment", "label_inventory"}:
            continue
        selected_reports.append(block)
        loaded_paths.add(block["_path"])
        result["records_selected"] += block["records_read"]
        dates.append(block["_date"])
        domains.update(block["domains"])
        unknown_domains.update(block["domains_unknown"])
        if block["_date"]:
            for baseline_key in block["_baseline_signals"]:
                baseline_dates[baseline_key] = max(baseline_dates.get(baseline_key, block["_date"]), block["_date"])
        detail_sites.update(block["_detail_sites"])
        detail_dates.update({key: block["_date"] for key in block["_detail_sites"]})
        workloads.add(block["workload"])
        for signal_key, count in block["_maxima"].items():
            if block['report_type'] == 'permission_snapshot' and block['_date'] and block['_date'] < newest_objects.get(signal_key[1:], block['_date']):
                continue
            maxima[signal_key] = max(maxima.get(signal_key, 0), count)
        for evidence_key, risk in block["_risks"].items():
            if block['report_type'] == 'permission_snapshot' and block['_date'] and block['_date'] < newest_objects.get(evidence_key[:3], block['_date']):
                continue
            site = risk["Site URL"] or risk["Site ID"] or risk["Site Name"]
            if site:
                sites.add(site)
            previous = evidence.get(evidence_key)
            if previous:
                source_files = sorted(set(previous["Source File"].split("; ") + [risk["Source File"]]))
                if risk["Signal Count"] > previous["Signal Count"]:
                    evidence[evidence_key] = risk
                evidence[evidence_key]["Source File"] = "; ".join(source_files)
            elif len(evidence) < MAX_EVIDENCE_ROWS:
                evidence[evidence_key] = risk
            else:
                result["evidence_truncated"] = True
    totals = defaultdict(int)
    for signal_key, count in maxima.items():
        baseline_date = baseline_dates.get(detail_sites.get(signal_key))
        detail_date = detail_dates.get(signal_key)
        if baseline_date and detail_date and baseline_date >= detail_date:
            # Keep item evidence, but do not add item rows already represented in
            # the same site's permission snapshot count.
            continue
        totals[signal_key[0]] += count
    result["signals"] = dict(totals)
    result["risk_rows"] = list(evidence.values())
    result["lifecycle_rows"] = list(lifecycle.values())[:MAX_EVIDENCE_ROWS]
    result["lifecycle_records_truncated"] = len(lifecycle) > MAX_EVIDENCE_ROWS
    result["lifecycle_summary"] = _summarize_lifecycle(list(lifecycle.values()), lifecycle_max_age, evaluation_date)
    result["affected_sites"] = sorted(sites, key=str.lower)
    result["files_loaded"] = len(loaded_paths)
    known_dates = [date for date in dates if date]
    result["latest_report_date"] = max(known_dates).isoformat() if known_dates else ""
    result["oldest_report_date"] = min(known_dates).isoformat() if known_dates else ""
    result["age_days"] = max((evaluated_at.date() - min(known_dates).date()).days, 0) if known_dates else None
    result["freshness"] = "missing" if not dates else (
        "stale" if any(b["freshness"] == "stale" for b in selected_reports) else (
            "unknown" if any(date is None or date.date() > evaluated_at.date() for date in dates) else "fresh"
        )
    )
    result["coverage"] = {
        "domains_observed": sorted(domains),
        "domains_missing": sorted((SAM_DOMAINS if source_kind == "sam" else DSPM_DOMAINS) - domains),
        "domains_unknown": sorted(unknown_domains),
        "workloads_observed": sorted(workloads),
        "workloads_missing": sorted({"SharePoint", "OneDrive"} - workloads) if source_kind == "sam" else [],
        "tenant_scope_verified": False,
        "scope_note": "Conclusions apply to the exported rows and fields; tenant-wide completeness is not established by a CSV alone.",
    }
    result["report_types"] = sorted({b["report_type"] for b in blocks})
    return result


def _summarize_lifecycle(rows, max_age, evaluation_date=None):
    evaluated_at = _parse_date(evaluation_date) if evaluation_date else datetime.now(timezone.utc)
    dates = [_parse_date(row["Report Date"]) for row in rows]
    known = [date for date in dates if date]
    freshness = "missing" if not rows else ("unknown" if any(date is None or date.date() > evaluated_at.date() for date in dates) else (
        "stale" if (evaluated_at.date() - min(known).date()).days > max_age else "fresh"
    ))
    date_bases = sorted({row.get("Report Date Basis", "Unknown") for row in rows})
    return {
        "available": bool(rows), "site_count": len(rows),
        "ownerless_site_count": sum(row["Is Ownerless"] == "Yes" for row in rows),
        "inactive_site_count": sum(row["Is Inactive"] == "Yes" for row in rows),
        "unknown_owner_status_count": sum(row["Is Ownerless"] == "Unknown" for row in rows),
        "missing_owner_contact_count": sum(not row["Owner"] for row in rows),
        "freshness": freshness,
        "max_age_days": max_age,
        "age_days": (evaluated_at.date() - min(known).date()).days if known else None,
        "date_basis": date_bases[0] if len(date_bases) == 1 else "Mixed report date sources" if date_bases else "Unknown",
        "operator_confirmed_dates": sorted({row["Report Date"] for row in rows
            if row.get("Report Date Basis") == "Operator-confirmed report date"}),
        "latest_report_date": max(known).isoformat() if known else "",
        "oldest_report_date": min(known).isoformat() if known else "",
    }


def _coverage_finding(feature, source_name, cli_option, env_name, documentation_url, instructions, errors=None):
    error_text = f" {len(errors)} supplied file(s) could not be read." if errors else ""
    return new_recommendation(
        "Data Exposure",
        feature,
        f"{source_name} evidence was not available, so this assessment did not verify this exposure domain.{error_text}",
        f"{instructions} Export the completed report, then rerun with {cli_option} PATH or the {env_name} environment variable. Enabling this capability is optional, but exposure remains unverified until recent evidence is supplied.",
        "Collection guidance",
        documentation_url,
        priority="High",
        status=NOT_ASSESSED_STATUS,
        category=CATEGORY_SCAN_COVERAGE,
        finding_key=f"data_exposure.coverage.{feature.lower().replace(' ', '_')}",
        disposition="Coverage",
        impact_area="Content access & grounding",
        ai_applicability="All AI using M365 data",
        evidence_basis="Not verified",
        confidence="Unknown",
    )


def _freshness_finding(feature, source_name, scan, cli_option, env_name, documentation_url, run_instructions):
    freshness = scan["freshness"]
    if freshness == "fresh":
        return None
    if freshness == "stale":
        observation = (
            f"The oldest selected {source_name} evidence is {scan['age_days']} day(s) old, "
            f"which exceeds the {scan['max_age_days']}-day freshness threshold."
        )
    else:
        observation = (
            f"The supplied {source_name} export did not contain a recognized report or scan date, "
            "so its freshness could not be verified."
        )
    return new_recommendation(
        "Data Exposure",
        feature,
        observation,
        f"{run_instructions} Export the completed result and rerun with {cli_option} PATH or {env_name}. The freshness threshold can be changed with {'SAM_REPORT_MAX_AGE_DAYS' if scan['kind'] == 'sam' else 'DSPM_REPORT_MAX_AGE_DAYS'}.",
        "Scan guidance",
        documentation_url,
        priority="High",
        status=NOT_ASSESSED_STATUS,
        category=CATEGORY_SCAN_COVERAGE,
        finding_key=f"data_exposure.freshness.{scan['kind']}",
        disposition="Coverage",
        impact_area="Content access & grounding",
        ai_applicability="All AI using M365 data",
        evidence_basis="Stale tenant evidence" if freshness == "stale" else "Freshness not verified",
        confidence="Low",
    )


def _risk_finding(feature, observation, recommendation, priority, key, url, confidence="High"):
    return new_recommendation(
        "Data Exposure",
        feature,
        observation,
        recommendation,
        "Remediation guidance",
        url,
        priority=priority,
        status="Action Required" if priority == "High" else "Attention Required",
        finding_key=f"data_exposure.{key}",
        evidence_key="data_exposure_detail",
        evidence_summary="See Data Exposure Detail for the exact sites or items and the report signals that produced this finding.",
        disposition="Action",
        impact_area="Content access & grounding",
        ai_applicability="All AI using M365 data",
        evidence_basis="Tenant evidence",
        confidence=confidence,
    )


def _findings_from_scan(scan):
    signals = scan["signals"]
    findings = []
    confidence = "High" if scan.get("freshness") == "fresh" else "Medium"

    def count_text(value, singular, plural=None):
        return f"{value:,} {singular if value == 1 else plural or singular + 's'}"

    def location_text(*signal_names):
        # The scan's affected_sites contains locations for every risk family.
        # Only detail rows for this finding can describe its affected scope.
        locations = defaultdict(set)
        for row in scan.get("risk_rows", []):
            if not set(signal_names).intersection(str(row.get("Risk Signals", "")).split("; ")):
                continue
            identity = row.get("Site URL") or row.get("Site ID") or row.get("Site Name")
            if identity:
                locations[row.get("Workload") or "Unknown"].add(str(identity).lower())
        parts = []
        for workload, identities in sorted(locations.items()):
            unit = "SharePoint site" if workload == "SharePoint" else "OneDrive location" if workload == "OneDrive" else "identified content location"
            parts.append(count_text(len(identities), unit))
        if not parts:
            return " The affected locations are not identified in the retained detail."
        scope = " The supporting detail identifies " + ", ".join(parts) + " for this finding."
        if scan.get("evidence_truncated"):
            scope += " The retained detail is truncated, so these are minimum location counts."
        return scope

    anyone = signals.get("Anyone links", 0)
    if anyone:
        findings.append(_risk_finding(
            "Review links that allow access without sign-in",
            f"The selected reports record {count_text(anyone, 'Anyone link')} allowing access without sign-in."
            + location_text("Anyone links"),
            "Validate business need, remove anonymous links that are not explicitly required, replace them with least-privilege links, and use sensitivity labels or DLP for sensitive content.",
            "High", "anonymous_links", "https://learn.microsoft.com/en-us/sharepoint/data-access-governance-reports", confidence,
        ))

    broad_parts = []
    for key, recipient in (("Everyone permissions", "Everyone"),
                           ("EEEU permissions", "Everyone except external users")):
        value = signals.get(key, 0)
        if value:
            broad_parts.append(count_text(value, "permission") + " granted to " + recipient)
    organization_links = signals.get("Organization links", 0)
    if organization_links:
        broad_parts.append(count_text(organization_links, "link") + " accessible to people in the organization")
    if broad_parts:
        findings.append(_risk_finding(
            "Review organization-wide links and group permissions",
            "The selected reports record " + "; ".join(broad_parts) + ". These measures can overlap and are not a count of distinct files or users."
            + location_text("Everyone permissions", "EEEU permissions", "Organization links"),
            "Have site owners validate the audience, remove broad permissions that are not required, and initiate site access reviews for the affected sites.",
            "High", "broad_internal_access", "https://learn.microsoft.com/en-us/sharepoint/site-access-review", confidence,
        ))

    external = signals.get("External exposure", 0)
    if external:
        findings.append(_risk_finding(
            "External sharing exposure",
            "The selected reports contain external-access permissions, recipients or links requiring review. Counts of these different measures do not establish a number of distinct exposed files or people."
            + location_text("External exposure"),
            "Review the external recipients and link purpose, remove expired or unjustified access, and apply expiration and least-privilege sharing defaults.",
            "High", "external_exposure", "https://learn.microsoft.com/en-us/sharepoint/turn-external-sharing-on-or-off", confidence,
        ))

    potential = signals.get("Potentially overshared items", 0)
    if potential:
        findings.append(_risk_finding(
            "Potentially overshared content",
            f"Purview DSPM identified {count_text(potential, 'potentially overshared item')}."
            + location_text("Potentially overshared items"),
            "Use the DSPM item-level results to resolve false positives, notify owners, apply labels, or remove excessive sharing links. Prioritize sensitive and unlabeled content first.",
            "High", "potentially_overshared", "https://learn.microsoft.com/en-us/purview/data-security-posture-management-oversharing", confidence,
        ))

    unlabeled = signals.get("Unlabeled sensitive items", 0)
    if unlabeled:
        findings.append(_risk_finding(
            "Sensitive content without labels",
            f"Purview DSPM identified {count_text(unlabeled, 'sensitive item')} without an adequate sensitivity label."
            + location_text("Unlabeled sensitive items"),
            "Validate the classifications and deploy an appropriate sensitivity label or auto-labeling policy before expanding AI access to the affected content.",
            "High", "unlabeled_sensitive", "https://learn.microsoft.com/en-us/purview/apply-sensitivity-label-automatically", confidence,
        ))

    ownerless = signals.get("Ownerless site", 0)
    if ownerless:
        findings.append(_risk_finding(
            "Ownerless SharePoint sites",
            f"SharePoint governance evidence identified {ownerless:,} site(s) without a recorded primary administrator or owner.",
            "Assign accountable owners and require them to attest that membership, sharing links, permissions, labels, and continued business need are correct.",
            "Medium", "ownerless_sites", "https://learn.microsoft.com/en-us/sharepoint/site-ownership-policy", confidence,
        ))

    return findings


def _scope_finding(scan):
    coverage = scan["coverage"]
    missing = coverage["domains_missing"] + [f"{name} (blank or invalid values)" for name in coverage["domains_unknown"]] + [f"{name} workload" for name in coverage["workloads_missing"]]
    if not scan["files_loaded"] or not missing:
        return None
    return new_recommendation(
        "Data Exposure", "Exported report coverage",
        f"The selected {scan['kind'].upper()} reports do not establish coverage for: " + ", ".join(missing) + ".",
        "Supply the relevant completed reports and confirm their scope. Conclusions are limited to the exported rows and fields; a narrow or filtered report does not establish tenant-wide coverage.",
        "Report guidance", "https://learn.microsoft.com/en-us/sharepoint/data-access-governance-reports" if scan["kind"] == "sam" else "https://learn.microsoft.com/en-us/purview/data-security-posture-management-oversharing",
        priority="High", status=NOT_ASSESSED_STATUS, category=CATEGORY_SCAN_COVERAGE,
        finding_key=f"data_exposure.scope.{scan['kind']}", disposition="Coverage",
        evidence_key="data_exposure_detail", evidence_basis="Exported report schema", confidence="High",
    )


def _lifecycle_findings(summary):
    if not summary["available"]:
        return []
    findings = []
    for field, feature, observation, action, key in (
        ("ownerless_site_count", "Ownerless SharePoint sites", "site(s) explicitly marked ownerless",
         "Assign accountable owners and have them review site membership, sharing and continued business need.", "ownerless_sites"),
        ("inactive_site_count", "Inactive SharePoint sites", "site(s) explicitly marked inactive",
         "Have site owners confirm continued business need and review permissions. Consider archiving after checking retention obligations and dependencies; inactivity alone does not establish oversharing.", "inactive_sites"),
    ):
        if summary[field]:
            finding = _risk_finding(feature,
                f"Content Management Assessment identified {summary[field]:,} {observation} in the supplied export.",
                action, "Medium", key, "https://learn.microsoft.com/en-us/sharepoint/content-management-assessment",
                "High" if summary["freshness"] == "fresh" else "Medium")
            finding["EvidenceKey"] = "sharepoint_lifecycle_detail"
            finding["EvidenceSummary"] = "See SharePoint Lifecycle Detail for the explicit Microsoft flags and owner contact fields."
            findings.append(finding)
    if summary["freshness"] != "fresh":
        findings.append(new_recommendation(
            "Data Exposure", "SharePoint lifecycle report freshness",
            "The Content Management Assessment export has no recognized report date." if summary["freshness"] == "unknown" else "The Content Management Assessment evidence is older than the accepted freshness window.",
            "Confirm the generation date in the SharePoint admin center and obtain a current export. Site creation and last activity dates describe individual sites; they do not establish when this report was generated.",
            "Assessment guidance", "https://learn.microsoft.com/en-us/sharepoint/content-management-assessment",
            priority="Medium", status=NOT_ASSESSED_STATUS, category=CATEGORY_SCAN_COVERAGE,
            finding_key="data_exposure.lifecycle.freshness", disposition="Coverage",
            evidence_key="sharepoint_lifecycle_detail", evidence_basis="Freshness not verified", confidence="Unknown",
        ))
    return findings


def build_data_exposure_assessment(sam_report_paths=None, dspm_report_paths=None, enabled=True, evaluation_date=None,
                                   sam_max_age_days=None, dspm_max_age_days=None,
                                   lifecycle_max_age_days=None, lifecycle_report_dates=None):
    """Load SAM/DSPM exports and create a dedicated, evidence-backed assessment result."""
    if not enabled and not sam_report_paths and not dspm_report_paths:
        return {"available": False, "recommendations": [], "evidence_rows": [], "lifecycle_evidence_rows": [], "lifecycle_summary": {}, "sources": {}, "operator_messages": []}

    sam_paths = _expand_paths(_split_env_paths("SAM_DAG_REPORT_PATHS") if sam_report_paths is None else sam_report_paths)
    dspm_paths = _expand_paths(_split_env_paths("DSPM_REPORT_PATHS") if dspm_report_paths is None else dspm_report_paths)
    sam = _scan_reports(sam_paths, "sam", evaluation_date, sam_max_age_days,
        lifecycle_max_age_days=lifecycle_max_age_days, lifecycle_report_dates=lifecycle_report_dates)
    dspm = _scan_reports(dspm_paths, "dspm", evaluation_date, dspm_max_age_days)

    recommendations = _lifecycle_findings(sam["lifecycle_summary"])
    if not sam["files_loaded"]:
        recommendations.append(_coverage_finding(
            "SharePoint oversharing evidence",
            "SharePoint Advanced Management Data Access Governance",
            "--sam-report",
            "SAM_DAG_REPORT_PATHS",
            "https://learn.microsoft.com/en-us/sharepoint/data-access-governance-reports",
            "First check SharePoint admin center > Reports > Data access governance for a completed report within the accepted freshness window and export it if one exists. If no current report exists, create the site-permissions baseline plus sharing-link and EEEU reports and let Microsoft complete them. If Data Access Governance is unavailable, optionally confirm SharePoint Advanced Management licensing and assign the SharePoint Administrator or SharePoint Advanced Management Administrator role.",
            errors=sam["errors"],
        ))
    else:
        recommendations.extend(_findings_from_scan(sam))
        freshness = _freshness_finding(
            "SharePoint oversharing scan freshness",
            "SharePoint Data Access Governance",
            sam,
            "--sam-report",
            "SAM_DAG_REPORT_PATHS",
            "https://learn.microsoft.com/en-us/sharepoint/data-access-governance-reports",
            "In the SharePoint admin center, open Reports > Data access governance and create or rerun the required reports; allow Microsoft to finish them.",
        )
        if freshness:
            recommendations.append(freshness)

    if not dspm["files_loaded"]:
        recommendations.append(_coverage_finding(
            "Sensitive-data exposure evidence",
            "Microsoft Purview DSPM data-risk assessment",
            "--dspm-report",
            "DSPM_REPORT_PATHS",
            "https://learn.microsoft.com/en-us/purview/data-security-posture-management-oversharing",
            "First check Microsoft Purview > Data Security Posture Management > Discover > Data risk assessments for a recent completed default or custom assessment and export it if one exists. If no current result exists, let the default assessment complete or run a custom assessment for the required SharePoint sites. If DSPM is not configured, optionally complete its setup tasks for Audit and analytics first.",
            errors=dspm["errors"],
        ))
    else:
        recommendations.extend(_findings_from_scan(dspm))
        freshness = _freshness_finding(
            "Sensitive-data exposure scan freshness",
            "Purview DSPM data-risk assessment",
            dspm,
            "--dspm-report",
            "DSPM_REPORT_PATHS",
            "https://learn.microsoft.com/en-us/purview/data-security-posture-management-oversharing",
            "In the Microsoft Purview portal, open Data Security Posture Management > Discover > Data risk assessments and wait for the default or custom assessment to complete.",
        )
        if freshness:
            recommendations.append(freshness)

    for scan in (sam, dspm):
        scope = _scope_finding(scan)
        if scope:
            recommendations.append(scope)

    total_signals = sum(sam["signals"].values()) + sum(dspm["signals"].values())
    complete_fields = all(not scan["coverage"]["domains_missing"] and not scan["coverage"]["domains_unknown"] and not scan["coverage"]["workloads_missing"] for scan in (sam, dspm))
    if (sam["freshness"] == "fresh" and dspm["freshness"] == "fresh" and total_signals == 0 and complete_fields and not sam["errors"] and not dspm["errors"]):
        recommendations.append(new_recommendation(
            "Data Exposure",
            "Supplied exposure report fields",
            f"No explicit exposure signal was present in the selected SAM and DSPM exported fields ({sam['records_selected'] + dspm['records_selected']:,} record(s)). This conclusion applies to the exported scope; tenant-wide completeness is not verified.",
            "Confirm export filters and assessment scope before using this result to support a tenant-wide readiness decision.",
            "Assessment guidance",
            "https://learn.microsoft.com/en-us/purview/data-security-posture-management-oversharing",
            status="Success",
            finding_key="data_exposure.assessed_no_signals",
            evidence_key="data_exposure_detail",
            evidence_summary="See Data Exposure Detail for source files, record counts, report dates, and freshness state.",
            disposition="Assurance",
            impact_area="Content access & grounding",
            ai_applicability="All AI using M365 data",
            evidence_basis="Tenant evidence",
            confidence="Medium",
        ))

    operator_messages = []
    for source_name, scan in (
        ("SharePoint Data Access Governance", sam),
        ("Purview DSPM data-risk assessment", dspm),
    ):
        if scan["freshness"] == "fresh":
            operator_messages.append(
                f"Using recent {source_name} evidence dated {scan['oldest_report_date'][:10]} to {scan['latest_report_date'][:10]} "
                f"({scan['records_selected']:,} selected record(s))."
            )
        elif scan["freshness"] == "stale":
            operator_messages.append(
                f"The oldest selected {source_name} evidence is {scan['age_days']} day(s) old. Confirm its continued relevance or supply a newer completed report for that scope."
            )
        elif scan["freshness"] == "unknown":
            operator_messages.append(
                f"{source_name} evidence was parsed, but its scan date could not be verified; "
                "the report requests a dated export."
            )
        else:
            operator_messages.append(
                f"No usable {source_name} export was supplied; check for a recent completed scan before "
                "starting a new one. The report records this as Coverage and includes optional enablement, "
                "scan, and export instructions."
            )

    if sam["lifecycle_summary"]["available"]:
        summary = sam["lifecycle_summary"]
        operator_messages.append(
            f"Content Management Assessment: {summary['site_count']:,} site(s), {summary['inactive_site_count']:,} explicitly inactive, "
            f"{summary['ownerless_site_count']:,} explicitly ownerless; freshness {summary['freshness']}. "
            "Lifecycle and ownership evidence does not replace permission or sensitive-data exposure reports."
        )
    for report in sam["reports"]:
        if report["status"] == "superseded":
            operator_messages.append(f"Excluded superseded {report['workload']} snapshot {report['source_file']} dated {report['report_date']} from current exposure counts.")
        elif report["status"] == "empty":
            operator_messages.append(f"Recognized {report['report_type']} export {report['source_file']} with zero data rows. Its scan date and tenant exposure cannot be established from the header or filename.")
        elif report['status'] == 'future_date':
            operator_messages.append(f"Excluded {report['source_file']} from exposure totals: report date {report['report_date']} is after the evaluation date. Confirm the date before using this snapshot.")

    for finding in recommendations:
        lifecycle = finding.get('EvidenceKey') == 'sharepoint_lifecycle_detail'
        text = ' '.join(str(finding.get(k, '')) for k in ('FindingKey', 'Feature')).lower()
        source = dspm if any(word in text for word in ('dspm', 'sensitive')) else sam
        reports = [r for r in source['reports'] if r['status'] == 'selected'
            and (r['report_type'] == 'content_management_assessment') == lifecycle]
        if not reports:
            continue
        dates = [r['report_date'] for r in reports if r.get('report_date')]
        finding.update(SourceType='portal_export', SourceFile='; '.join(sorted({r['source_file'] for r in reports})),
            ObservationDate=min(dates) if len(dates) == len(reports) else '',
            EvidenceScope='Exported objects and fields only', EvidenceComplete=False)
        if lifecycle:
            finding['EvidenceMaxAgeDays'] = sam['lifecycle_summary']['max_age_days']
            finding['EvidenceDateBasis'] = sam['lifecycle_summary']['date_basis']
            hashes = {r['source_hash'] for r in reports if r.get('source_hash')}
            if len(hashes) == 1:
                finding['SourceHash'] = next(iter(hashes))
        if dates:
            finding['EvidencePeriodStart'], finding['EvidencePeriodEnd'] = min(dates), max(dates)
    if any('populations differ' in r.get('scope_note', '') for r in sam['reports']):
        recommendations.append(new_recommendation(
            'Data Exposure', 'Confirm the scope of changed permission snapshots',
            'The permission snapshots cover different site populations. Older objects absent from a newer export remain dated observations; their removal or remediation has not been established.',
            'Ask the SharePoint owner to confirm export filters and the status of sites omitted from the newer snapshot. Review the dated records in the evidence workbook.',
            status='Not Assessed', disposition='Coverage', priority='Medium',
            finding_key='data_exposure.snapshot_scope', evidence_key='data_exposure_detail', confidence='Unknown'))

    return {
        "available": bool(sam["files_loaded"] or dspm["files_loaded"] or sam["lifecycle_summary"]["available"]),
        "recommendations": recommendations,
        "evidence_rows": sam["risk_rows"] + dspm["risk_rows"],
        "lifecycle_evidence_rows": sam["lifecycle_rows"],
        "lifecycle_summary": sam["lifecycle_summary"],
        "sources": {"sam": sam, "dspm": dspm},
        "evidence_truncated": sam["evidence_truncated"] or dspm["evidence_truncated"] or sam["lifecycle_records_truncated"],
        "operator_messages": operator_messages,
        "observations": sam['observations'] + dspm['observations'],
    }
