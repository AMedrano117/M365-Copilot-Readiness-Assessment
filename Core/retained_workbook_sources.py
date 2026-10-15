"""Recover native replay evidence from recorded, same-run technical workbooks.

Older snapshots retain native row references but not the source envelopes. Their
technical export does retain those envelopes. Accept them only when the existing
deterministic evidence hashes reproduce every exported native identifier.
"""
import json
from pathlib import Path

from .evidence_selection import _dataset_id, _evidence_id, evidence_record_ids


def _json_field(row, key):
    pieces = [row.get(key) or '']
    number = 2
    while f'{key} (continued {number})' in row:
        pieces.append(row[f'{key} (continued {number})'] or '')
        number += 1
    return json.loads(''.join(pieces))


def native_references(value):
    """Inspect explicit references, never infer associations from record text."""
    if isinstance(value, dict):
        if {'dataset', 'dataset_index', 'record_index'} <= value.keys():
            yield (value['dataset'], value['dataset_index'], value['record_index'])
        for child in value.values():
            yield from native_references(child)
    elif isinstance(value, list):
        for child in value:
            yield from native_references(child)


def read_native_sources(path, result, retained_sheets=None, retained_selection=None):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        manifest = dict((row[0], row[1]) for row in book['Run Manifest'].iter_rows(values_only=True)
                        if len(row) >= 2 and row[0])
        identity = result.get('identity') or {}
        for key in ('AssessmentId', 'RunId'):
            if not identity.get(key) or manifest.get(key) != identity[key]:
                raise ValueError('Retained technical workbook belongs to a different assessment run.')
        if manifest.get('Tenant ID') != result.get('tenant_id'):
            raise ValueError('Retained technical workbook belongs to a different tenant.')
        sources, blocks = {}, []
        required = {'Evidence Record ID', 'Source dataset', 'Source record', 'Source metadata', 'Raw record'}
        for sheet in book:
            if not sheet.title.startswith('Raw '):
                continue
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, ())
            if not required <= set(headers):
                continue
            block = None
            for values in rows:
                row = dict(zip(headers, values))
                if not any(values):
                    continue
                position = row['Source record']
                if not isinstance(position, int) or position < 1:
                    raise ValueError('Retained native source record numbering is invalid.')
                source, raw = _json_field(row, 'Source metadata'), _json_field(row, 'Raw record')
                if position == 1:
                    block = {'name': row['Source dataset'], 'source': source, 'records': [], 'ids': []}
                    blocks.append(block)
                if (block is None or block['name'] != row['Source dataset'] or block['source'] != source
                        or position != len(block['records']) + 1):
                    raise ValueError('Retained native source sequence or capture metadata changed.')
                block['records'].append(raw)
                block['ids'].append(row['Evidence Record ID'])
        refs = set(native_references(result))
        upper = max([index for _, index, _ in refs] + [len(blocks), 0]) + 1
        for block in blocks:
            name, source = block['name'], block['source']
            first = block['records'][0]
            # The raw-sheet writer wraps scalar/list records in a display-only
            # value column. Restore that shape only when the native hash agrees.
            shapes = [False, True] if isinstance(first, dict) and set(first) == {'value'} else [False]
            matches = [(index, unwrap) for index in range(upper) for unwrap in shapes
                       if _evidence_id(result['tenant_id'], _dataset_id(result['tenant_id'], name, index, source),
                                       0, first['value'] if unwrap else first) == block['ids'][0]]
            if not matches:
                # Parent-object collectors can retain thousands of empty captures
                # between populated captures. The hash, not sheet order or the
                # display name, proves the original index.
                for index in range(upper, 10000):
                    for unwrap in shapes:
                        if _evidence_id(result['tenant_id'], _dataset_id(result['tenant_id'], name, index, source),
                                        0, first['value'] if unwrap else first) == block['ids'][0]:
                            matches.append((index, unwrap))
                    if matches:
                        break
            if len(matches) != 1:
                raise ValueError(f'Retained native evidence identity cannot be reproduced for {name}.')
            index, unwrap = matches[0]
            if unwrap:
                if any(not isinstance(row, dict) or set(row) != {'value'} for row in block['records']):
                    raise ValueError('Retained native record shape changed within the source capture.')
                block['records'] = [row['value'] for row in block['records']]
            datasets = sources.setdefault(name, [])
            while len(datasets) <= index:
                datasets.append({'source': {}, 'records': []})
            if datasets[index]['records']:
                raise ValueError('Duplicate retained native source dataset.')
            dataset_id = _dataset_id(result['tenant_id'], name, index, source)
            if any(_evidence_id(result['tenant_id'], dataset_id, position, raw) != identifier
                   for position, (raw, identifier) in enumerate(zip(block['records'], block['ids']))):
                raise ValueError('Retained native evidence row integrity mismatch.')
            datasets[index] = {'source': source, 'records': block['records']}
        if refs - evidence_record_ids(sources, result['tenant_id']).keys():
            raise ValueError('Retained technical workbook does not contain every snapshot native reference.')
        if retained_sheets is not None:
            from .evidence_layer import SHEET_DEFINITIONS
            keys = {value['title']: key for key, value in SHEET_DEFINITIONS.items()}
            keys['Application Grant Detail'] = 'application_grant_detail'
            keys['Investigation Items'] = 'investigation_items'
            keys['Sharing Site Configuration'] = 'sharepoint_site_configuration_detail'
            rebuilt = {'Start Here', 'Action Plan', 'Findings', 'Findings Register', 'Findings Lineage',
                       'Recommendations', 'Evidence Index', 'Control Results', 'Collection Coverage',
                       'Run Manifest', 'Integrity Checks', 'Configuration Detail', 'Action Detail',
                       'Coverage Detail', 'Technical Recommendations', 'Raw Derived Records'}
            for sheet in book:
                if sheet.title in rebuilt or sheet.title.startswith(('Raw ', 'Records ', 'Evidence ')):
                    continue
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, ())
                records = [{key: value for key, value in zip(headers, values) if key is not None}
                           for values in rows if any(values)]
                if records:
                    key = keys.get(sheet.title, 'retained.' + sheet.title)
                    retained_sheets[key] = {'title': sheet.title, 'rows': records, 'restricted': True,
                                            'preview_columns': [], 'summary': 'Retained same-run supporting records.'}
        if retained_selection is not None:
            # Record IDs are authoritative in the completed export. Replaying
            # its display fields cannot recreate hashes of the earlier typed
            # materialization, so preserve IDs via exact finding/detail/EVD joins.
            from .finding_evidence import COMMON_COLUMNS
            detail_metadata, selected, derived, source_ids = {}, {}, {}, {}
            for sheet in book:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, ())
                if {'evidence_record_ids', 'dataset_id'} <= set(headers):
                    for values in sheet.iter_rows(min_row=2, values_only=True):
                        row = dict(zip(headers, values))
                        for identifier in str(row.get('evidence_record_ids') or '').split(';'):
                            if identifier.strip() and row.get('dataset_id'):
                                source_ids.setdefault(identifier.strip(), set()).add(row['dataset_id'])
                if {'finding_id', 'detail_record_id', 'evidence_record_ids'} <= set(headers):
                    for values in rows:
                        row = dict(zip(headers, values))
                        detail_metadata[(row['finding_id'], row['detail_record_id'])] = row
                elif sheet.title == 'Raw Derived Records':
                    for values in rows:
                        row = dict(zip(headers, values))
                        derived[row['Evidence Record ID']] = row
                elif (sheet.title.startswith(('Records ', 'Logs ')) and {'RecommendationId', 'Detail record ID'} <= set(headers)
                      and ('Evidence Record ID' in headers or 'Evidence ID' in headers)):
                    evidence_column = 'Evidence Record ID' if 'Evidence Record ID' in headers else 'Evidence ID'
                    for values in rows:
                        row = dict(zip(headers, values))
                        key = (row['RecommendationId'], row['Detail record ID'])
                        record = selected.setdefault(key, {'record_id': row['Detail record ID'],
                            'evidence_record_ids': [], 'fields': {}, 'source_refs': [], 'qualifications': []})
                        if row[evidence_column] not in record['evidence_record_ids']:
                            record['evidence_record_ids'].append(row[evidence_column])
                        record['fields'].update({(name[7:] if name.startswith('Record.') else name): value
                            for name, value in row.items() if name not in {'RecommendationId', 'Detail record ID', evidence_column} and value is not None})
            known_findings = {row['RecommendationId'] for row in result.get('recommendations') or []}
            native = set(evidence_record_ids(sources, result['tenant_id']).values())
            # Several selected details can cite one native row. The full-source
            # register stores that row once; its selected table retains each
            # distinct detail ID and qualification.
            for key, metadata in detail_metadata.items():
                if key not in selected:
                    selected[key] = {'record_id': key[1], 'fields': {}, 'source_refs': [], 'qualifications': [],
                        'evidence_record_ids': [value.strip() for value in str(metadata.get('evidence_record_ids') or '').split(';') if value.strip()]}
            for (finding_id, detail_id), record in selected.items():
                valid_detail = str(detail_id).startswith('DET-') or detail_id in native
                if finding_id not in known_findings or not valid_detail:
                    raise ValueError('Retained selected record has unresolved finding/detail ownership.')
                if set(record['evidence_record_ids']) - native - derived.keys():
                    raise ValueError('Retained selected record has unresolved evidence lineage.')
                metadata = detail_metadata.get((finding_id, detail_id)) or {}
                if metadata:
                    record['fields'] = {name.removesuffix(' (record field)'): value for name, value in metadata.items()
                                        if name not in COMMON_COLUMNS and value is not None}
                    if metadata.get('qualification'):
                        record['qualifications'] = [metadata['qualification']]
                    record['source_record_id'] = metadata.get('source_record_id')
                retained_selection.setdefault('records', {}).setdefault(finding_id, []).append(record)
            retained_selection['derived'] = derived
            retained_selection['selected_findings'] = sorted({finding for finding, _ in detail_metadata})
            retained_selection['derived_source_ids'] = {}
            for identifier in derived:
                candidates = source_ids.get(identifier, set())
                if len(candidates) != 1:
                    raise ValueError('Retained derived evidence has unresolved source-ID lineage.')
                retained_selection['derived_source_ids'][identifier] = next(iter(candidates))
        return sources
    finally:
        book.close()


def package_native_sources(folder, result, retained_sheets=None, retained_selection=None):
    """Select only technical exports recorded by a completed same-run receipt."""
    folder = Path(folder).resolve()
    log = folder / 'operator-log.jsonl'
    if not log.is_file():
        # Self-contained snapshots already retain their own selected detail.
        # Native adapters reject unresolved references when they consume them.
        return {}
    from .assessment_package import confined_path
    candidates = []
    for line in log.read_text(encoding='utf-8-sig').splitlines():
        receipt = json.loads(line)
        if receipt.get('status') != 'complete':
            continue
        if (receipt.get('semantic_run') or {}).get('RunId') != (result.get('identity') or {}).get('RunId'):
            continue
        for locator in receipt.get('deliverables') or []:
            if Path(locator).name.startswith('Technical Evidence') and Path(locator).suffix.lower() == '.xlsx':
                path = confined_path(folder, locator)
                if path.is_file() and path not in candidates:
                    candidates.append(path)
    if candidates:
        return read_native_sources(candidates[0], result, retained_sheets, retained_selection)
    if any(native_references(result)):
        raise ValueError('Snapshot native lineage requires the recorded same-run technical workbook.')
    return {}
