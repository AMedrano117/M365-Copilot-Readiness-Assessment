"""Workbook guides, visible destinations and restrained table formatting."""

from datetime import date, datetime
from pathlib import Path
import re


def apply_workbook_navigation(workbook, result, bundle, *, role='assessment', companion_path=None, tenant_name=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter
    from .export_recommendations import _append_dict_rows_to_sheet
    from .workbook_layout import ASSESSMENT_ORDER, COMPATIBILITY_TITLES

    hidden = set(COMPATIBILITY_TITLES) if role == 'assessment' else set()
    fill = PatternFill('solid', fgColor='102B40')
    font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    wrap = Alignment(vertical='top', wrap_text=True)
    for sheet in workbook:
        sheet.sheet_state = 'hidden' if sheet.title in hidden else 'visible'
        sheet.sheet_view.zoomScale = 85
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = ('007F85' if sheet.title in {'Action Plan', 'Findings', 'Findings Lineage'} else
                                          'C49943' if 'Coverage' in sheet.title else '102B40')
        sheet.freeze_panes = 'A2'
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.fill, cell.font, cell.alignment = fill, font, wrap
            cell.data_type = 's'
        for name, column in {str(cell.value or ''): cell.column for cell in sheet[1]}.items():
            letter = get_column_letter(column)
            narrative = name.lower() in {'what we found', 'what to do', 'recommended action', 'detail', 'description', 'purpose', 'issue',
                                         'source excerpt', 'follow-up', 'qualification'}
            longest = max((len(str(sheet.cell(number, column).value or '')) for number in range(1, sheet.max_row + 1)), default=12)
            sheet.column_dimensions[letter].width = 60 if narrative else min(max(longest + 2, 14), 45)
            for number in range(2, sheet.max_row + 1):
                cell = sheet.cell(number, column)
                cell.alignment = wrap
                value = cell.value
                is_id = bool(re.search(r'(^id$|\bid\b|identifier|guid|upn|url|principalname)', name, re.I))
                if isinstance(value, str):
                    cell.data_type = 's'
                if is_id:
                    if value is not None and not isinstance(value, str):
                        cell.value = str(value)
                        cell.data_type = 's'
                    cell.number_format = '@'
                elif isinstance(value, datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm'
                elif isinstance(value, date):
                    cell.number_format = 'yyyy-mm-dd'
                elif isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = '0.0%' if '%' in name or 'percentage' in name.lower() else '#,##0' if isinstance(value, int) else '0.0'
            states = ({'Critical': 'F7DEDA', 'High': 'F7DEDA', 'Medium': 'FFF1D2', 'Low': 'E4F1EA'} if name == 'Priority' else
                      {'Change': 'F7DEDA', 'Review': 'FFF1D2', 'Meets': 'E4F1EA'} if name == 'Status' else
                      {'available': 'E4F1EA', 'missing': 'F7DEDA', 'stale': 'FFF1D2', 'truncated': 'FFF1D2',
                       'incomplete': 'FFF1D2', 'owner review': 'FFF1D2'} if name == 'Evidence state' else {})
            if sheet.max_row > 1:
                for value, color in states.items():
                    sheet.conditional_formatting.add(f'{letter}2:{letter}{sheet.max_row}',
                        CellIsRule(operator='equal', formula=['"' + value + '"'], fill=PatternFill('solid', fgColor=color)))

    context = bundle.get('collection_context') or {}
    manifest = {row.get('Item'): row.get('Value') for row in (bundle.get('run_manifest') or {}).get('rows') or []}
    sources = bundle.get('assessment_sources') or {}
    timestamps = [str(entry['source'].get('collected_at') or entry['source'].get('refresh_date'))
                  for datasets in sources.values() for entry in datasets
                  if (entry.get('source') or {}).get('collected_at') or (entry.get('source') or {}).get('refresh_date')]
    if context.get('collected_at'):
        timestamps.append(str(context['collected_at']))
    companion = Path(companion_path).name if companion_path else ''
    entries = [
        {'Open': 'Purpose', 'Purpose': 'Findings, actions and the records supporting each finding.' if role == 'assessment' else
         'Complete retained source records, selected records by dataset, lineage and provenance for engineering review.', 'Rows': None, 'Type': 'Instructions'},
        {'Open': 'Tenant', 'Purpose': tenant_name or manifest.get('Tenant') or result.get('tenant_id') or 'Unknown', 'Rows': None, 'Type': 'Metadata'},
        {'Open': 'Evaluation date', 'Purpose': result.get('evaluation_date'), 'Rows': None, 'Type': 'Metadata'},
        {'Open': 'Latest collection timestamp', 'Purpose': max(timestamps, default='Unknown'), 'Rows': None, 'Type': 'Metadata'},
        {'Open': 'Run timestamp', 'Purpose': manifest.get('Generation Time (UTC)') or (bundle.get('run_manifest') or {}).get('generated_at') or 'Unknown', 'Rows': None, 'Type': 'Metadata'},
        {'Open': 'Methodology version', 'Purpose': result.get('methodology_version'), 'Rows': None, 'Type': 'Metadata'},
        {'Open': companion, 'Purpose': 'Technical workbook: complete raw exports, findings lineage, full selected records and historical context.' if role == 'assessment' else
         'Assessment workbook: prioritized actions, findings, coverage and readable supporting records.', 'Rows': None, 'Type': 'Companion'},
        {'Open': 'How to read the evidence', 'Purpose': 'Filter ID for one finding. Its evidence rows form a contiguous block. Evidence ID opens the exact technical source row. Records can support several findings; do not add counts across findings. Unknown values remain unknown.' if role == 'assessment' else
         'RecommendationId / finding_id joins findings. Detail record IDs join selected lineage; EVD evidence IDs join native Raw sheets. Derived records remain identified as derived. Raw inventories and selected populations are different scopes.', 'Rows': None, 'Type': 'Instructions'},
    ]
    if role == 'assessment':
        entries.append({'Open': 'Compatibility registers', 'Purpose': 'Hidden registers preserve prior-report, baseline and audit imports. Their links open visible evidence sheets in the technical workbook.', 'Rows': None, 'Type': 'Instructions'})
    source_sheets = {sheet.get('title'): sheet for sheet in (bundle.get('sheets') or {}).values()}
    for sheet in workbook:
        entries.append({'Open': sheet.title, 'Purpose': (source_sheets.get(sheet.title) or {}).get('summary') or
                       ('Compatibility register (hidden).' if sheet.sheet_state != 'visible' else
                        'Prioritized actions; numbers match the HTML report.' if sheet.title == 'Action Plan' else
                        'One row per finding with its supporting record block.' if sheet.title == 'Findings' else
                        'Required checks, evidence state and applicability.' if sheet.title == 'Coverage' else
                        'Page-backed PDF excerpts and follow-up; source context with separate capture and refresh dates.' if sheet.title == 'PDF Highlights' else
                        'Retained records and lineage.'), 'Rows': max(sheet.max_row - 1, 0),
                        'Type': 'Hidden register' if sheet.sheet_state != 'visible' else 'Worksheet'})
    guide = workbook.create_sheet('Start Here', 0)
    _append_dict_rows_to_sheet(guide, entries, fill, font, wrap, 'WorkbookGuide')
    for number, entry in enumerate(entries, 2):
        if entry['Open'] in workbook and workbook[entry['Open']].sheet_state == 'visible':
            guide.cell(number, 1).hyperlink = "#'" + entry['Open'].replace("'", "''") + "'!A1"
        elif entry['Type'] == 'Companion' and companion:
            guide.cell(number, 1).hyperlink = companion + "#'Start Here'!A1"
        if guide.cell(number, 1).hyperlink:
            guide.cell(number, 1).style = 'Hyperlink'
    for column, width in [('A', 36), ('B', 80), ('C', 14), ('D', 22)]:
        guide.column_dimensions[column].width = width
    guide.sheet_view.showGridLines = False
    guide.sheet_properties.tabColor = '007F85'
    leading = list(ASSESSMENT_ORDER) if role == 'assessment' else ['Start Here', 'Findings Lineage']
    if 'PDF Highlights' in workbook:
        leading.insert(4 if role == 'assessment' else 2, 'PDF Highlights')
    order = [title for title in leading if title in workbook] + [title for title in workbook.sheetnames if title not in leading]
    for index, title in enumerate(order):
        workbook.move_sheet(title, index - workbook.sheetnames.index(title))
    workbook.active = 0
