"""Read saved assessment results without rebuilding or invoking tenant collectors.

An assessment workbook/snapshot contains prior conclusions and exported evidence;
it is deliberately distinct from the raw collection saved by --save-collection.
"""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path


_RECOMMENDATION_FIELDS = (
    "RecommendationId", "ControlId", "MethodologyVersion", "FindingFingerprint",
    "FindingKey", "Service", "Disposition", "ReadinessStage", "ImpactArea",
    "AIApplicability", "Category", "Feature", "AlsoLicensedVia", "Status",
    "Priority", "Observation", "Recommendation", "LinkText", "LinkUrl",
    "EvidenceAvailable", "EvidenceSheet", "EvidenceKey", "EvidenceSummary",
    "EvidenceBasis", "Confidence", "ControlStatus",
)


def _key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


_FIELD_NAMES = {_key(name): name for name in _RECOMMENDATION_FIELDS}
_RESTRICTED_TITLES = {
    "copilotuserusage", "copilotreadinessusers", "copilotuseractivity",
    "copilotuserusagedetail", "copilotreadinessuserdetail",
}
_GUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)
_USER_DETAIL_KEYS = {"userdetail", "userdetails", "copilotuserusagedetail", "copilotreadinessuserdetail"}


def _restricted(title):
    normalized = _key(title)
    return normalized in _RESTRICTED_TITLES or (
        "copilot" in normalized and "user" in normalized
        and any(word in normalized for word in ("usage", "activity", "readiness"))
    )


def _cell(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _without_user_details(value):
    """Coverage objects can contain nested Copilot payloads in older snapshots."""
    if isinstance(value, list):
        return [_without_user_details(item) for item in value]
    if isinstance(value, dict):
        return {key: _without_user_details(item) for key, item in value.items()
                if _key(key) not in _USER_DETAIL_KEYS}
    return value


def _records(value, label):
    if not isinstance(value, list) or any(not isinstance(row, dict) for row in value):
        raise ValueError(f"Prior assessment {label} must be a list of records.")
    return [dict(row) for row in value]


def _recommendations(rows):
    result = []
    for row in _records(rows, "recommendations"):
        normalized = dict(row)
        for name, value in row.items():
            canonical = _FIELD_NAMES.get(_key(name))
            if canonical and canonical not in normalized:
                normalized[canonical] = value
        if not {"Service", "Feature", "Recommendation"}.issubset(normalized):
            raise ValueError("Prior assessment recommendations require Service, Feature, and Recommendation fields.")
        result.append(normalized)
    return result


def _worksheet_rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    first = next(iterator, ())
    if not first:
        return [], []
    headers = []
    for index, value in enumerate(first, 1):
        header = str(value).strip() if value is not None else f"Column {index}"
        if header in headers:
            header = f"{header} ({index})"
        headers.append(header)
    rows = [dict(zip(headers, (_cell(value) for value in row))) for row in iterator
            if any(value is not None for value in row)]
    return headers, rows


def _text(value, label):
    if value is None:
        return ""
    if not isinstance(value, (str, int, float, date, datetime)) or isinstance(value, bool):
        raise ValueError(f"Prior assessment {label} must be text.")
    return str(_cell(value)).strip()


def _coverage(value):
    if isinstance(value, list):
        return _records(value, "collection coverage")
    if not isinstance(value, dict):
        raise ValueError("Prior assessment collection coverage must be records or a source mapping.")
    result = []
    for source, state in value.items():
        if isinstance(state, str):
            state = {"availability_status": state}
        if not isinstance(state, dict):
            raise ValueError("Prior assessment collection coverage contains an invalid source state.")
        row = dict(state)
        row.setdefault("Source", source)
        for display, keys in {
            "State": ("availability_status", "status"),
            "Records": ("records_collected", "record_count"),
            "Pages": ("pages_collected",), "Truncated": ("truncated",),
            "Refresh Date": ("refresh_date",), "Freshness": ("freshness",),
            "Reason": ("reason",),
        }.items():
            for key in keys:
                if key in state:
                    row.setdefault(display, state[key])
                    break
        result.append(row)
    return result


def _read_workbook(path, include_user_details):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ValueError("Prior assessment workbook could not be read as XLSX.") from exc
    try:
        if "Recommendations" not in workbook.sheetnames:
            raise ValueError("Prior assessment workbook requires a Recommendations worksheet.")
        sheets = {}
        withheld = []
        recommendation_headers = []
        for sheet in workbook:
            if _restricted(sheet.title) and not include_user_details:
                withheld.append(sheet.title)
                continue
            headers, rows = _worksheet_rows(sheet)
            sheets[sheet.title] = {"rows": rows}
            if sheet.title == "Recommendations":
                recommendation_headers = headers
        canonical_headers = {_FIELD_NAMES.get(_key(header), header) for header in recommendation_headers}
        if not {"Service", "Feature", "Recommendation"}.issubset(canonical_headers):
            raise ValueError("Prior assessment Recommendations worksheet has an unsupported schema.")
        manifest = {str(row.get("Item", "")): row.get("Value", "")
                    for row in sheets.get("Run Manifest", {}).get("rows", [])}
        return {
            "source_kind": "assessment_workbook",
            "tenant_name": _text(manifest.get("Tenant", ""), "tenant name"),
            "tenant_id": _text(manifest.get("Tenant ID", ""), "tenant ID"),
            "generated_at": _text(manifest.get("Generation Time (UTC)", ""), "generation time"),
            "methodology_version": _text(manifest.get("Methodology Version", ""), "methodology version"),
            "assessment_version": _text(manifest.get("Assessment Version", ""), "assessment version"),
            "recommendations": _recommendations(sheets["Recommendations"]["rows"]),
            "collection_coverage": sheets.get("Collection Coverage", {}).get("rows", []),
            "control_results": sheets.get("Control Results", {}).get("rows", []),
            "sheets": sheets, "withheld_sheets": withheld, "warnings": [],
        }
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Prior assessment workbook contains unreadable worksheet data.") from exc
    finally:
        workbook.close()


def _read_snapshot(path, include_user_details):
    try:
        with path.open(encoding="utf-8-sig") as handle:
            payload = json.load(handle)
    except (OSError, UnicodeError, ValueError) as exc:
        raise ValueError("Prior assessment snapshot could not be read as JSON.") from exc
    if not isinstance(payload, dict):
        raise ValueError("Prior assessment snapshot must contain an assessment object.")
    if payload.get("format") == "m365-readiness-collection":
        raise ValueError("This JSON is a saved tenant collection; use --collection-input instead of --prior-report.")
    if "purview_data_json" in payload or "purview_data_json" in path.stem.lower():
        raise ValueError("This JSON is a Purview cache containing only a service subset; use --purview-cache for this file and --prior-report for a saved assessment workbook or snapshot.")
    if "recommendations" not in payload:
        raise ValueError("Prior assessment snapshot requires a recommendations array; a service cache is not a complete assessment.")
    recommendations = _recommendations(payload["recommendations"])
    manifest = payload.get("run_manifest", {})
    if not isinstance(manifest, dict):
        raise ValueError("Prior assessment run manifest must be an object.")
    metadata = payload.get("metadata", {})
    if not isinstance(metadata, dict):
        raise ValueError("Prior assessment metadata must be an object.")
    source_metadata = payload.get("source", {})
    if not isinstance(source_metadata, dict):
        source_metadata = {}
    sheets = {}
    withheld = []
    supplied_sheets = payload.get("sheets", {})
    if not isinstance(supplied_sheets, dict):
        raise ValueError("Prior assessment sheets must be a mapping.")
    for key, sheet in supplied_sheets.items():
        if not isinstance(sheet, dict):
            raise ValueError("Prior assessment sheets must contain worksheet records.")
        title = _text(sheet.get("title", key), "worksheet title")
        if _restricted(title) and not include_user_details:
            withheld.append(title)
            continue
        sheets[title] = {"rows": _records(sheet.get("rows", []), "worksheet rows")}
    coverage = _coverage(payload.get("collection_coverage", {}))
    controls = _records(payload.get("control_results", []), "control results")
    sheets.setdefault("Recommendations", {"rows": [dict(row) for row in recommendations]})
    if coverage:
        sheets.setdefault("Collection Coverage", {"rows": coverage})
    if controls:
        sheets.setdefault("Control Results", {"rows": controls})
    if manifest.get("rows"):
        sheets.setdefault("Run Manifest", {"rows": _records(manifest["rows"], "run manifest rows")})
    manifest_values = {str(row.get("Item", "")): row.get("Value", "")
                       for row in sheets.get("Run Manifest", {}).get("rows", [])}
    result = {
        "source_kind": "assessment_snapshot", "recommendations": recommendations,
        "collection_coverage": coverage, "control_results": controls,
        "sheets": sheets, "withheld_sheets": withheld,
        "warnings": [] if supplied_sheets else [
            "This assessment snapshot contains saved conclusions and coverage metadata, not the underlying tenant evidence tables."
        ],
    }
    for output, aliases in {
        "tenant_name": ("tenant_name", "tenant", "Tenant"), "tenant_id": ("tenant_id", "Tenant ID"),
        "generated_at": ("generated_at", "Generation Time (UTC)"),
        "methodology_version": ("methodology_version", "Methodology Version"),
        "assessment_version": ("assessment_version", "Assessment Version"),
    }.items():
        value = ""
        for source in (payload, metadata, source_metadata, manifest, manifest_values):
            value = next((source[key] for key in aliases if source.get(key) is not None), "")
            if value != "":
                break
        result[output] = _text(value, output.replace("_", " "))
    return result


def load_prior_report(path, include_user_details=False):
    """Return saved assessment rows and provenance, without refreshing conclusions.

    No dates, tenant identifiers, or collection coverage are inferred from names,
    file modification times, or the workstation environment. User-level Copilot
    sheets require a fresh opt-in for each import.
    """
    source = Path(path)
    if not source.is_file():
        raise ValueError("Prior assessment file does not exist or is not a file.")
    if source.suffix.lower() == ".xlsx":
        result = _read_workbook(source, include_user_details)
    elif source.suffix.lower() == ".json":
        result = _read_snapshot(source, include_user_details)
    else:
        raise ValueError("Prior assessment must be an XLSX workbook or assessment snapshot JSON.")
    result.update(available=True, source_file=source.name)
    if not include_user_details:
        original_coverage = result["collection_coverage"]
        result["collection_coverage"] = _without_user_details(original_coverage)
        for title, sheet in result["sheets"].items():
            if _key(title) == "collectioncoverage":
                sheet["rows"] = _without_user_details(sheet["rows"])
        if original_coverage != result["collection_coverage"]:
            result["warnings"].append("Nested user-level Copilot records were omitted from saved collection coverage.")
    if not result["generated_at"]:
        result["warnings"].append("The prior assessment has no recorded generation time; its evidence age is unknown.")
    if not _GUID_RE.fullmatch(result["tenant_id"]):
        result["warnings"].append("The prior assessment has no explicit tenant GUID; tenant identity cannot be verified automatically against supplied exports.")
    if not result["collection_coverage"]:
        result["warnings"].append("The prior assessment does not contain collection coverage metadata.")
    missing = set()
    present = {_key(title) for title in result["sheets"]}
    for row in result["recommendations"]:
        for title in str(row.get("EvidenceSheet", "") or "").split(";"):
            title = title.strip()
            if title and _key(title) not in present and not _restricted(title):
                missing.add(title)
    if missing:
        result["warnings"].append(
            "Saved recommendations reference evidence tabs absent from this file: " + "; ".join(sorted(missing)) + "."
        )
    if result["withheld_sheets"]:
        result["warnings"].append("User-level Copilot evidence tabs were omitted; include user details explicitly to import them.")
    return result
