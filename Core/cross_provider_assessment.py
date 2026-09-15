"""Control-based, cross-provider assessment and reproducibility helpers."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path


METHODOLOGY_VERSION = "2.1.0"
ASSESSMENT_VERSION = "2.1.0"
GUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)

CONTROL_CATALOG = (
    ("IDENTITY-001", "Identity and privileged access", {"Identity & access", "Privileged access"}),
    ("APPS-001", "Application consent and connected services", {"Apps, connectors & agents", "Application access"}),
    ("DATA-001", "SharePoint/OneDrive exposure and grounding permissions", {"Data exposure & oversharing", "Content access"}),
    ("GOV-001", "Classification, DLP, retention, audit, and investigation", {"Data governance", "Compliance & governance"}),
    ("EGRESS-001", "Endpoint, browser, and network data egress", {"Endpoint & device", "Network & browser"}),
    ("THREAT-001", "Threat and incident posture", {"Threat & incident posture", "Security operations"}),
    ("ACTIONS-001", "Connectors, agents, tools, and consequential actions", {"Apps, connectors & agents"}),
    ("ADOPTION-001", "Adoption evidence and outcome measurement", {"Adoption & value", "Platform capability"}),
)

STATUS_RANK = {"pass": 0, "not_applicable": 0, "not_assessed": 1, "fail": 2}


def _now_iso():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _count_phrase(count, singular, plural=None):
    return f"{count} {singular if count == 1 else (plural or singular + 's')}"


def _safe_basename(value):
    return Path(str(value)).name if value else ""


def _normalized_key(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _parse_date(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        try:
            return datetime.strptime(text[:10], "%Y-%m-%d").date()
        except ValueError:
            return None


def load_assessment_profile(path):
    if not path:
        return {"available": False, "status": "not_requested", "reason": "No assessment profile was supplied.", "products": [], "use_cases": []}
    source = Path(path)
    try:
        payload = json.loads(source.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        return {"available": False, "status": "unavailable", "reason": f"Assessment profile could not be read: {exc}", "filename": _safe_basename(path), "products": [], "use_cases": []}
    if not isinstance(payload, dict):
        return {"available": False, "status": "unavailable", "reason": "Assessment profile must be a JSON object.", "filename": _safe_basename(path), "products": [], "use_cases": []}
    version = str(payload.get("version") or payload.get("schemaVersion") or "").strip()
    products = payload.get("ai_products", payload.get("products", [])) or []
    use_cases = payload.get("use_cases", payload.get("useCases", [])) or []
    errors = []
    if not version:
        errors.append("version is required")
    if not isinstance(products, list):
        errors.append("products must be an array")
        products = []
    if not isinstance(use_cases, list):
        errors.append("use_cases must be an array")
        use_cases = []
    if "readiness_review" in payload:
        from .control_reviews import validate_readiness_review
        errors.extend(validate_readiness_review(payload))
    from .control_reviews import LoadedAssessmentProfile
    return LoadedAssessmentProfile({
        "available": not errors,
        "status": "available" if not errors else "partial",
        "reason": "; ".join(errors),
        "filename": _safe_basename(path),
        "version": version,
        "products": products,
        "use_cases": use_cases,
        "raw": payload,
    })


def _read_tabular(path):
    source = Path(path)
    if source.suffix.lower() == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if source.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values, [])]
        return [dict(zip(headers, row)) for row in values]
    raise ValueError("Provider evidence must be CSV or XLSX.")


def load_provider_evidence(path, max_age_days=None, evaluation_date=None):
    if not path:
        return {"available": False, "status": "not_requested", "reason": "No provider evidence register was supplied.", "rows": [], "max_age_days": int(max_age_days or os.getenv("PROVIDER_EVIDENCE_MAX_AGE_DAYS", "90"))}
    max_age = int(max_age_days or os.getenv("PROVIDER_EVIDENCE_MAX_AGE_DAYS", "90"))
    try:
        raw_rows = _read_tabular(path)
    except Exception as exc:
        return {"available": False, "status": "unavailable", "reason": f"Provider evidence could not be read: {exc}", "filename": _safe_basename(path), "rows": [], "max_age_days": max_age}
    today = _parse_date(evaluation_date) if evaluation_date else datetime.now(timezone.utc).date()
    rows = []
    for raw in raw_rows:
        normalized = {_normalized_key(key): value for key, value in (raw or {}).items()}
        review_date = normalized.get("reviewdate", "")
        expiration = normalized.get("expirationdate", "")
        reviewed = _parse_date(review_date)
        expires = _parse_date(expiration)
        stale = bool((reviewed and (today - reviewed).days > max_age) or (expires and expires < today))
        provider = str(normalized.get("provider", "") or "").strip()
        product = str(normalized.get("product", "") or "").strip()
        tier = str(normalized.get("tier", normalized.get("subscriptiontier", "")) or "").strip()
        approval = str(normalized.get("approvalstate", "") or "").strip()
        identity = str(normalized.get("enterpriseidentityandlifecyclecontrols", normalized.get("enterpriseidentity", "")) or "").strip()
        audit = str(normalized.get("auditinvestigationandincidentresponsecapabilities", normalized.get("auditandincidentresponse", "")) or "").strip()
        evidence_link = str(normalized.get("evidencelink", "") or "").strip()
        training = str(normalized.get("traininguseandretentionterms", "") or "").strip()
        privacy = str(normalized.get("residencydeletionsubprocessorsandprivacyreview", "") or "").strip()
        connectors = str(normalized.get("connectorsoauthapplicationsmcpserversagentsandactionpermissions", "") or "").strip()
        egress = str(normalized.get("dataegresscontrolsandapprovedclassifications", "") or "").strip()
        owner = str(normalized.get("owner", "") or "").strip()
        reviewer = str(normalized.get("reviewer", "") or "").strip()
        complete = all((provider, product, tier, approval, owner, identity, training, privacy, audit, connectors, egress, evidence_link, reviewer, review_date))
        rows.append({
            "Provider": provider,
            "Product": product,
            "Tier": tier,
            "Owner": owner,
            "Approval State": approval,
            "Enterprise Identity": identity,
            "Training Use and Retention": training,
            "Privacy and Residency": privacy,
            "Audit and Incident Response": audit,
            "Connectors and Actions": connectors,
            "Data Egress and Classifications": egress,
            "Evidence Link": evidence_link,
            "Reviewer": reviewer,
            "Review Date": review_date,
            "Expiration Date": expiration,
            "Freshness": "Stale" if stale else "Fresh",
            "Complete": "Yes" if complete else "No",
        })
    status = "available" if rows else "partial"
    return {"available": bool(rows), "status": status, "reason": "" if rows else "The provider register contained no records.", "filename": _safe_basename(path), "rows": rows, "max_age_days": max_age, "records_collected": len(rows)}


def _control_for(record):
    explicit = str(record.get("ControlId", "") or "")
    if explicit:
        return explicit
    area = str(record.get("ImpactArea", "") or "")
    text = " ".join(str(record.get(k, "") or "") for k in ("Service", "Feature", "Observation", "Recommendation")).lower()
    if "conditional access" in text or "mfa" in text or "privileged" in text or "admin role" in text:
        return "IDENTITY-001"
    if "sharepoint" in text or "onedrive" in text or "overshar" in text or "ground" in text:
        return "DATA-001"
    if "retention" in text or "purview" in text or "dlp" in text or "audit" in text or "label" in text:
        return "GOV-001"
    if "device" in text or "endpoint" in text or "browser" in text or "network access" in text:
        return "EGRESS-001"
    if "defender" in text or "incident" in text or "threat" in text:
        return "THREAT-001"
    if "agent" in text or "connector" in text or "power platform" in text or "copilot studio" in text:
        return "ACTIONS-001"
    if "app" in text or "oauth" in text or "consent" in text or "service principal" in text:
        return "APPS-001"
    for control_id, _, areas in CONTROL_CATALOG:
        if area in areas:
            return control_id
    return "ADOPTION-001"


def _fingerprint(control_id, record):
    stable_source = "|".join(str(record.get(key, "") or "").strip().lower() for key in (
        "FindingKey", "Service", "Feature", "EvidenceKey", "AffectedObjectIds"
    ))
    return hashlib.sha256(f"{METHODOLOGY_VERSION}|{control_id}|{stable_source}".encode("utf-8")).hexdigest()[:20]


def evaluate_controls(recommendations):
    catalog = {item[0]: item for item in CONTROL_CATALOG}
    grouped = {item[0]: [] for item in CONTROL_CATALOG}
    enriched = []
    for source in recommendations or []:
        row = dict(source)
        control_id = _control_for(row)
        row["ControlId"] = control_id
        row["MethodologyVersion"] = METHODOLOGY_VERSION
        row["FindingFingerprint"] = _fingerprint(control_id, row)
        grouped.setdefault(control_id, []).append(row)
        enriched.append(row)

    results = []
    for control_id, title, _ in CONTROL_CATALOG:
        rows = grouped.get(control_id, [])
        qualifying = [row for row in rows if str(row.get("EvidenceBasis", "")).lower() != "license signal"]

        def evidence_supported(row):
            return (
                str(row.get("EvidenceAvailable", "")).lower() == "yes"
                and str(row.get("Confidence", "")).lower() in {"medium", "high"}
            )

        # Legacy recommendation modules can still produce action/assurance wording from a
        # licensed capability alone or from an unverified observation. Those rows remain in
        # the compatibility register, but they cannot determine a control outcome.
        actions = [row for row in qualifying if row.get("Disposition") == "Action" and evidence_supported(row)]
        unsupported_actions = [row for row in qualifying if row.get("Disposition") == "Action" and not evidence_supported(row)]
        gaps = [row for row in qualifying if row.get("Disposition") == "Coverage"] + unsupported_actions
        strengths = [row for row in qualifying if row.get("Disposition") == "Assurance" and evidence_supported(row)]
        if actions:
            status, effect = "fail", "Limits M365 foundation readiness"
        elif gaps:
            status, effect = "not_assessed", "Limits confidence; no pass/fail conclusion"
        elif strengths:
            status, effect = "pass", "Supports M365 foundation readiness"
        else:
            status, effect = "not_assessed", "No qualifying tenant evidence collected"
        decision_rows = actions + gaps + strengths
        evidence_ids = [row.get("RecommendationId", "") for row in decision_rows if row.get("RecommendationId")]
        fingerprints = [row.get("FindingFingerprint", "") for row in decision_rows if row.get("FindingFingerprint")]
        results.append({
            "Control ID": control_id,
            "Control": title,
            "Methodology Version": METHODOLOGY_VERSION,
            "Applicability": "Applicable",
            "Applicability Reason": "General Microsoft 365 AI foundation control.",
            "Result": status,
            "Required Evidence": title,
            "Collected Evidence": "; ".join(evidence_ids) or "None",
            "Decision Effect": effect,
            "Severity Rationale": (
                f"{_count_phrase(len(actions), 'action')}, "
                f"{_count_phrase(len(gaps), 'evidence gap')}, and "
                f"{_count_phrase(len(strengths), 'confirmed strength')}."
            ),
            "Confidence": "High" if actions and all(row.get("EvidenceAvailable") == "Yes" for row in actions) else ("Medium" if qualifying else "Unknown"),
            "Freshness": "See Evidence Index",
            "Affected Object Identifiers": "; ".join(str(row.get("AffectedObjectIds", "") or "") for row in decision_rows if row.get("AffectedObjectIds")),
            "Remediation": " ".join(str(row.get("Recommendation", "") or "") for row in actions[:3]),
            "Authoritative References": "; ".join(str(row.get("LinkUrl", "") or "") for row in decision_rows if row.get("LinkUrl")),
            "Finding Fingerprints": "; ".join(fingerprints),
        })
    return enriched, results


def assess_provider_and_use_cases(profile, provider_evidence, foundation_decision):
    products = []
    provider_rows = provider_evidence.get("rows", []) if provider_evidence else []
    provider_index = {}
    for row in provider_rows:
        key = (_normalized_key(row.get("Provider")), _normalized_key(row.get("Product")), _normalized_key(row.get("Tier")))
        provider_index[key] = row
    if profile.get("available"):
        for item in profile.get("products", []):
            if not isinstance(item, dict):
                continue
            provider = item.get("provider", "")
            product = item.get("product", item.get("name", ""))
            tier = item.get("tier", item.get("subscription_tier", ""))
            evidence = provider_index.get((_normalized_key(provider), _normalized_key(product), _normalized_key(tier)))
            approved = bool(evidence and str(evidence.get("Approval State", "")).lower() == "approved" and evidence.get("Freshness") == "Fresh" and evidence.get("Complete") == "Yes")
            products.append({"Provider": provider, "Product": product, "Tier": tier, "Approval Status": "Approved" if approved else "Not approved", "Reason": "Current, complete provider review supplied." if approved else "A current, complete, approved provider-tier evidence record was not supplied."})

    use_cases = []
    product_lookup = {(_normalized_key(row["Provider"]), _normalized_key(row["Product"]), _normalized_key(row["Tier"])): row for row in products}
    if profile.get("available"):
        for item in profile.get("use_cases", []):
            if not isinstance(item, dict):
                continue
            provider = item.get("provider", "")
            product = item.get("product", "")
            tier = item.get("tier", item.get("subscription_tier", ""))
            provider_result = product_lookup.get((_normalized_key(provider), _normalized_key(product), _normalized_key(tier)))
            required_fields = (
                "name", "business_owner", "intended_users", "data_classifications",
                "approved_data", "prohibited_data", "required_outcome", "risk_measurements",
                "expand_stop_decision",
            )
            complete = all(item.get(field) not in (None, "", []) for field in required_fields)
            if item.get("can_write") or item.get("can_invoke_tools"):
                complete = complete and item.get("human_approval_required") is not None
            applicable_controls = ["IDENTITY-001", "EGRESS-001", "THREAT-001", "ADOPTION-001"]
            if item.get("can_read_m365_data"):
                applicable_controls.extend(["DATA-001", "GOV-001"])
            if item.get("uses_connectors"):
                applicable_controls.extend(["APPS-001", "ACTIONS-001"])
            if item.get("can_invoke_tools") or item.get("can_write"):
                applicable_controls.append("ACTIONS-001")
            applicable_controls = list(dict.fromkeys(applicable_controls))
            foundation_ok = foundation_decision in {"Ready for a controlled pilot", "Controlled pilot with conditions"}
            ready = bool(complete and foundation_ok and provider_result and provider_result["Approval Status"] == "Approved")
            use_cases.append({"Use Case": item.get("name", "Unnamed use case"), "Business Owner": item.get("business_owner", ""), "Provider": provider, "Product / Tier": f"{product} / {tier}".strip(" /"), "Applicable Controls": "; ".join(applicable_controls), "Readiness": "Ready for controlled pilot" if ready else "Not ready", "Reason": "Foundation, provider review, and required profile fields are satisfied." if ready else "Foundation, provider approval, or required use-case evidence is incomplete."})

    return {
        "foundation": {"Conclusion": foundation_decision, "Scope": "Microsoft 365 foundation"},
        "providers": products,
        "use_cases": use_cases,
        "provider_conclusion": "Not assessed" if not profile.get("available") else ("Approved" if products and all(row["Approval Status"] == "Approved" for row in products) else "Not approved"),
        "use_case_conclusion": "Not assessed" if not profile.get("available") else ("Ready" if use_cases and all(row["Readiness"].startswith("Ready") for row in use_cases) else "Not ready"),
    }


def build_run_manifest(tenant_name, inputs, collectors, source_statuses=None):
    rows = [
        {"Item": "Tenant", "Value": tenant_name or "Unknown"},
        {"Item": "Assessment Version", "Value": ASSESSMENT_VERSION},
        {"Item": "Methodology Version", "Value": METHODOLOGY_VERSION},
        {"Item": "Generation Time (UTC)", "Value": _now_iso()},
        {"Item": "Enabled Collectors", "Value": ", ".join(collectors or []) or "Default"},
    ]
    for label, value in (inputs or {}).items():
        if isinstance(value, (list, tuple)):
            sanitized = ", ".join(_safe_basename(item) for item in value if item)
        else:
            sanitized = _safe_basename(value)
        rows.append({"Item": f"Input: {label}", "Value": sanitized or "Not supplied"})
    for source, status in (source_statuses or {}).items():
        state = status.get("availability_status", status.get("status", "unknown")) if isinstance(status, dict) else str(status)
        count = status.get("records_collected", status.get("record_count", "")) if isinstance(status, dict) else ""
        pages = status.get("pages_collected", "") if isinstance(status, dict) else ""
        truncated = status.get("truncated", "") if isinstance(status, dict) else ""
        reason = status.get("reason", "") if isinstance(status, dict) else ""
        rows.append({"Item": f"Source: {source}", "Value": f"{state}; records={count}; pages={pages}; truncated={truncated}; {reason}".strip("; ")})
    return {"rows": rows, "generated_at": rows[3]["Value"], "assessment_version": ASSESSMENT_VERSION, "methodology_version": METHODOLOGY_VERSION}


def run_integrity_checks(recommendations, evidence_bundle):
    issues = []
    sheets = evidence_bundle.get("sheets", {})
    for key, sheet in sheets.items():
        for index, row in enumerate(sheet.get("rows", []), start=2):
            for field, value in row.items():
                if "name" in field.lower() and GUID_RE.match(str(value or "").strip()):
                    issues.append(f"{sheet.get('title', key)} row {index}: GUID appears in name field '{field}'.")
    for row in recommendations:
        if row.get("Confidence") == "High" and row.get("Disposition") in {"Action", "Assurance"} and row.get("EvidenceAvailable") != "Yes":
            issues.append(f"{row.get('RecommendationId', 'Unknown')}: high-confidence conclusion lacks linked evidence.")
        if row.get("Disposition") == "Assurance" and str(row.get("Status", "")).lower() in {"not assessed", "unavailable", "permission required"}:
            issues.append(f"{row.get('RecommendationId', 'Unknown')}: unavailable evidence is labeled as a strength.")
    for source, status in (evidence_bundle.get("source_statuses", {}) or {}).items():
        if isinstance(status, dict) and status.get("truncated") and status.get("required", True):
            issues.append(f"Required source '{source}' was truncated.")
    return {"valid": not issues, "issues": issues, "banner": "" if not issues else "Validation incomplete—do not use for deployment approval"}


def _load_baseline(path):
    source = Path(path)
    if source.suffix.lower() == ".json":
        return json.loads(source.read_text(encoding="utf-8-sig"))
    if source.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(source, read_only=True, data_only=True)
        payload = {"methodology_version": "", "recommendations": [], "control_results": []}
        if "Run Manifest" in workbook.sheetnames:
            rows = list(workbook["Run Manifest"].iter_rows(values_only=True))
            manifest = {str(row[0]): row[1] for row in rows[1:] if row and row[0]}
            payload["methodology_version"] = str(manifest.get("Methodology Version", ""))
        if "Recommendations" in workbook.sheetnames:
            rows = workbook["Recommendations"].iter_rows(values_only=True)
            headers = [str(value or "") for value in next(rows, [])]
            payload["recommendations"] = [dict(zip(headers, row)) for row in rows]
        return payload
    raise ValueError("Baseline must be an XLSX workbook or snapshot JSON.")


def compare_baseline(current_recommendations, current_controls, baseline_path):
    if not baseline_path:
        return {"available": False, "status": "not_requested", "rows": [], "reason": "No baseline was supplied."}
    try:
        baseline = _load_baseline(baseline_path)
    except Exception as exc:
        return {"available": False, "status": "unavailable", "rows": [], "reason": f"Baseline could not be read: {exc}", "filename": _safe_basename(baseline_path)}
    old_version = str(baseline.get("methodology_version", baseline.get("run_manifest", {}).get("methodology_version", "")) or "")
    if old_version and old_version != METHODOLOGY_VERSION:
        return {"available": True, "status": "available", "comparable": False, "rows": [{"Classification": "Unable to compare", "Reason": f"Methodology changed from {old_version} to {METHODOLOGY_VERSION}."}], "filename": _safe_basename(baseline_path)}
    old_rows = baseline.get("recommendations", []) or []
    old = {str(row.get("FindingFingerprint", row.get("Finding Fingerprint", "")) or ""): row for row in old_rows if row.get("FindingFingerprint", row.get("Finding Fingerprint", ""))}
    current = {str(row.get("FindingFingerprint", "") or ""): row for row in current_recommendations if row.get("FindingFingerprint")}
    rows = []
    for fingerprint, row in current.items():
        prior = old.get(fingerprint)
        if not prior:
            classification = "New"
        else:
            old_status = str(prior.get("ControlStatus", prior.get("Disposition", ""))).lower()
            new_status = str(row.get("ControlStatus", row.get("Disposition", ""))).lower()
            old_rank = 2 if old_status in {"fail", "action"} else 1 if old_status in {"not_assessed", "coverage"} else 0
            new_rank = 2 if new_status in {"fail", "action"} else 1 if new_status in {"not_assessed", "coverage"} else 0
            classification = "Improved" if new_rank < old_rank else "Regressed" if new_rank > old_rank else "Persistent"
        rows.append({"Finding Fingerprint": fingerprint, "Recommendation ID": row.get("RecommendationId", ""), "Control ID": row.get("ControlId", ""), "Classification": classification, "Feature": row.get("Feature", "")})
    for fingerprint, row in old.items():
        if fingerprint not in current:
            rows.append({"Finding Fingerprint": fingerprint, "Recommendation ID": row.get("RecommendationId", ""), "Control ID": row.get("ControlId", ""), "Classification": "Resolved", "Feature": row.get("Feature", "")})
    return {"available": True, "status": "available", "comparable": True, "rows": rows, "filename": _safe_basename(baseline_path)}


def write_snapshot(path, payload):
    if not path:
        return None
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
    return str(target.resolve())
