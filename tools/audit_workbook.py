"""Print privacy-safe QA counts from an assessment evidence workbook."""

import argparse
from collections import Counter
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook


def _records(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "") for value in next(rows, [])]
    for row in rows:
        yield dict(zip(headers, row))


def _compatibility_issues(path, workbook):
    issues = []
    try:
        with zipfile.ZipFile(path) as archive:
            corrupt_member = archive.testzip()
            if corrupt_member:
                issues.append(f"Corrupt ZIP member: {corrupt_member}")
            for name in archive.namelist():
                if name.endswith((".xml", ".rels")) or name == "[Content_Types].xml":
                    try:
                        ET.fromstring(archive.read(name))
                    except Exception as exc:
                        issues.append(f"Invalid Open XML part {name}: {type(exc).__name__}")
    except (OSError, zipfile.BadZipFile) as exc:
        return [f"Invalid XLSX package: {type(exc).__name__}"]

    table_names = []
    for sheet in workbook.worksheets:
        if sheet.tables and sheet.auto_filter.ref:
            issues.append(
                f"{sheet.title}: worksheet AutoFilter overlaps an Excel table ({sheet.auto_filter.ref})"
            )
        table_names.extend(table.displayName for table in sheet.tables.values())
    duplicates = [name for name, count in Counter(table_names).items() if count > 1]
    if duplicates:
        issues.append("Duplicate Excel table names: " + ", ".join(sorted(duplicates)))
    return issues


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=False, data_only=True)
    compatibility_issues = _compatibility_issues(args.workbook, workbook)
    print("Workbook compatibility:")
    if compatibility_issues:
        for issue in compatibility_issues:
            print(f"  FAIL: {issue}")
    else:
        print("  PASS: ZIP, Open XML, table names, and filter definitions are valid")

    print("Sheets:")
    for sheet in workbook.worksheets:
        count = max(sheet.max_row - 1, 0)
        print(f"  {sheet.title}: {count} {'data row' if count == 1 else 'data rows'}")

    for sheet_name, field in (
        ("Recommendations", "Disposition"),
        ("App Access Detail", "Flagged Because"),
        ("Admin Role Detail", "Assignment Type"),
        ("Collection Coverage", "State"),
    ):
        if sheet_name not in workbook.sheetnames:
            continue
        counts = Counter(str(record.get(field) or "Blank") for record in _records(workbook[sheet_name]))
        print(f"{sheet_name} by {field}:")
        for value, count in counts.most_common():
            print(f"  {value}: {count}")

    if "AI Adoption Usage" in workbook.sheetnames:
        print("AI Adoption Usage evidence:")
        for record in _records(workbook["AI Adoption Usage"]):
            print(
                "  {source} | {period} | {metric}: {value} | {availability} | {freshness} | {detail}".format(
                    source=record.get("Evidence Source") or "Unknown source",
                    period=record.get("Period") or "No period",
                    metric=record.get("Metric") or "Unknown metric",
                    value=record.get("Value") if record.get("Value") not in (None, "") else "Not provided",
                    availability=record.get("Availability") or "Unknown availability",
                    freshness=record.get("Freshness") or "Unknown freshness",
                    detail=record.get("Detail") or "",
                )
            )

    if compatibility_issues:
        sys.exit(1)


if __name__ == "__main__":
    main()
