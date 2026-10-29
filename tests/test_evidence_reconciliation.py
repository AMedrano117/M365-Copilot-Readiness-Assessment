"""Fictional regressions for component evidence, units, lineage and projections."""
from copy import deepcopy
from pathlib import Path
import json
import re
from tempfile import TemporaryDirectory
import unittest

from provenance_fixture import fictional_publication
from test_privileged_identity import build, sources, grant, dataset as role_dataset
from Core.assessment_result import build_assessment_result
from Core.evidence_reconciliation import customer_text, require_supported_confidence
from Core.coverage_summary import summarize_collection, coverage_text

DAY = '2026-10-27'
TENANT = '11111111-1111-4111-8111-111111111111'
PARENT = '22222222-2222-4222-8222-222222222222'


def dataset(rows, **metadata):
    return {'records':rows, 'source':{'available':True, 'availability_status':'available',
        'complete':True, 'truncated':False, 'collected_at':DAY, 'scope':'Fictional tenant scope',
        'tenant_id':TENANT, 'source_file':'fictional-evidence.json', **metadata}}


def assessment(data=None, **extra):
    bundle = dict(assessment_sources=data or {}, expected_tenant_id=TENANT,
        collection_context={'collected_at':DAY,'tenant_id':TENANT}, **extra)
    return build_assessment_result([], bundle, evaluation_date=DAY, expected_tenant_id=TENANT), bundle


def gap(result, control):
    return next(r for r in result['recommendations'] if r.get('ControlId') == control and r.get('SourceType') == 'assessment_requirement')


class ComponentEvidenceTests(unittest.TestCase):
    def licensing(self):
        usage = {'available':True, 'periods':{'D7':{'active_users':63},'D28':{'active_users':80},'D180':{'active_users':85}}}
        capacity = {'available':True, 'enabled_users':87,'copilot_license_coverage':64.93}
        return assessment(ai_usage={'copilot_usage':usage,'license_coverage':capacity})

    def test_limited_usage_and_capacity_are_retained_not_uncollected(self):
        result, _ = self.licensing(); row = gap(result,'LICENSE.ASSIGNMENT')
        self.assertIn('usage and limited license-capacity signals',row['Observation'])
        self.assertEqual(row['ObservedEvidence']['license_signal']['enabled_users'],87)
        self.assertEqual(row['ObservedEvidence']['license_signal']['copilot_license_coverage'],64.93)
        self.assertEqual([p['active_users'] for p in row['ObservedEvidence']['usage_signal']['periods'].values()],[63,80,85])

    def test_usage_does_not_prove_entitlement_or_apps(self):
        result, _ = self.licensing(); row = gap(result,'LICENSE.ASSIGNMENT')
        self.assertEqual(row['ObservedEvidence']['pilot_entitlement'],'not_established')
        self.assertEqual(row['ObservedEvidence']['apps_prerequisites'],'not_established')
        self.assertEqual(next(c for c in result['controls'] if c['control_id']=='LICENSE.ASSIGNMENT')['status'],'Not established')

    def test_unavailable_defaults_do_not_become_observed_signals(self):
        result,_=assessment(ai_usage={'copilot_usage':{'available':False,'periods':{},'enabled_users':0},
            'license_coverage':{'available':False,'enabled_users':0}})
        self.assertNotIn('ObservedEvidence',gap(result,'LICENSE.ASSIGNMENT'))

    def test_runtime_license_summary_retains_observed_capacity(self):
        result,_=assessment(ai_usage={'license_summary':{'enabled_users':87,'copilot_license_coverage':64.93}})
        self.assertEqual(gap(result,'LICENSE.ASSIGNMENT')['ObservedEvidence']['license_signal']['enabled_users'],87)

    def test_existing_module_gap_is_reconciled_with_current_components(self):
        _,bundle=assessment({'sensitivity_labels':[dataset([{'id':'label-one'}])]})
        legacy={'RecommendationId':'FICTIONAL-LABEL','ControlId':'DATA.LABELS','Disposition':'Coverage',
                'Service':'Purview','Feature':'Label evidence','Observation':'Label definitions were not collected.',
                'Recommendation':'Run --reports-dir PATH','EvidenceKey':'purview_policy_detail'}
        result=build_assessment_result([legacy],bundle,evaluation_date=DAY,expected_tenant_id=TENANT)
        row=next(r for r in result['recommendations'] if r['RecommendationId']=='FICTIONAL-LABEL')
        self.assertIn('1 sensitivity-label definition',row['Observation'])
        self.assertNotIn('not collected',row['Observation'])
        self.assertEqual(row['OriginalObservation'],legacy['Observation'])

    def test_one_label_definition_is_observed(self):
        result,_=assessment({'sensitivity_labels':[dataset([{'id':'label-one'}])], 'label_policies':[dataset([{'id':'policy-one'}])]})
        row=gap(result,'DATA.LABELS')
        self.assertIn('1 sensitivity-label definition',row['Observation'])
        self.assertNotIn('definitions were not collected',row['Observation'])

    def test_publication_presence_does_not_establish_scope_or_use(self):
        result,_=assessment({'sensitivity_labels':[dataset([{'id':'label-one'}])], 'label_policies':[dataset([{'id':'policy-one'}])]})
        row=gap(result,'DATA.PUBLISHING')
        self.assertEqual(row['ObservedEvidence']['publication_scope'],'not_established')
        self.assertEqual(row['ObservedEvidence']['content_coverage'],'not_established')
        self.assertEqual(row['ObservedEvidence']['operational_use'],'not_established')

    def audit(self):
        return assessment({'audit_config':[dataset([{'unified_audit_enabled':True,'admin_audit_enabled':False}])],
            'copilot_audit':[dataset([{'id':f'fictional-event-{i}'} for i in range(3554)],
                collection_window={'window_start':'2026-10-20','window_end':DAY})]})

    def test_audit_settings_and_events_remain_separate(self):
        result,_=self.audit(); row=gap(result,'DATA.AUDIT')
        counts={r['dataset']:r['record_count'] for r in row['ObservedEvidence']['components']}
        self.assertEqual(counts,{'audit_config':1,'copilot_audit':3554})
        self.assertIn('unified_audit_enabled=True',row['Observation'])
        self.assertIn('admin_audit_enabled=False',row['Observation'])
        self.assertIn('2026-10-20 through 2026-10-27',row['Observation'])

    def test_audit_events_do_not_pass_tenant_retention_or_review(self):
        result,_=self.audit(); row=gap(result,'DATA.AUDIT')
        self.assertEqual(row['ObservedEvidence']['retention_coverage'],'not_established')
        self.assertEqual(row['ObservedEvidence']['review_ownership'],'not_established')
        self.assertEqual(next(c for c in result['controls'] if c['control_id']=='DATA.AUDIT')['status'],'Not established')

    def test_legitimate_gaps_remain_unresolved(self):
        result,_=self.licensing()
        for control in ('DATA.EXPOSURE','LICENSE.APPS','CONTENT.PERMISSIONS'):
            self.assertEqual(gap(result,control)['Disposition'],'Coverage')
        self.assertEqual(result['decision'],'Tenant readiness unconfirmed')


class MetricReconciliationTests(unittest.TestCase):
    def test_assignment_and_schedule_units_are_not_users(self):
        model=build(sources([grant('a'),grant('b')]))
        rows={r['Metric']:r for r in model['metric_reconciliation']}
        self.assertEqual(rows['Assignment schedules']['Value'],2)
        self.assertEqual(rows['Assignment schedules']['Unit'],'schedule records')
        self.assertEqual(rows['Unique users with active privilege']['Value'],1)

    def test_human_group_and_workload_units_remain_separate(self):
        grants=[]
        for key,kind in [('person','user'),('team','group'),('app','servicePrincipal')]:
            row=grant(key,principal=key); row['principal']['@odata.type']='#microsoft.graph.'+kind
            grants.append(row)
        model=build(sources(grants,users=[{'id':'person','accountEnabled':True}]))
        rows={r['Metric']:r['Value'] for r in model['metric_reconciliation']}
        self.assertEqual([rows[n] for n in ('Unique human users','Unique groups','Unique service principals')],[1,1,1])
        self.assertEqual(rows['Unique effective privileged identities'],2)

    def test_unreconciled_140_user_claim_is_suppressed_not_normalized(self):
        from Core.privileged_presentation import qualify_privileged_recommendations
        model=build(); before=deepcopy(model['assignments'])
        row={'Observation':'140 users with privileged administrative roles','EvidenceKey':'admin_role_detail',
            'Disposition':'Assurance','Confidence':'High'}
        updated=qualify_privileged_recommendations([row],model)[0]
        self.assertNotIn('140 users',updated['Observation'])
        self.assertEqual(updated['PrivilegedMetricReconciliation']['status'],'unresolved_legacy_user_aggregate')
        self.assertEqual(updated['OriginalObservation'],row['Observation'])
        self.assertEqual(updated['Confidence'],'Low')
        self.assertEqual(model['assignments'],before)

    def mfa(self):
        from Core.authentication_methods import summarize_registrations
        rows=[{'id':f'fictional-user-{i}','isMfaRegistered':i<4,'isMfaCapable':i<6,
            'isPasswordlessCapable':i<6,'methodsRegistered':['fido2'] if i<6 else
            ['mobilePhone'] if i<9 else []} for i in range(586)]
        # One explicit MFA-registered phone-only record, preserving four true flags.
        rows[3]['methodsRegistered']=['mobilePhone']; rows[6]['methodsRegistered']=['fido2','mobilePhone']
        rows[8]['methodsRegistered']=[]
        return summarize_registrations(rows,{'availability_status':'available','scope':'Fictional returned registration report'})

    def test_mfa_fields_populations_and_denominators_remain_separate(self):
        report=self.mfa(); metrics=report['metrics']
        self.assertEqual(metrics['mfa_registered'],4)
        self.assertEqual(metrics['mfa_registered_known'],586)
        self.assertEqual(metrics['phishing_resistant_registered'],6)
        self.assertEqual(metrics['passwordless_capable'],6)
        self.assertEqual(metrics['phone_registered'],3)
        self.assertEqual(metrics['phone_only_mfa_registered'],1)
        rows={r['Metric ID']:r for r in report['metric_reconciliation']['metrics']}
        self.assertEqual(rows['mfa_registered']['Source field'],'isMfaRegistered')
        self.assertEqual(rows['phishing_resistant_registered']['Source field'],'methodsRegistered')
        self.assertEqual(rows['mfa_registered']['Denominator'],586)
        self.assertIn('Population',rows['passwordless_capable'])

    def test_conflicting_mfa_source_fields_are_qualified_not_normalized(self):
        report=self.mfa(); reconciliation=report['metric_reconciliation']
        self.assertFalse(reconciliation['directly_comparable'])
        self.assertGreater(reconciliation['conflicts']['users_with_cross_field_conflicts'],0)
        self.assertEqual(report['registration_population']['counts']['ExplicitlyNotRegistered'],582)
        self.assertIn('no values were adjusted',reconciliation['qualification'])

    def test_mfa_projection_uses_supplied_saved_result_over_bundle_report(self):
        from Core.customer_report import _authentication_methods
        report=self.mfa()
        saved={'authentication_assessment':{'registration':{'report':report}}}
        stale=deepcopy(report); stale['metrics']['phishing_resistant_registered']=999
        bundle={'authentication_methods':stale,'assessment_result':{'authentication_assessment':{'registration':{'report':stale}}}}
        html=_authentication_methods(bundle,saved)
        self.assertIn('<strong>6</strong>',html); self.assertNotIn('<strong>999</strong>',html)
        self.assertEqual(saved['authentication_assessment']['registration']['report'],report)


class IntegrityAndGuidanceTests(unittest.TestCase):
    def test_guid_fallback_with_valid_identifier_is_diagnostic(self):
        from Core.cross_provider_assessment import run_integrity_checks
        for field in ('Device Name','Group Display Name','userDisplayName','computerDnsName'):
            with self.subTest(field=field):
                result=run_integrity_checks([],{'sheets':{'source':{'rows':[{'id':PARENT,field:PARENT}]}}})
                self.assertTrue(result['valid']); self.assertEqual(result['diagnostics'][0]['kind'],'unresolved_display_name')

    def test_guid_in_name_without_valid_identifier_still_fails(self):
        from Core.cross_provider_assessment import run_integrity_checks
        result=run_integrity_checks([],{'sheets':{'source':{'rows':[{'id':'malformed','App Display Name':PARENT}]}}})
        self.assertFalse(result['valid'])

    def test_high_confidence_requires_real_support_even_if_flag_says_yes(self):
        row={'RecommendationId':'FICTIONAL-UNSUPPORTED','Confidence':'High','Disposition':'Action','EvidenceAvailable':'Yes','EvidenceId':'EV-fabricated'}
        with self.assertRaisesRegex(ValueError,'required supporting evidence'):
            require_supported_confidence({'recommendations':[row]})

    def test_endpoint_setting_finding_restores_selected_support(self):
        from Core.expanded_findings import expanded_findings
        data={'compliance_setting_states':[dataset([{'id':'setting-one','ParentId':'device-one','state':'noncompliant'}])]}
        rows=expanded_findings({'assessment_sources':data},evaluation_date=DAY,tenant_id=TENANT)
        row=next(r for r in rows if 'compliance_setting_states' in r['FindingKey'])
        self.assertEqual(row['EvidenceAvailable'],'Yes')
        self.assertEqual(row['EvidenceKey'],'compliance_setting_detail')
        require_supported_confidence({'recommendations':[row]})

    def test_customer_requests_have_no_internal_commands(self):
        samples=['Run --sam-report PATH using SAM_DAG_REPORT_PATHS','Run --dspm-report PATH using DSPM_REPORT_PATHS',
                 'Use --reports-dir PATH','For collection grant Policy.Read.All','Run --copilot-readiness-export PATH']
        for value in samples:
            projected=customer_text(value)
            self.assertNotRegex(projected,r'--[a-z]|SAM_DAG_REPORT_PATHS|DSPM_REPORT_PATHS|Policy.Read.All')
            self.assertIn('Provide',projected)

    def test_operator_import_instructions_are_retained(self):
        text=(Path(__file__).resolve().parents[1]/'docs'/'PORTAL_REPORTS_AND_OFFLINE.md').read_text(encoding='utf-8')
        self.assertIn('--sam-report',text); self.assertIn('--dspm-report',text); self.assertIn('--reports-dir',text)

    def test_customer_projection_preserves_facts_before_operator_instruction(self):
        value='87 enabled users were observed. Run --reports-dir PATH to add inputs.'
        self.assertIn('87 enabled users were observed.',customer_text(value))
        self.assertNotIn('--reports-dir',customer_text(value))

    def test_customer_projection_keeps_zero_and_false_values(self):
        self.assertEqual(customer_text(0),'0'); self.assertEqual(customer_text(False),'False')


class CoverageSummaryTests(unittest.TestCase):
    def coverage(self):
        rows=[]
        for index,state in enumerate(('available','available','unavailable','partial')):
            source=dataset([] if index<3 else [{'id':'returned'}])
            source['source'].update(availability_status=state,scope='Parent object '+str(index)+PARENT,
                source_api='https://graph.microsoft.com/v1.0/servicePrincipals/'+str(index)+PARENT+'/appRoleAssignments',
                complete=state=='available',truncated=state=='partial')
            rows.append(source)
        return summarize_collection({'application_permissions':rows})

    def test_repeated_requests_aggregate_by_logical_dataset(self):
        coverage=self.coverage(); row=coverage['summary'][0]
        self.assertEqual(len(coverage['summary']),1)
        self.assertEqual(row['Parent objects queried'],4)
        self.assertEqual(row['Successful queries'],2)
        self.assertEqual(row['Zero-result queries'],2)
        self.assertEqual(row['Returned records'],1)
        self.assertNotIn(PARENT,coverage_text(row)); self.assertNotIn('https://',coverage_text(row))

    def test_failed_and_partial_queries_stay_visible(self):
        row=self.coverage()['summary'][0]
        self.assertEqual(row['Failed queries'],1); self.assertEqual(row['Partial queries'],1)
        self.assertEqual(row['Completeness'],'Incomplete'); self.assertEqual(row['Truncation'],'Yes')

    def test_all_native_requests_and_empty_metadata_are_retained(self):
        coverage=self.coverage()
        self.assertEqual(len(coverage['requests']),4)
        self.assertEqual(coverage['requests'][0]['Returned records'],0)
        self.assertIn(PARENT,json.dumps(coverage['requests']))
        self.assertIn('source_api',coverage['requests'][2]['Source metadata'])

    def test_status_only_per_object_requests_are_not_dropped(self):
        states={f'app_permissions:{PARENT}':{'availability_status':'available','record_count':0},
                'app_permissions:33333333-3333-4333-8333-333333333333':{'availability_status':'failed'}}
        coverage=summarize_collection({},states)
        self.assertEqual(len(coverage['requests']),2)
        self.assertEqual(coverage['summary'][0]['Failed queries'],1)

    def test_not_requested_is_not_a_successful_zero_result(self):
        coverage=summarize_collection({}, {'specialized_purview':{'availability_status':'not_requested'}})
        row=coverage['summary'][0]
        self.assertEqual(row['Successful queries'],0); self.assertEqual(row['Zero-result queries'],0)
        self.assertIsNone(row['Returned records']); self.assertEqual(row['Completeness'],'Not requested')

    def test_customer_coverage_does_not_repeat_per_query_capture_timestamps(self):
        from Core.report_presentation import coverage_panels
        from Core.customer_report import table
        timestamps='; '.join(f'2026-10-27T12:00:{i:02d}Z' for i in range(30))
        check={'Check ID':'FICTIONAL.REQUESTS','Check':'Application requests','Domain':'Fictional domain',
               'State':'partial','Evidence dates':timestamps,'Required sources':['application_permissions']}
        result={'domain_coverage':[check],'assessment_domains':[{'id':'fictional','title':'Fictional domain'}],
                'dataset_coverage':self.coverage()}
        html=coverage_panels(result,table,'Technical.xlsx')
        self.assertNotIn('2026-10-27T12:00:29Z',html)
        self.assertIn('failed queries=1',html)
        self.assertIn('Collection Request Detail',html)


@fictional_publication
class EvidenceProjectionTests(unittest.TestCase):
    def fixture(self):
        from test_workbook_layout import fixture
        rows,bundle=fixture()
        bundle['evaluation_date']=DAY; bundle['expected_tenant_id']=TENANT
        risk=[{'id':f'fictional-risk-{i}','userDisplayName':f'Restricted person {i}',
               'userPrincipalName':f'restricted-{i}@example.invalid','riskLevel':'low','riskState':'atRisk',
               'riskLastUpdatedDateTime':DAY} for i in range(8)]
        bundle['assessment_sources']['risky_users']=[dataset(risk)]
        rows.append(dict(rows[0],RecommendationId='ENT-004',Service='Entra',DomainId='identity',ControlId='IDENTITY.RISK',
            Feature='Risky users requiring investigation',FindingKey='entra.identity_risk.users',EvidenceKey='identity_risk_detail',
            Observation='8 selected risky-user records require investigation.',
            InvestigationEvidence={'kind':'records','records':risk,'source':dataset([])['source']}))
        # Realistic empty parent queries retained in the technical register.
        bundle['assessment_sources']['application_permissions']=[dataset([],scope='Parent object '+PARENT,
            source_api='https://graph.microsoft.com/v1.0/servicePrincipals/'+PARENT+'/appRoleAssignments')]
        return rows,bundle

    def test_identity_risk_raw_ranges_and_exact_selected_evidence_ids_resolve(self):
        from Core.raw_evidence import prepare_raw_details
        from Core.evidence_selection import build_evidence_selection
        rows,bundle=self.fixture()
        result=build_assessment_result(rows,bundle,evaluation_date=DAY,expected_tenant_id=TENANT)
        prepare_raw_details(bundle,result)
        row=next(r for r in result['recommendations'] if r['RecommendationId']=='ENT-004')
        self.assertTrue(row['RawEvidenceRanges']); self.assertNotIn('not retained',row['RawEvidenceQualification'])
        selection=build_evidence_selection(result,bundle)
        finding=next(r for r in selection['findings'] if r['finding_id']=='ENT-004')
        self.assertEqual(finding['record_count'],8)
        self.assertTrue(all(r['evidence_record_ids'] for r in finding['records']))

    def test_excel_html_snapshot_replay_keep_semantics_and_technical_detail(self):
        from Core.export_recommendations import export_to_excel
        from Core.customer_report import render_customer_report
        from Core.workbook_layout import technical_workbook_path
        from Core.assessment_serialization import write_assessment_result,read_assessment_result
        from openpyxl import load_workbook
        rows,bundle=self.fixture()
        rows[0]['Recommendation']='Provide protection evidence; internal collector instructions: --reports-dir PATH'
        result=build_assessment_result(rows,bundle,evaluation_date=DAY,expected_tenant_id=TENANT)
        bundle['assessment_result']=result
        with TemporaryDirectory() as directory:
            path=Path(export_to_excel(rows,filename='Fictional.xlsx',tenant_name='Fictional',evidence_bundle=bundle,output_dir=directory))
            result=bundle['assessment_result']
            before=deepcopy(result)
            html=render_customer_report(result,bundle,'Fictional',workbook_path=path,technical_workbook_path=technical_workbook_path(path))
            self.assertEqual(result,before)
            from tools.audit_report_signal import public_report_source
            self.assertNotIn('--reports-dir',html); self.assertNotIn('restricted-0@example.invalid',public_report_source(html))
            self.assertNotIn(PARENT,html)
            self.assertIn('Technical appendix and evidence workbook',html)
            self.assertNotIn('Detailed request register',html)
            self.assertIn('Complete run, source, lifecycle and governance registers are retained in the workbooks',html)
            customer=load_workbook(path); technical=load_workbook(technical_workbook_path(path))
            try:
                customer_texts='\n'.join(str(c.value or '') for s in customer for row in s for c in row)
                technical_texts='\n'.join(str(c.value or '') for s in technical for row in s for c in row)
                self.assertNotIn('--reports-dir',customer_texts)
                self.assertNotIn('restricted-0@example.invalid',customer_texts)
                self.assertIn('restricted-0@example.invalid',technical_texts)
                self.assertIn('--reports-dir',technical_texts)
                self.assertIn(PARENT,technical_texts)
                self.assertIn('Collection Request Detail',technical.sheetnames)
                visible_coverage='\n'.join(str(c.value or '') for row in customer['Coverage'] for c in row)
                self.assertNotIn(PARENT,visible_coverage); self.assertNotIn('/appRoleAssignments',visible_coverage)
                headers=[c.value for c in customer['Users & Identity'][1]]
                identity_rows=[dict(zip(headers,row)) for row in customer['Users & Identity'].iter_rows(min_row=2,values_only=True)]
                selected=[r for r in identity_rows if r.get('ID')=='ENT-004']
                self.assertEqual(len(selected),8)
                self.assertTrue(all(r['User']=='Restricted identity-risk record' for r in selected))
                risk=next(r for r in result['recommendations'] if r['RecommendationId']=='ENT-004')
                self.assertTrue(risk['InvestigationRange']); self.assertTrue(risk['TechnicalEvidenceRanges'])
            finally:
                customer.close(); technical.close()
            snapshot=Path(directory)/'assessment-result.json'
            write_assessment_result(snapshot,result); replay=read_assessment_result(snapshot)
            for key in ('decision','controls','counts','dataset_coverage','authentication_assessment','privileged_assessment'):
                self.assertEqual(replay[key],json.loads(json.dumps(result[key],default=str)))
            replay_html=render_customer_report(replay,bundle,'Fictional',workbook_path=path,technical_workbook_path=technical_workbook_path(path))
            self.assertEqual(html,replay_html)


if __name__ == '__main__':
    unittest.main()
