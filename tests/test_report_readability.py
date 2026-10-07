"""Regression checks for concise presentation and shared, readable source logs."""

from pathlib import Path
import json
import tempfile
import unittest

from tests.workbook_test_helpers import load_workbook_pair as load_workbook

from Core.customer_report import render_customer_report
from Core.export_recommendations import export_to_excel
from Core.investigation_contract import validate_investigation_coverage
from Core.investigation_details import prepare_investigation_details
from Core.raw_evidence import split_evidence_sheets
from Core.assessment_result import build_assessment_result
from Core.tenant_baseline import assess_identity_baseline, conditional_access_facts
from tests.test_tenant_baseline import policy, entra, LEGACY_BLOCK


DAY = '2026-09-30'
TENANT = 'readability-fixture'


def source(records):
    return [{'source':{'scope':'Fixture tenant','complete':True,'collected_at':DAY,
                       'availability_status':'available'},'records':records}]


def bundle_and_result():
    bundle = {'collection_context':{'tenant_id':TENANT,'collected_at':DAY},'sheets':{},
              'assessment_sources':{'auth_methods':source([
                  {'id':'private-user-1','userPrincipalName':'private@example.test','isMfaRegistered':True,
                   'nested':{'methods':['fido2','phone']},'password':'exclude-this'},
                  {'id':'private-user-2','userPrincipalName':'other@example.test','isMfaRegistered':False}])}}
    records = [{'RecommendationId':f'TEST-{n}','Service':'Entra','DomainId':'identity',
                'Feature':f'Authentication registration {n}','Observation':'Registration details returned.',
                'Disposition':'Assurance','EvidenceKey':'authentication_detail','EvidenceLevel':'configuration',
                'ObservationDate':DAY,'EvidenceComplete':True,'EvidenceScope':'Fixture tenant','TenantId':TENANT,
                'EvidenceAvailable':'Yes','EvidenceBasis':'Tenant evidence','Status':'Success','Priority':'Low'}
               for n in (1,2)]
    return bundle,build_assessment_result(records,bundle,evaluation_date=DAY,expected_tenant_id=TENANT)


class ReadabilityTests(unittest.TestCase):
    def test_shared_inventory_is_retained_once_with_readable_original_fields(self):
        bundle,result = bundle_and_result()
        prepare_investigation_details(bundle,result)
        rows = [r for r in result['recommendations'] if r['RecommendationId'].startswith('TEST-')]
        self.assertEqual(rows[0]['InvestigationRange'],rows[1]['InvestigationRange'])
        self.assertEqual(rows[0]['InvestigationCount'],2)
        self.assertFalse(any(key.startswith('declared_investigation.TEST') for key in bundle['sheets']))
        raw = bundle['sheets']['raw_source.auth_methods']['rows']
        self.assertEqual(len(raw),2)
        self.assertEqual(list(raw[0])[0],'id')
        self.assertEqual(json.loads(raw[0]['nested']),{'methods':['fido2','phone']})
        self.assertEqual(json.loads(raw[0]['Raw record'])['id'],'private-user-1')
        self.assertNotIn('password',json.loads(raw[0]['Raw record']))
        self.assertEqual(validate_investigation_coverage(result,bundle),[])
        prepare_investigation_details(bundle,result)
        self.assertEqual(len(bundle['sheets']['raw_source.auth_methods']['rows']),2)

    def test_shared_inventory_links_survive_worksheet_splitting(self):
        bundle,result = bundle_and_result()
        bundle['assessment_sources']['auth_methods'][0]['records'].append({'id':'third','isMfaRegistered':True})
        prepare_investigation_details(bundle,result)
        split_evidence_sheets(bundle,result,max_rows=2)
        records = [r for r in result['recommendations'] if r['RecommendationId'].startswith('TEST-')]
        self.assertEqual(records[0]['InvestigationCount'],3)
        self.assertEqual(len(records[0]['InvestigationRanges']),2)
        self.assertEqual(validate_investigation_coverage(result,bundle),[])

    def test_label_finding_links_to_labels_without_duplicating_dlp_inventory(self):
        bundle,result = bundle_and_result()
        bundle['assessment_sources'].update(sensitivity_labels=source([{'id':'label'}]),
            label_policies=source([{'id':'publication'}]),dlp_policies=source([{'id':'dlp'}]))
        row=result['recommendations'][0]
        row.update(ControlId='DATA.PUBLISHING',Feature='Sensitivity labels',EvidenceKey='purview_policy_detail',
                   Observation='The Audit Labels publishing policy is configured.')
        prepare_investigation_details(bundle,result)
        self.assertEqual(row['InvestigationCount'],2)
        self.assertFalse(any('Dlp' in location for location in row['InvestigationRanges']))
        self.assertIn('raw_source.dlp_policies',bundle['sheets'])

    def test_html_keeps_every_catalog_check_but_omits_private_source_rows(self):
        bundle,result = bundle_and_result()
        prepare_investigation_details(bundle,result)
        html=render_customer_report(result,bundle,'Fixture','evidence.xlsx')
        self.assertEqual(html.count('class="domain-tile"'),9)
        self.assertEqual(html.count('class="check-detail"'),len(result['domain_coverage']))
        self.assertIn('Supported observations',html)
        self.assertNotIn('Verified strengths',html)
        self.assertNotIn('private@example.test',html)
        self.assertNotIn('private-user-1',html)
        self.assertLess(html.index('id="readiness-domains"'),html.index('id="action-plan"'))
        self.assertNotIn('action-detail" open',html)

    def test_workbook_guide_preserves_raw_logs_and_exact_links(self):
        bundle,result = bundle_and_result()
        bundle['assessment_result']=result
        with tempfile.TemporaryDirectory() as directory:
            filename=str(Path(directory)/'readable.xlsx')
            export_to_excel(result['recommendations'],filename,evidence_bundle=bundle)
            workbook=load_workbook(filename)
            self.addCleanup(workbook.close)
            self.assertEqual(workbook.active.title,'Start Here')
            self.assertEqual(workbook['Evidence Index'].sheet_state,'hidden')
            log=workbook.technical['Raw Auth Methods']
            self.assertEqual(log.cell(1,1).value,'id')
            self.assertEqual(log.cell(2,1).value,'private-user-1')
            self.assertFalse(log.column_dimensions['A'].hidden)
            headers={c.value:c.column for c in log[1]}
            from openpyxl.utils import get_column_letter
            self.assertFalse(log.column_dimensions[get_column_letter(headers['Raw record'])].hidden)
            register=workbook.technical['Findings Lineage']
            cols={c.value:c.column for c in register[1]}
            row=next(n for n in range(2,register.max_row+1) if register.cell(n,cols['RecommendationId']).value=='TEST-1')
            cell=register.cell(row,cols['Raw source evidence'])
            self.assertTrue(cell.hyperlink.target.startswith("#'Raw Auth Methods'!"))

    def test_device_code_block_does_not_hide_report_only_mfa_gap(self):
        unrelated=policy('Device code block',grant=('block',))
        unrelated['conditions']['authenticationFlows']={'transferMethods':'deviceCodeFlow'}
        policies=[policy('MFA report only',state='enabledForReportingButNotEnforced'),LEGACY_BLOCK,unrelated]
        self.assertNotIn('Device code block',conditional_access_facts(policies)['scope_unknown'])
        [row]=assess_identity_baseline(entra(policies),DAY)
        self.assertEqual((row['Disposition'],row['Priority']),('Action','High'))
        self.assertIn('report-only',row['Observation'])
        self.assertNotIn('No enforced policy blocks legacy',row['Observation'])

    def test_saved_dlp_summary_does_not_establish_operational_protection(self):
        bundle,result = bundle_and_result()
        bundle['verified_strengths']=[{'Key':'dlp-enforced','Area':'Data loss prevention',
            'Strength':'DLP policies and their protection rules are enabled.',
            'EvidenceKey':'purview_policy_detail','DomainId':'data_protection',
            'ObservationDate':DAY,'EvidenceComplete':True,'TenantId':TENANT,
            'EvidenceScope':'Fixture tenant'}]
        assessed=build_assessment_result([],bundle,evaluation_date=DAY,expected_tenant_id=TENANT)
        row=next(r for r in assessed['recommendations'] if r.get('FindingKey')=='dlp-enforced')
        self.assertEqual(row['OperationalResult'],'unknown')
        self.assertNotIn(row,assessed['strengths'])


if __name__=='__main__':
    unittest.main()
