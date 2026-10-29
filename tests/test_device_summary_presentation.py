"""Readable device populations and bounded technical examples from invented data."""
from copy import deepcopy
from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
import re
import unittest

from provenance_fixture import fictional_publication
from test_evidence_reconciliation import DAY, TENANT, dataset
from Core.device_compliance import build_device_compliance
from Core.expanded_findings import expanded_findings
from Core.assessment_result import build_assessment_result
from Core.evidence_selection import build_evidence_selection
from Core.finding_evidence import build_finding_evidence


def scenario(active=3):
    devices, settings = [], []
    old = (date.fromisoformat(DAY) - timedelta(days=60)).isoformat()
    overdue = (date.fromisoformat(DAY) - timedelta(days=31)).isoformat()
    def device(identifier, sync, state='noncompliant'):
        devices.append({'id': identifier, 'azureADDeviceId': 'entra-' + identifier,
            'deviceName': 'Example-' + identifier, 'lastSyncDateTime': sync, 'complianceState': state})
    def setting(identifier, label, state='noncompliant'):
        settings.append({'id': identifier + '-' + str(len(settings)), 'deviceId': identifier,
            'settingName': label, 'state': state, 'policyId': 'policy-1', 'lastReportedDateTime': DAY})
    for index in range(active):
        identifier = 'active-' + str(index)
        device(identifier, DAY)
        setting(identifier, 'Windows10CompliancePolicy.BitLockerEnabled')
    setting('active-0', 'Windows10CompliancePolicy.ActiveFirewallRequired')
    setting('active-1', 'Windows10CompliancePolicy.AntivirusRequired')
    for index in range(2):
        identifier = 'inactive-' + str(index)
        device(identifier, old)
        setting(identifier, 'Is active')
        setting(identifier, 'Windows10CompliancePolicy.BitLockerEnabled')
    device('overdue', overdue)
    setting('overdue', 'Windows10CompliancePolicy.BitLockerEnabled')
    device('unknown', None)
    setting('unknown', 'Windows10CompliancePolicy.BitLockerEnabled')
    device('no-policy', DAY, 'compliant')
    setting('no-policy', 'Has a compliance policy assigned')
    return {'evaluation_date': DAY, 'expected_tenant_id': TENANT,
        'collection_context': {'collected_at': DAY, 'tenant_id': TENANT}, 'sheets': {},
        'assessment_sources': {'managed_devices': [dataset(devices)],
            'compliance_setting_states': [dataset(settings)],
            'compliance_policies': [dataset([{'id': 'policy-1', 'displayName': 'Windows baseline'}])],
            'device_compliance_settings': [dataset([{'settings': {'secureByDefault': False,
                'deviceComplianceCheckinThresholdDays': 30}}])]}}


def assessment(bundle):
    rows = expanded_findings(bundle, evaluation_date=DAY)
    result = build_assessment_result(rows, bundle, evaluation_date=DAY, expected_tenant_id=TENANT)
    bundle['assessment_result'] = result
    return rows, result


class DeviceSummaryTests(unittest.TestCase):
    def test_unresolved_device_identity_is_not_displayed_as_zero_affected_devices(self):
        from Core.device_summary import finding_summary
        model = build_device_compliance(scenario(), evaluation_date=DAY)
        row = {'DeviceFindingType': 'ExplicitSecuritySettingFailure',
               'DeviceEvidenceCounts': {'resolved_device_count': 0, 'unresolved_identity_count': 2}}
        summary = finding_summary(row, model)
        self.assertIn('Device count unverified', summary['display_count'])
        self.assertIn('2 unresolved source record', summary['display_count'])
        self.assertNotIn('0 devices', summary['text'])

    def test_distinct_populations_partition_noncompliance_and_preserve_raw_states(self):
        bundle = scenario(); before = deepcopy(bundle)
        model = build_device_compliance(bundle, evaluation_date=DAY)
        self.assertEqual(model['summary']['counts'], {'active': 3, 'inactive': 2, 'overdue': 1, 'unknown': 1, 'retired': 0})
        self.assertEqual(model['summary']['total_noncompliant'], 7)
        self.assertEqual(model['summary']['validity_days'], 30)
        self.assertEqual(model['summary']['missing_policy_devices'], 1)
        self.assertTrue(any('marked compliant' in text for text in model['summary']['lines']))
        self.assertTrue(any('change it to Not compliant' in text for text in model['summary']['lines']))
        self.assertEqual(bundle, before)

    def test_policy_groups_count_devices_not_repeated_setting_records(self):
        bundle = scenario()
        rows = bundle['assessment_sources']['compliance_setting_states'][0]['records']
        rows.append(dict(rows[0]))
        model = build_device_compliance(bundle, evaluation_date=DAY)
        group = model['summary']['policy_groups'][0]
        self.assertEqual(group['policy'], 'Windows baseline')
        self.assertEqual(group['devices'], 3)
        self.assertEqual(next(item['devices'] for item in group['reasons'] if 'BitLocker' in item['reason']), 3)
        self.assertEqual(len(rows), len(bundle['assessment_sources']['compliance_setting_states'][0]['records']))
        self.assertEqual(model['summary']['counts']['active'], 3)

    def test_reported_inactivity_overrides_fresh_entra_sign_in(self):
        bundle = scenario()
        bundle['assessment_sources']['directory_devices'] = [dataset([{'id': 'directory-i0',
            'deviceId': 'entra-inactive-0', 'approximateLastSignInDateTime': DAY}])]
        model = build_device_compliance(bundle, evaluation_date=DAY)
        inactive = next(item for item in model['devices'] if item['device_id'] == 'inactive-0')
        self.assertEqual(inactive['compliance_activity'], 'inactive')
        self.assertTrue(all(item['triage'] != 'current_security_failure' for item in inactive['settings']))
        current = next(row for row in expanded_findings(bundle, evaluation_date=DAY)
                       if row.get('DeviceFindingType') == 'ExplicitSecuritySettingFailure')
        selected = current['InvestigationEvidence']['records']
        self.assertTrue(all(record['deviceId'].startswith('active-') for record in selected))

    def test_configured_validity_not_assessment_activity_window_controls_overdue(self):
        bundle = scenario()
        bundle['assessment_sources']['device_compliance_settings'][0]['records'][0]['settings']['deviceComplianceCheckinThresholdDays'] = 1
        bundle['assessment_sources']['managed_devices'][0]['records'][0]['lastSyncDateTime'] = (date.fromisoformat(DAY) - timedelta(days=2)).isoformat()
        model = build_device_compliance(bundle, evaluation_date=DAY)
        first = next(item for item in model['devices'] if item['device_id'] == 'active-0')
        self.assertEqual(first['compliance_activity'], 'overdue')
        self.assertEqual(model['summary']['counts']['active'], 2)

    def test_missing_or_conflicting_validity_does_not_assume_thirty_days(self):
        for conflict in [False, True]:
            bundle = scenario()
            if conflict:
                bundle['assessment_sources']['device_compliance_settings'].append(dataset([{'settings': {'deviceComplianceCheckinThresholdDays': 60}}]))
            else:
                bundle['assessment_sources'].pop('device_compliance_settings')
            summary = build_device_compliance(bundle, evaluation_date=DAY)['summary']
            self.assertIsNone(summary['validity_days'])
            self.assertFalse(any('30 days' in text for text in summary['lines']))
            self.assertEqual(summary['counts']['active'], 0)

    def test_duplicate_inventory_ids_are_counted_once_and_partial_scope_stays_visible(self):
        bundle = scenario()
        devices = bundle['assessment_sources']['managed_devices'][0]['records']
        devices.append(dict(devices[0]))
        bundle['assessment_sources']['managed_devices'][0]['source']['complete'] = False
        summary = build_device_compliance(bundle, evaluation_date=DAY)['summary']
        self.assertEqual(summary['total_noncompliant'], 7)
        self.assertTrue(any('inventory completeness' in line for line in summary['lines']))

    def test_missing_policy_name_is_explicit_and_reasons_remain_readable(self):
        bundle = scenario(); bundle['assessment_sources'].pop('compliance_policies')
        rows, result = assessment(bundle)
        current = next(row for row in result['recommendations'] if row.get('DeviceFindingType') == 'ExplicitSecuritySettingFailure')
        text = current['DeviceSummary']['text']
        self.assertIn('Policy name not retained', text)
        self.assertIn('BitLocker encryption requirement failed', text)
        self.assertNotIn('Windows10CompliancePolicy', text)

    def test_unknown_device_population_is_not_presented_as_zero_noncompliance(self):
        bundle = scenario(); bundle['assessment_sources'].pop('managed_devices')
        summary = build_device_compliance(bundle, evaluation_date=DAY)['summary']
        self.assertIn('could not be assessed', summary['lines'][0])


@fictional_publication
class TechnicalExamplesTests(unittest.TestCase):
    def test_evidence_pages_remain_linked_when_there_are_no_action_examples(self):
        from Core.technical_appendix import render_technical_appendix
        html = render_technical_appendix({}, {'html_evidence_folder': 'Technical Evidence'}, 'a.xlsx', 't.xlsx', {}, [])
        self.assertIn('href="Technical%20Evidence/index.html"', html)

    def test_html_examples_are_named_bounded_and_workbook_keeps_every_record(self):
        from Core.export_recommendations import export_to_excel, export_to_html
        from Core.workbook_layout import technical_workbook_path
        from tools.audit_report_signal import parse_report, summarize, public_report_source, technical_example_issues
        from openpyxl import load_workbook
        bundle = scenario(active=5); raw = deepcopy(bundle['assessment_sources'])
        rows, result = assessment(bundle)
        with TemporaryDirectory() as folder:
            workbook = Path(export_to_excel(rows, filename='Example.xlsx', evidence_bundle=bundle, output_dir=folder))
            html_path = Path(export_to_html(rows, filename='Example.html', evidence_bundle=bundle, output_dir=folder, excel_path=workbook))
            source = html_path.read_text(encoding='utf-8')
            self.assertTrue('<ul class="device-summary">' in source)
            self.assertTrue('Compliance policy' in source and 'Windows baseline' in source)
            self.assertIn('Device compliance at a glance', source)
            self.assertIn('2 noncompliant devices failed', source)
            self.assertNotIn('Example-active-0', public_report_source(source))
            self.assertTrue('Example-active-0' in source, 'Named technical examples should be available.')
            examples = re.findall(r'<details class="technical-example"[^>]*>(.*?)</details>', source, re.S)
            self.assertTrue(examples)
            self.assertTrue(all(fragment.count('<tr>') - 1 <= 3 for fragment in examples))
            self.assertEqual(technical_example_issues(source), [])
            audit = summarize(parse_report(html_path), html_source=source, workbook=workbook)
            self.assertEqual(audit['audit_issues'], [])
            book = load_workbook(technical_workbook_path(workbook))
            try:
                self.assertEqual(book['Raw Managed Devices'].max_row - 1, 10)
                self.assertEqual(book['Raw Compliance Setting States'].max_row - 1,
                                 len(raw['compliance_setting_states'][0]['records']))
            finally:
                book.close()
        self.assertEqual(bundle['assessment_sources'], raw)

    def test_audit_still_rejects_unbounded_examples_and_names_outside_appendix(self):
        from tools.audit_report_signal import summarize, technical_example_issues
        source = '<details class="technical-example"><table><tbody>' + '<tr><td>Example</td></tr>' * 4 + '</tbody></table></details>'
        self.assertTrue(technical_example_issues(source))
        audit = summarize([], html_source='<p>outside@example.invalid</p>')
        self.assertTrue(any('user identity' in issue for issue in audit['audit_issues']))


if __name__ == '__main__':
    unittest.main()
