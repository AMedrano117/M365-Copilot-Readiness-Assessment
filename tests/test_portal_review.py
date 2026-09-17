"""Manually reviewed portal visuals remain portable context, never scored evidence."""

import contextlib
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import struct
import tempfile
import unittest
from unittest.mock import patch
import zlib

from Core.assessment_package import restore_arguments
from Core.cli_parser import parse_arguments
from Core.offline_collection import empty_service_results, load_collection, save_collection
from Core.offline_report import run_offline_report
from Core.portal_report_import import route_portal_reports
from Core.portal_review import load_portal_review


TENANT = '44444444-4444-4444-8444-444444444444'


def reviewed_fixture(folder):
    folder = Path(folder)
    assets = folder / 'originals'
    previews = folder / 'images'
    assets.mkdir(parents=True, exist_ok=True)
    previews.mkdir(parents=True, exist_ok=True)
    pdf = b'%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\ntrailer\n<< /Root 1 0 R >>\n%%EOF\n'
    (assets / 'invented.pdf').write_bytes(pdf)
    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xffffffff)
    png = b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 1, 1, 8, 2, 0, 0, 0))
    png += chunk(b'IDAT', zlib.compress(b'\x00\x10\x20\x30')) + chunk(b'IEND', b'')
    (previews / 'invented.png').write_bytes(png)
    manifest = {'schema_version': 1, 'tenant_id': TENANT, 'captures': [{
        'id': 'invented-review', 'title': 'Fictional portal review', 'domain_id': 'data_protection',
        'captured_at': '2026-09-14', 'report_date': '2026-09-12',
        'source_file': 'originals/invented.pdf', 'source_sha256': hashlib.sha256(pdf).hexdigest(),
        'previews': ['images/invented.png'], 'summary': '=A1+A2',
        'limitations': ['The pictured page does not establish tenant-wide policy enforcement.'],
        'review_notes': ['A reviewer read the supplied fictional page.'],
        'coverage': [{'area': 'Policy overview', 'assessment_coverage': 'Visual context only',
                      'next_step': 'Obtain scoped policy configuration evidence.'}],
    }]}
    path = folder / 'review.json'
    path.write_text(json.dumps(manifest), encoding='utf-8')
    return path, manifest


class PortalReviewTests(unittest.TestCase):
    def parse(self, *arguments):
        with patch('sys.argv', ['main.py', *map(str, arguments)]):
            return parse_arguments(None, [])

    def test_review_contains_standalone_assets_and_metadata_without_scoring(self):
        with tempfile.TemporaryDirectory() as directory:
            path, _ = reviewed_fixture(directory)
            with patch('socket.create_connection', side_effect=AssertionError('No external resource fetching')):
                review = load_portal_review(path, TENANT, '2026-09-15')
            self.assertTrue(review['available'])
            self.assertTrue(review['captures'][0]['source_data_uri'].startswith('data:application/pdf;base64,'))
            self.assertTrue(review['captures'][0]['previews'][0]['data_uri'].startswith('data:image/png;base64,'))
            self.assertNotIn('data_uri', json.dumps(review['rows']))
            self.assertNotIn('observations', review)
            self.assertNotIn('recommendations', review)
            self.assertEqual(review['receipt'][0]['Status'], 'reviewed_reference')

    def test_strict_schema_tenant_date_hash_and_asset_boundaries(self):
        with tempfile.TemporaryDirectory() as directory:
            path, original = reviewed_fixture(directory)
            cases = [
                ('tenant', lambda value: value.update(tenant_id='55555555-5555-4555-8555-555555555555')),
                ('schema', lambda value: value.update(schema_version=2)),
                ('unsupported', lambda value: value.update(auto_pass=True)),
                ('future', lambda value: value['captures'][0].update(captured_at='2099-01-01')),
                ('later', lambda value: value['captures'][0].update(report_date='2026-09-15')),
                ('hash', lambda value: value['captures'][0].update(source_sha256='0' * 64)),
                ('escapes', lambda value: value['captures'][0].update(previews=['../outside.png'])),
                ('local relative', lambda value: value['captures'][0].update(previews=['https://example.invalid/picture.png'])),
            ]
            for reason, change in cases:
                value = json.loads(json.dumps(original))
                change(value)
                path.write_text(json.dumps(value), encoding='utf-8')
                with self.subTest(reason=reason), self.assertRaisesRegex(ValueError, reason):
                    load_portal_review(path, TENANT)
            path.write_text(json.dumps(original), encoding='utf-8')
            (Path(directory) / 'images' / 'invented.png').write_text('<svg onload="alert(1)"></svg>', encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'valid PNG or JPEG'):
                load_portal_review(path, TENANT)

    def test_pdf_directory_discovery_reports_reference_without_importing(self):
        with tempfile.TemporaryDirectory() as directory:
            path, _ = reviewed_fixture(directory)
            routed = route_portal_reports([path.parent / 'originals'])
            self.assertFalse(routed['sam'] or routed['dspm'] or routed['readiness'])
            self.assertIn('--portal-review', routed['reference_warnings'][0])

    def test_collection_and_recipe_package_keep_all_review_assets_after_move(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory).resolve()
            original, _ = reviewed_fixture(root / 'source')
            collection = save_collection(root / 'collection.json', tenant_id=TENANT, tenant_name='Invented review',
                                         service_results=empty_service_results(),
                                         supplemental_inputs={'portal_review': original})
            initial = load_collection(collection)
            portable = root / 'portable'
            shutil.copytree(initial['package_directory'], portable)
            original.rename(original.with_name('unavailable.json'))
            moved = load_collection(portable / 'collection.json')
            args = restore_arguments(self.parse('--collection-input', portable / 'collection.json'), moved)
            review = load_portal_review(args.portal_review, TENANT)
            self.assertEqual(len(review['asset_paths']), 2)
            self.assertTrue(all(Path(asset).is_relative_to(portable) for asset in review['asset_paths']))
            with contextlib.redirect_stdout(io.StringIO()), \
                 patch('Core.processor.process_and_print_all_information', return_value={'html_path': 'example.html'}), \
                 patch('Core.offline_collection.save_collection', side_effect=AssertionError('Offline saved raw collection')):
                self.assertEqual(run_offline_report(args), 0)
            recipe = json.loads((portable / 'rebuild.json').read_text(encoding='utf-8'))
            self.assertEqual({Path(item['path']).suffix for item in recipe['files']}, {'.json', '.pdf', '.png'})
            (Path(review['asset_paths'][1])).write_bytes(b'changed preview')
            with self.assertRaisesRegex(ValueError, 'has changed'):
                load_collection(portable / 'collection.json')

    def test_processor_visual_context_does_not_change_decision_and_workbook_is_safe(self):
        from Core.processor import process_and_print_all_information
        from Core.export_recommendations import export_to_excel
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path, _ = reviewed_fixture(root / 'review')
            with contextlib.redirect_stdout(io.StringIO()), \
                 patch('Core.processor.export_tabular_reports', return_value=(None, None)), \
                 patch('Core.processor.export_to_html', return_value='example.html'), \
                 patch('Core.processor.print_recommendations_summary'):
                arguments = dict(**empty_service_results(), expected_tenant_id=TENANT,
                                 collection_context={'mode': 'offline', 'freshness': 'missing', 'evaluation_date': '2026-09-15'})
                baseline = process_and_print_all_information(**arguments)
                reviewed = process_and_print_all_information(**arguments, portal_review=path)
            self.assertEqual(baseline['evidence_bundle']['assessment_result'], reviewed['evidence_bundle']['assessment_result'])
            bundle = reviewed['evidence_bundle']
            workbook_path = export_to_excel(bundle['recommendations'], filename=str(root / 'review.xlsx'), evidence_bundle=bundle)
            workbook = load_workbook(workbook_path, read_only=True, data_only=False)
            try:
                rows = list(workbook['Portal Review'].values)
                self.assertIn('SHA-256', rows[0])
                summary_column = rows[0].index('Summary')
                self.assertEqual(rows[1][summary_column], '=A1+A2')
                self.assertEqual(workbook['Portal Review'].cell(2, summary_column + 1).data_type, 's')
                self.assertNotIn('base64,', str(rows))
                self.assertIn('Rollout Progress', workbook.sheetnames)
            finally:
                workbook.close()

    def test_portal_only_offline_build_packages_review_and_its_tenant(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, _ = reviewed_fixture(root / 'review')
            previous = Path.cwd()
            try:
                os.chdir(root)
                with contextlib.redirect_stdout(io.StringIO()), \
                     patch('Core.processor.export_tabular_reports', return_value=(None, None)), \
                     patch('Core.processor.export_to_html', return_value='example.html'), \
                     patch('Core.processor.print_recommendations_summary'), \
                     patch('socket.create_connection', side_effect=AssertionError('Offline network access')), \
                     patch('subprocess.Popen', side_effect=AssertionError('Offline process launch')):
                    args = self.parse('--mode', 'offline', '--portal-review', manifest)
                    self.assertEqual(run_offline_report(args), 0)
                recipe = next((root / 'output' / 'assessments').glob('*/rebuild.json'))
                payload = load_collection(recipe)
                self.assertEqual(payload['tenant_id'], TENANT)
                self.assertEqual(len(payload['package']['files']), 3)
                replay = restore_arguments(self.parse('--collection-input', recipe), payload)
                self.assertEqual(load_portal_review(replay.portal_review)['tenant_id'], TENANT)
            finally:
                os.chdir(previous)

    def test_snapshot_cannot_overwrite_review_manifest_or_original_assets(self):
        with tempfile.TemporaryDirectory() as directory:
            manifest, _ = reviewed_fixture(directory)
            for target in (manifest, manifest.parent / 'originals' / 'invented.pdf', manifest.parent / 'images' / 'invented.png'):
                before = target.read_bytes()
                args = self.parse('--mode', 'offline', '--portal-review', manifest, '--snapshot-json', target)
                with self.subTest(path=target.name), self.assertRaisesRegex(ValueError, 'cannot overwrite'):
                    run_offline_report(args)
                self.assertEqual(target.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
