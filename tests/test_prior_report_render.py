import os
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from Core.export_recommendations import export_to_excel, export_to_html


class PriorReportRenderTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.original_cwd = Path.cwd()
        os.chdir(self.directory.name)
        self.current = [{
            'Service': 'Assessment', 'Feature': 'Offline tenant evidence coverage',
            'Observation': 'No current tenant collection is available.',
            'Recommendation': 'Review historical evidence and exported reports.',
            'Disposition': 'Coverage', 'Status': 'Not Assessed', 'Priority': 'High',
            'Category': 'Scan Coverage',
        }]
        self.prior = {
            'available': True, 'source_file': 'prior-tenant-assessment.xlsx',
            'tenant_name': 'Example tenant', 'generated_at': '2026-07-16T20:20:02+00:00',
            'methodology_version': 'historical-methodology', 'assessment_version': 'earlier-version',
            'recommendations': [{
                'Service': 'Entra', 'Feature': 'Historical authentication finding',
                'Observation': 'Earlier authentication policy gap.',
                'Recommendation': 'Review the authentication policies.',
                'Priority': 'High', 'Status': 'Action Required', 'Disposition': 'Action',
                'LinkUrl': 'javascript:alert(1)',
            }, {
                'Service': 'M365', 'Feature': 'Historical verified control',
                'Observation': 'Original control observation.', 'Disposition': 'Assurance',
            }],
            'collection_coverage': [{'Source': 'Graph at original run', 'Freshness': 'Fresh'}],
            'sheets': {
                'AI Adoption Usage': {'rows': [{'Metric': 'Active users', 'Value': 5, 'Period': 'D28', 'Refresh Date': '2026-07-15'}]},
                'M365 Activity Detail': {'rows': [{'Application': 'Teams', 'Active users': 8}]},
                'Copilot User Usage': {'rows': [{'UPN': 'private-user@example.invalid'}]},
                'Copilot Readiness Users': {'rows': [{'UPN': 'private-readiness@example.invalid'}]},
            },
            'warnings': ['An earlier report does not establish the current tenant state.'],
        }
        self.bundle = {
            'evaluation_date': '2026-09-15',
            'prior_report': self.prior,
            'collection_context': {
                'mode': 'offline', 'freshness': 'missing',
                'scope': 'Historical tenant assessment plus portal exports',
                'historical_sources': [{
                    'source_file': self.prior['source_file'], 'source_type': 'Previous workbook',
                    'reported_at': self.prior['generated_at'],
                }],
            },
        }

    def tearDown(self):
        os.chdir(self.original_cwd)
        self.directory.cleanup()

    def test_historical_findings_join_domain_actions_and_keep_original_dates(self):
        path = export_to_html(self.current, evidence_bundle=self.bundle)
        html = Path(path).read_text(encoding='utf-8')
        result = self.bundle['assessment_result']
        self.assertNotIn('Previous tenant assessment', html)
        self.assertNotIn('Full historical recommendation register', html)
        self.assertIn(self.prior['generated_at'], html)
        self.assertIn(self.prior['source_file'], html)
        self.assertIn('Historical verified control', html)
        identity = next(domain for domain in result['domains'] if domain['id'] == 'identity')
        historical = [row for row in identity['actions'] if row.get('Historical') == 'Yes']
        self.assertEqual(len(historical), 1)
        self.assertEqual(historical[0]['ActionType'], 'Confirmation')
        self.assertEqual(historical[0]['ObservationDate'], '')
        self.assertEqual(historical[0]['PriorReportDate'], self.prior['generated_at'])
        self.assertIn('Earlier authentication policy gap.', html)
        self.assertIn('Historical aggregate usage', html)
        historical_usage = [row for row in result['adoption_metrics'] if row.get('source_type') == 'prior_assessment']
        self.assertTrue(any(row['metric_id'] == 'copilot.active_users' and row['value'] == 5 and row['observed_at'] == '2026-07-15' for row in historical_usage))
        self.assertEqual(result['counts']['confirmation'], 1)
        self.assertEqual(result['counts']['strengths'], 0)
        self.assertIn('Readiness unconfirmed', html)
        self.assertIn('Its conclusion has not been revalidated from saved facts', html)
        self.assertNotIn('supplied portal exports only', html)
        self.assertNotIn('private-user@example.invalid', html)
        self.assertNotIn('private-readiness@example.invalid', html)
        self.assertNotIn('href="javascript:', html)

    def test_imported_html_is_escaped_and_missing_date_is_explicit(self):
        self.prior['generated_at'] = ''
        self.prior['recommendations'][0]['Observation'] = '<img src=x onerror=alert(1)>'
        self.prior['source_file'] = '<script>bad()</script>.xlsx'
        path = export_to_html(self.current, evidence_bundle=self.bundle)
        html = Path(path).read_text(encoding='utf-8')
        self.assertIn('&lt;img src=x onerror=alert(1)&gt;', html)
        self.assertNotIn('<img src=x', html)
        self.assertNotIn('<script>bad()', html)
        self.assertRegex(html, r'Original report date: <strong>(?:<span[^>]*>)?Not established')
        self.assertIn('The original observation date is not recorded', html)

    def test_workbook_preserves_historical_rows_and_maps_unique_tab_names(self):
        self.prior['sheets'] = {
            'A very long original evidence title one': {'rows': [{'=Header': '=HYPERLINK("bad")'}]},
            'A very long original evidence title two': {'rows': [{'=Header': 'Second'}]},
            'Report Sources': {'rows': [{'Detail': 'Existing source metadata'}]},
            'AI Adoption Usage': {'rows': [{'Metric': 'Original usage', 'Value': 5}]},
            'AI_Adoption_Usage': {'rows': [{'Metric': 'Similar table name', 'Value': 6}]},
            'Empty evidence': {'rows': []},
        }
        path = export_to_excel(self.current, evidence_bundle=self.bundle)
        workbook = load_workbook(path, data_only=False)
        try:
            self.assertIn('Prior Report Sources', workbook.sheetnames)
            mapping = list(workbook['Prior Report Sources'].values)
            rows = [dict(zip(mapping[0], row)) for row in mapping[1:]]
            by_original = {row['Original Sheet']: row['New Workbook Tab'] for row in rows}
            self.assertEqual(len(rows), 8)
            self.assertEqual(len({title.casefold() for title in workbook.sheetnames}), len(workbook.sheetnames))
            self.assertTrue(all(len(title) <= 31 for title in workbook.sheetnames))
            self.assertNotEqual(by_original['A very long original evidence title one'],
                                by_original['A very long original evidence title two'])
            self.assertNotEqual(by_original['Report Sources'], 'Prior Report Sources')
            self.assertIn(by_original['Empty evidence'], workbook.sheetnames)
            sheet = workbook[by_original['A very long original evidence title one']]
            self.assertEqual(sheet['A1'].value, '=Header')
            self.assertEqual(sheet['A1'].data_type, 's')
            self.assertEqual(sheet['A2'].value, '=HYPERLINK("bad")')
            self.assertEqual(sheet['A2'].data_type, 's')
            self.assertTrue(all(row['Original Report Date'] == self.prior['generated_at'] for row in rows))
            self.assertTrue(all(row['Source File'] == self.prior['source_file'] for row in rows))
            table_names = [name.casefold() for sheet in workbook for name in sheet.tables]
            self.assertEqual(len(set(table_names)), len(table_names))
        finally:
            workbook.close()


if __name__ == '__main__':
    unittest.main()
