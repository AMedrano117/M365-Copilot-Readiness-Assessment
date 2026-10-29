"""Fictional production-bundle and cross-output regressions for PR 24 follow-up."""
from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest

from provenance_fixture import fictional_publication
from test_evidence_reconciliation import DAY, TENANT, PARENT, dataset, gap
from Core.assessment_result import build_assessment_result, _action
from Core.coverage_summary import summarize_collection
from Core.evidence_reconciliation import customer_text, blocker_control_mapping


def runtime_fixture():
    from Core.evidence_layer import build_evidence_bundle
    registrations=[{'id':f'fictional-{i}','isMfaRegistered':i<4,
        'isPasswordlessCapable':i<6,'isMfaCapable':i<6,'userType':'Member',
        'methodsRegistered':['fido2'] if i<5 else ['microsoftAuthenticatorPasswordless'] if i==5 else []}
        for i in range(586)]
    m365=SimpleNamespace(users_summary={},license_coverage={'available':True,'refresh_date':DAY,
        'copilot_subscription_capacity':{'available':True,'enabled_seats':106,'assigned_seats':87,
            'qualification':'Fictional paid subscription seats; not unique people.'}},
        copilot_usage={'available':True,'selected_period':'D28','refresh_date':DAY,
            'periods':{'D28':{'enabled_users':87,'active_users':80}},'user_detail':[]},
        copilot_interaction_audit={'available':True,'availability_status':'available','complete':True,
            'window_start':'2026-10-20','window_end':DAY,'records':[{'id':f'event-{i}'} for i in range(3620)]})
    source=dataset([])['source']
    entra=SimpleNamespace(auth_methods_registration=registrations,
        users=[{'id':r['id'],'accountEnabled':True,'userType':'Member'} for r in registrations],
        collection_status={'auth_methods':source,'users':source})
    purview=SimpleNamespace(sensitivity_labels={'available':True,'labels':[{'id':'label-one','displayName':'Fictional label'}]},
        label_policies={'available':True,'policies':[{'id':'policy-one','Enabled':True}]},
        audit_config={'available':True,'unified_audit_enabled':True},
        collection_status={name:source for name in ('sensitivity_labels','label_policies','audit_config')})
    bundle=build_evidence_bundle([],({'_client':m365},[]),{'_client':entra},{'_client':purview},{},{},{})
    from Core.assessment_catalog import collect_assessment_sources
    bundle['assessment_sources']=collect_assessment_sources([m365,entra,purview],collected_at=DAY)
    bundle.update(expected_tenant_id=TENANT,evaluation_date=DAY,
        collection_context={'tenant_id':TENANT,'collected_at':DAY,'source_file':'fictional-collection.json'})
    return build_assessment_result([],bundle,evaluation_date=DAY,expected_tenant_id=TENANT),bundle


class RuntimeEvidenceTests(unittest.TestCase):
    def test_unavailable_response_envelopes_are_not_label_or_audit_records(self):
        from test_evidence_reconciliation import assessment
        result,_=assessment({name:[dataset([{'available':False,'total_labels':0}],available=False,availability_status='unavailable')]
            for name in ('sensitivity_labels','label_policies','audit_config')})
        for control in ('DATA.LABELS','DATA.PUBLISHING','DATA.AUDIT'):
            self.assertNotIn('ObservedEvidence',gap(result,control))

    def test_actual_bundle_keeps_paid_capacity_and_usage_but_not_pilot_entitlement(self):
        result,bundle=runtime_fixture(); row=gap(result,'LICENSE.ASSIGNMENT')
        self.assertEqual(bundle['ai_usage']['license_coverage']['copilot_subscription_capacity']['enabled_seats'],106)
        self.assertEqual(row['ObservedEvidence']['subscription_capacity']['unassigned_seats'],19)
        for text in ('106 seats','87 assigned','19 unassigned','87 enabled users','not proof of assigned entitlement'):
            self.assertIn(text,row['Observation'])
        self.assertEqual(row['ObservedEvidence']['pilot_entitlement'],'not_established')

    def test_actual_label_and_active_policy_remain_incomplete(self):
        result,_=runtime_fixture(); row=gap(result,'DATA.LABELS')
        self.assertIn('1 sensitivity-label definition',row['Observation'])
        self.assertEqual(row['ObservedEvidence']['active_publishing_policies'],1)
        for key in ('publication_scope','intended_population','broader_taxonomy','operational_use'):
            self.assertEqual(row['ObservedEvidence'][key],'not_established')

    def test_actual_audit_config_and_3620_bounded_events_do_not_pass(self):
        result,_=runtime_fixture(); row=gap(result,'DATA.AUDIT')
        self.assertIn('1 audit-configuration record',row['Observation'])
        self.assertIn('3620 Copilot audit-event',row['Observation'])
        self.assertIn('2026-10-20 through 2026-10-27',row['Observation'])
        self.assertEqual(next(c for c in result['controls'] if c['control_id']=='DATA.AUDIT')['status'],'Not established')

    def test_registered_passwordless_inventory_is_not_capability_or_mfa_flag(self):
        result,_=runtime_fixture(); reconciliation=result['authentication_assessment']['registration']['metric_reconciliation']
        self.assertIn('4 of 586',reconciliation['comparison_explanation'])
        self.assertIn('6 of 586',reconciliation['comparison_explanation'])
        metrics={r['Metric ID']:r for r in reconciliation['metrics']}
        self.assertEqual(metrics['passwordless_registered']['Source field'],'methodsRegistered')
        self.assertEqual(metrics['passwordless_registered']['Denominator'],586)
        self.assertEqual(metrics['passwordless_capable']['Source field'],'isPasswordlessCapable')
        self.assertGreater(reconciliation['conflicts']['passwordless_method_with_false_mfa_flag'],0)


class PresentationTests(unittest.TestCase):
    def test_device_conditions_and_reporting_ages_have_distinct_confirmation_titles(self):
        from Core.evidence_reconciliation import qualify_device_action_titles
        actions=[]
        for condition in ('real-time protection turned off','overdue security intelligence updates'):
            for age in ('current','older'):
                row={'Feature':'Review Windows device protection','Observation':condition,'Recommendation':'Validate the affected device population.',
                    'FindingKey':'expanded.windows_protection.'+('disabled' if condition.startswith('real') else 'overdue')+'.'+age,
                    'Disposition':'Action','EvidenceStatus':'limited','Status':'Not Assessed','DomainId':'endpoints'}
                action=_action(row); actions.append(action)
                self.assertEqual(action['ActionType'],'Confirmation')
                self.assertEqual(row['Observation'],condition)
        titles=[r['Feature'] for r in qualify_device_action_titles(actions)]
        self.assertEqual(len(set(titles)),4)
        self.assertTrue(any('real-time' in title for title in titles))
        self.assertTrue(any('security intelligence' in title for title in titles))

    def test_four_failed_controls_map_to_three_actions_without_changing_decisions(self):
        actions=[{'RecommendationId':f'F-{i}','PilotImpact':'Blocks pilot'} for i in range(3)]
        controls=[{'control_id':f'C-{i}','security_gate':True,'status':'Action required',
                   'recommendation_ids':[f'F-{min(i,2)}']} for i in range(4)]
        result={'actions':actions,'controls':controls}; before=deepcopy(result)
        mapping=blocker_control_mapping(result)
        self.assertTrue(mapping['mapping_complete'])
        self.assertIn('4 failed controls map to 3 blocker actions',mapping['explanation'])
        self.assertEqual(mapping['actions'][-1]['control_ids'],['C-2','C-3'])
        self.assertEqual(result,before)

    def test_unproven_mapping_is_qualified(self):
        result={'actions':[{'RecommendationId':'F-1','PilotImpact':'Blocks pilot'}],
            'controls':[{'control_id':'C-1','security_gate':True,'status':'Action required','recommendation_ids':[]}]}
        self.assertFalse(blocker_control_mapping(result)['mapping_complete'])

    def test_failed_controls_with_later_actions_are_not_invented_as_pilot_blockers(self):
        actions=[{'RecommendationId':f'F-{i}','PilotImpact':'Blocks pilot' if i<3 else 'Fix before broad rollout'} for i in range(5)]
        controls=[{'control_id':f'C-{i}','security_gate':True,'status':'Action required',
                   'recommendation_ids':[['F-0','F-1'],['F-2'],['F-3'],['F-4']][i]} for i in range(4)]
        result={'actions':actions,'controls':controls}; before=deepcopy(result)
        mapping=blocker_control_mapping(result)
        self.assertTrue(mapping['mapping_complete'])
        self.assertEqual(mapping['blocker_action_count'],3)
        self.assertIn('2 failed controls link to pilot blockers',mapping['explanation'])
        self.assertIn('remaining 2',mapping['explanation'])
        self.assertNotIn('4 failed controls map to 3',mapping['explanation'])
        self.assertEqual(result,before)

    def test_customer_requests_remove_procedures_without_erasing_observation_windows(self):
        for instruction in ('Wait 24 hours for the scan and rerun collection.','The scan may take 48 hours.',
                            'Import-Module Microsoft.Graph. Rerun the assessment.',
                            'Required permissions: AuditLog.Read.All; retry collection.'):
            projected=customer_text('3620 audit events were observed over seven days. '+instruction,'DATA.AUDIT')
            self.assertIn('3620 audit events were observed over seven days.',projected)
            self.assertNotRegex(projected,r'24 hours|48 hours|Import-Module|AuditLog.Read.All|rerun')
            self.assertIn('Provide dated audit configuration',projected)


class CoverageAliasTests(unittest.TestCase):
    def test_prefixed_mirrors_are_aggregated_and_missing_failed_request_is_retained(self):
        native=[]; states={}
        for i in range(5):
            parent=f'{i:08d}-2222-4222-8222-222222222222'
            request=dataset([],scope='Parent object '+parent,source_api='https://graph.microsoft.com/v1.0/servicePrincipals/'+parent+'/appRoleAssignments')
            native.append(request); states['entra_application_permissions:'+parent]=deepcopy(request['source'])
        failed=dataset([])['source']; failed.update(availability_status='unavailable',available=False,complete=False,error='Fictional denied request')
        states['entra_application_permissions:'+PARENT]=failed
        before=deepcopy((native,states)); coverage=summarize_collection({'application_permissions':native},states)
        self.assertEqual(len(coverage['summary']),1)
        row=coverage['summary'][0]
        self.assertEqual(row['Parent objects queried'],6)
        self.assertEqual(row['Successful queries'],5); self.assertEqual(row['Zero-result queries'],5)
        self.assertEqual(row['Failed queries'],1); self.assertEqual(len(coverage['requests']),6)
        self.assertEqual(sum(len(r['Mirrored status metadata']) for r in coverage['requests']),5)
        self.assertEqual((native,states),before)

    def test_false_available_flag_cannot_hide_a_failed_request(self):
        coverage=summarize_collection({'application_permissions':[dataset([],available=False)]})
        self.assertEqual(coverage['summary'][0]['Failed queries'],1)
        self.assertEqual(coverage['summary'][0]['Successful queries'],0)

    def test_planned_parent_is_not_counted_as_queried_or_returned_zero(self):
        coverage=summarize_collection({'application_permissions':[dataset([],availability_status='not_requested',available=False,
            scope='Parent object '+PARENT)]})
        row=coverage['summary'][0]
        self.assertEqual(row['Parent objects queried'],0)
        self.assertEqual(row['Successful queries'],0)
        self.assertIsNone(row['Returned records'])
        self.assertEqual(coverage['requests'][0]['Parent object'],PARENT)


class PreservedPrivilegeTests(unittest.TestCase):
    def test_14_active_3_eligible_128_assignments_and_nonhuman_counts_are_preserved(self):
        from test_privileged_identity import grant,sources,dataset as role_dataset,build
        active=[]
        for i in range(124):
            kind='user' if i%37<14 else 'group' if i%37<17 else 'servicePrincipal'
            record=grant(f'active-{i}',principal=f'principal-{i%37}',appScopeId=f'/fictional/{i}')
            record['principal']['@odata.type']='#microsoft.graph.'+kind; active.append(record)
        eligible=[grant(f'eligible-{i}',principal=f'principal-{i%3}',appScopeId=f'/eligible/{i}') for i in range(4)]
        data=sources(active,users=[{'id':f'principal-{i}','accountEnabled':True} for i in range(14)],
            role_eligibility_schedules=[role_dataset(eligible)])
        original=deepcopy(data); model=build(data); counts=model['counts']
        self.assertEqual((counts['active_user_identities'],counts['eligible_user_identities'],counts['assignments'],
                          counts['active_assignments'],counts['eligible_assignments']),(14,3,128,124,4))
        metrics={r['Metric']:r['Value'] for r in model['metric_reconciliation']}
        self.assertEqual(metrics['Unique groups'],3); self.assertEqual(metrics['Unique service principals'],20)
        self.assertEqual(data,original)


@fictional_publication
class CrossOutputTests(unittest.TestCase):
    def test_workbooks_html_summary_snapshot_and_replay_agree(self):
        from test_evidence_reconciliation import EvidenceProjectionTests
        from Core.export_recommendations import export_to_excel
        from Core.customer_report import render_customer_report
        from Core.pilot_summary import render_pilot_summary
        from Core.assessment_serialization import write_assessment_result,read_assessment_result
        from Core.workbook_layout import technical_workbook_path
        from openpyxl import load_workbook
        _,runtime=runtime_fixture(); rows,bundle=EvidenceProjectionTests().fixture()
        for key in ('ai_usage','authentication_methods'):
            bundle[key]=runtime[key]
        for key,values in runtime['assessment_sources'].items():
            if key not in bundle['assessment_sources']: bundle['assessment_sources'][key]=values
        rows[0]['Recommendation']='Wait 24 hours, then rerun collection with --reports-dir PATH.'
        result=build_assessment_result(rows,bundle,evaluation_date=DAY,expected_tenant_id=TENANT)
        bundle['assessment_result']=result
        with TemporaryDirectory() as folder:
            path=Path(export_to_excel(rows,filename='Fictional.xlsx',evidence_bundle=bundle,output_dir=folder))
            result=bundle['assessment_result']; before=deepcopy(result)
            html=render_customer_report(result,bundle,'Fictional',path,technical_workbook_path=technical_workbook_path(path))
            summary=render_pilot_summary(result,bundle,'Fictional',workbook_path=path)
            for output in (html,summary):
                self.assertIn('4 of 586',output); self.assertIn('6 of 586',output)
                self.assertNotIn('--reports-dir',output); self.assertNotIn('24 hours',output)
                from tools.audit_report_signal import public_report_source
                self.assertNotIn('restricted-0@example.invalid',public_report_source(output))
                if result['blocker_control_mapping']['explanation']:
                    self.assertIn(result['blocker_control_mapping']['explanation'],output)
            self.assertEqual(result,before)
            customer=load_workbook(path); technical=load_workbook(technical_workbook_path(path))
            try:
                texts='\n'.join(str(c.value or '') for s in customer for row in s for c in row)
                self.assertIn('106 seats',texts); self.assertIn('4 of 586',texts); self.assertIn('6 of 586',texts)
                self.assertNotIn('--reports-dir',texts); self.assertNotIn('24 hours',texts)
                for phrase in ('Copilot license information was not collected','Sensitivity label definitions were not collected','The audit configuration was not collected'):
                    leaked=[(s.title,c.coordinate,str(s.cell(1,c.column).value)) for s in customer for row in s for c in row
                            if phrase in str(c.value or '')]
                    self.assertEqual(leaked,[])
                technical_text='\n'.join(str(c.value or '') for s in technical for row in s for c in row)
                self.assertIn('--reports-dir',technical_text)
                risk=next(r for r in result['recommendations'] if r['RecommendationId']=='ENT-004')
                self.assertTrue(risk['RawEvidenceRanges']); self.assertTrue(risk['InvestigationRange'])
                self.assertTrue(risk['TechnicalEvidenceRanges']); self.assertEqual(risk['InvestigationCount'],8)
                self.assertIn('Collection Request Detail',technical.sheetnames)
                index=list(technical['Evidence Index'].values)
                self.assertTrue(any('ENT-004' in row for row in index[1:]))
            finally:
                customer.close(); technical.close()
            snapshot=Path(folder)/'result.json'; write_assessment_result(snapshot,result); replay=read_assessment_result(snapshot)
            for key in ('decision','controls','counts','authentication_assessment','dataset_coverage','blocker_control_mapping'):
                self.assertEqual(replay[key],json.loads(json.dumps(result[key],default=str)))
            self.assertEqual(summary,render_pilot_summary(replay,bundle,'Fictional',workbook_path=path))


if __name__=='__main__':
    unittest.main()
