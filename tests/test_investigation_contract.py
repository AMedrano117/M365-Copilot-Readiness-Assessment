"""A new service exports useful investigation evidence without a service adapter."""

from copy import deepcopy
import json
import os
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from Core.export_recommendations import export_to_excel, export_to_html
from Core.investigation_contract import materialize_declared_evidence, recommended_actions, validate_investigation_coverage
from Core.investigation_details import prepare_investigation_details


DAY = '2026-09-30T08:00:00Z'
ABSENT = object()


def recommendation(identifier='NOVEL-001', declaration=ABSENT, disposition='Action'):
    row = {'RecommendationId': identifier, 'Service': 'Fictional Observatory',
           'Feature': 'Review suspended widgets', 'Observation': '2 widget records have a suspended state.',
           'Recommendation': 'Investigate the suspended widget records.', 'Disposition': disposition,
           'Status': 'Action Required', 'Priority': 'High', 'DomainId': 'applications',
           'EvidenceBasis': 'Tenant evidence', 'EvidenceAvailable': 'Yes', 'Confidence': 'High',
           'ObservationDate': DAY, 'TenantId': 'fixture-tenant'}
    if declaration is not ABSENT:
        row['InvestigationEvidence'] = declaration
    return row


def declaration(**values):
    return {'kind': 'records', 'sheet_name': 'Widget Investigation',
            'records': [{'id': 'private-widget-1', 'name': 'Private widget one', 'state': 'suspended',
                         'timestamp': DAY, 'attributes': {'region': 'fixture-region', 'codes': ['x', 'y']}},
                        {'id': 'private-widget-2', 'name': 'Private widget two', 'state': 'suspended',
                         'timestamp': DAY, 'attributes': {'region': 'fixture-region', 'codes': ['z']}}],
            'source': {'api': 'https://example.test/v1/widgets', 'file': 'fixture-widgets.json',
                       'collected_at': DAY, 'window': '2026-09-01 through 2026-09-30',
                       'filter': 'state=suspended', 'scope': 'Fixture widgets', 'pages': 1,
                       'complete': True, 'truncated': False, 'permissions': 'Widget.Read',
                       'licensing': 'Fixture entitlement', 'retention': '30 days'},
            'reconciliation': {'operation': 'count', 'expected': 2}, **values}


class InvestigationContractTests(unittest.TestCase):
    def materialize(self, contract):
        bundle = {'sheets': {}}
        output = materialize_declared_evidence(bundle, [recommendation(declaration=contract)])
        return bundle, output['NOVEL-001']

    def test_count_distinct_and_sum_reproduce_different_measures_without_collapsing_rows(self):
        records = [{'id': 'one', 'entity': 'same', 'amount': 2}, {'id': 'two', 'entity': 'same', 'amount': 4},
                   {'id': 'three', 'entity': 'other', 'amount': 1}]
        for operation, field, expected in [('count', None, 3), ('distinct', 'entity', 2), ('sum', 'amount', 7)]:
            with self.subTest(operation=operation):
                bundle, info = self.materialize(declaration(records=records,
                    reconciliation={'operation': operation, 'field': field, 'expected': expected}))
                self.assertEqual(info['InvestigationStatus'], 'Records available')
                self.assertEqual(info['InvestigationCount'], 3)
                self.assertIn(f'{operation} reproduces {expected}', info['InvestigationQualification'])
                self.assertEqual(len(next(iter(bundle['sheets'].values()))['rows']), 3)

    def test_incomplete_source_cannot_establish_absence(self):
        _, info = self.materialize({'kind': 'absence', 'reason': 'No schedules were returned.',
                                   'source': {'complete': False, 'truncated': True}})
        self.assertEqual(info['InvestigationStatus'], 'Evidence unavailable')
        self.assertIn('not established', info['InvestigationPublicQualification'])

    def test_each_measure_mismatch_is_visible_without_padding_or_dropping_evidence(self):
        records = [{'id': 'one', 'entity': 'same', 'amount': 2}, {'id': 'two', 'entity': 'same', 'amount': 4}]
        for operation, field, expected in [('count', None, 8), ('distinct', 'entity', 2), ('sum', 'amount', 9)]:
            with self.subTest(operation=operation):
                _, info = self.materialize(declaration(records=records,
                    reconciliation={'operation': operation, 'field': field, 'expected': expected}))
                self.assertEqual(info['InvestigationStatus'], 'Count mismatch')
                self.assertEqual(info['InvestigationCount'], 2)
                self.assertIn('Count mismatch', info['InvestigationQualification'])
                self.assertTrue(info['InvestigationRange'])

    def test_unmeasurable_fields_and_invalid_sum_values_do_not_claim_reconciliation(self):
        for record, operation in [({'id': 'one'}, 'distinct'), ({'id': 'one', 'amount': True}, 'sum'),
                                  ({'id': 'one', 'amount': float('inf')}, 'sum')]:
            with self.subTest(operation=operation, record=record):
                _, info = self.materialize(declaration(records=[record],
                    reconciliation={'operation': operation, 'field': 'amount', 'expected': 1}))
                self.assertEqual(info['InvestigationStatus'], 'Count mismatch')
                self.assertIn('could not be established', info['InvestigationQualification'])

    def test_absence_unavailable_and_planning_require_precise_explanations(self):
        for kind, status, reason in [
            ('absence', 'Configuration evidence', 'The completed widget query returned zero configured schedules.'),
            ('unavailable', 'Evidence unavailable', 'The widget API returned HTTP 403; Widget.Read was not granted.'),
            ('planning', 'Planning decision', 'A widget operating schedule must be agreed with the business owner.'),
        ]:
            with self.subTest(kind=kind):
                bundle, info = self.materialize({'kind': kind, 'reason': reason})
                self.assertEqual(info['InvestigationStatus'], status)
                self.assertEqual(info['InvestigationQualification'], reason)
                self.assertEqual(info['InvestigationCount'], 0)
                self.assertFalse(info['InvestigationRange'])
                self.assertFalse(bundle['sheets'])
                _, unexplained = self.materialize({'kind': kind})
                self.assertEqual(unexplained['InvestigationStatus'], 'Detail mapping missing')

    def test_malformed_declarations_are_visible_mapping_gaps_instead_of_export_errors(self):
        for contract in ['aggregate only', {'records': 2}, {'records': ['record']}, {'kind': []},
                         declaration(record_id_field=['id']), {'kind': 'unknown'}]:
            with self.subTest(contract=contract):
                _, info = self.materialize(contract)
                self.assertEqual(info['InvestigationStatus'], 'Detail mapping missing')
                self.assertTrue(info['InvestigationQualification'])
                self.assertEqual(info['InvestigationCount'], 0)

    def test_reserved_source_fields_do_not_replace_provenance_and_evidence_ids_are_stable(self):
        records = [{'id': 'source-one', 'RecommendationId': 'source-original', 'Evidence ID': 'source-evidence',
                    'Record ID': 'source-record', 'attributes': {'nested': [1, 2]}}]
        bundle, _ = self.materialize(declaration(records=records, reconciliation={'expected': 1}))
        row = next(iter(bundle['sheets'].values()))['rows'][0]
        self.assertEqual(row['RecommendationId'], 'NOVEL-001')
        self.assertEqual(row['Record.RecommendationId'], 'source-original')
        self.assertEqual(row['Record.Evidence ID'], 'source-evidence')
        self.assertEqual(row['Record.Record ID'], 'source-record')
        attributes = json.loads(row['attributes']) if isinstance(row['attributes'], str) else row['attributes']
        self.assertEqual(attributes, {'nested': [1, 2]})
        self.assertTrue(row['Evidence ID'].startswith('INV-'))
        second, _ = self.materialize(declaration(records=records, reconciliation={'expected': 1}))
        self.assertEqual(row['Evidence ID'], next(iter(second['sheets'].values()))['rows'][0]['Evidence ID'])

    def test_omitted_declaration_without_legacy_mapping_is_a_visible_gap(self):
        rec = recommendation()
        bundle = {'sheets': {}, 'recommendations': [deepcopy(rec)]}
        result = {'actions': [rec], 'recommendations': [deepcopy(rec)]}
        prepare_investigation_details(bundle, result)
        self.assertEqual(result['actions'][0]['InvestigationStatus'], 'Detail mapping missing')
        self.assertTrue(result['actions'][0]['InvestigationQualification'])
        self.assertEqual(validate_investigation_coverage(result, bundle), [])
        self.assertEqual(bundle['investigation_validation'], [])

    def test_opportunity_with_a_recommendation_has_coverage_without_becoming_a_primary_action(self):
        rec = recommendation(declaration=declaration(), disposition='Opportunity')
        bundle = {'sheets': {}, 'recommendations': [deepcopy(rec)]}
        result = {'actions': [], 'recommendations': [rec]}
        prepare_investigation_details(bundle, result)
        self.assertEqual(result['actions'], [])
        self.assertEqual(recommended_actions(result), [rec])
        self.assertEqual(rec['InvestigationStatus'], 'Records available')
        self.assertEqual(rec['InvestigationCount'], 2)
        self.assertEqual(validate_investigation_coverage(result, bundle), [])
        self.assertEqual(len(bundle['investigation_coverage']), 1)


class InvestigationContractExportTests(unittest.TestCase):
    def setUp(self):
        self.previous = os.getcwd()
        self.directory = tempfile.TemporaryDirectory()
        os.chdir(self.directory.name)

    def tearDown(self):
        os.chdir(self.previous)
        self.directory.cleanup()

    def export(self, rec):
        bundle = {'sheets': {}, 'evaluation_date': '2026-09-30',
                  'collection_context': {'collected_at': DAY, 'tenant_id': 'fixture-tenant', 'mode': 'offline'}}
        path = export_to_excel([rec], filename='generic-investigation.xlsx', evidence_bundle=bundle)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        return path, workbook, bundle

    def row(self, workbook, title, identifier):
        sheet = workbook[title]
        headers = {cell.value: cell.column for cell in sheet[1]}
        position = next(number for number in range(2, sheet.max_row + 1)
                        if sheet.cell(number, headers['RecommendationId']).value == identifier)
        return {header: sheet.cell(position, column) for header, column in headers.items()}

    def test_nested_source_timestamps_survive_workbook_serialization(self):
        from datetime import datetime, timezone
        timestamp = datetime(2026, 9, 30, 8, tzinfo=timezone.utc)
        rec = recommendation(declaration=declaration(records=[{'id': 'dated', 'timestamp': timestamp,
              'attributes': {'lastSeen': timestamp}}], reconciliation={'expected': 1}))
        _, workbook, _ = self.export(rec)
        row = dict(zip([cell.value for cell in workbook['Widget Investigation'][1]],
                       [cell.value for cell in workbook['Widget Investigation'][2]]))
        self.assertEqual(row['Observed At'], timestamp.isoformat())
        self.assertEqual(json.loads(row['attributes'])['lastSeen'], timestamp.isoformat())

    def test_fictional_service_exports_retained_records_and_exact_links_without_an_adapter(self):
        rec = recommendation(declaration=declaration())
        path, workbook, bundle = self.export(rec)
        sheet = workbook['Widget Investigation']
        values = list(sheet.values)
        records = [dict(zip(values[0], row)) for row in values[1:]]
        self.assertEqual([row['Record ID'] for row in records], ['private-widget-1', 'private-widget-2'])
        self.assertEqual(records[0]['Entity'], 'Private widget one')
        self.assertEqual(json.loads(records[0]['attributes']), {'region': 'fixture-region', 'codes': ['x', 'y']})
        self.assertEqual(records[0]['Collection Window'], '2026-09-01 through 2026-09-30')
        self.assertEqual(records[0]['Selection / Scope'], 'state=suspended')
        expected = f"'Widget Investigation'!A2:{get_column_letter(sheet.max_column)}3"
        for title, field in [('Action Plan', 'Investigation Details'), ('Evidence Index', 'Workbook Tab'),
                             ('Recommendations', 'Evidence Sheet')]:
            row = self.row(workbook, title, 'NOVEL-001')
            self.assertEqual(row[field].hyperlink.target, '#' + expected)
        action = self.row(workbook, 'Action Plan', 'NOVEL-001')
        self.assertIn('2', action['Investigation Details'].value)
        self.assertIn('count reproduces 2', action['Qualification'].value)
        self.assertEqual(bundle['investigation_validation'], [])
        html_path = export_to_html([rec], filename='generic-investigation.html', evidence_bundle=bundle, excel_path=path)
        html = Path(html_path).read_text(encoding='utf-8')
        for value in ('private-widget-1', 'private-widget-2', 'Private widget one', 'fixture-region'):
            self.assertNotIn(value, html)

    def test_long_original_attributes_survive_and_exact_range_includes_continuation_columns(self):
        original = 'X' * 70000
        contract = declaration(records=[{'id': 'long-record', 'investigationAttributes': original}],
                               reconciliation={'operation': 'count', 'expected': 1})
        _, workbook, _ = self.export(recommendation(declaration=contract))
        sheet = workbook['Widget Investigation']
        headers = [cell.value for cell in sheet[1]]
        pieces = [sheet.cell(2, column).value for column, header in enumerate(headers, 1)
                  if header == 'investigationAttributes' or str(header).startswith('investigationAttributes (continued')]
        self.assertEqual(''.join(pieces), original)
        expected = f"#'Widget Investigation'!A2:{get_column_letter(sheet.max_column)}2"
        self.assertEqual(self.row(workbook, 'Action Plan', 'NOVEL-001')['Investigation Details'].hyperlink.target, expected)

    def test_missing_mapping_is_visible_in_workbook_action_and_registers(self):
        _, workbook, _ = self.export(recommendation())
        action = self.row(workbook, 'Action Plan', 'NOVEL-001')
        self.assertEqual(action['Investigation Details'].value, 'Detail mapping missing')
        self.assertTrue(action['Qualification'].value)
        for title in ('Evidence Index', 'Recommendations'):
            row = self.row(workbook, title, 'NOVEL-001')
            self.assertIn('Detail mapping missing', [cell.value for cell in row.values()])

    def test_nonrecord_decisions_keep_their_precise_reason_visible_in_the_workbook(self):
        for kind, status, reason in [
            ('absence', 'Configuration evidence', 'The completed widget query returned zero schedules.'),
            ('unavailable', 'Evidence unavailable', 'Widget.Read was denied by the source API with HTTP 403.'),
            ('planning', 'Planning decision', 'The business owner has not supplied the widget operating schedule.'),
        ]:
            with self.subTest(kind=kind):
                _, workbook, _ = self.export(recommendation(declaration={'kind': kind, 'reason': reason}))
                self.assertNotIn('Widget Investigation', workbook.sheetnames)
                action = self.row(workbook, 'Action Plan', 'NOVEL-001')
                self.assertEqual(action['Investigation Details'].value, status)
                self.assertIn(reason, action['Qualification'].value)
                for title in ('Evidence Index', 'Recommendations'):
                    row = self.row(workbook, title, 'NOVEL-001')
                    self.assertIn(status, [cell.value for cell in row.values()])
                workbook.close()

    def test_opportunity_records_export_and_link_from_both_registers(self):
        _, workbook, bundle = self.export(recommendation(declaration=declaration(), disposition='Opportunity'))
        self.assertIn('Widget Investigation', workbook.sheetnames)
        self.assertFalse(any(row['RecommendationId'] == 'NOVEL-001' for row in bundle['assessment_result']['actions']))
        for title, column in (('Evidence Index', 'Workbook Tab'), ('Recommendations', 'Evidence Sheet')):
            cell = self.row(workbook, title, 'NOVEL-001')[column]
            self.assertTrue(cell.hyperlink.target.startswith("#'Widget Investigation'!A2:"))


if __name__ == '__main__':
    unittest.main()
