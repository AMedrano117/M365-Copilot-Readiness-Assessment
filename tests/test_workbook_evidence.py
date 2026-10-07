"""Workbook evidence sheets, Findings Register links and technical recommendations."""

import os
from pathlib import Path
import tempfile
import unittest

from tests.workbook_test_helpers import load_workbook_pair as load_workbook
from openpyxl.utils.cell import range_boundaries

from Core.assessment_result import build_assessment_result
from Core.export_recommendations import export_to_excel
from Core.investigation_details import prepare_investigation_details
from Core.workbook_evidence import add_evidence_sheets
from tests.test_finding_evidence import build, event, legacy_finding, source


class WorkbookEvidenceTests(unittest.TestCase):
    def workbook(self, max_rows=None):
        logs = [event(number, code=[0, 53003, 50126][number % 3], id=f'{number:020d}') for number in range(1, 9)]
        row = legacy_finding('ENT-018', Status='Action Required', ControlId='IDENTITY.AUTH', ObservationDate='2026-09-28')
        bundle = {'assessment_sources': {'signin_logs': [{'records': logs, 'source': source()}]}, 'sheets': {}}
        if max_rows:
            bundle['worksheet_row_capacity'] = max_rows
        result = build_assessment_result([row], bundle, evaluation_date='2026-09-30',
                                         expected_tenant_id='11111111-1111-4111-8111-111111111111')
        bundle['assessment_result'] = result
        prepare_investigation_details(bundle, result)
        _, model = build(result['recommendations'], bundle['assessment_sources'])
        add_evidence_sheets(bundle, result, model, html_folder='report_evidence', app_builder_files={})
        cwd = os.getcwd()
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        os.chdir(directory.name)
        try:
            path = export_to_excel(result['recommendations'], filename='evidence.xlsx', evidence_bundle=bundle)
            path = Path(path).resolve()
        finally:
            os.chdir(cwd)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        return workbook, model, result['recommendations'][0]['RecommendationId']

    def test_findings_register_links_to_the_exact_evidence_block(self):
        workbook, model, identifier = self.workbook()
        register = workbook.technical['Findings Lineage']
        headers = {cell.value: cell.column for cell in register[1]}
        number = next(n for n in range(2, register.max_row + 1) if register.cell(n, headers['RecommendationId']).value == identifier)
        link = register.cell(number, headers['Evidence records']).hyperlink.target.lstrip('#')
        title, cells = link.rsplit('!', 1)
        sheet = workbook.technical[title.strip("'")]
        self.assertTrue(sheet.tables, 'Evidence sheets are filterable Excel tables.')
        _, first, _, last = range_boundaries(cells)
        columns = [cell.value for cell in sheet[1]]
        block = [dict(zip(columns, (cell.value for cell in row))) for row in sheet.iter_rows(min_row=first, max_row=last)]
        evidence = model['findings'][0]['evidence']
        self.assertEqual(len(block), evidence['record_count'])
        self.assertTrue(all(row['finding_id'] == identifier for row in block))
        self.assertEqual(register.cell(number, headers['Record unit']).value, 'sign-in events')
        self.assertEqual(register.cell(number, headers['Affected entities']).value, 3)
        id_cell = sheet.cell(first, columns.index('eventId') + 1)
        self.assertEqual(id_cell.data_type, 's')
        self.assertEqual(id_cell.value, '00000000000000000001')
        self.assertIn('Technical Recommendations', workbook.technical.sheetnames)
        fix = register.cell(number, headers['Technical fix'])
        self.assertTrue(fix.hyperlink.target.startswith("#'Technical Recommendations'!A"))

    def test_technical_recommendations_link_to_verified_documentation(self):
        workbook, _, identifier = self.workbook()
        sheet = workbook.technical['Technical Recommendations']
        headers = {cell.value: cell.column for cell in sheet[1]}
        row = next(n for n in range(2, sheet.max_row + 1) if sheet.cell(n, headers['RecommendationId']).value == identifier)
        self.assertEqual(sheet.cell(row, headers['Guidance status']).value, 'verified')
        self.assertIn('Conditional Access', sheet.cell(row, headers['Where to configure']).value)
        self.assertTrue(sheet.cell(row, headers['Documentation 1']).hyperlink.target.startswith('https://learn.microsoft.com/'))
        self.assertTrue(sheet.cell(row, headers['RecommendationId']).hyperlink.target.startswith("#'Findings Lineage'!A"))
        start = workbook.technical['Start Here']
        self.assertIn('Technical Recommendations', [cell.value for cell in start['A']])

    def test_split_evidence_sheets_remap_finding_links(self):
        workbook, model, identifier = self.workbook(max_rows=3)
        register = workbook.technical['Findings Lineage']
        headers = {cell.value: cell.column for cell in register[1]}
        number = next(n for n in range(2, register.max_row + 1) if register.cell(n, headers['RecommendationId']).value == identifier)
        link = register.cell(number, headers['Evidence records']).hyperlink.target.lstrip('#')
        title, cells = link.rsplit('!', 1)
        _, first, _, last = range_boundaries(cells)
        self.assertLessEqual(last, workbook.technical[title.strip("'")].max_row)
        parts = [name for name in workbook.technical.sheetnames if name.startswith('Evidence Legacy Auth Sign-ins')]
        self.assertGreater(len(parts), 1)
        total = sum(workbook.technical[name].max_row - 1 for name in parts)
        self.assertEqual(total, model['findings'][0]['evidence']['record_count'])


    def test_finding_across_tables_links_largest_block_and_lists_all(self):
        columns = ['finding_id', 'detail_record_id']
        model = {'findings': [], 'tables': {
            'small': {'role': 'selected', 'concern_id': 'data-protection', 'short_title': 'small', 'title': 'Small',
                      'columns': columns, 'rows': [{'finding_id': 'PUR-014', 'detail_record_id': 'DET-a'}]},
            'large': {'role': 'selected', 'concern_id': 'data-protection', 'short_title': 'large', 'title': 'Large',
                      'columns': columns, 'rows': [{'finding_id': 'PUR-014', 'detail_record_id': f'DET-{n}'} for n in range(3)]}}}
        result = {'recommendations': [{'RecommendationId': 'PUR-014'}]}
        bundle = {'sheets': {}}
        add_evidence_sheets(bundle, result, model)
        row = result['recommendations'][0]
        large = model['tables']['large']['sheet_title']
        self.assertEqual(row['EvidenceRecordsRange'], f"'{large}'!A2:B4")
        self.assertEqual(len(row['EvidenceRecordsRanges']), 2)
        self.assertNotIn('_largest_evidence_block', row)

if __name__ == '__main__':
    unittest.main()
