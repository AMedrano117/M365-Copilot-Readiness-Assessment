"""Keep offline release checks aligned with rendered reports without weakening gates."""

from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tools.audit_report_signal import parse_report, summarize


def modern_html(*, count=1, linked=True, kind='Remediation'):
    action = (f'<article class="action" data-priority="high" id="action-1">'
              f'<span class="action-kind">{kind}</span><h3>Review fictional incident</h3>'
              '<p><strong>What we found.</strong> An invented finding.</p>'
              '<p><strong>What to do.</strong> Review invented evidence.</p>'
              + ('<a href="#evidence-test-001">Evidence and qualifications</a>' if linked else '')
              + '</article>') if count else ''
    return (f'<dt>Open actions</dt><dd>{count} to resolve or verify</dd>'
            f'<section id="action-plan">{action}</section>'
            '<table><tr id="evidence-test-001"><td>TEST-001</td></tr></table>')


class ReportAuditTests(unittest.TestCase):
    def setUp(self):
        self.directory = self.enterContext(tempfile.TemporaryDirectory())
        self.root = Path(self.directory)

    def report(self, source):
        target = self.root / 'report.html'
        target.write_text(source, encoding='utf-8')
        return parse_report(target)

    def workbook(self, *, action=True, action_id='TEST-001', disposition='Action',
                 source='Offline report build', state='offline', truncated=False, integrity='Passed',
                 empty_text='No deployment actions were identified from the evidence collected.'):
        book = Workbook()
        book.remove(book.active)
        rows = {
            'Action Plan': ([['RecommendationId', 'What We Found'], [action_id, 'An invented finding.']]
                            if action else [['What We Found'], [empty_text]]),
            'Recommendations': [['RecommendationId', 'Disposition', 'Observation', 'Control ID'],
                                *([['TEST-001', disposition, 'An invented finding.', 'THREAT.INCIDENTS']] if action else [])],
            'Evidence Index': [['RecommendationId'], *([['TEST-001']] if action else [])],
            'Control Results': [['Control ID'], ['THREAT.INCIDENTS']],
            'Collection Coverage': [['Source', 'State', 'Truncated'], [source, state, truncated]],
            'Run Manifest': [['Item', 'Value'], ['Input: collection', 'collection.json']],
            'Integrity Checks': [['Status', 'Issue'], [integrity, 'Invented integrity failure' if integrity == 'Failed' else '']],
        }
        for title, sheet_rows in rows.items():
            sheet = book.create_sheet(title)
            for row in sheet_rows:
                sheet.append(row)
        path = self.root / 'evidence.xlsx'
        book.save(path)
        book.close()
        return path

    def test_current_action_cards_match_workbook_by_linked_identifier(self):
        source = modern_html()
        records = self.report(source)
        self.assertEqual(records[0]['RecommendationId'], 'TEST-001')
        self.assertEqual(records[0]['Observation'], 'An invented finding.')
        self.assertEqual(summarize(records, source, self.workbook())['audit_issues'], [])

    def test_current_evidence_actions_remain_in_cross_output_checks(self):
        source = modern_html(kind='Evidence')
        records = self.report(source)
        self.assertEqual(records[0]['Disposition'], 'Coverage')
        self.assertEqual(summarize(records, source, self.workbook(disposition='Coverage'))['audit_issues'], [])

    def test_incorrect_action_plan_identifier_fails(self):
        source = modern_html()
        audit = summarize(self.report(source), source, self.workbook(action_id='OTHER-001'))
        self.assertIn('Action Plan and Recommendations action identifiers do not match.', audit['audit_issues'])

    def test_zero_action_placeholder_is_allowed_only_for_empty_register(self):
        source = modern_html(count=0)
        self.assertTrue(summarize(self.report(source), source, self.workbook(action=False))['valid'])
        audit = summarize(self.report(source), source, self.workbook(action=False, empty_text='Unexpected unregistered finding'))
        self.assertIn('Action Plan findings do not match the Recommendations action register.', audit['audit_issues'])

    def test_metadata_states_do_not_make_invalid_source_states_valid(self):
        source = modern_html()
        audit = summarize(self.report(source), source, self.workbook(source='incidents', state='offline'))
        self.assertTrue(any('invalid state' in issue for issue in audit['audit_issues']))
        migration = self.workbook(source='Methodology compatibility migration', state='migrated')
        self.assertTrue(summarize(self.report(source), source, migration)['valid'])

    def test_workbook_integrity_failure_is_still_fatal(self):
        source = modern_html()
        audit = summarize(self.report(source), source, self.workbook(integrity='Failed'))
        self.assertIn('Workbook integrity gate reports failure: Invented integrity failure', audit['audit_issues'])

    def test_incomplete_source_cannot_be_marked_available(self):
        source = modern_html()
        audit = summarize(self.report(source), source,
                          self.workbook(source='incidents', state='available', truncated=True))
        self.assertIn('Truncated source incidents is incorrectly marked available.', audit['audit_issues'])

    def test_checkpoint_pipeline_coverage_remains_valid_when_unavailable(self):
        source = modern_html()
        audit = summarize(self.report(source), source,
                          self.workbook(source='pipeline_entra', state='unavailable'))
        self.assertTrue(audit['valid'])

    def test_missing_html_evidence_link_and_wrong_headline_fail(self):
        source = modern_html(linked=False).replace('<dd>1 to', '<dd>2 to')
        audit = summarize(self.report(source), source, self.workbook())
        self.assertIn('A rendered HTML action has no linked evidence identifier.', audit['audit_issues'])
        self.assertIn('HTML action headline does not match rendered action cards.', audit['audit_issues'])
        self.assertIn('HTML and workbook action identifiers do not match.', audit['audit_issues'])

    def test_read_only_workbook_is_closed_after_audit(self):
        source = modern_html()
        path = self.workbook()
        book = load_workbook(path, read_only=True, data_only=True)
        with patch('openpyxl.load_workbook', return_value=book):
            summarize(self.report(source), source, path)
        self.assertIsNone(book._archive.fp)
        path.unlink()


if __name__ == '__main__':
    unittest.main()
