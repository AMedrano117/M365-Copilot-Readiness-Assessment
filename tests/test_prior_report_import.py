import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from Core.prior_report_import import load_prior_report


class PriorReportImportTests(unittest.TestCase):
    def workbook(self, path, *, manifest=True):
        workbook = Workbook()
        recommendations = workbook.active
        recommendations.title = "Recommendations"
        recommendations.append([
            "Recommendation ID", "Service", "Feature", "Recommendation", "Control ID",
            "Methodology Version", "Evidence Sheet", "Evidence Available", "Readiness Stage",
            "Link URL", "Observation", "Custom field",
        ])
        recommendations.append([
            "ENTRA-001", "Entra", "Example control", "Review example policy", "EXAMPLE-01",
            "2.0", "Conditional Access Detail", "Yes", "Foundation", "https://example.test",
            "Original saved observation", "preserved",
        ])
        detail = workbook.create_sheet("Conditional Access Detail")
        detail.append(["Policy Name", "State", "Observed"])
        detail.append(["Example policy", "enabled", datetime(2026, 9, 8, 12, 30)])
        coverage = workbook.create_sheet("Collection Coverage")
        coverage.append(["Source", "State", "Truncated", "Records", "Refresh Date"])
        coverage.append(["entra_conditional_access", "partial", True, 7, "2026-09-08"])
        controls = workbook.create_sheet("Control Results")
        controls.append(["Control ID", "Status"])
        controls.append(["EXAMPLE-01", "Fail"])
        if manifest:
            sheet = workbook.create_sheet("Run Manifest")
            sheet.append(["Item", "Value"])
            for row in (
                ["Tenant", "Example tenant"],
                ["Tenant ID", "11111111-1111-1111-1111-111111111111"],
                ["Generation Time (UTC)", "2026-09-08T12:30:00+00:00"],
                ["Assessment Version", "3.0"], ["Methodology Version", "2.0"],
            ):
                sheet.append(row)
        for title in ("Copilot User Usage", "Copilot Readiness Users"):
            sheet = workbook.create_sheet(title)
            sheet.append(["User Principal Name", "Last Activity Date"])
            sheet.append(["private-user@example.test", "2026-09-08"])
        aggregate = workbook.create_sheet("AI Adoption Usage")
        aggregate.append(["Metric", "Value"])
        aggregate.append(["Active users", 3])
        workbook.save(path)
        workbook.close()
        return path

    def json_file(self, path, payload):
        path.write_text(json.dumps(payload), encoding="utf-8")
        return path

    def snapshot(self):
        return {
            "tenant": "Example tenant", "generated_at": "2026-09-08T12:30:00+00:00",
            "methodology_version": "2.0", "assessment_version": "3.0",
            "recommendations": [{
                "RecommendationId": "M365-001", "Service": "M365", "Feature": "Example",
                "Recommendation": "Review example", "EvidenceSheet": "SharePoint Governance",
                "Disposition": "Action", "EvidenceAvailable": "Yes",
                "Nested original field": {"data": [1, 2]},
            }],
            "collection_coverage": {"m365_sharepoint": {
                "availability_status": "partial", "records_collected": 4,
                "truncated": True, "refresh_date": "2026-09-07",
            }},
            "control_results": [{"ControlId": "EXAMPLE-01", "Status": "Fail"}],
        }

    def test_workbook_preserves_provenance_and_original_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.workbook(Path(directory) / "renamed_report.xlsx")
            original = path.read_bytes()
            with patch("socket.create_connection", side_effect=AssertionError("network forbidden")):
                loaded = load_prior_report(path)
            self.assertTrue(loaded["available"])
            self.assertEqual(loaded["source_file"], "renamed_report.xlsx")
            self.assertEqual(loaded["tenant_name"], "Example tenant")
            self.assertEqual(loaded["tenant_id"], "11111111-1111-1111-1111-111111111111")
            self.assertEqual(loaded["generated_at"], "2026-09-08T12:30:00+00:00")
            self.assertEqual(loaded["methodology_version"], "2.0")
            self.assertEqual(loaded["assessment_version"], "3.0")
            row = loaded["recommendations"][0]
            self.assertEqual(row["RecommendationId"], "ENTRA-001")
            self.assertEqual(row["ControlId"], row["Control ID"])
            self.assertEqual(row["LinkUrl"], "https://example.test")
            self.assertEqual(row["Custom field"], "preserved")
            self.assertEqual(loaded["sheets"]["Conditional Access Detail"]["rows"][0]["Observed"], "2026-09-08T12:30:00")
            self.assertEqual(loaded["collection_coverage"][0]["State"], "partial")
            self.assertTrue(loaded["collection_coverage"][0]["Truncated"])
            self.assertEqual(path.read_bytes(), original)

    def test_restricted_copilot_tabs_require_new_opt_in(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.workbook(Path(directory) / "assessment.xlsx")
            private = load_prior_report(path)
            self.assertNotIn("private-user@example.test", json.dumps(private))
            self.assertIn("AI Adoption Usage", private["sheets"])
            opted_in = load_prior_report(path, include_user_details=True)
            self.assertIn("private-user@example.test", json.dumps(opted_in))
            self.assertIn("Copilot Readiness Users", opted_in["sheets"])
            self.assertFalse(opted_in["withheld_sheets"])

    def test_missing_manifest_never_uses_filename_or_modified_date(self):
        with tempfile.TemporaryDirectory() as directory:
            loaded = load_prior_report(self.workbook(
                Path(directory) / "m365_recommendations_inventedtenant_20260915_170000.xlsx",
                manifest=False,
            ))
            self.assertEqual(loaded["generated_at"], "")
            self.assertEqual(loaded["tenant_name"], "")
            self.assertEqual(loaded["tenant_id"], "")
            self.assertTrue(any("age is unknown" in warning for warning in loaded["warnings"]))
            self.assertTrue(any("tenant GUID" in warning for warning in loaded["warnings"]))

    def test_snapshot_retains_findings_but_flags_missing_underlying_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            payload = self.snapshot()
            loaded = load_prior_report(self.json_file(Path(directory) / "snapshot.json", payload))
            self.assertEqual(loaded["recommendations"], payload["recommendations"])
            self.assertEqual(loaded["tenant_name"], "Example tenant")
            self.assertEqual(loaded["collection_coverage"][0]["Source"], "m365_sharepoint")
            self.assertEqual(loaded["collection_coverage"][0]["State"], "partial")
            self.assertTrue(loaded["collection_coverage"][0]["Truncated"])
            self.assertEqual(loaded["control_results"], payload["control_results"])
            self.assertTrue(any("not the underlying" in warning for warning in loaded["warnings"]))
            self.assertTrue(any("SharePoint Governance" in warning for warning in loaded["warnings"]))

    def test_snapshot_embedded_tabs_and_coverage_rows_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            payload = self.snapshot()
            payload["collection_coverage"] = [{"Source": "Example", "State": "unavailable", "Reason": "Denied"}]
            payload["sheets"] = {
                "governance": {"title": "SharePoint Governance", "rows": [{"Property": "Saved", "Value": "Original"}]},
                "users": {"title": "Copilot Readiness Users", "rows": [{"User Principal Name": "private-user@example.test"}]},
            }
            loaded = load_prior_report(self.json_file(Path(directory) / "snapshot.json", payload))
            self.assertEqual(loaded["sheets"]["SharePoint Governance"]["rows"], payload["sheets"]["governance"]["rows"])
            self.assertNotIn("private-user@example.test", json.dumps(loaded))
            self.assertEqual(loaded["collection_coverage"], payload["collection_coverage"])
            self.assertFalse(any("absent" in warning for warning in loaded["warnings"]))

    def test_snapshot_nested_usage_records_require_opt_in(self):
        with tempfile.TemporaryDirectory() as directory:
            payload = self.snapshot()
            payload["collection_coverage"]["copilot_usage"] = {
                "availability_status": "available", "active_users": 3,
                "user_detail": [{"userPrincipalName": "private-user@example.test"}],
                "nested": {"user_details": [{"userPrincipalName": "private-user@example.test"}]},
            }
            path = self.json_file(Path(directory) / "snapshot.json", payload)
            loaded = load_prior_report(path)
            self.assertNotIn("private-user@example.test", json.dumps(loaded))
            self.assertEqual(loaded["collection_coverage"][1]["active_users"], 3)
            loaded = load_prior_report(path, include_user_details=True)
            self.assertIn("private-user@example.test", json.dumps(loaded))

    def test_snapshot_manifest_rows_supply_only_explicit_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            payload = {"recommendations": [], "run_manifest": {"rows": [
                {"Item": "Tenant", "Value": "Explicit tenant"},
                {"Item": "Generation Time (UTC)", "Value": "2026-09-08T12:30:00Z"},
                {"Item": "Methodology Version", "Value": "2.0"},
            ]}}
            loaded = load_prior_report(self.json_file(Path(directory) / "snapshot.json", payload))
            self.assertEqual(loaded["tenant_name"], "Explicit tenant")
            self.assertEqual(loaded["generated_at"], "2026-09-08T12:30:00Z")
            self.assertEqual(loaded["methodology_version"], "2.0")

    def test_purview_cache_and_raw_collection_have_clear_guidance(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "purview_data_json_example.json"
            self.json_file(path, {"recommendations": []})
            with self.assertRaisesRegex(ValueError, "Purview cache.*subset"):
                load_prior_report(path)
            path = Path(directory) / "collection.json"
            self.json_file(path, {"format": "m365-readiness-collection"})
            with self.assertRaisesRegex(ValueError, "collection-input"):
                load_prior_report(path)

    def test_rejects_invalid_json_shapes_and_unrelated_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            for index, payload in enumerate((
                [], {"recommendations": {}}, {"recommendations": ["not a row"]},
                {"recommendations": [{"unrelated": "data"}]},
                {"recommendations": [], "collection_coverage": "bad"},
                {"recommendations": [], "collection_coverage": {"source": []}},
                {"recommendations": [], "generated_at": {}},
            )):
                with self.subTest(payload=payload):
                    path = self.json_file(Path(directory) / f"bad{index}.json", payload)
                    with self.assertRaises(ValueError):
                        load_prior_report(path)
            workbook = Workbook()
            workbook.active.title = "Portal Export"
            workbook.active.append(["Site URL", "Anyone link count"])
            path = Path(directory) / "portal.xlsx"
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "Recommendations worksheet"):
                load_prior_report(path)

    def test_formulas_are_not_loaded_as_executable_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.workbook(Path(directory) / "assessment.xlsx")
            from openpyxl import load_workbook
            workbook = load_workbook(path)
            sheet = workbook.create_sheet("Formula Evidence")
            sheet.append(["Value"])
            sheet.append(['=HYPERLINK("https://example.test", "do not evaluate")'])
            workbook.save(path)
            workbook.close()
            loaded = load_prior_report(path)
            self.assertNotIn("HYPERLINK", json.dumps(loaded))

    def test_invalid_files_return_actionable_value_errors(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "broken.xlsx"
            path.write_text("not a workbook", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "could not be read as XLSX"):
                load_prior_report(path)
            path = Path(directory) / "broken.json"
            path.write_text("{broken", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "could not be read as JSON"):
                load_prior_report(path)
            with self.assertRaisesRegex(ValueError, "does not exist"):
                load_prior_report(Path(directory) / "missing.xlsx")


if __name__ == "__main__":
    unittest.main()
