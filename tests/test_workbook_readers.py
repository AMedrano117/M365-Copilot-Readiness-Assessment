"""Reader layout acceptance: exact lineage, complete records and unchanged decisions."""
from contextlib import redirect_stdout
from copy import deepcopy
import io
import json
from pathlib import Path
import tempfile
import unittest

from provenance_fixture import fictional_publication
from openpyxl import Workbook, load_workbook
from Core.workbook_readers import (CONFIGURATION_COLUMNS, ACTION_COLUMNS, TOPIC_COLUMNS, CONTROL_COLUMNS,
    configuration_rows, action_rows, coverage_rows, evidence_discrepancies, state_label, next_step)


def group(identifier='CFG-A', count=6, **changes):
    return {'Group ID': identifier, 'Setting': 'Sharing', 'Scope': 'site', 'Current': 'Anyone links',
        'Recommended': 'Existing guests or stricter', 'Status': 'Change', 'Total objects assessed': count,
        'Affected objects': count, 'Observed': None, **changes}


def member(identifier, number, **changes):
    return {'Group ID': identifier, 'Object ID': f'site-{number}', 'Object name': f'Site {number}',
        'Locator': f'https://example.invalid/sites/{number}', 'Configuration key': 'SharingCapability', **changes}


class ConfigurationReaderTests(unittest.TestCase):
    def test_small_groups_name_all_distinct_sites_and_large_groups_label_examples(self):
        details = [member('CFG-small', i) for i in range(5)] + [member('CFG-large', i) for i in range(6)]
        output = configuration_rows([group('CFG-small', 5), group('CFG-large', 6)], details)
        self.assertEqual(output[0]['Applies to / named locations'].splitlines(), [f'Site {i}' for i in range(5)])
        self.assertEqual(output[1]['Applies to / named locations'], '6 sites; examples only: Site 0; Site 1; Site 2')
        self.assertEqual(configuration_rows([group('CFG-large', 6)], list(reversed(details)))[0], output[1])

    def test_exact_group_join_retains_repeated_setting_names_and_distinct_counts(self):
        rows = [group('CFG-one', 3), group('CFG-two', 3, Current='Existing guests')]
        details = [member('CFG-one', 1), member('CFG-one', 1), member('CFG-one', 2), member('CFG-two', 3)]
        before = deepcopy((rows, details))
        output = configuration_rows(rows, details)
        self.assertEqual(len(output), 2)
        self.assertEqual(output[0]['Applies to / named locations'], 'Site 1\nSite 2')
        self.assertEqual(output[1]['Applies to / named locations'], 'Site 3')
        self.assertEqual(output[0]['Group size / objects assessed'], '2 objects in this configuration-value group / 3 distinct objects assessed')
        self.assertEqual((rows, details), before)

    def test_duplicate_names_use_locator_suffix_and_missing_names_use_retained_identity(self):
        details = [member('CFG-A', i, **{'Object name': 'Same name'}) for i in (1, 2)]
        details += [member('CFG-A', 3, **{'Object name': None}), member('CFG-A', 4, **{'Object name': None, 'Locator': None})]
        applies = configuration_rows([group()], details)[0]['Applies to / named locations']
        self.assertIn('Same name [https://example.invalid/sites/1]', applies)
        self.assertIn('Same name [https://example.invalid/sites/2]', applies)
        self.assertIn('https://example.invalid/sites/3', applies)
        self.assertIn('site-4', applies)

    def test_missing_group_identity_and_counts_stay_missing(self):
        for details in ([], [member('CFG-unrelated', 1)], [member('CFG-A', 1, **{'Object ID': None})]):
            output = configuration_rows([group()], details)[0]
            self.assertEqual(output['Applies to / named locations'], 'Site identities not available')
            self.assertIsNone(output['Group size / objects assessed'])
            self.assertIsNone(output['Observed'])

    def test_unknown_members_do_not_become_extra_distinct_sites(self):
        output = configuration_rows([group()], [member('CFG-A', 1), member('CFG-A', 2, **{'Object ID': None})])[0]
        self.assertIn('Additional site identities not available', output['Applies to / named locations'])
        self.assertTrue(output['Group size / objects assessed'].startswith('1 objects'))

    def test_tenant_labels_and_original_codes_and_stricter_qualifier_are_retained(self):
        for key, label in [('SharingCapability', 'Tenant-level SharePoint sharing setting'),
                ('OneDriveSharingCapability', 'Tenant-level OneDrive sharing setting'),
                ('DefaultSharingLinkType', 'Tenant-level default sharing link')]:
            row = group(Scope='tenant', Status='Meets')
            output = configuration_rows([row], [member('CFG-A', 1, **{'Configuration key': key})])[0]
            self.assertEqual(output['Applies to / named locations'], label)
            self.assertEqual(output['Assessment'], 'Meets recommendation')
            self.assertEqual(output['Status'], 'Meets')
            self.assertIn('or stricter', output['Recommended configuration'])


class ActionReaderTests(unittest.TestCase):
    def fixture(self):
        findings = [{'RecommendationId': 'A', 'PilotImpact': 'Confirm before pilot', 'Priority': 'Critical'},
            {'RecommendationId': 'B', 'PilotImpact': 'Blocks pilot', 'Priority': 'Low'},
            {'RecommendationId': 'C', 'PilotImpact': 'Planning input', 'Priority': 'High'}]
        actions = [dict(row, Recommendation='Review scoped results. Never revoke access based only on counts.',
            CustomerAction='Review results.', Status='Awaiting validation', OwnerRole='Security administrator') for row in findings]
        legacy = [dict(RecommendationId=row['RecommendationId'], Finding='Same finding name', Evidence="'Users & Identity'!A2:H2") for row in findings]
        return actions, findings, legacy

    def test_pilot_effect_joins_exact_identifier_and_orders_blockers_before_priority(self):
        actions, findings, legacy = self.fixture()
        before = deepcopy((actions, findings))
        rows, details = action_rows(actions, findings, legacy, {})
        self.assertEqual([row['RecommendationId'] for row in rows], ['B', 'A', 'C'])
        self.assertEqual([row['_detail_row'] for row in rows], [3, 2, 4])
        self.assertEqual([row['Priority'] for row in rows], ['Low', 'Critical', 'High'])
        self.assertEqual((actions, findings), before)
        self.assertEqual(details[1]['Original action record'], actions[1])

    def test_absent_or_ambiguous_relationship_remains_unresolved(self):
        actions, findings, legacy = self.fixture()
        for records in (findings[1:], findings + [findings[0]]):
            rows, details = action_rows(actions, records, legacy, {})
            row = next(row for row in rows if row['RecommendationId'] == 'A')
            self.assertIsNone(row['Pilot impact'])
            self.assertIn('Unresolved', details[0]['Finding relationship'])

    def test_status_role_dates_and_count_based_access_warning_are_preserved(self):
        actions, findings, legacy = self.fixture()
        rows, details = action_rows(actions, findings, legacy, {})
        for row in rows:
            self.assertEqual(row['Status'], 'Awaiting validation')
            self.assertEqual(row['Responsible role'], 'Security administrator')
            self.assertIsNone(row['Due date'])
            self.assertIn('Never revoke access based only on counts.', row['Next step'])
        self.assertEqual(details[0]['Full recommendation'], actions[0]['Recommendation'])

    def test_conditions_and_stricter_qualifiers_are_not_lost(self):
        row = {'Recommendation': 'Apply controls or stricter. If exceptions are approved, restrict the pilot scope first.',
            'CustomerAction': 'Apply controls.'}
        text = next_step(row, {})
        self.assertIn('or stricter', text)
        self.assertIn('If exceptions are approved, restrict the pilot scope first.', text)


class CoverageReaderTests(unittest.TestCase):
    def test_explicit_collections_preserve_overlapping_ids_and_blank_outcomes(self):
        result = {'domain_coverage': [{'Check ID': 'SAME', 'Check': 'Topic', 'State': 'collected'}],
            'control_results': [{'Control ID': 'SAME', 'Control': 'Control', 'Result': 'not_assessed',
                'Configuration result': 'not_established', 'Operational result': 'not_applicable'}]}
        topics, controls, details, _ = coverage_rows(result, {}, {})
        self.assertEqual(len(topics), 1)
        self.assertEqual(len(controls), 1)
        self.assertEqual(topics[0]['Configuration'], 'Not a control result')
        self.assertEqual(topics[0]['Operation'], 'Not a control result')
        self.assertEqual(topics[0]['Evidence'], 'Available')
        self.assertEqual(controls[0]['Overall result'], 'Not assessed')
        self.assertEqual(controls[0]['Configuration'], 'Not established')
        self.assertEqual([row['Row type'] for row in details], ['Assessment topic', 'Readiness control'])
        self.assertIsNone(topics[0]['Owner'])
        self.assertIn('Owner assignment required', topics[0]['Next evidence / interpretation'])

    def test_all_requested_state_translations_and_unknown_inputs(self):
        for raw in ('available', 'missing', 'incomplete', 'not_established', 'supported', 'not_applicable', 'not_assessed', 'unknown', 'fail', 'pass'):
            self.assertEqual(state_label(raw), raw.replace('_', ' ').capitalize())
        self.assertIsNone(state_label(None))
        self.assertEqual(state_label('custom state'), 'custom state')

    def test_specific_requests_and_retained_scope_identity_surface_without_inventory_join(self):
        row = {'Check ID': 'TOPIC', 'Check': 'Label configuration', 'State': 'missing',
            'Required sources': ['sensitivity_labels'], 'Missing evidence': 'sensitivity_labels; dated scope review',
            'Source metadata': [{'dataset': 'sensitivity_labels', 'scope': 'Pilot policy', 'object_name': 'Retained policy'}], 'Owner': 'Assigned owner'}
        topics, _, details, _ = coverage_rows({'domain_coverage': [row]}, {}, {})
        text = topics[0]['Next evidence / interpretation']
        self.assertIn('Sensitivity label definitions', text)
        self.assertIn('Pilot policy', text)
        self.assertIn('Retained policy', text)
        self.assertEqual(topics[0]['Owner'], 'Assigned owner')
        self.assertEqual(details[0]['Original record'], row)

    def test_apparent_discrepancies_keep_decisions_and_both_sources(self):
        for control, topic, dataset in [('LICENSE.ASSIGNMENT', 'SCOPE.LICENSES', 'licenses'),
                ('DATA.AUDIT', 'GOVERNANCE.AUDIT', 'audit_config'), ('DATA.LABELS', 'CLASSIFICATION.LABELS', 'sensitivity_labels')]:
            gap = {'RecommendationId': 'GAP', 'ControlId': control, 'Disposition': 'Coverage', 'EvidenceAvailable': 'No'}
            result = {'tenant_id': 'tenant-a', 'recommendations': [gap], 'domain_coverage': [
                {'Check ID': topic, 'Required sources': [dataset], 'Source metadata': [
                    {'dataset': dataset, 'availability_status': 'available', 'tenant_id': 'tenant-a'}]}]}
            before = deepcopy(result)
            discrepancy = evidence_discrepancies(result)['GAP']
            self.assertIn('Evidence reconciliation required', discrepancy['note'])
            self.assertEqual(discrepancy['gap_reference'], 'GAP')
            self.assertEqual(discrepancy['collected_reference'], topic)
            self.assertEqual(result, before)
            result['domain_coverage'][0]['Source metadata'][0]['tenant_id'] = 'other-tenant'
            self.assertFalse(evidence_discrepancies(result))


@fictional_publication
class WorkbookReaderIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from tests.test_configuration_outputs import shared_fixture
        from Core.export_recommendations import export_to_excel
        from Core.workbook_layout import technical_workbook_path
        cls.temp = tempfile.TemporaryDirectory()
        cls.bundle, cls.result = shared_fixture()
        cls.decisions = deepcopy({key: cls.result[key] for key in ('control_results', 'decision', 'counts')})
        cls.original_actions = deepcopy(cls.result['actions'])
        with redirect_stdout(io.StringIO()):
            cls.path = Path(export_to_excel(cls.result['recommendations'], filename='Assessment.xlsx', evidence_bundle=cls.bundle, output_dir=cls.temp.name))
        cls.assessment = load_workbook(cls.path)
        cls.technical = load_workbook(technical_workbook_path(cls.path))

    @classmethod
    def tearDownClass(cls):
        cls.assessment.close(); cls.technical.close(); cls.temp.cleanup()

    def test_primary_columns_and_hidden_technical_fields_and_no_freeze_or_example_tabs(self):
        for title, columns in [('Action Plan', ACTION_COLUMNS), ('Configuration', CONFIGURATION_COLUMNS)]:
            sheet = self.assessment[title]
            self.assertEqual(tuple(cell.value for cell in sheet[1][:8]), columns)
            for column in range(9, sheet.max_column + 1):
                self.assertTrue(sheet.column_dimensions[sheet.cell(1, column).column_letter].hidden)
        for book in (self.assessment, self.technical):
            for sheet in book:
                self.assertIsNone(sheet.freeze_panes)
                self.assertNotIn('Example', sheet.title)
        self.assertIn('operational effectiveness is not proven', self.assessment['Configuration']['E1'].comment.text)

    def test_coverage_has_two_filterable_sections_and_complete_counts(self):
        sheet = self.assessment['Coverage']
        self.assertEqual(len(sheet.tables), 2)
        from openpyxl.utils.cell import range_boundaries
        counts = []
        for table in sheet.tables.values():
            _, start, _, end = range_boundaries(table.ref)
            columns = TOPIC_COLUMNS if table.displayName == 'AssessmentTopics' else CONTROL_COLUMNS
            self.assertEqual(tuple(cell.value for cell in sheet[start][:7]), columns)
            counts.append(end - start)
        self.assertEqual(counts, [len(self.result['domain_coverage']), len(self.result['control_results'])])
        self.assertIn('Available means evidence collected', ' '.join(str(row[0].value) for row in sheet))

    def test_action_source_references_remain_exact_after_sorting_and_live_rule_exists(self):
        from tests.test_workbook_layout import records
        from Core.workbook_layout import RANGE
        actions = records(self.assessment['Action Plan'])
        self.assertEqual(len(actions), len(self.original_actions))
        headers = {cell.value: cell.column for cell in self.assessment['Action Plan'][1]}
        detail = records(self.technical['Action Detail'])
        for number, row in enumerate(actions, 2):
            target = self.assessment['Action Plan'].cell(number, headers['Source record']).hyperlink.target.split('#')[1]
            source = detail[int(RANGE.fullmatch(target)[3]) - 2]
            self.assertEqual(row['RecommendationId'], source['RecommendationId'])
            matching = [rec for rec in self.result['recommendations'] if rec['RecommendationId'] == row['RecommendationId']]
            self.assertEqual(row['Pilot impact'], matching[0].get('PilotImpact') or None)
        rules = [rule for rules in self.assessment['Action Plan'].conditional_formatting._cf_rules.values() for rule in rules]
        self.assertTrue(any(rule.type == 'expression' and '$B2' in str(rule.formula) for rule in rules))

    def test_all_configuration_members_and_original_decisions_remain(self):
        from tests.test_workbook_layout import records
        detail = records(self.technical['Configuration Detail'])
        self.assertEqual(len(detail), len(self.result['configuration_groups']['members']))
        self.assertEqual(len(records(self.assessment['Configuration'])), len(self.result['configuration_groups']['groups']))
        self.assertEqual({key: self.result[key] for key in self.decisions}, self.decisions)
        for row in detail:
            for column in ('Group ID', 'Setting', 'Scope', 'Object name', 'Locator', 'Object ID', 'Current value', 'Source reference'):
                self.assertIn(column, row)

    def test_explicit_freeze_configuration_and_actual_formulas_are_preserved(self):
        from Core.workbook_navigation import apply_workbook_navigation
        workbook = Workbook()
        self.addCleanup(workbook.close)
        sheet = workbook.active; sheet.title = 'Calculation'
        sheet.append(['Input', 'Derived'])
        sheet.append([2, '=A2*2'])
        apply_workbook_navigation(workbook, {}, {'workbook_settings': {'freeze_panes': 'A2'}}, role='technical')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet['B2'].data_type, 'f')
        self.assertEqual(sheet['B2'].value, '=A2*2')
