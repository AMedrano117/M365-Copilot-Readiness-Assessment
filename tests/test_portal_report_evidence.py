"""Portal evidence provenance, offline presentation, and restricted workbook checks."""

import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from Core.copilot_readiness_import import FLAG_COLUMNS
from Core.evidence_layer import _build_data_exposure_sheet, build_evidence_bundle
from Core.export_recommendations import export_to_excel, export_to_html
from Core.new_recommendation import new_recommendation


class PortalReportEvidenceTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.original_cwd = os.getcwd()
        os.chdir(self.directory.name)

    def tearDown(self):
        os.chdir(self.original_cwd)
        self.directory.cleanup()

    def readiness(self, include_users=False):
        result = {
            "available": True, "status": "available", "source_file": "readiness.csv",
            "report_date": "2026-09-12", "report_period": "30", "total_rows": 3,
            "summary": "Readiness flags describe exported rows only.", "warnings": [], "error": "",
            "metrics": {key: {"true": 1, "false": 1, "unknown": 1} for key in FLAG_COLUMNS},
        }
        if include_users:
            result["user_details"] = [{
                "user_principal_name": "restricted.person@example.test", "report_date": "2026-09-12",
                "report_period": "30", **{key: None for key in FLAG_COLUMNS},
            }]
        return result

    def exposure(self):
        return {
            "sources": {"sam": {
                "files_loaded": 1, "records_read": 2, "signals": {}, "freshness": "unknown",
                "max_age_days": 35, "latest_report_date": "2026-09-12",
                "reports": [{
                    "source_file": "sites.csv", "source_sheet": "", "report_type": "inactive_sites",
                    "workload": "SharePoint", "report_date": "2026-09-12", "records_read": 2,
                    "status": "selected", "domains": ["lifecycle"], "age_days": 3,
                    "freshness": "fresh", "tenant_id": "tenant-placeholder",
                }, {
                    "source_file": "older-sites.csv", "report_type": "inactive_sites",
                    "report_date": "2026-08-01", "records_read": 1, "status": "superseded",
                    "domains": ["lifecycle"], "freshness": "stale",
                }],
            }},
            "evidence_rows": [{
                "Source File": "sharing.csv", "Report Date": "2026-09-11",
                "Risk Signals": "Anyone links", "Freshness": "fresh",
            }],
            "lifecycle_summary": {
                "available": True, "site_count": 2, "ownerless_site_count": 1,
                "inactive_site_count": 1, "unknown_owner_status_count": 0,
                "missing_owner_contact_count": 1, "freshness": "fresh",
                "latest_report_date": "2026-09-12", "oldest_report_date": "2026-09-12",
            },
            "lifecycle_evidence_rows": [{
                "Source File": "sites.csv", "Report Date": "2026-09-12",
                "Site Name": "Restricted Internal Site", "Owner": "private.owner@example.test",
                "Is Ownerless": "No", "Is Inactive": "Yes", "Last Activity Date": "2025-01-01",
            }],
        }

    def bundle(self, readiness=None, exposure=None):
        client = SimpleNamespace(copilot_readiness_export=readiness or {})
        recommendation = new_recommendation(
            "M365", "Review assessment scope", "Review the supplied readiness evidence.",
            recommendation="Confirm the pilot scope with the tenant owner.", priority="Medium",
        )
        bundle = build_evidence_bundle([recommendation], ({"_client": client}, []), {}, {}, {}, {}, {}, exposure)
        bundle["evaluation_date"] = "2026-09-15"
        return bundle

    def html(self, bundle):
        return Path(export_to_html(bundle["recommendations"], filename="test.html", evidence_bundle=bundle)).read_text(encoding="utf-8")

    def test_readiness_summary_is_scoped_and_identity_free_with_user_detail_enabled(self):
        bundle = self.bundle(self.readiness(include_users=True))
        self.assertNotIn("user_details", bundle["copilot_readiness_export"])
        self.assertIn("copilot_readiness_user_detail", bundle["sheets"])
        self.assertNotIn("copilot_readiness_user_detail", [section["key"] for section in bundle["appendix_sections"]])
        html = self.html(bundle)
        section = html.split('id="copilot-readiness-export"', 1)[1].split("</section>", 1)[0]
        self.assertIn("3 exported rows", section)
        self.assertIn("2026-09-12", section)
        self.assertIn("30 days", section)
        self.assertIn("do not establish tenant-wide license totals", section)
        self.assertIn("do not measure actual Copilot usage", section)
        self.assertIn("Uses eligible update channel", section)
        self.assertIn("<td>Uses Teams chat</td><td>1</td><td>1</td><td>1</td>", section)
        self.assertNotIn("restricted.person@example.test", html)
        self.assertNotIn('<script src=', html)
        self.assertNotIn('<link rel="stylesheet"', html)

    def test_user_readiness_sheet_requires_explicit_details(self):
        bundle = self.bundle(self.readiness())
        self.assertIn("copilot_readiness_detail", bundle["sheets"])
        self.assertNotIn("copilot_readiness_user_detail", bundle["sheets"])

    def test_workbook_preserves_unknown_flags_and_restricted_details(self):
        bundle = self.bundle(self.readiness(include_users=True), self.exposure())
        path = export_to_excel(bundle["recommendations"], filename="test.xlsx", evidence_bundle=bundle)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        self.assertIn("Copilot Readiness Export", workbook.sheetnames)
        self.assertIn("Copilot Readiness Users", workbook.sheetnames)
        self.assertIn("SharePoint Lifecycle Detail", workbook.sheetnames)
        users = list(workbook["Copilot Readiness Users"].values)
        self.assertIn("restricted.person@example.test", users[1])
        self.assertIn("Unknown", users[1])
        lifecycle = list(workbook["SharePoint Lifecycle Detail"].values)
        self.assertIn("private.owner@example.test", lifecycle[1])
        self.assertIn("2026-09-12", lifecycle[1])
        source_rows = list(workbook["Data Exposure Detail"].values)
        self.assertIn("Source Status", source_rows[0])
        self.assertIn("Report Type", source_rows[0])
        status_index = source_rows[0].index("Source Status")
        self.assertIn("superseded", [row[status_index] for row in source_rows[1:]])

    def test_imported_workbook_strings_remain_text_instead_of_formulas(self):
        readiness = self.readiness(include_users=True)
        readiness["user_details"][0]["user_principal_name"] = '=HYPERLINK("https://example.test","contact")'
        bundle = self.bundle(readiness)
        path = export_to_excel(bundle["recommendations"], filename="literal.xlsx", evidence_bundle=bundle)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        cell = workbook["Copilot Readiness Users"]["A2"]
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, readiness["user_details"][0]["user_principal_name"])

    def test_lifecycle_html_is_aggregate_and_does_not_claim_oversharing(self):
        html = self.html(self.bundle(exposure=self.exposure()))
        section = html.split('id="sharepoint-lifecycle-summary"', 1)[1].split("</section>", 1)[0]
        self.assertIn("Reported ownerless", section)
        self.assertIn("Reported inactive", section)
        self.assertIn("do not establish oversharing", section)
        self.assertIn("missing owner contact does not establish", section)
        self.assertNotIn("private.owner@example.test", html)
        self.assertNotIn("Restricted Internal Site", html)

    def test_source_metadata_preserves_per_report_and_risk_row_dates(self):
        sheet = _build_data_exposure_sheet(self.exposure())
        reports = [row for row in sheet["rows"] if row["Detail Type"] == "Source Report"]
        self.assertEqual([row["Source Status"] for row in reports], ["selected", "superseded"])
        self.assertEqual(reports[0]["Report Date"], "2026-09-12")
        self.assertEqual(reports[0]["Coverage Domains"], "lifecycle")
        evidence = next(row for row in sheet["rows"] if row["Detail Type"] == "Risk Evidence")
        self.assertEqual(evidence["Report Date"], "2026-09-11")
        self.assertEqual(evidence["Freshness"], "fresh")
        html = self.html(self.bundle(exposure=self.exposure()))
        self.assertIn("Imported source reports and dates", html)
        self.assertIn("older-sites.csv", html)
        self.assertIn("superseded", html)

    def test_offline_provenance_follows_decision_and_preserves_original_collection(self):
        bundle = self.bundle(self.readiness())
        bundle["collection_context"] = {
            "mode": "offline", "source_file": "snapshot.json", "collected_at": "2026-06-01T12:00:00Z",
            "age_days": 106, "freshness": "Stale", "scope": "Saved tenant evidence plus portal exports",
        }
        html = self.html(bundle)
        executive = html.split('id="executive"', 1)[1].split("</section>", 1)[0]
        self.assertNotIn("Built offline", executive)
        self.assertNotIn("snapshot.json", executive)
        self.assertIn("Original tenant collection: 2026-06-01T12:00:00Z", html)
        self.assertIn("Report execution: offline", html)
        self.assertLess(html.index('id="executive"'), html.index('id="engineer-appendix"'))
        self.assertEqual(bundle["assessment_result"]["decision"], "Readiness unconfirmed")
        self.assertTrue(bundle["assessment_result"]["decision_coverage"])
        self.assertNotIn("Current report assessed", html)
        path = export_to_excel(bundle["recommendations"], filename="offline.xlsx", evidence_bundle=bundle)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        rows = list(workbook["Collection Coverage"].values)
        self.assertIn("2026-06-01T12:00:00Z", rows[1])
        self.assertIn("Stale", rows[1])

    def test_portal_only_domains_mark_tenant_controls_unassessed(self):
        bundle = self.bundle(self.readiness())
        bundle["collection_context"] = {"mode": "offline", "scope": "Portal exports only"}
        html = self.html(bundle)
        result = bundle["assessment_result"]
        self.assertEqual(result["decision"], "Readiness unconfirmed")
        missing = {row["control_id"] for row in result["controls"] if row["status"] == "Not established"}
        self.assertIn("IDENTITY.AUTH", missing)
        self.assertIn("DATA.DLP", missing)
        self.assertIn("THREAT.INCIDENTS", missing)
        self.assertIn("Confirm sign-in policy coverage for the pilot", html)
        self.assertIn("Confirm data loss prevention coverage and enforcement", html)
        self.assertIn("The supplied evidence does not establish this check", html)

    def test_invalid_readiness_export_shows_validation_without_zero_counts(self):
        readiness = self.readiness()
        readiness.update(available=False, metrics={}, status="invalid_headers", error="Expected readiness column headers.")
        html = self.html(self.bundle(readiness))
        section = html.split('id="copilot-readiness-export"', 1)[1].split("</section>", 1)[0]
        self.assertIn("could not be assessed", section)
        self.assertIn("Expected readiness column headers", section)
        self.assertNotIn("<td>0</td>", section)

    def test_legacy_restricted_appendices_are_filtered_before_rendering(self):
        bundle = self.bundle(self.readiness())
        bundle["appendix_sections"].append({
            "key": "copilot_user_usage_detail", "title": "secret.person@example.test",
            "summary": "secret.person@example.test", "preview_rows": [],
        })
        self.assertNotIn("secret.person@example.test", self.html(bundle))


if __name__ == "__main__":
    unittest.main()
