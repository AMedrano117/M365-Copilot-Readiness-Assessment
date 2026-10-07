"""PDF context is visible and traceable without becoming scored measurements."""

from copy import deepcopy
from html.parser import HTMLParser
import json
from pathlib import Path
import tempfile
import unittest

from Core.portal_insights import featured_highlights, report_highlights, workbook_highlights
from Core.portal_review import load_portal_review
from tests.test_portal_review import TENANT, reviewed_fixture


TEXT = {
    'Health': ["Last month, 17% of Copilot\nusers' devices were not on a\nsupported update channel\n"
               "Last month, 42% of Copilot\nusers' devices were multiple\nversions behind\n"
               'Last month, connected\nexperiences were enabled for\n97% of Copilot users\n'
               'Last month, OneDrive was\nenabled for 98% of Copilot\nusers'],
    'Overview': ['Copilot adoption score\n45/100\nNo feedback from users yet\nCopilot Chat not pinned\n'
                 'License requests pending\nCopilot responses are limited',
                 "29%\nof licensed Copilot users haven't tried Copilot yet"],
    'Security': ['Policy is monitoring only — not enforced\nPrompts with sensitive info types\n2.5K',
                 'File referenced with labeling protection\n4/90\nSharepoint sites referenced\n6\n'
                 'Improve permissions\nFiles referenced\n90\nCompleted by your organization\n5/20'],
    'Usage': ['Enabled users\n4\nActive users\n2\nActive users\n19\nAvg daily active users\n3',
              'Agent usage\nActive agents\n3\nActive users\n2\nCopilot search usage\nLoading data...\nNo credit usage available'],
}


def fixture(directory):
    path, manifest = reviewed_fixture(directory)
    template = manifest['captures'][0]
    manifest['captures'] = []
    for title, pages in TEXT.items():
        capture = deepcopy(template)
        capture.update(id='PDF_' + title + '.1', title='Fictional ' + title,
                       domain_id='licensing' if title == 'Health' else 'data_protection' if title == 'Security' else 'adoption',
                       review_method='automated', captured_at='2026-09-10', report_date='2026-09-09',
                       extracted_pages=[{'page': index, 'method': 'Windows OCR', 'text': text}
                                        for index, text in enumerate(pages, 1)],
                       previews=template['previews'] * len(pages))
        manifest['captures'].append(capture)
    path.write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
    return load_portal_review(path, TENANT, '2026-09-15')


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.ids, self.anchors = set(), []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if 'id' in attrs:
            self.ids.add(attrs['id'])
        if tag == 'a' and attrs.get('href', '').startswith('#'):
            self.anchors.append(attrs['href'][1:])


class PortalInsightsTests(unittest.TestCase):
    def test_excerpts_are_exact_page_text_with_stable_provenance_and_separate_dates(self):
        with tempfile.TemporaryDirectory() as directory:
            review = fixture(directory)
            original = deepcopy(review)
            rows = report_highlights(review)
            self.assertEqual(len(rows), 20)
            self.assertEqual(len({row['highlight_id'] for row in rows}), 20)
            captures = {capture['id']: capture for capture in review['captures']}
            for row in rows:
                capture = captures[row['capture_id']]
                text = capture['extracted_pages'][row['page'] - 1]['text']
                self.assertIn(row['excerpt'], text)
                self.assertEqual(row['captured_at'], '2026-09-10')
                self.assertEqual(row['report_date'], '2026-09-09')
                self.assertEqual(row['evidence_kind'], 'portal_context')
                self.assertEqual(row['source_sha256'], capture['source_sha256'])
                self.assertNotIn('value', row)
                self.assertNotIn('data_uri', json.dumps(row))
            self.assertEqual(report_highlights(review), rows)
            self.assertEqual(review, original)
            featured = featured_highlights(rows)
            self.assertEqual(len(featured), 8)
            self.assertIn('Label protection on referenced files', [row['topic'] for row in featured])

    def test_unmatched_layout_and_missing_values_do_not_invent_card_measurements(self):
        review = {'captures': [{'id': 'unknown', 'title': 'Unknown layout', 'summary': 'Review the captured layout.',
                               'extracted_pages': [{'page': 1, 'method': 'Windows OCR',
                                                    'text': 'Active users\nEnabled users\n19\n7\nLoading data...'}]}]}
        rows = report_highlights(review)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['excerpt'], 'Review the captured layout.')
        self.assertIsNone(rows[0]['page'])
        self.assertIsNone(rows[0]['captured_at'])
        self.assertIsNone(rows[0]['report_date'])
        self.assertNotIn('value', rows[0])

    def test_repeated_portal_statements_are_not_counted_twice(self):
        review = {'captures': [{'id': 'repeated', 'title': 'Repeated', 'extracted_pages': [
            {'page': index, 'method': 'PDF text', 'text': 'No feedback from users yet'} for index in (1, 2)]}]}
        rows = report_highlights(review)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['page'], 1)

    def test_html_and_summary_surface_context_without_changing_the_assessment(self):
        from Core.assessment_result import build_assessment_result
        from Core.customer_report import render_customer_report
        from Core.pilot_summary import render_pilot_summary
        with tempfile.TemporaryDirectory() as directory:
            review = fixture(directory)
            bundle = {'portal_review': review}
            result = build_assessment_result([], bundle, expected_tenant_id=TENANT, evaluation_date='2026-09-15')
            before = deepcopy(result)
            html = render_customer_report(result, bundle, 'Fictional customer')
            self.assertLess(html.index('id="pdf-report-highlights"'), html.index('id="readiness-domains"'))
            for row in report_highlights(review):
                self.assertIn(row['topic'], html)
            parser = Links()
            parser.feed(html)
            self.assertFalse(set(parser.anchors) - parser.ids)
            self.assertIn('portal-pdf-security-1-page-2', parser.ids)
            summary = render_pilot_summary(result, bundle, 'Fictional customer', 'Report.html')
            self.assertIn('id="pdf-summary"', summary)
            for row in featured_highlights(report_highlights(review)):
                self.assertIn(row['topic'], summary)
            self.assertIn('Report.html#portal-pdf-security-1', summary)
            self.assertEqual(result, before)

    def test_reviewed_notes_remain_escaped_and_unscored(self):
        from Core.assessment_result import build_assessment_result
        from Core.customer_report import render_customer_report
        review = {'captures': [{'id': 'manual', 'title': 'Manual', 'summary': 'Manual context',
                               'review_notes': ['=SUM(1,2) <script>alert(1)</script>'], 'previews': []}]}
        bundle = {'portal_review': review}
        result = build_assessment_result([], bundle, evaluation_date='2026-09-15')
        html = render_customer_report(result, bundle, 'Fictional customer')
        self.assertIn('=SUM(1,2) &lt;script&gt;alert(1)&lt;/script&gt;', html)
        self.assertNotIn('<script>alert(1)</script>', html)
        self.assertIsNone(workbook_highlights(review)[0]['Page'])

    def test_both_workbooks_keep_readable_excerpts_and_links_to_exact_source_pages(self):
        from Core.export_recommendations import export_to_excel
        from openpyxl import load_workbook
        from tests.test_workbook_layout import fixture as workbook_fixture
        from Core.workbook_layout import technical_workbook_path
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            recommendations, bundle = workbook_fixture()
            bundle['portal_review'] = fixture(root / 'review')
            path = Path(export_to_excel(recommendations, filename=str(root / 'Assessment.xlsx'), evidence_bundle=bundle))
            books = [load_workbook(path), load_workbook(technical_workbook_path(path))]
            try:
                for book in books:
                    self.assertEqual(book['PDF Highlights'].sheet_state, 'visible')
                    sheet = book['PDF Highlights']
                    self.assertEqual(sheet.max_row - 1, 20)
                    header = [cell.value for cell in sheet[1]]
                    columns = {value: index + 1 for index, value in enumerate(header)}
                    for number, expected in enumerate(workbook_highlights(bundle['portal_review']), 2):
                        self.assertEqual(sheet.cell(number, columns['Source excerpt']).value, expected['Source excerpt'])
                        cell = sheet.cell(number, columns['Source detail'])
                        target = cell.hyperlink.target
                        self.assertIn("#'PDF Extracted Text'!A", target)
                        row = int(target.rsplit('A', 1)[-1])
                        source = list(books[1]['PDF Extracted Text'].values)
                        record = dict(zip(source[0], source[row - 1]))
                        self.assertEqual(record['Capture ID'], expected['Capture ID'])
                        self.assertEqual(record['Page'], expected['Page'])
                    self.assertTrue(sheet.tables)
                visible = [sheet.title for sheet in books[0] if sheet.sheet_state == 'visible']
                self.assertEqual(visible[4], 'PDF Highlights')
            finally:
                for book in books:
                    book.close()

    def test_json_package_and_app_builder_preserve_all_pdf_context_as_context(self):
        from Core.assessment_result import build_assessment_result
        from Core.dashboard_export import build_dashboard_export
        from Core.dashboard_package import read_dashboard_package, write_dashboard_package
        from Core.finding_evidence import build_finding_evidence
        from Core.app_builder_export import write_app_builder_export
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bundle = {'portal_review': fixture(root / 'review')}
            result = build_assessment_result([], bundle, expected_tenant_id=TENANT, evaluation_date='2026-09-15')
            payload = build_dashboard_export(result, bundle, tenant_name='Fictional customer')
            rows = report_highlights(bundle['portal_review'])
            self.assertEqual(payload['portal_report_highlights'], rows)
            path = write_dashboard_package(root / 'JSON', payload)
            restored = read_dashboard_package(path)
            self.assertEqual(restored['portal_report_highlights'], rows)
            index = json.loads(Path(path).read_text(encoding='utf-8'))
            self.assertEqual(index['entry_points']['portal_reports'], 'portal/index.json')
            summary = json.loads((root / 'JSON/summary.json').read_text(encoding='utf-8'))
            self.assertEqual(summary['portal_reports']['highlight_count'], 20)
            self.assertEqual(len(summary['portal_reports']['highlights']), 8)
            model = build_finding_evidence(payload)
            output = write_app_builder_export(model, root / 'App Builder')
            overview = json.loads((root / 'App Builder/01-overview.json').read_text(encoding='utf-8'))
            files = overview['portal_report_context']['files']
            retained = []
            for filename in files:
                retained.extend(json.loads((root / 'App Builder' / filename).read_text(encoding='utf-8'))['highlights'])
            self.assertEqual(retained, rows)
            self.assertEqual(overview['assessment_counts'], result['counts'])
            for upload in output['upload_sets'].values():
                self.assertTrue(set(files) <= set(upload['supporting_context']))

    def test_reviewed_note_is_inert_in_excel_and_links_to_capture_metadata(self):
        from Core.export_recommendations import export_to_excel
        from Core.workbook_layout import technical_workbook_path
        from openpyxl import load_workbook
        from tests.test_workbook_layout import fixture as workbook_fixture
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            manifest_path, manifest = reviewed_fixture(root / 'review')
            manifest['captures'][0]['review_notes'] = ['=SUM(1,2)']
            manifest_path.write_text(json.dumps(manifest), encoding='utf-8')
            recommendations, bundle = workbook_fixture()
            bundle['portal_review'] = load_portal_review(manifest_path, TENANT, '2026-09-15')
            path = export_to_excel(recommendations, filename=str(root / 'Assessment.xlsx'), evidence_bundle=bundle)
            for filename in (path, technical_workbook_path(path)):
                book = load_workbook(filename)
                try:
                    sheet = book['PDF Highlights']
                    columns = {cell.value: cell.column for cell in sheet[1]}
                    cell = sheet.cell(2, columns['Source excerpt'])
                    self.assertEqual(cell.value, '=SUM(1,2)')
                    self.assertEqual(cell.data_type, 's')
                    self.assertIn("#'Portal Review'!A", sheet.cell(2, columns['Source detail']).hyperlink.target)
                    self.assertIsNone(sheet.cell(2, columns['Page']).value)
                finally:
                    book.close()


if __name__ == '__main__':
    unittest.main()
