"""Open the two real workbook files for tests that inspect both export layers."""

from openpyxl import load_workbook
from Core.workbook_layout import technical_workbook_path


def load_workbook_pair(path, *args, **kwargs):
    assessment = load_workbook(path, *args, **kwargs)
    technical_path = technical_workbook_path(path)
    assessment.technical_path = technical_path
    if technical_path.is_file():
        assessment.technical = load_workbook(technical_path, *args, **kwargs)
        close = assessment.close

        def close_pair():
            close()
            assessment.technical.close()

        assessment.close = close_pair
    else:
        assessment.technical = assessment
    return assessment
