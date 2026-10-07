import unittest

from openpyxl import Workbook

from tools.validate_offline_report import _investigation_issues


class InvestigationWorkbookAuditTests(unittest.TestCase):
    def workbook(self, location="'Novel Evidence'!A2:B2", count=1, qualification='One retained record supports the action.'):
        workbook = Workbook()
        workbook.active.title = 'Recommendations'
        workbook.active.append(['RecommendationId', 'Recommendation', 'Investigation Status', 'Supporting Records',
                                'Evidence Sheet', 'Investigation Qualification'])
        workbook.active.append(['NEW-001', 'Review the retained record.', 'Records available', count, location, qualification])
        for title, field in [('Evidence Index', 'Workbook Tab'), ('Action Plan', 'Evidence')]:
            sheet = workbook.create_sheet(title)
            sheet.append(['RecommendationId', field])
            sheet.append(['NEW-001', location])
            if location:
                sheet.cell(2, 2).hyperlink = '#' + location
        if location:
            workbook['Recommendations'].cell(2, 5).hyperlink = '#' + location
        detail = workbook.create_sheet('Novel Evidence')
        detail.append(['ID', 'State'])
        detail.append(['record-id', 'Suspended'])
        self.addCleanup(workbook.close)
        return workbook

    def test_generic_valid_ranges_and_zero_record_configuration_context_are_accepted(self):
        self.assertEqual(_investigation_issues(self.workbook()), [])
        self.assertEqual(_investigation_issues(self.workbook(count=0)), [])
        self.assertEqual(_investigation_issues(self.workbook(location='', count=0,
                        qualification='The source returned HTTP 403; no supporting records could be read.')), [])

    def test_missing_status_or_reason_and_generic_or_out_of_bounds_locations_are_rejected(self):
        workbook = self.workbook(location='', count=0, qualification='')
        workbook['Recommendations'].cell(2, 3).value = ''
        self.assertEqual(len(_investigation_issues(workbook)), 2)
        self.assertTrue(_investigation_issues(self.workbook(location='Evidence Index')))
        self.assertTrue(_investigation_issues(self.workbook(location="'Novel Evidence'!A2:B7")))

    def test_action_link_must_open_the_actual_supporting_records(self):
        workbook = self.workbook()
        workbook['Action Plan'].cell(2, 2).value = 'Evidence Index'
        workbook['Action Plan'].cell(2, 2).hyperlink = "#'Evidence Index'!A2"
        self.assertIn('Action Plan does not link', ' '.join(_investigation_issues(workbook)))


if __name__ == '__main__':
    unittest.main()
