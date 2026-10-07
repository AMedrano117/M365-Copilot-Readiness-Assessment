"""Synthetic acceptance checks for the assessment/technical workbook pair."""

from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import unquote
import re

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from Core.export_recommendations import export_to_excel, export_to_html
from Core.workbook_layout import ASSESSMENT_ORDER, FINDING_COLUMNS, WorkbookLayoutError, technical_workbook_path


def records(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def fixture():
    source = {'complete': True, 'available': True, 'availability_status': 'available',
              'collected_at': '2026-10-05T12:00:00Z', 'refresh_date': '2026-10-05'}
    data = {
        'sharepoint_tenant_settings': [{'SharingCapability': 2, 'OneDriveSharingCapability': 2,
                                       'DefaultSharingLinkType': 3, 'RequireAnonymousLinksExpireInDays': -1,
                                       'FileAnonymousLinkType': 2, 'FolderAnonymousLinkType': 2}],
        'windows_protection': [{'id': 'protection-001', 'ParentId': 'device-001', 'realTimeProtectionEnabled': False,
                                'isTamperProtected': False, 'signatureUpdateOverdue': False,
                                'lastReportedDateTime': '2026-10-05T12:00:00Z'}],
        'managed_devices': [{'id': 'device-001', 'deviceName': 'LAPTOP-0142', 'userPrincipalName': 'j.smith@example.invalid',
                             'operatingSystem': 'Windows 11', 'osVersion': '23H2', 'complianceState': 'compliant'}],
        'compliance_setting_states': [{'id': 'setting-001', 'ParentId': 'device-001', 'settingName': 'Minimum OS version',
                                       'state': 'noncompliant'}, {'id': 'setting-002', 'ParentId': 'device-001',
                                                                'settingName': 'Require encryption', 'state': 'noncompliant'}],
    }
    bundle = {'evaluation_date': '2026-10-06', 'collection_context': {'collected_at': source['collected_at']},
              'assessment_sources': {name: [{'source': dict(source, dataset=name), 'records': rows}]
                                     for name, rows in data.items()}, 'sheets': {}}
    common = {'Service': 'Entra', 'DomainId': 'endpoints', 'Disposition': 'Action', 'Status': 'Action Required',
              'Priority': 'High', 'ControlId': 'ENDPOINT.POSTURE', 'EvidenceAvailable': 'Yes', 'EvidenceBasis': 'Tenant evidence',
              'ObservationDate': '2026-10-05', 'TenantId': '11111111-1111-4111-8111-111111111111',
              'Observation': 'Synthetic condition requires a change.', 'Recommendation': 'Apply and verify the policy.'}
    rows = [dict(common, RecommendationId='WIN-1', FindingKey='expanded.windows_protection.disabled', Feature='Restore real-time protection',
                 EvidenceKey='entra_device_detail', InvestigationEvidence={'records': data['windows_protection'],
                 'source': dict(source, dataset='windows_protection')}),
            dict(common, RecommendationId='WIN-2', FindingKey='expanded.windows_protection.disabled.second', Feature='Review endpoint protection',
                 EvidenceKey='entra_device_detail', InvestigationEvidence={'records': data['windows_protection'],
                 'source': dict(source, dataset='windows_protection')}),
            dict(common, RecommendationId='SETTING-1', FindingKey='expanded.compliance_setting_states.noncompliant',
                 Feature='Resolve failed device compliance settings', EvidenceKey='entra_device_detail',
                 InvestigationEvidence={'records': data['compliance_setting_states'],
                                       'source': dict(source, dataset='compliance_setting_states')}),
            dict(common, RecommendationId='M365-116', Service='Microsoft 365', DomainId='content', ControlId='CONTENT.SHARING',
                 FindingKey='sharepoint.sharing.permissive_anonymous_defaults', Feature='Tighten sharing defaults',
                 EvidenceKey='sharepoint_governance_detail', InvestigationEvidence={'kind': 'configuration',
                 'records': data['sharepoint_tenant_settings'], 'source': dict(source, dataset='sharepoint_tenant_settings')})]
    # Generic producers exercise all five entity projections without invoking
    # live collectors or reading any customer exports.
    for identifier, key, record in [
        ('USER-1', 'mfa_registration_detail', {'id': '000000000000000001', 'userPrincipalName': 'user@example.invalid',
                                            'displayName': 'Fictional user', 'isMfaRegistered': False}),
        ('APP-1', 'app_access_detail', {'id': 'app-001', 'appId': 'app-client-001', 'displayName': 'Fictional app',
                                     'consentType': 'AllPrincipals', 'scope': 'Mail.Read'}),
        ('SITE-1', 'data_exposure_detail', {'id': 'site-001', 'siteUrl': 'https://example.invalid/site',
                                         'siteName': 'Fictional site', 'Metric': 'Anyone links', 'value': 4}),
    ]:
        rows.append(dict(common, RecommendationId=identifier, Feature='Review ' + identifier, EvidenceKey=key,
                         ControlId='IDENTITY.MFA' if identifier.startswith('USER') else 'APPS.CONSENT' if identifier.startswith('APP') else 'CONTENT.SHARING',
                         DomainId='identity' if identifier.startswith('USER') else 'applications' if identifier.startswith('APP') else 'content',
                         InvestigationEvidence={'records': [record], 'source': dict(source, dataset=key)}))
    return rows, bundle


class WorkbookLayoutTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.rows, self.bundle = fixture()
        self.path = Path(export_to_excel(self.rows, filename='Assessment.xlsx', tenant_name='Synthetic',
                                       evidence_bundle=self.bundle, output_dir=self.directory.name))
        self.technical_path = technical_workbook_path(self.path)
        self.assessment = load_workbook(self.path)
        self.technical = load_workbook(self.technical_path)
        self.addCleanup(self.assessment.close)
        self.addCleanup(self.technical.close)

    def test_visible_order_and_compatibility_registers(self):
        self.assertEqual([sheet.title for sheet in self.assessment if sheet.sheet_state == 'visible'], list(ASSESSMENT_ORDER))
        self.assertEqual([cell.value for cell in self.assessment['Findings'][1]], list(FINDING_COLUMNS))
        for title in ('Recommendations', 'Evidence Index', 'Control Results', 'Collection Coverage', 'Run Manifest', 'Integrity Checks'):
            self.assertEqual(self.assessment[title].sheet_state, 'hidden')
        self.assertTrue(all(sheet.sheet_state == 'visible' for sheet in self.technical))
        self.assertEqual(self.technical.sheetnames[:2], ['Start Here', 'Findings Lineage'])
        self.assertFalse(any(sheet.title.startswith(('Raw ', 'Logs ', 'Evidence ')) for sheet in self.assessment if sheet.sheet_state == 'visible'))

    def test_evidence_blocks_contiguous_and_links_resolve_with_ids(self):
        for row in self.bundle['assessment_result']['recommendations']:
            identifier = row['RecommendationId']
            for location in row['AssessmentEvidenceRanges']:
                title, cells = location.rsplit('!', 1)
                sheet = self.assessment[title.strip("'")]
                _, first, _, last = range_boundaries(cells)
                self.assertTrue(all(sheet.cell(number, 1).value == identifier for number in range(first, last + 1)))
                self.assertEqual(sum(sheet.cell(number, 1).value == identifier for number in range(2, sheet.max_row + 1)), last - first + 1)
                for number in range(first, last + 1):
                    cell = sheet.cell(number, sheet.max_column)
                    filename, target = cell.hyperlink.target.split('#', 1)
                    self.assertEqual(filename, self.technical_path.name)
                    title, cells = target.rsplit('!', 1)
                    destination = self.technical[title.strip("'")]
                    _, source_row, _, last_row = range_boundaries(cells)
                    self.assertEqual(source_row, last_row)
                    self.assertLessEqual(source_row, destination.max_row)
                    self.assertIn(cell.value, [entry.value for entry in destination[source_row]])

    def test_device_names_and_setting_units_are_kept_separate(self):
        findings = {row['ID']: row for row in records(self.assessment['Findings'])}
        self.assertEqual(findings['WIN-1']['Affected'], '1 devices')
        self.assertEqual(findings['SETTING-1']['Affected'], '2 settings')
        devices = records(self.assessment['Devices'])
        protection = next(row for row in devices if row['ID'] == 'WIN-1')
        self.assertEqual(protection['Device'], 'LAPTOP-0142')
        self.assertEqual(protection['Primary user'], 'j.smith@example.invalid')
        self.assertEqual(protection['OS'], 'Windows 11 23H2')
        self.assertEqual(protection['Intune compliance'], 'Compliant')
        self.assertEqual(protection['Issue'], 'Real-time protection off')
        settings = [row for row in devices if row['ID'] == 'SETTING-1']
        self.assertEqual({row['Issue'] for row in settings}, {'Minimum OS version', 'Require encryption'})
        self.assertEqual(sum(row['ID'] == 'WIN-2' for row in devices), 1, 'One source record appears under both findings.')

    def test_sharing_values_and_recommendations_are_readable(self):
        values = {row['Setting']: row for row in records(self.assessment['Configuration']) if row['ID'] == 'M365-116'}
        self.assertEqual(values['Default sharing link']['Current'], 'Anyone link')
        self.assertEqual(values['Default sharing link']['Recommended'], 'Specific people')
        self.assertEqual(values['Anyone link expiry']['Current'], 'No expiry requirement (not set)')
        self.assertEqual(values['Anyone link expiry']['Recommended'], '30 days or fewer')
        self.assertEqual(values['Anyone link permission (files)']['Current'], 'View and edit')
        self.assertEqual(values['Anyone link permission (files)']['Recommended'], 'View only')
        self.assertEqual(values['Default sharing link']['Status'], 'Change')
        self.assertEqual(next(row for row in records(self.assessment['Findings']) if row['ID'] == 'M365-116')['Affected'], '6 settings')

    def test_all_raw_and_source_rows_reconcile_without_dropped_columns(self):
        for name, datasets in self.bundle['assessment_sources'].items():
            raw = [sheet for sheet in self.technical if sheet.title.startswith('Raw ' + name.replace('_', ' ').title())]
            self.assertEqual(sum(sheet.max_row - 1 for sheet in raw), sum(len(dataset['records']) for dataset in datasets))
            for original in [row for dataset in datasets for row in dataset['records']]:
                retained = next(row for sheet in raw for row in records(sheet) if row.get('id') == original.get('id'))
                for field, value in original.items():
                    self.assertEqual(retained[field], value)
        checks = records(self.assessment['Integrity Checks'])
        self.assertTrue(all(row['Status'] == 'Passed' for row in checks))
        self.assertTrue(any(row.get('Check') == 'Source worksheet rows' for row in checks))
        self.assertEqual(records(self.assessment['Run Manifest'])[-1]['Value'], self.technical_path.name)

    def test_every_workbook_link_targets_visible_cells_and_tables_are_unique(self):
        books = {self.path.name: self.assessment, self.technical_path.name: self.technical}
        for book in books.values():
            names = [name for sheet in book for name in sheet.tables]
            self.assertEqual(len(names), len(set(names)))
            for sheet in book:
                self.assertFalse(sheet.merged_cells)
                self.assertEqual(sheet.freeze_panes, 'A2')
                for row in sheet:
                    for cell in row:
                        if not cell.hyperlink or '#' not in cell.hyperlink.target:
                            continue
                        name, target = cell.hyperlink.target.split('#', 1)
                        destination = books[name] if name else book
                        title = target.rsplit('!', 1)[0].strip("'")
                        self.assertEqual(destination[title].sheet_state, 'visible')

    def test_html_and_summary_link_to_both_workbooks_and_signal_audit_passes(self):
        from tools.audit_report_signal import parse_report, summarize
        html_path = Path(export_to_html(self.rows, filename='Assessment.html', tenant_name='Synthetic',
                        evidence_bundle=self.bundle, excel_path=str(self.path), output_dir=self.directory.name))
        html = html_path.read_text(encoding='utf-8')
        cards = html.split('id="action-plan"', 1)[1].split('</section>', 1)[0]
        self.assertIn('Full source records (technical workbook)', cards)
        for location in re.findall(r'href="([^"]+)"[^>]*>\d+ supporting records? in workbook', cards):
            decoded = unquote(location)
            self.assertIn("Assessment.xlsx#'", decoded)
            self.assertNotIn('Investigation Items', decoded)
        summary = Path(self.bundle['summary_html_path']).read_text(encoding='utf-8')
        self.assertIn(self.path.name, unquote(summary))
        self.assertIn(self.technical_path.name, unquote(summary))
        audit = summarize(parse_report(html_path), html_source=html, workbook=self.path)
        self.assertEqual(audit['audit_issues'], [])

    def test_prior_report_and_baseline_still_import_assessment_workbook(self):
        from Core.prior_report_import import load_prior_report
        from Core.cross_provider_assessment import _load_baseline
        prior = load_prior_report(self.path, include_user_details=True)
        self.assertTrue(prior['available'])
        self.assertTrue(prior['recommendations'])
        baseline = _load_baseline(self.path)
        self.assertEqual(len(baseline['recommendations']), len(self.bundle['assessment_result']['recommendations']))

    def test_layout_mismatch_fails_excel_build_instead_of_csv_fallback(self):
        from Core.processor import export_tabular_reports
        with patch('Core.processor.export_to_excel', side_effect=WorkbookLayoutError('Synthetic mismatch')):
            with self.assertRaisesRegex(WorkbookLayoutError, 'Synthetic mismatch'):
                export_tabular_reports(self.rows, 'Synthetic', 'excel', self.bundle, output_dir=self.directory.name)

    def test_package_deliverables_include_the_technical_workbook(self):
        from contextlib import redirect_stdout
        import io
        from Core.assessment_package import record_package_run
        folder = Path(self.directory.name) / 'Package'
        folder.mkdir()
        with redirect_stdout(io.StringIO()):
            deliverables = record_package_run(folder, mode='offline', tenant_id='synthetic',
                          collected_at='2026-10-05', evaluation_date='2026-10-06',
                          outputs={'excel_path': str(self.path), 'technical_excel_path': str(self.technical_path),
                                   'evidence_bundle': self.bundle})
        self.assertTrue(any(Path(path).name == self.path.name for path in deliverables))
        self.assertTrue(any(Path(path).name == self.technical_path.name for path in deliverables))
        self.assertTrue(all((folder / path).is_file() for path in deliverables))

    def test_declared_logs_are_consolidated_by_dataset_without_dropped_fields(self):
        logs = [sheet for sheet in self.technical if sheet.title.startswith('Logs windows protection')]
        self.assertEqual(len(logs), 1)
        self.assertFalse(any(sheet.title.startswith('Records Windows Protection') for sheet in self.technical))
        data = records(logs[0])
        self.assertEqual({row['RecommendationId'] for row in data}, {'WIN-1', 'WIN-2'})
        self.assertEqual(len(data), 2)
        self.assertTrue(all(row['ParentId'] == 'device-001' for row in data))
        self.assertTrue(all(row['realTimeProtectionEnabled'] is False for row in data))

    def test_reconciliation_failure_is_recorded_and_stops_generation(self):
        from Core import workbook_layout
        original = workbook_layout._clone

        def lose_source_row(source, workbook, title=None):
            copied = original(source, workbook, title)
            if source.title == 'Raw Windows Protection':
                copied.delete_rows(2, 1)
            return copied

        rows, bundle = fixture()
        with tempfile.TemporaryDirectory() as directory:
            with patch('Core.workbook_layout._clone', side_effect=lose_source_row):
                with self.assertRaisesRegex(WorkbookLayoutError, 'source rows do not reconcile'):
                    export_to_excel(rows, filename='Failed.xlsx', evidence_bundle=bundle, output_dir=directory)
            failed = load_workbook(Path(directory) / 'Failed.xlsx')
            try:
                self.assertTrue(any(row['Status'] == 'Failed' for row in records(failed['Integrity Checks'])))
                self.assertFalse(bundle['integrity']['valid'])
            finally:
                failed.close()

    def test_empty_evidence_sheets_are_omitted_and_csv_does_not_write_excel(self):
        from Core.processor import export_tabular_reports
        with tempfile.TemporaryDirectory() as directory:
            export_tabular_reports(self.rows, 'Synthetic', 'csv', deepcopy(self.bundle), output_dir=directory)
            self.assertFalse(list(Path(directory).glob('*.xlsx')))
        with tempfile.TemporaryDirectory() as directory:
            path = export_to_excel([{'RecommendationId': 'EMPTY', 'Service': 'Custom', 'Feature': 'Planning',
                                    'Disposition': 'Coverage', 'Recommendation': 'Obtain evidence.'}], output_dir=directory)
            book = load_workbook(path)
            try:
                self.assertEqual([sheet.title for sheet in book if sheet.sheet_state == 'visible'], list(ASSESSMENT_ORDER[:4]))
            finally:
                book.close()


if __name__ == '__main__':
    unittest.main()
