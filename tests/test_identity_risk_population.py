"""Fictional regression for a positive legacy risk count and its affected list."""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest

from provenance_fixture import fictional_publication
from test_evidence_reconciliation import DAY, TENANT, dataset
from Core.assessment_result import build_assessment_result
from Core.evidence_selection import build_evidence_selection
from Core.identity_investigation import qualify_identity_recommendation
from Core.saved_control_checks import qualify_saved_recommendations


def fixture(rows=None, **source_overrides):
    if rows is None:
        rows = [{'id': 'medium-' + str(i), 'riskLevel': 'medium', 'riskState': 'atRisk'} for i in range(7)]
        rows += [{'id': 'low-' + str(i), 'riskLevel': 'low', 'riskState': 'atRisk'} for i in range(2)]
        rows += [{'id': 'dismissed-' + str(i), 'riskLevel': 'none', 'riskState': 'dismissed'} for i in range(4)]
        rows += [{'id': 'remediated-' + str(i), 'riskLevel': 'none', 'riskState': 'remediated'} for i in range(3)]
    state = dataset([], **source_overrides)['source']
    client = SimpleNamespace(risky_users=rows, collection_status={'risky_users': state})
    original = {'RecommendationId': 'RISK-POSITIVE', 'Service': 'Entra', 'Feature': 'Microsoft Entra ID P2',
        'FindingKey': '', 'EvidenceKey': 'identity_risk_detail', 'ControlId': 'IDENTITY-001', 'DomainId': 'identity',
        'Observation': '7 risky users detected in the tenant (0 high-risk, 7 medium-risk)',
        'Recommendation': 'Investigate the detected users.', 'Priority': 'Medium', 'Disposition': 'Action',
        'Status': 'Attention Required', 'EvidenceComplete': False, 'Confidence': 'Unknown'}
    bundle = {'assessment_sources': {'risky_users': [{'source': state, 'records': rows}]},
        'source_statuses': {'risky_users': state}, 'expected_tenant_id': TENANT,
        'collection_context': {'tenant_id': TENANT, 'collected_at': DAY}}
    return original, client, bundle


def selected(original, client, bundle):
    row = qualify_saved_recommendations([original], 'Entra', client)[0]
    result = build_assessment_result([row], bundle, evaluation_date=DAY, expected_tenant_id=TENANT)
    finding = next(item for item in result['recommendations'] if item.get('FindingKey') == 'entra.identity_risk.users')
    detail = next(item for item in build_evidence_selection(result, bundle)['findings']
                  if item['finding_id'] == finding['RecommendationId'])
    return result, finding, detail


class IdentityRiskPopulationTests(unittest.TestCase):
    def test_positive_legacy_count_selects_all_nine_unresolved_users(self):
        original, client, bundle = fixture()
        before = deepcopy(bundle)
        result, finding, detail = selected(original, client, bundle)
        self.assertEqual(finding['ReportedRiskyUserCount'], 9)
        self.assertEqual(finding['RiskReviewRecordCount'], 9)
        self.assertEqual(finding['EvidenceSource'], 'risky_users')
        self.assertTrue(finding['EvidenceComplete'])
        self.assertTrue(finding['ObservationDate'].startswith(DAY))
        self.assertEqual(detail['record_type'], 'risky_user')
        self.assertEqual(len(detail['records']), 9)
        self.assertEqual({row['fields']['userId'] for row in detail['records']},
                         {'medium-' + str(i) for i in range(7)} | {'low-0', 'low-1'})
        self.assertTrue(all(row['evidence_record_ids'] for row in detail['records']))
        self.assertEqual(finding['OriginalObservation'], original['Observation'])
        self.assertEqual(bundle['assessment_sources'], before['assessment_sources'])
        self.assertEqual(len(client.risky_users), 16)

    def test_complete_empty_or_resolved_only_does_not_retain_a_positive_aggregate(self):
        for rows in [[], [{'id': 'closed', 'riskLevel': 'medium', 'riskState': 'dismissed'}],
                     [{'id': 'safe', 'riskLevel': 'high', 'riskState': 'confirmedSafe'}]]:
            with self.subTest(rows=rows):
                original, client, bundle = fixture(rows)
                original['ReportedRiskyUserCount'] = 99
                row = qualify_identity_recommendation(original, client)
                self.assertEqual(row['ReportedRiskyUserCount'], 0)
                self.assertEqual(row['RiskReviewRecordCount'], 0)
                self.assertEqual(row['Disposition'], 'Reference')
                self.assertEqual(row['Recommendation'], '')
                self.assertTrue(row['EvidenceComplete'])

    def test_partial_results_are_a_retained_population_not_a_tenant_total(self):
        original, client, bundle = fixture([{'id': 'retained', 'riskLevel': 'low', 'riskState': 'atRisk'}],
            availability_status='partial', available=False, complete=False, truncated=True)
        row = qualify_identity_recommendation(original, client)
        self.assertEqual(row['ReportedRiskyUserCount'], 1)
        self.assertFalse(row['EvidenceComplete'])
        self.assertEqual(row['SourceAvailability'], 'partial')
        self.assertIn('source is incomplete', row['Observation'])
        self.assertNotIn('7 risky users detected', row['Observation'])
        detail = build_evidence_selection({'tenant_id': TENANT, 'recommendations': [row]}, bundle)['findings'][0]
        self.assertEqual(len(detail['records']), 1)

    def test_unavailable_empty_source_cannot_establish_absence(self):
        original, client, bundle = fixture([], availability_status='unavailable', available=False, complete=False)
        row = qualify_identity_recommendation(original, client)
        self.assertEqual(row['Disposition'], 'Coverage')
        self.assertFalse(row['EvidenceComplete'])
        self.assertIn('absence of unresolved risk cannot be established', row['Observation'])
        self.assertNotIn('0 high/medium', row['Observation'])

    def test_unknown_state_high_medium_records_are_qualified_separately(self):
        rows = [{'id': 'unknown-' + str(i), 'riskLevel': 'medium', 'riskState': state}
                for i, state in enumerate([None, 'unknownFutureValue', 'unknown'])]
        original, client, bundle = fixture(rows)
        row = qualify_identity_recommendation(original, client)
        self.assertEqual(row['ReportedRiskyUserCount'], 0)
        self.assertEqual(row['RiskReviewRecordCount'], 3)
        self.assertEqual(row['Disposition'], 'Coverage')
        self.assertIn('not counted as unresolved', row['Observation'])
        detail = build_evidence_selection({'tenant_id': TENANT, 'recommendations': [row]}, bundle)['findings'][0]
        self.assertEqual(len(detail['records']), 3)
        self.assertTrue(all(any('not an established unresolved' in note for note in item['qualifications'])
                            for item in detail['records']))
        from Core.investigation_details import _select
        for record in rows:
            item = _select('identity_risk_detail', {'Record Type': 'Risky User', 'Object ID': record['id'],
                'Risk State': record['riskState'], 'Risk Level': record['riskLevel']}, row)
            self.assertIsNotNone(item)
            self.assertIn('confirm before classifying', item['Why Review'])

    def test_detection_event_finding_is_not_reclassified_as_a_user_cohort(self):
        original, client, bundle = fixture()
        original.update(EvidenceSource='risk_detections', Observation='2 risk detections associated with risky users.')
        self.assertEqual(qualify_identity_recommendation(original, client), original)

    def test_combined_entitlement_text_is_not_replaced_by_one_risk_cohort(self):
        original, client, bundle = fixture()
        original['Observation'] = 'Microsoft Entra ID P2 is active in a subscription; capabilities include monitoring risky users.'
        self.assertEqual(qualify_identity_recommendation(original, client), original)

    def test_at_risk_legacy_user_count_uses_the_same_unresolved_population(self):
        original, client, bundle = fixture()
        original['Observation'] = '7 at-risk user(s) detected with suspicious sign-in patterns.'
        row = qualify_identity_recommendation(original, client)
        self.assertEqual(row['ReportedRiskyUserCount'], 9)
        self.assertEqual(row['FindingKey'], 'entra.identity_risk.users')

    def test_state_normalization_and_repeat_qualification_preserve_the_cohort(self):
        original, client, bundle = fixture([{'id': 'retained', 'riskLevel': 'LOW', 'riskState': ' AtRisk '}])
        once = qualify_identity_recommendation(original, client)
        twice = qualify_identity_recommendation(once, client)
        self.assertEqual(twice, once)
        self.assertEqual(twice['ReportedRiskyUserCount'], 1)


@fictional_publication
class IdentityRiskOutputTests(unittest.TestCase):
    def test_customer_summary_technical_detail_and_raw_inventory_reconcile(self):
        from openpyxl import load_workbook
        from Core.export_recommendations import export_to_excel
        from Core.workbook_layout import technical_workbook_path
        from Core.customer_report import render_customer_report
        from Core.assessment_serialization import write_assessment_result, read_assessment_result
        original, client, bundle = fixture()
        result, finding, detail = selected(original, client, bundle)
        bundle['assessment_result'] = result
        with TemporaryDirectory() as folder:
            path = Path(export_to_excel(result['recommendations'], filename='Fictional.xlsx',
                        tenant_name='Fictional', evidence_bundle=bundle, output_dir=folder))
            technical = load_workbook(technical_workbook_path(path))
            customer = load_workbook(path)
            try:
                self.assertEqual(technical['Raw Risky Users'].max_row - 1, 16)
                rows = customer['Users & Identity'].iter_rows(values_only=True)
                headers = next(rows)
                cohort = [dict(zip(headers, row)) for row in rows if row[0] == finding['RecommendationId']]
                self.assertEqual(len(cohort), 9)
                self.assertTrue(all(row['User'] == 'Restricted identity-risk record' for row in cohort))
                risk = next(row for row in bundle['assessment_result']['recommendations']
                            if row['RecommendationId'] == finding['RecommendationId'])
                self.assertEqual(risk['AssessmentEvidenceCount'], 9)
                html = render_customer_report(bundle['assessment_result'], bundle, 'Fictional', path,
                                               technical_workbook_path=technical_workbook_path(path))
                self.assertTrue('9 unresolved risky user' in html, 'HTML must show the corrected unresolved-user population.')
                snapshot = Path(folder) / 'snapshot.json'
                write_assessment_result(snapshot, bundle['assessment_result'])
                saved = read_assessment_result(snapshot)
                saved_risk = next(row for row in saved['recommendations'] if row['RecommendationId'] == finding['RecommendationId'])
                self.assertEqual(saved_risk['ReportedRiskyUserCount'], 9)
                self.assertTrue(saved_risk['EvidenceComplete'])
            finally:
                customer.close()
                technical.close()


if __name__ == '__main__':
    unittest.main()
