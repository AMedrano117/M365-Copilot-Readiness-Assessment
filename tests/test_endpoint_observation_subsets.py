"""One condition across reporting dates, with dated subsets and intact lineage."""
from copy import deepcopy
from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from provenance_fixture import fictional_publication
from test_device_compliance_reasons import fixture
from test_device_example_details import render
from test_evidence_reconciliation import DAY, TENANT, dataset
from Core.expanded_findings import expanded_findings
from Core.assessment_result import build_assessment_result
from Core.evidence_selection import build_evidence_selection
from Core.finding_evidence import build_finding_evidence


OLD = (date.fromisoformat(DAY)-timedelta(days=60)).isoformat()


def bundle_for(dates=(DAY, OLD, None, '0001-01-01'), *, same_device=False, signatures=False):
    bundle = fixture()
    bundle['assessment_sources'].pop('compliance_setting_states')
    devices, health = [], []
    for index, observed in enumerate(dates):
        identifier = 'intune-0' if same_device else f'intune-{index}'
        if not same_device or index == 0:
            devices.append({'id': identifier, 'deviceName': f'Fictional-{identifier}',
                'azureADDeviceId': 'entra-'+identifier, 'lastSyncDateTime': DAY, 'complianceState': 'compliant'})
        health.append({'id': f'health-{index}', 'ParentId': identifier, 'realTimeProtectionEnabled': False,
                       'signatureUpdateOverdue': signatures, 'lastReportedDateTime': observed})
    bundle['assessment_sources']['managed_devices'] = [dataset(devices)]
    bundle['assessment_sources']['windows_protection'] = [dataset(health)]
    return bundle


def rtp(rows):
    return [row for row in rows if row.get('DeviceFindingType') == 'DefenderRealTimeProtectionObservation']


class ObservationSubsetTests(unittest.TestCase):
    def test_same_condition_has_one_finding_and_four_explicit_subsets(self):
        bundle = bundle_for(); before = deepcopy(bundle)
        rows = rtp(expanded_findings(bundle, evaluation_date=DAY))
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row['FindingKey'], 'expanded.windows_protection.disabled')
        subsets = row['EndpointObservationSubsets']
        self.assertEqual([subset['counts']['source_record_count'] for subset in subsets.values()], [1, 1, 1, 1])
        ids = [key for subset in subsets.values() for key in subset['evidence_ids']]
        self.assertCountEqual(ids, row['DeviceEvidenceIds'])
        self.assertEqual(len(set(ids)), 4)
        self.assertTrue(all(subset['requires_confirmation'] for key, subset in subsets.items() if key != 'current'))
        self.assertEqual(bundle, before)

    def test_unique_total_and_subset_device_counts_are_not_added(self):
        row = rtp(expanded_findings(bundle_for((DAY, OLD), same_device=True), evaluation_date=DAY))[0]
        self.assertEqual(row['DeviceEvidenceCounts']['resolved_device_count'], 1)
        self.assertEqual(row['DeviceEvidenceCounts']['source_record_count'], 2)
        self.assertEqual([row['EndpointObservationSubsets'][key]['counts']['resolved_device_count'] for key in ['current', 'stale']], [1, 1])
        self.assertIn('do not add subset device counts', row['Observation'])

    def test_stale_reporting_does_not_establish_inactivity(self):
        row = rtp(expanded_findings(bundle_for((OLD,)), evaluation_date=DAY))[0]
        self.assertEqual(row['EndpointObservationSubsets']['stale']['counts']['source_record_count'], 1)
        self.assertEqual(row['DeviceEvidenceCounts']['active_device_count'], 1)
        self.assertEqual(row['DeviceEvidenceCounts']['stale_device_count'], 0)

    def test_stale_or_invalid_only_findings_do_not_become_current(self):
        for dates, expected in [((OLD,), 'stale'), ((None, '0001-01-01'), 'unknown')]:
            bundle = bundle_for(dates)
            result = build_assessment_result(expanded_findings(bundle, evaluation_date=DAY), bundle,
                                            evaluation_date=DAY, expected_tenant_id=TENANT)
            row = rtp(result['recommendations'])[0]
            self.assertEqual(row['Freshness'], expected)
            self.assertNotEqual(row['EvidenceStatus'], 'supported')
            if expected == 'unknown':
                self.assertEqual(row['ObservationDate'], '')

    def test_current_support_does_not_refresh_other_observation_dates(self):
        bundle = bundle_for()
        result = build_assessment_result(expanded_findings(bundle, evaluation_date=DAY), bundle,
                                        evaluation_date=DAY, expected_tenant_id=TENANT)
        row = rtp(result['recommendations'])[0]
        self.assertEqual(row['ObservationDate'], DAY)
        self.assertEqual(row['EndpointObservationSubsets']['stale']['counts']['evidence_period']['end'], OLD)
        self.assertIn('support date applies only to the current reporting subset', row['Qualification'])
        self.assertFalse(row['DevicePilotBlockerEligible'])

    def test_distinct_conditions_stay_separate_and_signature_current_support_survives(self):
        bundle = bundle_for((DAY, OLD), signatures=True)
        rows = expanded_findings(bundle, evaluation_date=DAY)
        health = [row for row in rows if row.get('EndpointObservationSubsets')]
        self.assertEqual({row['FindingKey'] for row in health}, {'expanded.windows_protection.disabled', 'expanded.windows_protection.overdue'})
        result = build_assessment_result(rows, bundle, evaluation_date=DAY, expected_tenant_id=TENANT)
        signature = next(row for row in result['recommendations'] if row.get('DeviceFindingType') == 'DefenderSignatureObservation')
        self.assertEqual(signature['Freshness'], 'current')
        self.assertEqual(signature['EvidenceStatus'], 'supported')
        self.assertEqual(signature['EndpointObservationSubsets']['stale']['counts']['source_record_count'], 1)

    def test_configured_reporting_window_controls_subset_membership(self):
        recent = (date.fromisoformat(DAY)-timedelta(days=4)).isoformat()
        bundle = bundle_for((recent,))
        row = rtp(expanded_findings(bundle, profile={'assessment_context': {'endpoint_reporting_days': 3}}, evaluation_date=DAY))[0]
        self.assertEqual(row['EndpointReportingDays'], 3)
        self.assertEqual(row['EndpointObservationSubsets']['stale']['counts']['source_record_count'], 1)

    def test_offline_reassessment_replaces_old_date_split_native_findings(self):
        bundle = bundle_for((DAY, OLD))
        fresh = rtp(expanded_findings(bundle, evaluation_date=DAY))[0]
        previous = [dict(deepcopy(fresh), FindingKey=fresh['FindingKey']+'.'+age) for age in ['current', 'older']]
        result = build_assessment_result(previous, bundle, evaluation_date=DAY, expected_tenant_id=TENANT)
        rows = rtp(result['recommendations'])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['FindingKey'], fresh['FindingKey'])
        self.assertEqual(rows[0]['DeviceEvidenceIds'], fresh['DeviceEvidenceIds'])

    def test_rebuild_recalculates_age_for_a_new_evaluation_date_or_reporting_window(self):
        for age, new_day, profile in [
            (7, (date.fromisoformat(DAY)+timedelta(days=1)).isoformat(), {}),
            (4, DAY, {'assessment_context': {'endpoint_reporting_days': 3}})]:
            reported = (date.fromisoformat(DAY)-timedelta(days=age)).isoformat()
            bundle = bundle_for((reported,))
            previous = expanded_findings(bundle, evaluation_date=DAY)
            self.assertEqual(rtp(previous)[0]['EndpointObservationSubsets']['current']['counts']['source_record_count'], 1)
            bundle['assessment_profile'] = profile
            result = build_assessment_result(previous, bundle, evaluation_date=new_day, expected_tenant_id=TENANT)
            row = rtp(result['recommendations'])[0]
            self.assertEqual(row['EndpointObservationSubsets']['current']['counts']['source_record_count'], 0)
            self.assertEqual(row['EndpointObservationSubsets']['stale']['counts']['source_record_count'], 1)
            self.assertEqual(row['Freshness'], 'stale')

    def test_html_examples_include_current_and_stale_with_one_heading(self):
        bundle = bundle_for((DAY, DAY, DAY, OLD))
        _, html = render(bundle)
        self.assertEqual(html.count('<summary>Validate Defender real-time protection observations</summary>'), 1)
        fragment = html.split('<summary>Validate Defender real-time protection observations</summary>')[1].split('</details>')[0]
        self.assertIn('Fictional-intune-3', fragment)
        self.assertIn('Stale observations', fragment)
        self.assertIn('Current observations', fragment)
        self.assertIn('Device activity', fragment)
        self.assertIn('Evidence recency', fragment)
        self.assertIn('Reported time', fragment)
        self.assertIn('data-label="Device activity">Current', fragment)
        self.assertIn('data-label="Evidence recency">Stale', fragment)
        self.assertEqual(fragment.count('data-label="Example"'), 3)

    def test_subsets_and_example_order_survive_sorted_snapshot_serialization(self):
        import json
        from Core.technical_appendix import render_technical_appendix
        bundle = bundle_for()
        result, before = render(bundle)
        replay = json.loads(json.dumps(result, sort_keys=True, default=str))
        model = build_finding_evidence(build_evidence_selection(replay, bundle))
        after = render_technical_appendix(replay, dict(bundle, finding_evidence=model), 'assessment.xlsx', 'technical.xlsx',
                    {row['finding_id']: row for row in model['findings']}, [])
        self.assertEqual(after, before)


@fictional_publication
class ObservationSubsetWorkbookTests(unittest.TestCase):
    def test_workbook_and_shared_html_evidence_keep_subset_membership_and_raw_dates(self):
        from Core.export_recommendations import export_to_excel
        from Core.workbook_layout import technical_workbook_path
        from openpyxl import load_workbook
        bundle = bundle_for(); raw = deepcopy(bundle['assessment_sources'])
        result, _ = render(bundle); bundle['assessment_result'] = result
        row = rtp(result['recommendations'])[0]
        model = build_finding_evidence(build_evidence_selection(result, bundle))
        records = [record for table in model['tables'].values() for record in table['rows'] if record.get('finding_id') == row['RecommendationId']]
        self.assertEqual(len(records), 4)
        self.assertEqual(sum(record['confirmation_required'] for record in records), 3)
        with TemporaryDirectory() as folder:
            path = Path(export_to_excel(result['recommendations'], filename='Fictional.xlsx', evidence_bundle=bundle, output_dir=folder))
            book = load_workbook(technical_workbook_path(path))
            try:
                values = list(book['Device Detail'].values)
                rows = [dict(zip(values[0], cells)) for cells in values[1:]]
                selected = [item for item in rows if item['ID'] == row['RecommendationId']]
                self.assertEqual(len(selected), 4)
                self.assertEqual(sum(item['Confirmation required'] == 'Yes' for item in selected), 3)
                self.assertTrue(any(item['Source report date'] == 'Not retained' for item in selected))
                self.assertTrue(any(item['Source report date'] == OLD for item in selected))
                self.assertEqual(book['Raw Windows Protection'].max_row-1, 4)
            finally:
                book.close()
        self.assertEqual(bundle['assessment_sources'], raw)


if __name__ == '__main__':
    unittest.main()
