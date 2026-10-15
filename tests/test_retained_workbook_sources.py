"""Older completed snapshots replay exact native evidence, never recollect."""
from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from Core.evidence_selection import evidence_record_ids
from Core.retained_workbook_sources import package_native_sources, read_native_sources


class RetainedWorkbookSourcesTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / 'Technical Evidence.xlsx'
        self.result = {'tenant_id': 'tenant-a', 'identity': {'AssessmentId': 'AST-a', 'RunId': 'RUN-a'},
                       'recommendations': [{'source_refs': [{'dataset': 'sites', 'dataset_index': 300, 'record_index': 1}]}]}
        source = {'tenant_id': 'tenant-a', 'scope': 'Retained sites', 'collected_at': '2026-10-13T10:00:00Z'}
        self.sources = {'sites': [{'source': {}, 'records': []} for _ in range(300)] +
                                  [{'source': source, 'records': [{'id': 'site-a'}, {'id': 'site-b'}]}],
                        'counts': [{'source': source, 'records': [7]}]}
        ids = evidence_record_ids(self.sources, 'tenant-a')
        book = Workbook(); book.active.title = 'Run Manifest'
        for row in [('Item', 'Value'), ('AssessmentId', 'AST-a'), ('RunId', 'RUN-a'), ('Tenant ID', 'tenant-a')]:
            book.active.append(row)
        for name, datasets in self.sources.items():
            sheet = book.create_sheet('Raw ' + name)
            sheet.append(['Evidence Record ID', 'Source dataset', 'Source record', 'Source metadata', 'Raw record'])
            for index, dataset in enumerate(datasets):
                for position, raw in enumerate(dataset['records']):
                    displayed = raw if isinstance(raw, dict) else {'value': raw}
                    sheet.append([ids[(name, index, position)], name, position + 1,
                                  json.dumps(dataset['source']), json.dumps(displayed)])
        sheet = book.create_sheet('MFA Registration Review')
        sheet.append(['User ID', 'Qualification']); sheet.append(['user-a', 'Registration is distinct from enforcement.'])
        book.save(self.path); book.close()

    def test_exact_sparse_indices_and_scalar_shape_restore_every_evidence_id(self):
        restored = read_native_sources(self.path, self.result)
        self.assertEqual(restored, self.sources)
        self.assertEqual(evidence_record_ids(restored, 'tenant-a'), evidence_record_ids(self.sources, 'tenant-a'))

    def test_retained_investigation_records_restore_framework_keys_and_qualifications(self):
        sheets = {}
        read_native_sources(self.path, self.result, sheets)
        self.assertEqual(sheets['mfa_registration_detail']['title'], 'MFA Registration Review')
        self.assertEqual(sheets['mfa_registration_detail']['rows'], [
            {'User ID': 'user-a', 'Qualification': 'Registration is distinct from enforcement.'}])
        self.assertFalse(any(s['title'].startswith('Raw ') for s in sheets.values()))

    def test_replay_preserves_materialized_selection_and_qualification_without_reselecting(self):
        from Core.investigation_details import prepare_investigation_details
        from unittest.mock import patch
        result = {'recommendations': [{'RecommendationId': 'FIND-a', 'Feature': 'Retained finding',
                   'InvestigationCount': 4, 'InvestigationQualification': 'Counts are not proof of exposure.',
                   'InvestigationRange': "'Investigation Items'!A2:L5"}]}
        retained = {'title': 'Investigation Items', 'rows': [{'RecommendationId': 'FIND-a', 'Item': 'Original work item'}]}
        bundle = {'retained_investigations': True, 'sheets': {'investigation_items': deepcopy(retained)}}
        before = deepcopy(result)
        with patch('Core.investigation_details._sharing_items', side_effect=AssertionError('Must not reselect')):
            prepare_investigation_details(bundle, result)
        for key, value in before['recommendations'][0].items():
            self.assertEqual(result['recommendations'][0][key], value)
        self.assertEqual(bundle['sheets']['investigation_items'], retained)

    def test_original_detail_derived_evidence_and_source_ids_preserved_by_exact_joins(self):
        from Core.evidence_selection import build_evidence_selection
        result = {'tenant_id': 'tenant-a', 'identity': self.result['identity'],
                  'recommendations': [{'RecommendationId': 'FIND-a', 'Feature': 'Retained finding',
                                       'EvidenceKind': 'supporting_detail'}]}
        detail_id, evidence_id, source_id = 'DET-' + 'a' * 24, 'EVD-' + 'b' * 24, 'SRC-' + 'c' * 24
        book = load_workbook(self.path)
        sheet = book.create_sheet('Records Retained Details')
        sheet.append(['RecommendationId', 'Evidence ID', 'Detail record ID', 'Detail'])
        sheet.append(['FIND-a', evidence_id, detail_id, 'Original selected detail'])
        sheet = book.create_sheet('Raw Derived Records')
        sheet.append(['Evidence Record ID', 'Source dataset', 'Detail'])
        sheet.append([evidence_id, 'investigation.FIND-a', 'Original selected detail'])
        sheet = book.create_sheet('Evidence Retained Detail')
        sheet.append(['finding_id', 'detail_record_id', 'evidence_record_ids', 'dataset_id', 'qualification', 'Detail'])
        sheet.append(['FIND-a', detail_id, evidence_id, source_id, 'Retained scope only.', 'Original selected detail'])
        second_detail = 'DET-' + 'd' * 24
        sheet.append(['FIND-a', second_detail, evidence_id, source_id, 'A distinct selection of the same native row.', 'Second selection'])
        book.save(self.path); book.close()
        retained = {}
        sources = read_native_sources(self.path, result, retained_selection=retained)
        selection = build_evidence_selection(result, {'assessment_sources': sources, 'retained_record_selection': retained})
        record = selection['recommendations'][0]['records'][0]
        self.assertEqual({r['record_id'] for r in selection['recommendations'][0]['records']}, {detail_id, second_detail})
        self.assertEqual(record['record_id'], detail_id)
        self.assertEqual(record['evidence_record_ids'], [evidence_id])
        self.assertEqual(record['qualifications'], ['Retained scope only.'])
        self.assertEqual(selection['recommendations'][0]['record_status'], 'records_available')
        derived = next(row for row in selection['evidence_records'] if row['record_id'] == evidence_id)
        self.assertEqual(derived['dataset_id'], source_id)
        self.assertEqual(derived['raw']['Detail'], 'Original selected detail')

    def test_changed_native_record_rejected(self):
        book = load_workbook(self.path)
        book['Raw sites']['E3'] = json.dumps({'id': 'changed'})
        book.save(self.path); book.close()
        with self.assertRaisesRegex(ValueError, 'integrity mismatch'):
            read_native_sources(self.path, self.result)

    def test_foreign_run_and_missing_reference_rejected(self):
        foreign = deepcopy(self.result); foreign['identity']['RunId'] = 'RUN-other'
        with self.assertRaisesRegex(ValueError, 'different assessment run'):
            read_native_sources(self.path, foreign)
        missing = deepcopy(self.result); missing['recommendations'][0]['source_refs'][0]['record_index'] = 8
        with self.assertRaisesRegex(ValueError, 'every snapshot native reference'):
            read_native_sources(self.path, missing)

    def test_only_completed_same_run_receipt_selects_sources(self):
        receipt = {'status': 'complete', 'semantic_run': {'RunId': 'RUN-a'}, 'deliverables': [self.path.name]}
        log = self.root / 'operator-log.jsonl'; log.write_text(json.dumps(receipt), encoding='utf-8')
        self.assertEqual(package_native_sources(self.root, self.result), self.sources)
        receipt['semantic_run']['RunId'] = 'RUN-other'; log.write_text(json.dumps(receipt), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'same-run technical workbook'):
            package_native_sources(self.root, self.result)

    def test_missing_package_does_not_invent_sources_and_native_adapter_rejects_missing_reference(self):
        from Core.evidence_selection import _Registry
        self.assertEqual(package_native_sources(self.root, self.result), {})
        with self.assertRaisesRegex(ValueError, 'original completed package'):
            _Registry('tenant-a', {}).link([{'source_refs': self.result['recommendations'][0]['source_refs']}])


if __name__ == '__main__':
    unittest.main()
