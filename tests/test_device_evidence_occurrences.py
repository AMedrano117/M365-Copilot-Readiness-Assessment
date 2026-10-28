"""Fictional regression: equal payloads are separate native source occurrences."""

from collections import Counter
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from provenance_fixture import fictional_publication
from test_device_compliance_reasons import fixture
from test_evidence_reconciliation import DAY,TENANT,dataset
from Core.assessment_result import build_assessment_result
from Core.expanded_findings import expanded_findings
from Core.evidence_selection import build_evidence_selection,evidence_record_ids
from Core.finding_evidence import build_finding_evidence


def cohort():
    bundle=fixture()
    records=bundle['assessment_sources']['compliance_setting_states'][0]['records']
    seed=dict(records[0])
    records[:]=[dict(seed,id='activity-'+str(i)) for i in range(167)]
    records.extend(dict(records[i]) for i in range(5))
    errors=[dict(seed,id='error-'+str(i),settingName='Fictional setting',state='error') for i in range(255)]
    records.extend(errors+[dict(errors[i]) for i in range(43)])
    rows=expanded_findings(bundle,evaluation_date=DAY)
    result=build_assessment_result(rows,bundle,evaluation_date=DAY,expected_tenant_id=TENANT)
    return bundle,rows,result


def device_findings(result):
    return [row for row in result['recommendations'] if row.get('DeviceFindingType') in
            {'ActivityOrEnrollmentFailure','ReportingErrorOrConflict'}]


class DeviceOccurrenceSelectionTests(unittest.TestCase):
    def test_exact_reported_populations_keep_repeated_payloads_and_native_ids(self):
        bundle,rows,result=cohort()
        original=deepcopy(bundle)
        selection=build_evidence_selection(result,bundle)
        native=evidence_record_ids(bundle['assessment_sources'],TENANT)
        self.assertEqual(bundle,original)
        for finding in device_findings(result):
            expected=172 if finding['DeviceFindingType']=='ActivityOrEnrollmentFailure' else 298
            detail=next(row for row in selection['findings'] if row['finding_id']==finding['RecommendationId'])
            self.assertEqual(finding['DeviceEvidenceCounts']['source_record_count'],expected)
            self.assertEqual(len(detail['records']),expected)
            self.assertEqual([row['evidence_record_ids'][0] for row in detail['records']],finding['DeviceEvidenceIds'])
            self.assertEqual(len(set(finding['DeviceEvidenceIds'])),expected)
            self.assertTrue(all(len(row['source_refs'])==1 for row in detail['records']))
            self.assertTrue(set(finding['DeviceEvidenceIds'])<=set(native.values()))
            self.assertEqual(Counter(str(row['fields']) for row in detail['records']),
                             Counter(str(row) for row in finding['InvestigationEvidence']['records']))

    def test_identical_rows_in_separate_reads_link_to_their_own_dataset(self):
        bundle=fixture()
        raw=bundle['assessment_sources']['compliance_setting_states'][0]['records'][0]
        bundle['assessment_sources']['compliance_setting_states']=[
            dataset([dict(raw)],source_api='https://example.invalid/first-read'),
            dataset([dict(raw)],source_api='https://example.invalid/second-read')]
        result=build_assessment_result(expanded_findings(bundle,evaluation_date=DAY),bundle,
                                      evaluation_date=DAY,expected_tenant_id=TENANT)
        finding=device_findings(result)[0]
        detail=next(row for row in build_evidence_selection(result,bundle)['findings'] if row['finding_id']==finding['RecommendationId'])
        self.assertEqual(len(detail['records']),2)
        self.assertEqual([row['source_refs'][0]['dataset_index'] for row in detail['records']],[0,1])
        self.assertNotEqual(detail['records'][0]['evidence_record_ids'],detail['records'][1]['evidence_record_ids'])

    def test_finding_tables_preserve_source_occurrence_counts(self):
        bundle,rows,result=cohort()
        model=build_finding_evidence(build_evidence_selection(result,bundle))
        for finding in device_findings(result):
            detail=next(row for row in model['findings'] if row['finding_id']==finding['RecommendationId'])
            selected=[row for table in model['tables'].values() if table['role']=='selected'
                      for row in table['rows'] if row['finding_id']==finding['RecommendationId']]
            self.assertEqual(len(selected),finding['DeviceEvidenceCounts']['source_record_count'])
            self.assertEqual(len({row['detail_record_id'] for row in selected}),len(selected))

    def test_missing_or_repeated_occurrence_ids_fail_closed(self):
        for invalid in ('missing','repeated'):
            with self.subTest(invalid=invalid):
                bundle,rows,result=cohort()
                finding=device_findings(result)[0]
                if invalid=='missing':finding['DeviceEvidenceIds'][0]='EVD-fictional-missing'
                else:finding['DeviceEvidenceIds'][1]=finding['DeviceEvidenceIds'][0]
                with self.assertRaisesRegex(ValueError,'occurrence'):
                    build_evidence_selection(result,bundle)

    def test_old_collapsed_retained_detail_is_repaired_from_frozen_native_cohort(self):
        bundle,rows,result=cohort()
        selection=build_evidence_selection(result,bundle)
        retained={'records':{},'selected_findings':[]}
        for finding in device_findings(result):
            detail=next(row for row in selection['findings'] if row['finding_id']==finding['RecommendationId'])
            unique={str(row['fields']):deepcopy(row) for row in detail['records']}
            retained['records'][finding['RecommendationId']]=list(unique.values())
            retained['selected_findings'].append(finding['RecommendationId'])
        self.assertEqual(sorted(len(value) for value in retained['records'].values()),[167,255])
        bundle['retained_record_selection']=retained
        before=deepcopy(result)
        repaired=build_evidence_selection(result,bundle)
        for finding in device_findings(result):
            detail=next(row for row in repaired['findings'] if row['finding_id']==finding['RecommendationId'])
            self.assertEqual(len(detail['records']),finding['DeviceEvidenceCounts']['source_record_count'])
            self.assertEqual([row['evidence_record_ids'][0] for row in detail['records']],finding['DeviceEvidenceIds'])
        self.assertEqual(result,before)
        self.assertEqual(sorted(len(value) for value in retained['records'].values()),[167,255])

    def test_valid_retained_detail_keeps_its_existing_ids_and_qualifications(self):
        bundle,rows,result=cohort()
        selection=build_evidence_selection(result,bundle)
        finding=device_findings(result)[0]
        detail=next(row for row in selection['findings'] if row['finding_id']==finding['RecommendationId'])
        records=deepcopy(detail['records'])
        records[0]['qualifications']=['Retained fictional qualification.']
        bundle['retained_record_selection']={'records':{finding['RecommendationId']:records},
                                             'selected_findings':[finding['RecommendationId']]}
        replay=build_evidence_selection(result,bundle)
        actual=next(row for row in replay['findings'] if row['finding_id']==finding['RecommendationId'])
        self.assertEqual(actual['records'],records)


@fictional_publication
class DeviceOccurrenceWorkbookTests(unittest.TestCase):
    def test_workbook_native_detail_and_summary_all_reconcile(self):
        from openpyxl import load_workbook
        from Core.export_recommendations import export_to_excel
        from Core.workbook_layout import technical_workbook_path
        from Core.customer_report import render_customer_report
        bundle,rows,result=cohort()
        bundle['assessment_result']=result
        original=deepcopy(bundle['assessment_sources'])
        with TemporaryDirectory() as folder:
            path=Path(export_to_excel(rows,filename='Fictional.xlsx',evidence_bundle=bundle,output_dir=folder))
            result=bundle['assessment_result']
            customer=load_workbook(path);technical=load_workbook(technical_workbook_path(path))
            try:
                raw=technical['Raw Compliance Setting States']
                self.assertEqual(raw.max_row-1,470)
                columns=[cell.value for cell in raw[1]]
                native_ids={raw.cell(index,columns.index('Evidence Record ID')+1).value for index in range(2,raw.max_row+1)}
                self.assertEqual(len(native_ids),470)
                detail=technical['Device Detail']
                columns=[cell.value for cell in detail[1]]
                counts=Counter(detail.cell(index,columns.index('ID')+1).value for index in range(2,detail.max_row+1))
                for finding in device_findings(result):
                    expected=finding['DeviceEvidenceCounts']['source_record_count']
                    self.assertEqual(counts[finding['RecommendationId']],expected)
                    self.assertEqual(finding['AssessmentEvidenceCount'],expected)
                    self.assertTrue(set(finding['DeviceEvidenceIds'])<=native_ids)
                    self.assertEqual(finding['DeviceEvidenceCounts']['distinct_intune_devices'],1)
                    if finding['DeviceFindingType']=='ActivityOrEnrollmentFailure':
                        selected=[index for index in range(2,detail.max_row+1)
                                  if detail.cell(index,columns.index('ID')+1).value==finding['RecommendationId']]
                        self.assertTrue(all('Default Device Compliance Policy' in str(detail.cell(index,columns.index('Detail')+1).value)
                                            for index in selected))
                integrity=list(customer['Integrity Checks'].values)
                checks=[dict(zip(integrity[0],row)) for row in integrity[1:]]
                self.assertTrue(all(row['Status']=='Passed' for row in checks))
                self.assertEqual(len(list(customer['Devices'].values))-1,4)
                html=render_customer_report(result,bundle,'Fictional',path,technical_workbook_path=technical_workbook_path(path))
                self.assertIn('172',html)
                self.assertIn('298',html)
            finally:
                customer.close();technical.close()
        self.assertEqual(bundle['assessment_sources'],original)

    def test_incorrect_summary_count_still_blocks_workbook(self):
        from Core.export_recommendations import export_to_excel
        from Core.workbook_layout import WorkbookLayoutError
        bundle,rows,result=cohort()
        device_findings(result)[0]['DeviceEvidenceCounts']['source_record_count']+=1
        bundle['assessment_result']=result
        with TemporaryDirectory() as folder,self.assertRaisesRegex(WorkbookLayoutError,'selected device records do not reconcile'):
            export_to_excel(rows,filename='Fictional.xlsx',evidence_bundle=bundle,output_dir=folder)


if __name__=='__main__':
    unittest.main()
