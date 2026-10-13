from provenance_fixture import fictional_publication
"""Actual fictional workbooks, HTML, snapshots and immutable lifecycle results."""
from contextlib import redirect_stdout
from copy import deepcopy
import io
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from Core.assessment_result import build_assessment_result
from Core.export_recommendations import export_to_excel, export_to_html
from Core.assessment_serialization import write_assessment_result, read_assessment_result
from Core.workbook_layout import technical_workbook_path, RANGE
from tests.test_workbook_layout import fixture, records
from tests.test_configuration_grouping import selection
from openpyxl import load_workbook


def shared_fixture():
    recommendations,bundle=fixture()
    row=next(r for r in recommendations if r.get('FindingKey','').startswith('sharepoint.'))
    row['FindingKey']='sharepoint.sharing.partial_settings'
    row.pop('InvestigationEvidence',None)
    payload=selection()
    source=dict(payload['sources'][0]['source'],tenant_id=row['TenantId'])
    bundle['assessment_sources']['sharepoint_site_settings']=[{'source':source,'records':[r['raw'] for r in payload['evidence_records']]}]
    result=build_assessment_result([row],bundle,evaluation_date='2026-10-09',expected_tenant_id=source['tenant_id'])
    bundle['assessment_result']=result
    return bundle,result


@fictional_publication
class ConfigurationOutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=TemporaryDirectory()
        cls.bundle,cls.result=shared_fixture()
        with redirect_stdout(io.StringIO()):
            cls.path=export_to_excel(cls.result['recommendations'],filename='fiction.xlsx',evidence_bundle=cls.bundle,output_dir=cls.temp.name)
            cls.html=export_to_html(cls.result['recommendations'],filename='fiction.html',evidence_bundle=cls.bundle,output_dir=cls.temp.name,excel_path=cls.path)
        cls.workbook=load_workbook(cls.path); cls.technical=load_workbook(technical_workbook_path(cls.path))

    @classmethod
    def tearDownClass(cls):
        cls.workbook.close(); cls.technical.close(); cls.temp.cleanup()

    def test_four_site_groups_have_distinct_counts_and_scope(self):
        rows=[r for r in records(self.workbook['Configuration']) if r['Scope']=='site']
        self.assertEqual(len(rows),4)
        self.assertEqual(sorted(r['Affected objects'] for r in rows),[1,1,2,2])
        self.assertTrue(all(r['Total objects assessed']==3 for r in rows))
        self.assertTrue(all(r['Group ID'].startswith('CFG-') for r in rows))

    def test_group_links_cover_complete_contiguous_membership(self):
        sheet=self.workbook['Configuration']; headers={c.value:c.column for c in sheet[1]}
        for number,row in enumerate(records(sheet),2):
            target=sheet.cell(number,headers['Affected objects link']).hyperlink.target
            filename,location=target.split('#'); self.assertEqual(filename,technical_workbook_path(self.path).name)
            match=RANGE.fullmatch(location); self.assertIsNotNone(match)
            detail=self.technical[match[1].strip("'")]
            members=records(detail)[int(match[3])-2:int(match[5])-1]
            self.assertEqual({m['Group ID'] for m in members},{row['Group ID']})
            self.assertEqual(len({m['Object ID'] for m in members if m['Object ID']}),row['Affected objects'])

    def test_members_have_all_native_evidence_and_precision(self):
        model=self.result['configuration_groups']; detail=records(self.technical['Configuration Detail'])
        self.assertEqual(len(detail),len(model['members']))
        sites=[m for m in model['members'] if m['scope']=='site']
        self.assertEqual({m['object_id'] for m in sites},{'site-a','site-b','site-c'})
        for m in sites: self.assertIn('1234567Z',str(m['source_capture']))
        evidence={i for m in model['members'] for i in m['evidence_ids']}
        all_values={str(c.value) for s in self.technical for row in s for c in row}
        self.assertTrue(evidence <= all_values)

    def test_html_summary_is_shared_and_has_no_object_locators(self):
        from Core.configuration_groups import summary_html
        from Core.customer_report import _settings
        html=_settings(self.bundle)
        self.assertIn(summary_html(self.result['configuration_groups']),html)
        self.assertNotIn('https://example.invalid/a',html)
        self.assertNotIn('site-a',html)
        self.assertNotIn('EVD-',html)

    def test_snapshot_replay_preserves_model_without_recollection(self):
        from Core.assessment_replay import render_snapshot
        with TemporaryDirectory() as folder,redirect_stdout(io.StringIO()):
            source=Path(folder)/'snapshot.json'; write_assessment_result(source,self.result); before=source.read_bytes()
            loaded=read_assessment_result(source)
            self.assertEqual(loaded['configuration_groups'],self.result['configuration_groups'])
            rendered=render_snapshot(source,Path(folder)/'outputs')
            workbook=load_workbook(rendered['excel_path'])
            try:
                self.assertEqual([r['Group ID'] for r in records(workbook['Configuration'])],
                                 [r['Group ID'] for r in records(self.workbook['Configuration'])])
            finally: workbook.close()
            self.assertEqual(source.read_bytes(),before)

    def test_html_membership_pages_paginate_every_retained_member(self):
        from Core.evidence_selection import build_evidence_selection
        from Core.finding_evidence import build_finding_evidence
        from Core.html_evidence_pages import write_html_evidence_pages
        model=build_finding_evidence(build_evidence_selection(self.result,self.bundle))
        with TemporaryDirectory() as folder:
            output=write_html_evidence_pages(model,Path(folder)/'Evidence',page_rows=1)
            self.assertEqual(set(output['configuration_pages']),{g['group_id'] for g in model['configuration_groups']['groups']})
            self.assertEqual(sum(len(p) for p in output['configuration_pages'].values()),len(model['configuration_groups']['members']))


@fictional_publication
class ConfigurationLifecycleTests(unittest.TestCase):
    def test_display_grouping_does_not_change_real_delta(self):
        from test_delta_lifecycle import pair,compare
        from Core.configuration_groups import group_observations
        from tests.test_configuration_grouping import observation
        baseline,current=pair(); before=deepcopy(current['identity'])
        first=compare(baseline,current)
        current['configuration_groups']=group_observations([observation(),observation('site-b')])
        second=compare(baseline,current)
        self.assertEqual(first,second)
        self.assertEqual(before,current['identity'])
