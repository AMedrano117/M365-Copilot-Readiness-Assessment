"""Retained deliverables work independently of the discontinued export pipeline."""

import contextlib
import copy
import io
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

from Core.cli_parser import parse_arguments
from Core.offline_collection import load_collection
import test_offline_report as offline_tests
from synthetic_package_fixture import create_synthetic_package


RETIRED = ('Core.dashboard_export', 'Core.dashboard_package', 'Core.dashboard_records', 'Core.app_builder_export')
ROOT = Path(__file__).resolve().parents[1]


class DashboardRetirementTests(unittest.TestCase):
    def parse(self, *arguments):
        with patch('sys.argv', ['main.py', *arguments]), contextlib.redirect_stderr(io.StringIO()):
            return parse_arguments(None, None)

    def test_retired_export_choices_are_rejected(self):
        for choice in ('dashboard-json', 'app-builder'):
            with self.subTest(choice=choice):
                with self.assertRaises(SystemExit) as error:
                    self.parse('--extra-exports', choice)
                self.assertEqual(error.exception.code, 2)
        self.assertEqual(self.parse('--extra-exports', 'evidence-pages').extra_exports, ['evidence-pages'])

    def test_excel_html_and_evidence_pages_work_with_retired_modules_blocked(self):
        from openpyxl import load_workbook
        from Core.workbook_layout import technical_workbook_path
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            collection = Path(create_synthetic_package(root / 'inputs', active_incident=True))
            original = collection.read_bytes()
            with patch.dict(sys.modules, dict.fromkeys(RETIRED)):
                html = offline_tests.OfflineReportTests().run_cli([
                    '--collection-input', str(collection), '--evaluation-date', '2026-09-15',
                    '--extra-exports', 'evidence-pages'], root)
            offline_tests.OfflineReportTests().assertLocalHtmlLinksExist(html.parent)
            self.assertEqual(collection.read_bytes(), original)
            for path in (html.with_suffix('.xlsx'), technical_workbook_path(html.with_suffix('.xlsx'))):
                workbook = load_workbook(path, read_only=True)
                try:
                    self.assertTrue(workbook.sheetnames)
                finally:
                    workbook.close()
            self.assertTrue((html.parent / 'Evidence' / 'index.html').is_file())
            self.assertFalse(any(path.suffix in {'.json', '.zip'} for path in html.parent.rglob('*')))
            self.assertFalse((html.parent / 'App Builder').exists())
            self.assertFalse((html.parent / 'JSON').exists())
            self.assertNotIn('App Builder', html.read_text(encoding='utf-8'))

    def test_snapshot_serializes_shared_result_without_dashboard_projection(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            collection = create_synthetic_package(root / 'inputs', active_incident=True)
            snapshot = root / 'assessment.json'
            offline_tests.OfflineReportTests().run_cli(['--collection-input', str(collection), '--evaluation-date',
                                         '2026-09-15', '--snapshot-json', str(snapshot)], root)
            result = json.loads(snapshot.read_text(encoding='utf-8'))
            for key in ('recommendations', 'controls', 'domains', 'counts', 'decision'):
                self.assertIn(key, result)
            for key in ('dashboard_schema_version', 'findings', 'sources', 'evidence_records', 'assessment_result'):
                self.assertNotIn(key, result)
            self.assertEqual(result['counts']['actions'], len(result['actions']))
            self.assertNotIn('finding_uid', json.dumps(result))

    def test_moved_package_replays_without_original_inputs_and_preserves_hashes(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            collection = create_synthetic_package(root / 'original')
            loaded = load_collection(collection)
            package = Path(loaded['package_directory'])
            before = {path.relative_to(package): path.read_bytes() for path in package.rglob('*') if path.is_file()}
            moved = root / 'moved'
            shutil.move(str(package), moved)
            shutil.rmtree(root / 'original')
            html = offline_tests.OfflineReportTests().run_cli(['--collection-input', str(moved / 'collection.json'),
                                                '--evaluation-date', '2026-09-15'], root)
            self.assertTrue(html.is_relative_to(moved))
            for relative, original in before.items():
                self.assertEqual((moved / relative).read_bytes(), original)
            self.assertTrue(html.with_suffix('.xlsx').is_file())
            self.assertFalse(list(html.parent.rglob('*.zip')))

    def test_shared_selection_keeps_native_occurrences_and_references(self):
        from Core.evidence_selection import build_evidence_selection, evidence_record_ids
        from test_finding_evidence import legacy_finding, source, LOGS, TENANT
        sources = {'signin_logs': [{'records': copy.deepcopy(LOGS), 'source': source()}]}
        result = {'tenant_id': TENANT, 'evaluation_date': '2026-09-30', 'recommendations': [legacy_finding()]}
        selected = build_evidence_selection(result, {'assessment_sources': sources})
        ids = evidence_record_ids(sources, TENANT)
        self.assertEqual(set(ids.values()), {row['record_id'] for row in selected['evidence_records']})
        for record in selected['findings'][0]['records']:
            self.assertTrue(record['evidence_record_ids'])
            self.assertTrue(set(record['evidence_record_ids']).issubset(set(ids.values())))
        self.assertNotIn('finding_uid', selected['findings'][0])

    def test_retired_source_modules_and_helper_are_removed(self):
        for name in RETIRED:
            self.assertFalse((ROOT / Path(*name.split('.')).with_suffix('.py')).exists(), name)
        self.assertFalse((ROOT / 'tools' / 'build_app_builder_export.py').exists())

    def test_retained_runtime_has_no_retired_imports_or_output_keys(self):
        forbidden = ('from .dashboard_', 'from Core.dashboard_', 'dashboard_schema_version',
                     'dashboard_json_path', 'dashboard_json_folder', 'app_builder_folder')
        for path in (ROOT / 'Core').glob('*.py'):
            if path.stem in {name.rsplit('.', 1)[1] for name in RETIRED}:
                continue
            with self.subTest(path=path.name):
                text = path.read_text(encoding='utf-8')
                for marker in forbidden:
                    self.assertNotIn(marker, text)


if __name__ == '__main__':
    unittest.main()
