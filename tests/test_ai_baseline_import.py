"""Compliance Manager reference evidence must never substitute for technical tests."""
import contextlib
import copy
from datetime import datetime
import io
from pathlib import Path
import shutil
import tempfile
import unittest
from provenance_fixture import fictional_publication

from openpyxl import Workbook, load_workbook

from Core.ai_baseline_import import load_ai_baseline_export, load_ai_baseline_exports
from Core.customer_report import _ai_baseline_review
from Core.portal_report_import import route_portal_reports

TENANT = '11111111-1111-1111-1111-111111111111'
OTHER = '22222222-2222-2222-2222-222222222222'


def make_export(path, tenant=TENANT, export_date='2026-10-12', assessment='AI Baseline',
                test_date=None, regulation=None):
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    summary.append(['Assessment name', assessment])
    if regulation:
        summary.append(['Regulation name', regulation])
    if tenant:
        summary.append(['Tenant ID', tenant])
    if export_date:
        summary.append(['Export date', export_date])
    for title, action, status, result in (
            ('Customer actions', 'Restrict <script> access', 'Implemented', 'Passed'),
            ('Microsoft actions', 'Protect service infrastructure', 'Implemented', 'Passed')):
        sheet = book.create_sheet(title)
        sheet.append(['Improvement action', 'Implementation status', 'Test result', 'Test date', 'Action ID', 'Owner'])
        sheet.append([action, status, result, test_date or datetime(2026, 10, 11, 12), title, 'owner@example.invalid'])
    book.save(path)
    return path


@fictional_publication
class AIBaselineImportTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def test_native_dates_original_fields_and_separate_responsibilities(self):
        path = make_export(self.root / 'unrelated filename.xlsx')
        original = path.read_bytes()
        result = load_ai_baseline_export(path, TENANT)
        self.assertTrue(result['identity_verified'])
        self.assertEqual(result['export_date'], '2026-10-12')
        self.assertEqual([row['Responsibility'] for row in result['rows']], ['Customer', 'Microsoft'])
        self.assertEqual(result['rows'][0]['Test date'], '2026-10-11T12:00:00')
        self.assertEqual(result['rows'][0]['Original fields']['Owner'], 'owner@example.invalid')
        self.assertEqual(result['rows'][0]['Source row'], 2)
        self.assertEqual(path.read_bytes(), original)

    def test_missing_tenant_and_dates_remain_unverified_without_filename_inferences(self):
        path = make_export(self.root / f'AI Baseline {TENANT} 2026-10-12.xlsx', tenant='',
                           export_date='', test_date='11/10/2026')
        result = load_ai_baseline_export(path, TENANT)
        self.assertFalse(result['identity_verified'])
        self.assertEqual(result['tenant_id'], '')
        self.assertEqual(result['export_date'], '')
        self.assertEqual(result['rows'][0]['Test date'], '')
        self.assertEqual(result['rows'][0]['Original fields']['Test date'], '11/10/2026')

    def test_tenant_mismatch_and_multiple_tenants_are_rejected(self):
        first = make_export(self.root / 'a.xlsx')
        second = make_export(self.root / 'b.xlsx', tenant=OTHER)
        with self.assertRaisesRegex(ValueError, 'does not match'):
            load_ai_baseline_export(first, OTHER)
        with self.assertRaisesRegex(ValueError, 'multiple tenants'):
            load_ai_baseline_exports([first, second])
        book = load_workbook(first)
        book['Summary'].append(['Tenant ID', OTHER])
        book.save(first)
        book.close()
        with self.assertRaisesRegex(ValueError, 'multiple tenants'):
            load_ai_baseline_export(first)

    def test_custom_assessment_name_with_baseline_regulation_is_accepted(self):
        path = make_export(self.root / 'custom.xlsx', assessment='Customer pilot', regulation='AI Baseline')
        result = load_ai_baseline_export(path)
        self.assertEqual(result['rows'][0]['Assessment'], 'Customer pilot')
        self.assertEqual(result['rows'][0]['Regulation'], 'AI Baseline')

    def test_content_deduplication_preserves_distinct_dated_snapshots(self):
        first = make_export(self.root / 'a.xlsx', export_date='2026-10-10')
        duplicate = self.root / 'copy.xlsx'
        shutil.copyfile(first, duplicate)
        second = make_export(self.root / 'b.xlsx')
        result = load_ai_baseline_exports([first, duplicate, second], TENANT)
        self.assertEqual(len(result['rows']), 4)
        self.assertEqual(len(result['sources']), 2)
        self.assertEqual(result['receipt'][1]['Status'], 'duplicate')
        self.assertEqual(result['test_results'], {'Passed': 4})

    def test_rejects_unrelated_or_ambiguous_exports_and_invalid_workbooks(self):
        wrong = make_export(self.root / 'AI Baseline.xlsx', assessment='ISO 27001')
        with self.assertRaisesRegex(ValueError, 'does not identify'):
            load_ai_baseline_export(wrong)
        ambiguous = make_export(self.root / 'ambiguous.xlsx')
        book = load_workbook(ambiguous)
        book['Customer actions'].cell(1, 7, 'Implementation state')
        book.save(ambiguous)
        book.close()
        with self.assertRaisesRegex(ValueError, 'ambiguous'):
            load_ai_baseline_export(ambiguous)
        invalid = self.root / 'bad.xlsx'
        invalid.write_text('Not an Excel workbook', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'readable Excel'):
            load_ai_baseline_export(invalid)

    def test_empty_export_does_not_become_assurance(self):
        path = make_export(self.root / 'empty.xlsx')
        book = load_workbook(path)
        for name in ('Customer actions', 'Microsoft actions'):
            book[name].delete_rows(2)
        book.save(path)
        book.close()
        result = load_ai_baseline_exports([path], TENANT)
        self.assertFalse(result['available'])
        self.assertEqual(result['receipt'][0]['Status'], 'no_rows')

    def test_sheet_local_metadata_does_not_inherit_another_assessment_identity(self):
        path = make_export(self.root / 'mixed.xlsx')
        book = load_workbook(path)
        sheet = book.create_sheet('Other assessment')
        sheet.append(['Assessment name', 'ISO 27001'])
        sheet.append(['Improvement action', 'Test result'])
        sheet.append(['Unrelated compliance action', 'Passed'])
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'does not identify an AI Baseline'):
            load_ai_baseline_export(path)

    def test_ai_only_plan_registers_input_without_reading_or_writing_exports(self):
        from unittest.mock import patch
        from Core.cli_parser import parse_arguments
        from Core.run_plan import build_run_plan
        from provenance_fixture import fictional_revision
        export = self.root / 'not-created.xlsx'
        with patch('sys.argv', ['main.py', '--offline', '--ai-baseline-export', str(export)]):
            args = parse_arguments(None, [])
        with patch('Core.ai_baseline_import.load_ai_baseline_exports', side_effect=AssertionError('preflight cannot import evidence')):
            plan = build_run_plan(args, revision=fictional_revision())
        self.assertNotIn('offline_evidence_required', plan['validation_errors'])
        self.assertEqual(plan['sources']['imports_supplied']['ai_baseline_export'], 1)
        self.assertFalse(export.exists())

    def test_cli_accepts_explicit_baseline_input_without_live_authentication(self):
        from tests.test_offline_report import OfflineReportTests
        export = make_export(self.root / 'explicit.xlsx')
        report = OfflineReportTests().run_cli(['--mode', 'offline', '--run-type', 'Standalone',
                   '--tenant-id', TENANT,
                   '--ai-baseline-export', str(export), '--tenant-name', 'Fictional customer'], self.root)
        self.assertIn('Compliance Manager AI Baseline reference', report.read_text(encoding='utf-8'))

    def test_routing_and_html_preserve_statuses_escape_content_and_do_not_mutate(self):
        path = make_export(self.root / 'arbitrary.xlsx')
        self.assertEqual(route_portal_reports([self.root])['ai_baseline'], [str(path)])
        review = load_ai_baseline_exports([path], TENANT)
        before = copy.deepcopy(review)
        html = _ai_baseline_review({'ai_baseline_review': review})
        self.assertIn('Implemented', html)
        self.assertIn('Passed', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('<script>', html)
        self.assertIn('do not pass technical controls', html)
        self.assertEqual(review, before)

    def test_explicit_input_survives_portable_package_and_restore(self):
        from unittest.mock import patch
        from Core.assessment_package import restore_arguments
        from Core.cli_parser import parse_arguments
        from Core.offline_collection import empty_service_results, load_collection, save_collection
        export = make_export(self.root / 'original.xlsx')
        collection = Path(save_collection(self.root / 'portable' / 'collection.json', tenant_id=TENANT,
                         tenant_name='Example', service_results=empty_service_results(),
                         supplemental_inputs={'ai_baseline_export': [str(export)]}))
        with patch('sys.argv', ['main.py', '--mode', 'offline', '--collection-input', str(collection)]):
            args = parse_arguments(None, [])
        export.unlink()
        restored = restore_arguments(args, load_collection(collection))
        self.assertEqual(len(restored.ai_baseline_export), 1)
        self.assertEqual(len(load_ai_baseline_export(restored.ai_baseline_export[0], TENANT)['rows']), 2)

    def test_report_build_does_not_change_readiness_and_both_workbooks_include_actions(self):
        from Core.offline_collection import empty_service_results
        from Core.processor import process_and_print_all_information
        export = make_export(self.root / 'baseline.xlsx')
        options = dict(tenant_name='Fictional customer', expected_tenant_id=TENANT,
                       collection_context={'mode': 'offline', 'evaluation_date': '2026-10-12'})
        with contextlib.redirect_stdout(io.StringIO()):
            without = process_and_print_all_information(**empty_service_results(), **options,
                                                       output_dir=self.root / 'without')
            with_import = process_and_print_all_information(**empty_service_results(), **options,
                              output_dir=self.root / 'with', ai_baseline_exports=[str(export)])
        before = without['evidence_bundle']['assessment_result']
        after = with_import['evidence_bundle']['assessment_result']
        for key in ('control_results', 'recommendations', 'decision', 'rationale', 'readiness_review'):
            self.assertEqual(after.get(key), before.get(key), key)
        self.assertEqual(len(after['supplemental_reviews']['ai_baseline_review']['rows']), 2)
        for key in ('excel_path', 'technical_excel_path'):
            book = load_workbook(with_import[key], read_only=True, data_only=True)
            try:
                self.assertIn('AI Baseline Actions', book.sheetnames)
                self.assertIn('AI Baseline Sources', book.sheetnames)
                self.assertIn('Evidence Requests', book.sheetnames)
                self.assertIn('DSPM Billing Guidance', book.sheetnames)
                self.assertEqual(book['AI Baseline Actions'].max_row, 3)
            finally:
                book.close()
        body = Path(with_import['html_path']).read_text(encoding='utf-8')
        self.assertIn('Compliance Manager AI Baseline reference', body)
        self.assertIn('Restrict &lt;script&gt; access', body)
        from Core.assessment_serialization import write_assessment_result
        from Core.assessment_replay import render_snapshot
        snapshot = self.root / 'snapshot.json'
        write_assessment_result(snapshot, after)
        with contextlib.redirect_stdout(io.StringIO()):
            rendered = render_snapshot(snapshot, self.root / 'render')
        replay = Path(rendered['html_path']).read_text(encoding='utf-8')
        self.assertIn('Compliance Manager AI Baseline reference', replay)
        self.assertIn('DSPM scope and billing', replay)
        self.assertEqual(rendered['evidence_bundle']['ai_baseline_review']['rows'],
                         after['supplemental_reviews']['ai_baseline_review']['rows'])


if __name__ == '__main__':
    unittest.main()
