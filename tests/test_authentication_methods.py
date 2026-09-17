"""Registered credentials, preferred second factors and enforced strengths are distinct."""

import contextlib
import io
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

from Core.authentication_methods import authentication_method_report, legacy_registration_summary, summarize_registrations
from Core.customer_report import _authentication_methods
from Core.offline_collection import empty_service_results, load_collection, save_collection


def registration(identifier, methods, **values):
    return {'id': identifier, 'userPrincipalName': identifier + '@example.com', 'userDisplayName': 'Private fixture name',
            'methodsRegistered': methods, 'isMfaRegistered': True, 'isMfaCapable': True,
            'isPasswordlessCapable': False, 'isAdmin': False, 'userType': 'member',
            'isSystemPreferredAuthenticationMethodEnabled': False,
            'userPreferredMethodForSecondaryAuthentication': 'sms',
            'lastUpdatedDateTime': '2026-09-14T10:00:00Z', **values}


class AuthenticationMethodsTests(unittest.TestCase):
    def test_email_push_passwordless_and_phishing_resistance_are_distinct(self):
        rows = [registration('email-only', ['email'], isMfaRegistered=False, isMfaCapable=False),
                registration('sms-only', ['mobilePhone', 'email']),
                registration('push', ['microsoftAuthenticatorPush', 'softwareOneTimePasscode', 'mobilePhone']),
                registration('phone-signin', ['microsoftAuthenticatorPasswordless']),
                registration('key', ['passKeyDeviceBoundAuthenticator', 'fido2']),
                registration('hello', ['windowsHelloForBusiness'], isAdmin=True),
                registration('cert', ['certificateBasedAuthentication', 'mobilePhone'])]
        report = summarize_registrations(rows, {'available': True})
        self.assertEqual(report['metrics']['phishing_resistant_registered'], 2)
        self.assertEqual(report['metrics']['passwordless_registered'], 3)
        self.assertEqual(report['metrics']['phone_only_mfa_registered'], 1)
        self.assertEqual(report['metrics']['email_registered'], 2)
        self.assertEqual(report['metrics']['methods_need_review'], 1)
        self.assertEqual(report['metrics']['resistant_inventory_known'], 6)
        self.assertIn('not workforce MFA', next(row for row in report['method_rows'] if row['Method'] == 'Email')['Strength / purpose'])

    def test_system_preference_overrides_user_default_and_aliases_count_once(self):
        rows = [registration('preferred-app', ['mobilePhone', 'microsoftAuthenticatorPush'],
                    isSystemPreferredAuthenticationMethodEnabled=True, systemPreferredAuthenticationMethods=['phoneAppNotification', 'push']),
                registration('preferred-phone', ['mobilePhone'], isSystemPreferredAuthenticationMethodEnabled=True,
                    systemPreferredAuthenticationMethods=['sms', 'voice']),
                registration('manual-app', ['mobilePhone', 'softwareOneTimePasscode'], userPreferredMethodForSecondaryAuthentication='oath')]
        report = summarize_registrations(rows)
        self.assertEqual(report['metrics']['user_phone_preferred'], 2)
        self.assertEqual(report['metrics']['phone_preferred'], 1)
        self.assertEqual(report['metrics']['current_preference_needs_review'], 0)
        push = next(row for row in report['preference_rows'] if row['Method'] == 'Authenticator push')
        self.assertEqual(push['System-preferred users'], 1)
        self.assertEqual(push['Preferred users in report'], 1)
        self.assertEqual(sum(row['Preferred users in report'] for row in report['preference_rows']), 4)

    def test_missing_preference_is_unknown_and_never_falls_back_from_system_mode(self):
        rows = [registration('no-system-detail', ['mobilePhone'], isSystemPreferredAuthenticationMethodEnabled=True),
                registration('mode-unknown', None, isSystemPreferredAuthenticationMethodEnabled=None),
                registration('no-default', [], userPreferredMethodForSecondaryAuthentication=None)]
        report = summarize_registrations(rows)
        metric = next(row for row in report['summary_rows'] if row['Metric'] == 'SMS/voice preferred in report')
        self.assertEqual(metric['Value'], 'Unknown')
        self.assertEqual(metric['Unknown users'], 3)
        self.assertEqual(report['population_rows'][0]['SMS/voice preferred'], 'Unknown')
        self.assertEqual(report['metrics']['method_inventory_known'], 2)

    def test_unknown_and_legacy_methods_are_not_silently_safe(self):
        rows = [registration('unknown', ['mobilePhone', 'unknownFutureValue']),
                registration('legacy', ['microsoftAuthenticator']),
                registration('future', ['<script>alert(1)</script>'], userPreferredMethodForSecondaryAuthentication='futureMethod')]
        report = summarize_registrations(rows)
        self.assertEqual(report['metrics']['phone_only_mfa_registered'], 0)
        self.assertEqual(report['metrics']['resistant_inventory_known'], 0)
        html = _authentication_methods({'authentication_methods': report})
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertIn('Unknown', html)

    def test_sdk_shape_matches_raw_http_shape(self):
        raw = registration('example', ['windowsHelloForBusiness'], isAdmin=True)
        import re
        sdk = SimpleNamespace(**{re.sub(r'(?<!^)(?=[A-Z])', '_', key).lower(): value for key,value in raw.items()})
        self.assertEqual(summarize_registrations([raw]), summarize_registrations([sdk]))

    def test_future_preferences_and_unavailable_columns_remain_unknown(self):
        rows = [registration('future', ['mobilePhone'], userPreferredMethodForSecondaryAuthentication='unknownFutureValue')]
        report = summarize_registrations(rows)
        self.assertEqual(report['metrics']['phone_preference_known'], 0)
        metric = next(row for row in report['summary_rows'] if row['Metric'] == 'SMS/voice preferred in report')
        self.assertEqual(metric['Value'], 'Unknown')
        self.assertEqual(report['population_rows'][0]['SMS/voice preferred'], 'Unknown')
        self.assertEqual(next(row for row in report['summary_rows'] if row['Metric'] == 'User-selected SMS/voice default')['Value'], 'Unknown')
        self.assertEqual(report['preference_rows'][0]['System-preferred users'], 'Unknown')
        self.assertIn('preferred methods needing classification review', _authentication_methods({'authentication_methods': report}))

    def test_duplicate_and_conflicting_users_are_not_counted_twice(self):
        row = registration('same', ['mobilePhone'])
        report = summarize_registrations([row, dict(row), registration('conflict', ['mobilePhone']),
                                          registration('conflict', ['fido2'])], {'available': True})
        self.assertEqual(report['total_users'], 2)
        self.assertEqual(report['duplicates_removed'], 1)
        self.assertEqual(report['conflicting_users'], 1)
        self.assertEqual(report['metrics']['method_inventory_known'], 1)
        self.assertFalse(report['complete'])

    def test_partial_source_and_unknown_dates_remain_qualified(self):
        report = summarize_registrations([registration('dated', [], lastUpdatedDateTime=None)],
                                          {'availability_status': 'available', 'truncated': True})
        self.assertEqual(report['source_state'], 'partial')
        self.assertFalse(report['complete'])
        self.assertEqual(report['dates_unknown'], 1)
        self.assertEqual(report['updated_from'], '')

    def test_collector_summary_uses_current_names_and_unique_passwordless_users(self):
        rows = [registration('both', ['microsoftAuthenticatorPush', 'microsoftAuthenticatorPasswordless', 'windowsHelloForBusiness']),
                registration('push-only', ['microsoftAuthenticatorPush']), registration('key', ['fido2', 'passKeySynced'])]
        summary = legacy_registration_summary(rows)
        self.assertEqual(summary['passwordless_enabled'], 2)
        self.assertEqual(summary['methods']['microsoftAuthenticator'], 2)
        self.assertEqual(summary['methods']['microsoftAuthenticatorPasswordless'], 1)
        self.assertEqual(summary['methods']['fido2'], 1)
        self.assertEqual(summary['methods']['windowsHello'], 1)

    def test_legacy_summary_alone_does_not_invent_methods_or_defaults(self):
        client = SimpleNamespace(auth_summary={'total_users': 100, 'methods': {'phone': 0}},
                                 collection_status={'auth_methods': {'available': True}})
        report = authentication_method_report(client)
        self.assertFalse(report['available'])
        self.assertIn('were not available', _authentication_methods({'authentication_methods': report}))

    def test_report_and_workbook_keep_aggregate_values_without_user_identities(self):
        from Core.evidence_layer import build_evidence_bundle
        from Core.export_recommendations import export_to_excel
        from openpyxl import load_workbook
        results = empty_service_results()
        client = SimpleNamespace(auth_methods_registration=[registration('do-not-render-identity', ['mobilePhone'], isAdmin=True)],
                                 collection_status={'auth_methods': {'availability_status': 'available'}})
        results['entra_info']['_client'] = client
        results['entra_info']['available'] = True
        row = {'Service': 'Entra', 'Feature': 'MFA registration', 'Observation': 'Registration was collected.', 'Recommendation': '', 'Status': 'Reference'}
        bundle = build_evidence_bundle([row], **results)
        html = _authentication_methods(bundle)
        self.assertNotIn('do-not-render-identity', json.dumps(bundle))
        self.assertNotIn('Private fixture name', html)
        self.assertIn('Phone-only MFA registration', html)
        with tempfile.TemporaryDirectory() as directory, contextlib.redirect_stdout(io.StringIO()):
            output = export_to_excel(bundle['recommendations'], filename=str(Path(directory)/'methods.xlsx'), evidence_bundle=bundle)
            book = load_workbook(output, read_only=True)
            try:
                for sheet in ('Authentication Coverage', 'Authentication Methods', 'MFA Preferences', 'MFA Populations'):
                    self.assertIn(sheet, book.sheetnames)
                    self.assertNotIn('do-not-render-identity', str(list(book[sheet].values)))
            finally:
                book.close()
            collection = save_collection(Path(directory)/'collection.json', tenant_id='11111111-1111-4111-8111-111111111111',
                tenant_name='Synthetic tenant', service_results=results)
            saved = load_collection(collection)['service_results']['entra_info']['_client']
            self.assertEqual(authentication_method_report(client), authentication_method_report(saved))

    def test_saved_passwordless_claim_is_recomputed_and_not_treated_as_enforcement(self):
        from Core.saved_control_checks import qualify_saved_recommendations
        client = SimpleNamespace(auth_methods_registration=[registration('push', ['microsoftAuthenticatorPush'])],
                                 collection_status={'auth_methods': {'available': True}})
        old = {'Observation': '100% of users use passwordless authentication, providing phishing-resistant protection',
               'Status': 'Success', 'Disposition': 'Assurance'}
        row = qualify_saved_recommendations([old], 'Entra', client)[0]
        self.assertIn('0 of 1', row['Observation'])
        self.assertEqual(row['Disposition'], 'Reference')


class AuthenticationCollectorTests(unittest.IsolatedAsyncioTestCase):
    async def test_existing_registration_request_keeps_preferences_and_normalizes_methods(self):
        from Core.get_entra_client import get_entra_client
        rows = [registration('fictional', ['windowsHelloForBusiness', 'microsoftAuthenticatorPush'],
                             isSystemPreferredAuthenticationMethodEnabled=True, systemPreferredAuthenticationMethods=['phoneAppNotification'])]
        async def collection(path, **options):
            return {'available': True, 'availability_status': 'available', 'value': rows if 'userRegistrationDetails' in path else []}
        with patch('Core.get_entra_client._fetch_graph_collection_via_http', side_effect=collection), \
             patch('Core.get_entra_client._fetch_graph_object_via_http', AsyncMock(return_value={})), \
             patch('Core.get_entra_client._get_graph_token_roles', return_value=set()), \
             patch('socket.create_connection', side_effect=AssertionError('No tenant access')), contextlib.redirect_stdout(io.StringIO()):
            result = await get_entra_client(SimpleNamespace(), 'fictional')
        self.assertEqual(result.auth_summary['passwordless_enabled'], 1)
        self.assertEqual(result.auth_methods_registration[0]['systemPreferredAuthenticationMethods'], ['phoneAppNotification'])
        self.assertEqual(authentication_method_report(result)['metrics']['phone_preferred'], 0)


if __name__ == '__main__':
    unittest.main()
