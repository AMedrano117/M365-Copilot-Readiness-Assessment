"""Execution, portable snapshots, CLI, and shared renderer integration."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from test_assessment_run_workflow import RunFixture,TENANT,DAY,result
from test_assessment_run_outputs import parse


class DeltaWorkflowTests(RunFixture):
    def test_initial_and_standalone_record_not_evaluated(self):
        from Core.assessment_runs import prepare_run,complete_run
        for mode in ('Initial','Standalone'):
            with self.subTest(mode=mode):
                run=prepare_run(mode,TENANT,evaluated_at=DAY,history_path=self.root/mode/'history.json' if mode=='Initial' else None)
                current=result(run.identity);complete_run(current,run,self.root/mode/'package')
                self.assertEqual(current['lifecycle']['Outcome'],'NotEvaluated')
                self.assertEqual(current['lifecycle']['Records'],[])

    def test_reassessment_history_appends_delta_without_changing_baseline(self):
        from Core.assessment_runs import prepare_run,complete_run
        from Core.assessment_history import read_history
        initial,baseline=self.initial();history_before=read_history(self.history)
        source=next((self.root/'initial'/'assessment-results').glob('*.json'));bytes_before=source.read_bytes()
        run=prepare_run('Reassessment',TENANT,evaluated_at='2026-09-16',history_path=self.history,
            assessment_id=initial.identity['AssessmentId'],baseline_run_id=initial.identity['RunId'],catalog_version='catalog-1')
        current=result(run.identity);path=complete_run(current,run,self.root/'current')
        self.assertTrue(current['lifecycle']['Records']);self.assertEqual(source.read_bytes(),bytes_before)
        history=read_history(self.history);self.assertEqual(history['Runs'][0],history_before['Runs'][0])
        self.assertEqual(history['Runs'][1]['Lifecycle']['Summary'],current['lifecycle']['Summary'])
        self.assertTrue(Path(path).is_file());self.assertEqual(history['SchemaVersion'],'1.0.0')

    def test_explicit_disabled_reassessment_is_recorded_and_persisted(self):
        from Core.assessment_runs import prepare_run,complete_run
        initial,_=self.initial()
        run=prepare_run('Reassessment',TENANT,evaluated_at='2026-09-16',history_path=self.history,
            assessment_id=initial.identity['AssessmentId'],baseline_run_id=initial.identity['RunId'],catalog_version='catalog-1',delta_enabled=False)
        current=result(run.identity);complete_run(current,run,self.root/'current')
        self.assertFalse(current['run_context']['DeltaEnabled']);self.assertEqual(current['lifecycle']['Records'],[])

    def test_legacy_snapshot_roundtrip_does_not_create_lifecycle(self):
        from Core.assessment_runs import prepare_run
        from Core.assessment_serialization import write_assessment_result,read_assessment_result
        current=result(prepare_run('Standalone',TENANT,evaluated_at=DAY).identity)
        path=self.root/'old.json';write_assessment_result(path,current)
        self.assertNotIn('lifecycle',read_assessment_result(path))

    def test_package_tree_movement_preserves_baseline_selection_and_lifecycle(self):
        from Core.assessment_runs import prepare_run,complete_run
        from Core.assessment_history import select_baseline
        initial,_=self.initial()
        run=prepare_run('Reassessment',TENANT,evaluated_at='2026-09-16',history_path=self.history,
            assessment_id=initial.identity['AssessmentId'],baseline_run_id=initial.identity['RunId'],catalog_version='catalog-1')
        current=result(run.identity);complete_run(current,run,self.root/'current')
        moved=self.root/'moved';moved.mkdir()
        # Move individually named synthetic children, never a computed recursive tree.
        for name in ('initial','current','assessment-history.json','run-seeds'):
            (self.root/name).rename(moved/name)
        restored=select_baseline(moved/'assessment-history.json',assessment_id=initial.identity['AssessmentId'],
            environment_id=initial.identity['PrimaryEnvironmentId'],baseline_run_id=run.identity['RunId'])
        self.assertEqual(restored['lifecycle'],current['lifecycle'])

    def test_renderer_range_remapping_does_not_mutate_lifecycle_references(self):
        from Core.raw_evidence import split_evidence_sheets
        snapshot={'lifecycle':{'Records':[{'SourceFile':"'Raw'!A2:B4"}]}}
        original=copy.deepcopy(snapshot)
        sheets={'raw':{'title':'Raw','rows':[{'id':n} for n in range(3)]}}
        split_evidence_sheets({'sheets':sheets},snapshot,max_rows=1)
        self.assertEqual(snapshot,original)


class DeltaCliTests(unittest.TestCase):
    def test_disable_requires_explicit_reassessment_and_has_no_interactive_prompt(self):
        valid=['--mode','offline','--run-type','Reassessment','--assessment-id','AST-fiction',
            '--baseline-run-id','RUN-fiction','--assessment-history','history.json']
        self.assertEqual(parse(*valid,'--delta-mode','disabled').delta_mode,'disabled')
        for options in [('--delta-mode','enabled'),('--run-type','Initial','--delta-mode','disabled'),
                ('--replay-snapshot','result.json','--delta-mode','enabled')]:
            with self.subTest(options=options),self.assertRaises(SystemExit):parse(*options)


class DeltaRendererTests(unittest.TestCase):
    def test_technical_fields_and_long_values_are_preserved(self):
        from Core.lifecycle_presentation import detail_rows
        from Core.workbook_layout import _write
        from openpyxl import Workbook
        records=[{'EntityId':'PCT-fiction','EntityType':'control','State':'Unchanged'},
            {'EntityId':'POB-fiction','EntityType':'metric','State':'Changed','MetricId':'metric-explicit',
             'Qualifications':['x'*40000]}]
        book=Workbook();self.addCleanup(book.close)
        sheet=_write(book,'Lifecycle Records',detail_rows({'lifecycle':{'Records':records}}))
        columns=[cell.value for cell in sheet[1]]
        self.assertIn('MetricId',columns);self.assertIn('Qualifications (continued 2)',columns)
        values=dict(zip(columns,[cell.value for cell in sheet[3]]))
        self.assertEqual(json.loads(values['Qualifications']+values['Qualifications (continued 2)']),records[1]['Qualifications'])

    def test_replay_preserves_lifecycle_and_renders_same_records_without_calculation(self):
        from Core.assessment_runs import prepare_run,complete_run
        from Core.assessment_result import build_assessment_result
        from Core.assessment_replay import render_snapshot
        from Core.assessment_serialization import read_assessment_result
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            initial=prepare_run('Initial',TENANT,evaluated_at=DAY,history_path=root/'assessment-history.json')
            def built(seed):return build_assessment_result([],{'collection_context':{'identity':seed,'evaluation_date':DAY}},evaluation_date=DAY,expected_tenant_id=TENANT)
            complete_run(built(initial.identity),initial,root/'initial')
            run=prepare_run('Reassessment',TENANT,evaluated_at=DAY,history_path=root/'assessment-history.json',
                assessment_id=initial.identity['AssessmentId'],baseline_run_id=initial.identity['RunId'])
            current=built(run.identity);source=Path(complete_run(current,run,root/'current'))
            original=source.read_bytes(); lifecycle=copy.deepcopy(current['lifecycle'])
            with patch('Core.assessment_delta.evaluate_delta',side_effect=AssertionError('Renderer recalculation')):
                outputs=render_snapshot(source,root/'renders',snapshot_output=root/'copy.json')
            self.assertEqual(source.read_bytes(),original);self.assertEqual(read_assessment_result(root/'copy.json')['lifecycle'],lifecycle)
            from openpyxl import load_workbook
            for name in ('excel_path','technical_excel_path'):
                book=load_workbook(outputs[name]);self.addCleanup(book.close)
                self.assertIn('Reassessment',book.sheetnames);self.assertIn('Lifecycle Records',book.sheetnames)
                self.assertEqual(book['Lifecycle Records'].max_row,len(lifecycle['Records'])+1)
            for name in ('html_path','summary_html_path'):
                html=Path(outputs[name]).read_text(encoding='utf-8')
                self.assertIn('Missing current evidence does not indicate improvement',html)
                self.assertIn(initial.identity['RunId'],html)
            html=Path(outputs['html_path']).read_text(encoding='utf-8')
            self.assertIn('Lifecycle records',html)
