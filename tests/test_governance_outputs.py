"""Governance persistence and deliverables under the OS-temp-only test guard."""
import copy
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
from test_governance_decisions import draft,advance,NOW,LATER,fixture


class SerializationTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)

    def test_log_atomic_roundtrip_and_compare_and_swap(self):
        from Core.governance_history import write_log,read_log
        r,log,key,p=draft();path=self.root/'decisions.json';write_log(path,log)
        before=path.read_bytes();updated=advance(r,log,key,p)
        with self.assertRaises(ValueError):write_log(path,updated,expected_hash='stale')
        self.assertEqual(path.read_bytes(),before)
        write_log(path,updated,expected_hash=log['Integrity'])
        self.assertEqual(read_log(path),updated)
        with self.assertRaises(ValueError):write_log(path,log,expected_hash=updated['Integrity'])

    def test_integrity_failure_blocks_read(self):
        from Core.governance_history import write_log,read_log
        r,log,key,p=draft();path=self.root/'decisions.json';write_log(path,log)
        bad=json.loads(path.read_text());bad['Events'][0]['Actor']='changed';path.write_text(json.dumps(bad))
        with self.assertRaises(ValueError):read_log(path)

    def test_failed_atomic_update_and_lock_preserve_existing_log(self):
        from Core.governance_history import write_log
        r,log,key,p=draft();path=self.root/'decisions.json';write_log(path,log);before=path.read_bytes();updated=advance(r,log,key,p)
        with patch('Core.governance_history.atomic_write',side_effect=OSError('Synthetic disk failure')):
            with self.assertRaises(OSError):write_log(path,updated,expected_hash=log['Integrity'])
        self.assertEqual(path.read_bytes(),before);self.assertFalse(path.with_name(path.name+'.lock').exists())
        lock=path.with_name(path.name+'.lock');lock.write_text('synthetic lock')
        with self.assertRaises(ValueError):write_log(path,updated,expected_hash=log['Integrity'])
        self.assertEqual(path.read_bytes(),before)

    def test_snapshot_preserves_separate_states_and_decision_history(self):
        from Core.governance import attach
        from Core.assessment_serialization import write_assessment_result,read_assessment_result
        r,log,key,p=draft();log=advance(r,log,key,p);out=attach(r,log,as_of=NOW,locator='governance/decisions.json')
        path=self.root/'snapshot.json';write_assessment_result(path,out)
        restored=read_assessment_result(path)
        self.assertEqual(restored['governance'],out['governance']);self.assertEqual(restored['identity'],r['identity'])

    def test_projected_summary_or_state_edits_block_publication(self):
        from Core.governance import attach
        from Core.assessment_serialization import write_assessment_result
        r,log,key,p=draft();out=attach(r,advance(r,log,key,p),as_of=NOW,locator='governance/decisions.json')
        for field in ('Summary','Records'):
            bad=copy.deepcopy(out);bad['governance'][field]={} if field=='Summary' else []
            with self.subTest(field=field),self.assertRaises(ValueError):write_assessment_result(self.root/'invalid.json',bad)

    def test_portable_tree_copy_preserves_log_and_snapshot(self):
        from Core.governance import attach
        from Core.governance_history import write_log,read_log
        from Core.assessment_serialization import write_assessment_result,read_assessment_result
        r,log,key,p=draft();log=advance(r,log,key,p);package=self.root/'package';package.mkdir()
        write_log(package/'governance/decisions.json',log)
        write_assessment_result(package/'snapshot.json',attach(r,log,as_of=NOW,locator='governance/decisions.json'))
        moved=self.root/'moved';shutil.copytree(package,moved)
        self.assertEqual(read_log(moved/'governance/decisions.json'),log)
        self.assertEqual(read_assessment_result(moved/'snapshot.json')['governance']['Records'][0]['DecisionId'],key)

    def test_legacy_fields_remain_unresolved_without_invented_authority(self):
        from Core.assessment_serialization import write_assessment_result,read_assessment_result
        r,_,_=fixture();r['recommendations'][0].update({'AcceptedRisk':'yes','Status':'Closed','Exception':'approved'})
        path=self.root/'legacy.json';write_assessment_result(path,r);out=read_assessment_result(path)
        self.assertNotIn('governance',out)
        self.assertTrue(any(d['severity']=='unresolved_legacy_governance_reference' for d in out['governance_diagnostics']))
        self.assertEqual(out['recommendations'],r['recommendations'])

    def test_history_reference_update_preserves_runs_and_snapshot_bytes(self):
        from test_assessment_run_workflow import result,TENANT,DAY
        from Core.assessment_runs import prepare_run,complete_run
        from Core.assessment_history import read_history,select_baseline
        from Core.governance import new_log,project
        from Core.governance_history import write_log,link_history
        history=self.root/'assessment-history.json';run=prepare_run('Initial',TENANT,evaluated_at=DAY,history_path=history)
        r=result(run.identity);source=Path(complete_run(r,run,self.root/'initial'));before=source.read_bytes();old=read_history(history)
        log=new_log(r);path=self.root/'governance/decisions.json';write_log(path,log)
        link_history(history,path,as_of=NOW,expected_hash=old['Integrity'])
        updated=read_history(history);self.assertEqual(updated['Runs'],old['Runs']);self.assertEqual(updated['SchemaVersion'],'1.0.0')
        self.assertEqual(source.read_bytes(),before);self.assertEqual(updated['Governance']['Summary'],project(log,as_of=NOW)['Summary'])
        select_baseline(history,assessment_id=r['identity']['AssessmentId'],environment_id=r['identity']['PrimaryEnvironmentId'],baseline_run_id=r['identity']['RunId'])


class RendererTests(unittest.TestCase):
    def test_recorded_projection_replay_never_transitions(self):
        from Core.governance import attach
        from Core.assessment_serialization import write_assessment_result,read_assessment_result
        from Core.assessment_replay import render_snapshot
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);r,log,key,p=draft()
            from Core.assessment_result import build_assessment_result
            built=build_assessment_result([],{'collection_context':{'identity':r['identity']}},expected_tenant_id=r['tenant_id'],evaluation_date='2026-10-08')
            built.update(r);r=built
            out=attach(r,advance(r,log,key,p),as_of=NOW,locator='governance/decisions.json')
            source=root/'source.json';write_assessment_result(source,out);before=source.read_bytes()
            with patch('Core.governance.transition',side_effect=AssertionError('Renderer decision transition')):
                with patch('Core.governance.create_draft',side_effect=AssertionError('Renderer draft creation')):
                    outputs=render_snapshot(source,root/'outputs',snapshot_output=root/'copy.json')
            self.assertEqual(source.read_bytes(),before);self.assertEqual(read_assessment_result(root/'copy.json')['governance'],out['governance'])
            for name in ('excel_path','technical_excel_path'):
                book=load_workbook(outputs[name]);self.addCleanup(book.close)
                self.assertIn('Governance',book.sheetnames);self.assertIn('Decision Register',book.sheetnames)
                self.assertEqual(book['Decision Register'].max_row,2)
                if name=='technical_excel_path':self.assertIn('Decision Audit',book.sheetnames)
                else:self.assertIn('Governance treatment',[cell.value for cell in book['Findings'][1]])
            for name in ('html_path','summary_html_path'):
                text=Path(outputs[name]).read_text(encoding='utf-8');self.assertIn('Governance',text)
                self.assertIn('Risk acceptance is not remediation',text)
            self.assertIn('Governance audit history',Path(outputs['html_path']).read_text(encoding='utf-8'))

    def test_html_escaping_and_full_long_workbook_fields(self):
        from Core.governance_presentation import technical_html,detail_rows
        from Core.workbook_layout import _write
        from openpyxl import Workbook
        r,log,key,p=draft();from Core.governance import attach
        result=attach(r,log,as_of=NOW,locator='governance/decisions.json')
        result['governance']['Records'][0]['Rationale']='<script>'+('x'*40000)
        self.assertNotIn('<script>',technical_html(result))
        book=Workbook();self.addCleanup(book.close);sheet=_write(book,'Decision Register',detail_rows(result))
        self.assertIn('Rationale (continued 2)',[cell.value for cell in sheet[1]])

    def test_governance_locators_are_not_remapped_as_workbook_ranges(self):
        from Core.raw_evidence import split_evidence_sheets
        r={'governance':{'LogReference':{'Locator':"'Raw'!A2:B4"}}};before=copy.deepcopy(r)
        split_evidence_sheets({'sheets':{'raw':{'title':'Raw','rows':[{'id':n} for n in range(3)]}}},r,max_rows=1)
        self.assertEqual(r,before)


class CliTests(unittest.TestCase):
    def test_cli_separates_draft_approval_activation_and_exports(self):
        from Core.governance_cli import main
        from Core.assessment_serialization import write_assessment_result
        from Core.governance_history import read_log
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);r,d,p=fixture();snapshot=root/'snapshot.json';write_assessment_result(snapshot,r)
            data=root/'draft.json';data.write_text(json.dumps(d));policy=root/'policy.json';policy.write_text(json.dumps(p));log=root/'governance.json'
            common=['--snapshot',str(snapshot),'--log',str(log),'--at',NOW]
            self.assertEqual(main(['draft',*common,'--actor','operator','--input',str(data)]),0)
            key=read_log(log)['Events'][0]['DecisionId']
            for op in ('submit','approve','activate'):
                self.assertEqual(main([op,*common,'--actor','operator' if op=='submit' else 'approver','--decision-id',key,'--policy',str(policy)]),0)
            exported=root/'export.json'
            self.assertEqual(main(['export',*common,'--output',str(exported)]),0)
            self.assertTrue(exported.exists());self.assertEqual(len(read_log(log)['Events']),4)

    def test_unknown_bulk_approval_operation_rejected(self):
        from Core.governance_cli import main
        with self.assertRaises(SystemExit):main(['approve-all'])

    def test_cli_default_denies_approval_and_original_snapshot_is_immutable(self):
        from Core.governance_cli import main
        from Core.assessment_serialization import write_assessment_result
        from Core.governance_history import write_log
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);r,log,key,p=draft();log=advance(r,log,key,p,('submit',))
            snapshot=root/'snapshot.json';write_assessment_result(snapshot,r);path=root/'log.json';write_log(path,log)
            before=(snapshot.read_bytes(),path.read_bytes())
            self.assertEqual(main(['approve','--snapshot',str(snapshot),'--log',str(path),'--at',NOW,'--actor','operator','--decision-id',key]),2)
            self.assertEqual((snapshot.read_bytes(),path.read_bytes()),before)
