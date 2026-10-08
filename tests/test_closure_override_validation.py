"""Closure hardening uses synthetic, valid lifecycle graphs and explicit grants."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from test_governance_decisions import fixture, advance, NOW
from test_governance_semantics import lifecycle_pair


STATES = ('ResolvedByCurrentEvidence', 'Improved', 'Continuing', 'Indeterminate', 'NotReassessed')
INVALID_STATEMENTS = (True, False, 1, 0, 1.5, None, '', ' \t\n ', [], {}, ['text'], {'text': 'value'})
REQUIRED = ('ResidualRisk', 'Conditions', 'ReviewAt', 'Rationale', 'DecisionAuthority', 'AuthorityRole',
            'ValidationResult', 'EvidenceReferences', 'Population', 'ResourceScope',
            'ClosureCriteriaEvaluated', 'AccountableOwner', 'EffectiveAt', 'Period')


def closure_fixture(state='ResolvedByCurrentEvidence'):
    from Core.assessment_delta import evaluate_delta
    from Core.assessment_references import require_valid_assessment
    result, baseline, record, policy = lifecycle_pair(state != 'Continuing')
    if state == 'Improved':
        result['evidence'][0]['value'] = 2
        result['control_results'][0]['Status'] = 'Fail'
    elif state == 'Indeterminate':
        result['evidence'][0]['availability'] = 'partial'
    elif state == 'NotReassessed':
        result['evidence'][0]['availability'] = 'unavailable'
    result['lifecycle'] = evaluate_delta(result, baseline)
    require_valid_assessment(result)
    finding = next(row for row in result['lifecycle']['Records'] if row['EntityType'] == 'finding')
    assert finding['State'] == state
    record['ClosureOverride'] = True
    return result, baseline, record, policy


def start(result, record):
    from Core.governance import new_log, create_draft
    log = create_draft(new_log(result), result, record, actor='operator', at=NOW)
    return log, log['Events'][0]['DecisionId']


def omit(record, field):
    if field == 'Period':
        record['EvidenceReferences'][0].pop('Period', None)
    else:
        record.pop(field, None)


def edited_history(log, field, value=None, *, remove=False):
    """Model a retained older log, including a valid chain, without editing files."""
    from Core.assessment_history import seal
    from Core.assessment_identity import digest
    changed = copy.deepcopy(log)
    previous = None
    for event in changed['Events']:
        if event['Operation'] == 'draft':
            if remove:
                omit(event['Data'], field)
            else:
                event['Data'][field] = value
        if remove:
            omit(event['Record'], field)
        else:
            event['Record'][field] = value
        event['PreviousHash'] = previous
        event['Hash'] = digest({key: item for key, item in event.items() if key != 'Hash'})
        previous = event['Hash']
    return seal(changed)


class TypedStatementTests(unittest.TestCase):
    def test_required_statements_reject_nontext_and_blank_values(self):
        from Core.governance import transition, project
        for kind, field in (('ClosedByRemediation', 'ResidualRisk'), ('ClosedByRemediation', 'ValidationResult'),
                            ('AcceptedRisk', 'ResidualRisk')):
            for value in INVALID_STATEMENTS:
                with self.subTest(kind=kind, field=field, value=value):
                    result, record, policy = fixture(kind)
                    record[field] = copy.deepcopy(value)
                    log, key = start(result, record)
                    before = copy.deepcopy(log)
                    with self.assertRaises(ValueError):
                        transition(log, result, key, 'submit', actor='operator', at=NOW)
                    self.assertEqual(log, before)
                    view = project(log, as_of=NOW)
                    self.assertEqual(view['Records'][0]['WorkflowState'], 'Open')
                    self.assertFalse(view['Records'][0]['EffectiveActive'])

    def test_meaningful_statements_normalize_without_coercion_or_input_mutation(self):
        from Core.governance import project
        result, record, policy = fixture('ClosedByRemediation')
        record.update(ResidualRisk=' \tRetain monitoring.\n ', ValidationResult='  Validated conditionally.  ')
        before = copy.deepcopy((result, record))
        log, key = start(result, record)
        log = advance(result, log, key, policy)
        row = project(log, as_of=NOW)['Records'][0]
        self.assertEqual(row['ResidualRisk'], 'Retain monitoring.')
        self.assertEqual(row['ValidationResult'], 'Validated conditionally.')
        self.assertEqual(log['Events'][0]['Data']['ResidualRisk'], 'Retain monitoring.')
        self.assertEqual((result, record), before)

    def test_amend_normalizes_new_statements_and_preserves_previous_events(self):
        from Core.governance import transition
        result, record, policy = fixture('ClosedByRemediation')
        log, key = start(result, record)
        before = copy.deepcopy(log)
        material = {'ResidualRisk': '  Monitor the remaining risk.  ', 'ValidationResult': '\nReviewed.\t'}
        changed = transition(log, result, key, 'amend', actor='operator', at=NOW, data=material)
        self.assertEqual(log, before)
        self.assertEqual(changed['Events'][:-1], before['Events'])
        self.assertEqual(changed['Events'][-1]['Record']['ResidualRisk'], 'Monitor the remaining risk.')
        self.assertEqual(changed['Events'][-1]['Record']['ValidationResult'], 'Reviewed.')
        self.assertEqual(material['ValidationResult'], '\nReviewed.\t')
        advance(result, changed, key, policy)

    def test_other_required_override_narratives_already_reject_invalid_values(self):
        from Core.governance import transition
        for field in ('Rationale', 'AccountableOwner', 'DecisionAuthority', 'AuthorityRole'):
            for value in INVALID_STATEMENTS:
                with self.subTest(field=field, value=value):
                    result, record, _ = fixture('ClosedByRemediation')
                    record[field] = copy.deepcopy(value)
                    log, key = start(result, record)
                    with self.assertRaises(ValueError):
                        transition(log, result, key, 'submit', actor='operator', at=NOW)


class OverrideRequirementsTests(unittest.TestCase):
    def test_every_lifecycle_state_requires_complete_override_at_submission(self):
        from Core.governance import transition
        for state in STATES:
            result, _, original, _ = closure_fixture(state)
            for field in REQUIRED:
                with self.subTest(state=state, field=field):
                    record = copy.deepcopy(original)
                    omit(record, field)
                    log, key = start(result, record)
                    with self.assertRaises(ValueError):
                        transition(log, result, key, 'submit', actor='operator', at=NOW)

    def test_approval_and_activation_independently_validate_all_required_fields(self):
        from Core.governance import _evolve
        for state in STATES:
            result, _, record, policy = closure_fixture(state)
            log, key = start(result, record)
            log = advance(result, log, key, policy)
            for index in (2, 3):  # Independently exercise approve and activate, beyond submit.
                for field in REQUIRED:
                    with self.subTest(state=state, operation=log['Events'][index]['Operation'], field=field):
                        prior = copy.deepcopy(log['Events'][index - 1]['Record'])
                        omit(prior, field)
                        with self.assertRaises(ValueError):
                            _evolve({key: prior}, copy.deepcopy(log['Events'][index]), log)

    def test_approval_and_activation_reject_invalid_statement_types(self):
        from Core.governance import _evolve
        result, _, record, policy = closure_fixture()
        log, key = start(result, record)
        log = advance(result, log, key, policy)
        for index in (2, 3):
            for field in ('ResidualRisk', 'ValidationResult'):
                for value in INVALID_STATEMENTS:
                    with self.subTest(operation=log['Events'][index]['Operation'], field=field, value=value):
                        prior = copy.deepcopy(log['Events'][index - 1]['Record'])
                        prior[field] = copy.deepcopy(value)
                        with self.assertRaises(ValueError):
                            _evolve({key: prior}, copy.deepcopy(log['Events'][index]), log)

    def test_fully_justified_authorized_overrides_preserve_state_separation(self):
        from Core.governance import attach
        for state in STATES:
            with self.subTest(state=state):
                result, baseline, record, policy = closure_fixture(state)
                before = copy.deepcopy((result, baseline))
                log, key = start(result, record)
                output = attach(result, advance(result, log, key, policy), as_of=NOW, locator='governance/decisions.json')
                row = output['governance']['Records'][0]
                self.assertEqual(row['WorkflowState'], 'Active')
                self.assertEqual(row['GovernanceState'], 'ClosedByRemediation')
                self.assertEqual(row['CurrentLifecycleState'], state)
                self.assertEqual(output['lifecycle'], before[0]['lifecycle'])
                self.assertEqual((result, baseline), before)

    def test_override_requires_exact_policy_permission_even_for_resolved_targets(self):
        from Core.governance import transition
        for state in STATES:
            for permission in (None, False, 'true', 1):
                with self.subTest(state=state, permission=permission):
                    result, _, record, policy = closure_fixture(state)
                    policy['Roles']['risk-authority']['AllowClosureOverride'] = permission
                    log, key = start(result, record)
                    log = advance(result, log, key, policy, ('submit',))
                    with self.assertRaises(ValueError):
                        transition(log, result, key, 'approve', actor='approver', at=NOW, policy=policy)

    def test_activation_rechecks_override_permission(self):
        from Core.governance import transition
        result, _, record, policy = closure_fixture()
        log, key = start(result, record)
        log = advance(result, log, key, policy, ('submit', 'approve'))
        before = copy.deepcopy(log)
        policy['Roles']['risk-authority']['AllowClosureOverride'] = False
        with self.assertRaises(ValueError):
            transition(log, result, key, 'activate', actor='approver', at=NOW, policy=policy)
        self.assertEqual(log, before)

    def test_normal_resolved_closure_does_not_require_override_only_fields(self):
        from Core.governance import new_log, project
        result, baseline, record, policy = closure_fixture()
        record['ClosureOverride'] = False
        for field in ('ResidualRisk', 'Conditions', 'ReviewAt'):
            record.pop(field)
        policy['Roles']['risk-authority']['AllowClosureOverride'] = False
        before = copy.deepcopy((result, baseline))
        self.assertEqual(new_log(result)['Events'], [])
        log, key = start(result, record)
        row = project(advance(result, log, key, policy), as_of=NOW)['Records'][0]
        self.assertEqual(row['WorkflowState'], 'Active')
        self.assertFalse(row['ClosureOverride'])
        self.assertEqual((result, baseline), before)

    def test_normal_unresolved_closure_is_not_a_substitute_for_an_override(self):
        from Core.governance import transition
        for state in STATES[1:]:
            with self.subTest(state=state):
                result, _, record, _ = closure_fixture(state)
                record['ClosureOverride'] = False
                log, key = start(result, record)
                with self.assertRaises(ValueError):
                    transition(log, result, key, 'submit', actor='operator', at=NOW)

    def test_invalid_conditions_and_review_dates_are_blocked(self):
        from Core.governance import transition
        for field, values in (('Conditions', (None, [], True, 'Monitor', [' '], [True])),
                              ('ReviewAt', (None, '', ' ', '2026-10-08', 'invalid'))):
            for value in values:
                with self.subTest(field=field, value=value):
                    result, _, record, _ = closure_fixture()
                    record[field] = value
                    try:
                        log, key = start(result, record)
                    except ValueError:
                        continue
                    with self.assertRaises(ValueError):
                        transition(log, result, key, 'submit', actor='operator', at=NOW)

    def test_override_flag_cannot_use_truthy_nonboolean_values(self):
        from Core.governance import transition
        for value in (1, 'true', 'false', None, [], {}):
            with self.subTest(value=value):
                result, _, record, _ = closure_fixture()
                record['ClosureOverride'] = value
                log, key = start(result, record)
                with self.assertRaises(ValueError):
                    transition(log, result, key, 'submit', actor='operator', at=NOW)

    def test_replay_diagnoses_invalid_retained_overrides_without_repair(self):
        from Core.governance import validate_log, project
        result, _, record, policy = closure_fixture()
        log, key = start(result, record)
        log = advance(result, log, key, policy)
        variants = [edited_history(log, field, value) for field in ('ResidualRisk', 'ValidationResult')
                    for value in (True, ' ')]
        variants += [edited_history(log, field, remove=True) for field in ('ResidualRisk', 'Conditions', 'ReviewAt')]
        for index, bad in enumerate(variants):
            with self.subTest(index=index):
                before = copy.deepcopy(bad)
                self.assertTrue(any(row['severity'] == 'error' for row in validate_log(bad)))
                with self.assertRaises(ValueError):
                    project(bad, as_of=NOW)
                self.assertEqual(bad, before)

    def test_valid_historical_whitespace_is_not_rewritten_during_replay(self):
        from Core.governance import project
        result, _, record, policy = closure_fixture()
        log, key = start(result, record)
        log = advance(result, log, key, policy)
        historical = edited_history(log, 'ResidualRisk', '  Recorded prior justification.  ')
        before = copy.deepcopy(historical)
        self.assertEqual(project(historical, as_of=NOW)['Records'][0]['ResidualRisk'], '  Recorded prior justification.  ')
        self.assertEqual(historical, before)


class ClosurePersistenceTests(unittest.TestCase):
    def test_valid_normal_and_override_closures_render_without_decision_operations(self):
        from Core.governance import attach
        from Core.assessment_result import build_assessment_result
        from Core.assessment_serialization import write_assessment_result, read_assessment_result
        from Core.assessment_replay import render_snapshot
        from openpyxl import load_workbook
        for override in (False, True):
            with self.subTest(override=override), tempfile.TemporaryDirectory() as directory:
                result, _, record, policy = closure_fixture()
                record['ClosureOverride'] = override
                if not override:
                    for field in ('ResidualRisk', 'Conditions', 'ReviewAt'):
                        record.pop(field)
                built = build_assessment_result([], {'collection_context': {'identity': result['identity']}},
                                                expected_tenant_id=result['tenant_id'], evaluation_date='2026-10-08')
                built.update(result)
                log, key = start(built, record)
                output = attach(built, advance(built, log, key, policy), as_of=NOW, locator='governance/decisions.json')
                root = Path(directory)
                source = root/'source.json'
                write_assessment_result(source, output)
                before = source.read_bytes()
                with patch('Core.governance.transition', side_effect=AssertionError('Renderer transition')), \
                     patch('Core.governance.create_draft', side_effect=AssertionError('Renderer draft')):
                    paths = render_snapshot(source, root/'renders', snapshot_output=root/'copy.json')
                self.assertEqual(source.read_bytes(), before)
                self.assertEqual(read_assessment_result(root/'copy.json')['governance'], output['governance'])
                for key in ('excel_path', 'technical_excel_path'):
                    book = load_workbook(paths[key])
                    try:
                        self.assertIn('Decision Register', book.sheetnames)
                        headers = [cell.value for cell in book['Decision Register'][1]]
                        self.assertEqual(book['Decision Register'].cell(2, headers.index('ClosureOverride') + 1).value, override)
                    finally:
                        book.close()
                self.assertIn('Governance audit history', Path(paths['html_path']).read_text(encoding='utf-8'))
                self.assertIn('Governance', Path(paths['summary_html_path']).read_text(encoding='utf-8'))

    def test_invalid_retained_overrides_cannot_serialize_or_render(self):
        from Core.governance import attach
        from Core.governance_validation import validate_governance
        from Core.assessment_serialization import write_assessment_result, read_assessment_result
        from Core.assessment_replay import render_snapshot
        result, _, record, policy = closure_fixture()
        log, key = start(result, record)
        log = advance(result, log, key, policy)
        output = attach(result, log, as_of=NOW, locator='governance/decisions.json')
        variants = [edited_history(log, 'ResidualRisk', True), edited_history(log, 'ValidationResult', ' '),
                    edited_history(log, 'ReviewAt', remove=True)]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for index, bad_log in enumerate(variants):
                with self.subTest(index=index):
                    bad = copy.deepcopy(output)
                    bad['governance']['DecisionLog'] = bad_log
                    before = copy.deepcopy(bad)
                    diagnostics = validate_governance(bad)
                    self.assertTrue(any(row['severity'] == 'error' and row['code'] in
                                        {'governance_closure', 'governance_closure_support'} for row in diagnostics))
                    with self.assertRaises(ValueError):
                        write_assessment_result(root/'forbidden.json', bad)
                    self.assertFalse((root/'forbidden.json').exists())
                    source = root/f'legacy-{index}.json'
                    source.write_text(json.dumps(bad), encoding='utf-8')
                    original_bytes = source.read_bytes()
                    with self.assertRaises(ValueError):
                        read_assessment_result(source)
                    with self.assertRaises(ValueError):
                        render_snapshot(source, root/'renders', snapshot_output=root/'copy.json')
                    self.assertEqual(source.read_bytes(), original_bytes)
                    self.assertFalse((root/'renders').exists())
                    self.assertFalse((root/'copy.json').exists())
                    self.assertEqual(bad, before)

    def test_normalized_override_roundtrips_without_changing_completed_source(self):
        from Core.governance import attach
        from Core.assessment_serialization import write_assessment_result, read_assessment_result
        result, _, record, policy = closure_fixture()
        record.update(ResidualRisk='  Retain explicit monitoring.  ', ValidationResult='\tReviewed criteria.\n')
        log, key = start(result, record)
        output = attach(result, advance(result, log, key, policy), as_of=NOW, locator='governance/decisions.json')
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory)/'snapshot.json'
            write_assessment_result(source, output)
            before = source.read_bytes()
            restored = read_assessment_result(source)
            self.assertEqual(restored['governance'], output['governance'])
            self.assertEqual(restored['governance']['Records'][0]['ResidualRisk'], 'Retain explicit monitoring.')
            self.assertEqual(restored['governance']['Records'][0]['ValidationResult'], 'Reviewed criteria.')
            self.assertEqual(source.read_bytes(), before)
