"""Fictional retained evidence: default-policy activity is not BitLocker risk."""

from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest

from provenance_fixture import fictional_publication
from test_evidence_reconciliation import dataset, DAY, TENANT
from Core.device_compliance import build_device_compliance
from Core.expanded_findings import expanded_findings
from Core.assessment_result import build_assessment_result
from Core.identity_investigation import build_intune_device_investigation, qualify_identity_recommendation


def fixture(sync=DAY, setting='Is active', **extra):
    device = {'id': 'intune-1', 'azureADDeviceId': 'entra-1', 'deviceName': 'Fictional laptop',
              'complianceState': 'noncompliant', 'lastSyncDateTime': sync}
    failed = {'id': 'setting-1', 'deviceId': 'intune-1', 'ParentId': 'summary-1',
              'settingName': setting, 'state': 'noncompliant', **extra}
    return {'evaluation_date': DAY, 'collection_context': {'collected_at': DAY, 'tenant_id': TENANT},
            'expected_tenant_id': TENANT, 'sheets': {}, 'assessment_sources': {
                'managed_devices': [dataset([device])],
                'compliance_setting_states': [dataset([failed], source_api='https://graph.microsoft.com/v1.0/deviceManagement/deviceCompliancePolicySettingStateSummaries/summary-1/deviceComplianceSettingStates')]}}


class ComplianceReasonTests(unittest.TestCase):
    def model(self, bundle):
        return build_device_compliance(bundle, evaluation_date=DAY)

    def test_screenshot_activity_failure_preserves_other_compliant_checks(self):
        bundle = fixture('2026-08-01')
        raw = bundle['assessment_sources']['compliance_setting_states'][0]['records']
        raw.extend([dict(raw[0], id='assigned', settingName='Has a compliance policy assigned', state='compliant'),
                    dict(raw[0], id='user', settingName='Enrolled user exists', state='compliant')])
        before = deepcopy(bundle)
        model = self.model(bundle)
        self.assertEqual(bundle, before)
        device = model['devices'][0]
        self.assertEqual(device['activity'], 'stale')
        self.assertEqual(device['reason'], 'Is active: noncompliant; policy: Default Device Compliance Policy')
        self.assertEqual(sum(item['failed'] for item in model['settings']), 1)
        self.assertEqual(device['device']['complianceState'], 'noncompliant')
        rows = expanded_findings(bundle, evaluation_date=DAY)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['Priority'], 'Medium')
        self.assertEqual(rows[0]['Disposition'], 'Opportunity')
        self.assertNotIn('BitLocker', rows[0]['Observation'])
        self.assertIn('Is active', rows[0]['Observation'])

    def test_recent_default_policy_failure_is_still_reporting_review(self):
        model = self.model(fixture())
        self.assertEqual(model['settings'][0]['triage'], 'reporting_review')
        self.assertEqual(model['devices'][0]['activity'], 'active')

    def test_real_setting_and_explicit_policy_are_named_and_supported(self):
        bundle = fixture(setting='Require BitLocker', policyId='policy-1')
        bundle['assessment_sources']['compliance_policies'] = [dataset([{'id': 'policy-1', 'displayName': 'Windows encryption baseline'}])]
        row = expanded_findings(bundle, evaluation_date=DAY)[0]
        self.assertEqual(row['Priority'], 'High')
        self.assertEqual(row['Disposition'], 'Action')
        self.assertIn('Require BitLocker', row['Observation'])
        self.assertIn('Windows encryption baseline', row['Observation'])
        self.assertEqual(row['ObservationDate'], DAY)

    def test_old_security_failure_is_retained_without_current_high_risk_claim(self):
        bundle = fixture('2026-07-01', 'Require BitLocker')
        row = expanded_findings(bundle, evaluation_date=DAY)[0]
        self.assertEqual(row['ObservationDate'], '2026-07-01')
        self.assertEqual(row['Priority'], 'Medium')
        self.assertEqual(row['InvestigationEvidence']['records'][0]['state'], 'noncompliant')
        self.assertIn('Require BitLocker', row['Observation'])
        self.assertIn('current security risk is not established', row['Observation'])

    def test_fresh_defender_signal_does_not_make_old_intune_state_current(self):
        bundle = fixture('2026-07-01', 'Require BitLocker')
        bundle['assessment_sources']['machines'] = [dataset([{'id': 'defender-1', 'aadDeviceId': 'entra-1', 'lastSeen': DAY}])]
        model = self.model(bundle)
        self.assertEqual(model['devices'][0]['activity'], 'active')
        self.assertEqual(model['settings'][0]['triage'], 'historical_or_unknown_review')
        self.assertEqual(model['settings'][0]['last_reported'], '2026-07-01')

    def test_missing_and_future_dates_remain_unknown(self):
        for date in (None, '2027-01-01'):
            with self.subTest(date=date):
                model = self.model(fixture(date, 'Require BitLocker'))
                self.assertEqual(model['devices'][0]['activity'], 'unknown')
                self.assertEqual(model['settings'][0]['triage'], 'historical_or_unknown_review')

    def test_setting_summary_parent_cannot_be_used_as_device_or_policy_id(self):
        bundle = fixture(setting='Require BitLocker')
        record = bundle['assessment_sources']['compliance_setting_states'][0]['records'][0]
        record.pop('deviceId')
        record['ParentId'] = 'intune-1'
        bundle['assessment_sources']['compliance_policies'] = [dataset([{'id': 'intune-1', 'displayName': 'Wrong policy'}])]
        model = self.model(bundle)
        self.assertIsNone(model['settings'][0]['device_id'])
        self.assertEqual(model['settings'][0]['policy'], 'not retained')
        self.assertEqual(model['devices'][0]['reason'], 'Reason not retained')

    def test_duplicate_ids_and_matching_names_never_establish_device_join(self):
        bundle = fixture(setting='Require BitLocker')
        devices = bundle['assessment_sources']['managed_devices'][0]['records']
        devices.append(dict(devices[0], azureADDeviceId='other-entra'))
        self.assertIsNone(self.model(bundle)['settings'][0]['device_id'])
        record = bundle['assessment_sources']['compliance_setting_states'][0]['records'][0]
        record.update(deviceId='unrelated', deviceName='Fictional laptop')
        self.assertIsNone(self.model(bundle)['settings'][0]['device_id'])

    def test_mixed_cohorts_keep_counts_and_exact_source_rows(self):
        bundle = fixture(setting='Require BitLocker')
        data = bundle['assessment_sources']['compliance_setting_states'][0]['records']
        data.append(dict(data[0], id='activity', settingName='Is active'))
        rows = expanded_findings(bundle, evaluation_date=DAY)
        self.assertEqual(len(rows), 2)
        self.assertEqual(sorted(row['Priority'] for row in rows), ['High', 'Medium'])
        self.assertTrue(all('1 distinct managed device' in row['Observation'] for row in rows))
        selected = [record for row in rows for record in row['InvestigationEvidence']['records']]
        self.assertCountEqual(selected, data)

    def test_missing_reason_and_old_plan_label_are_qualified(self):
        source = SimpleNamespace(managed_devices=[{'id': 'one', 'complianceState': 'noncompliant'}], collection_status={})
        row = qualify_identity_recommendation({'Feature': 'Microsoft Intune (Plan A)', 'Priority': 'High',
            'Observation': '1 of 1 managed devices are non-compliant'}, source)
        self.assertEqual(row['Priority'], 'Medium')
        self.assertEqual(row['Disposition'], 'Opportunity')
        self.assertEqual(row['OriginalFeature'], 'Microsoft Intune (Plan A)')
        self.assertEqual(row['Feature'], 'Review Intune device noncompliance reasons')
        detail = build_intune_device_investigation(source)['rows'][0]
        self.assertEqual(detail['Compliance reason'], 'Reason not retained')
        self.assertEqual(detail['Device activity'], 'unknown')

    def test_summary_can_supply_setting_name_without_implying_policy(self):
        bundle = fixture(setting='')
        bundle['assessment_sources']['compliance_setting_summaries'] = [dataset([{'id': 'summary-1', 'settingName': 'Is active'}])]
        self.assertEqual(self.model(bundle)['settings'][0]['setting'], 'Is active')

    def test_conflicting_policy_names_remain_unknown(self):
        bundle = fixture(setting='Require BitLocker', policyId='duplicate')
        bundle['assessment_sources']['compliance_policies'] = [dataset([
            {'id': 'duplicate', 'displayName': 'First'}, {'id': 'duplicate', 'displayName': 'Second'}])]
        self.assertEqual(self.model(bundle)['settings'][0]['policy'], 'not retained')

    def test_graph_builtin_setting_identifier_matches_portal_activity_label(self):
        bundle = fixture(setting='DefaultDeviceCompliancePolicy.RequireRemainContact')
        item = self.model(bundle)['settings'][0]
        self.assertEqual(item['setting'], 'Is active')
        self.assertEqual(item['triage'], 'reporting_review')

    def test_summary_parent_can_supply_builtin_identifier_without_device_join(self):
        bundle = fixture(setting='')
        row = bundle['assessment_sources']['compliance_setting_states'][0]['records'][0]
        row['ParentId'] = 'DefaultDeviceCompliancePolicy.RequireRemainContact'
        row.pop('deviceId')
        item = self.model(bundle)['settings'][0]
        self.assertEqual(item['setting'], 'Is active')
        self.assertIsNone(item['device_id'])

    def test_decision_model_does_not_restore_generic_high_priority_from_plan_name(self):
        bundle = fixture('2026-07-01')
        original = {'Feature': 'Microsoft Intune (Plan A)', 'Service': 'Entra', 'Priority': 'High',
                    'Observation': '1 of 1 managed devices are non-compliant',
                    'Recommendation': 'Block all devices', 'FindingKey': 'entra.devices.noncompliant'}
        before = deepcopy(original)
        result = build_assessment_result([original], bundle, evaluation_date=DAY, expected_tenant_id=TENANT)
        row = next(row for row in result['recommendations'] if row.get('FindingKey') == 'entra.devices.noncompliant')
        self.assertEqual(original, before)
        self.assertEqual(row['Priority'], 'Medium')
        self.assertEqual(row['Disposition'], 'Opportunity')
        self.assertEqual(row['OriginalFeature'], 'Microsoft Intune (Plan A)')
        self.assertEqual(result['device_compliance']['devices'][0]['activity'], 'stale')


@fictional_publication
class ComplianceOutputTests(unittest.TestCase):
    def test_workbook_html_and_replay_keep_reason_and_reporting_qualification(self):
        from openpyxl import load_workbook
        from Core.export_recommendations import export_to_excel
        from Core.customer_report import render_customer_report
        from Core.pilot_summary import render_pilot_summary
        from Core.assessment_serialization import write_assessment_result, read_assessment_result
        from Core.workbook_layout import technical_workbook_path
        bundle = fixture('2026-07-01')
        rows = expanded_findings(bundle, evaluation_date=DAY)
        result = build_assessment_result(rows, bundle, evaluation_date=DAY, expected_tenant_id=TENANT)
        bundle['assessment_result'] = result
        with TemporaryDirectory() as folder:
            path = Path(export_to_excel(rows, filename='Fictional.xlsx', evidence_bundle=bundle, output_dir=folder))
            result = bundle['assessment_result']
            before = deepcopy(result)
            html = render_customer_report(result, bundle, 'Fictional', path, technical_workbook_path=technical_workbook_path(path))
            summary = render_pilot_summary(result, bundle, 'Fictional', workbook_path=path)
            self.assertIn('Is active', html)
            workbook = load_workbook(path)
            technical = load_workbook(technical_workbook_path(path))
            try:
                # Every existing reason/date assertion now checks its technical
                # destination, with an additional customer privacy assertion.
                values = list(technical['Device Detail'].values)
                cells = '\n'.join(str(cell or '') for row in values for cell in row)
                self.assertNotIn('Fictional laptop', '\n'.join(str(cell or '') for row in workbook['Devices'].values for cell in row))
                self.assertIn('Is active', cells)
                self.assertIn('Default Device Compliance Policy', cells)
                self.assertIn('2026-07-01', cells)
                self.assertIn('activity: stale', cells)
                self.assertNotIn('Microsoft Intune (Plan A)', cells)
                finding = next(row for row in result['recommendations'] if row.get('DeviceComplianceTriage') == 'review')
                self.assertEqual(finding['Priority'], 'Medium')
                self.assertEqual(finding['Disposition'], 'Opportunity')
                self.assertTrue(finding['RawEvidenceRanges'])
                self.assertTrue(finding['TechnicalEvidenceRanges'])
            finally:
                workbook.close()
                technical.close()
            self.assertEqual(result, before)
            snapshot = Path(folder) / 'result.json'
            write_assessment_result(snapshot, result)
            replay = read_assessment_result(snapshot)
            self.assertEqual(replay['device_compliance'], json.loads(json.dumps(result['device_compliance'], default=str)))
            self.assertEqual(summary, render_pilot_summary(replay, bundle, 'Fictional', workbook_path=path))


if __name__ == '__main__':
    unittest.main()
