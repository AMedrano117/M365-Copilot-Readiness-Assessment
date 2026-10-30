"""Fictional endpoint populations: retained rows are not devices or pilot scope."""

from copy import deepcopy
from datetime import date
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from provenance_fixture import fictional_publication
from test_device_compliance_reasons import fixture
from test_evidence_reconciliation import DAY, TENANT, dataset
from Core.device_compliance import build_device_compliance
from Core.device_evidence import (CA_QUALIFICATION, build_device_identities, interpret_records,
                                  population, timestamp, pilot_blocker_eligible)
from Core.expanded_findings import expanded_findings
from Core.assessment_result import build_assessment_result


def roster(*ids, complete=False):
    return {'readiness_review':{'pilot_scope':{'id':'fictional-pilot', 'description':'Fictional intended pilot',
        'reviewed_at':DAY, 'reviewer_role':'Endpoint owner', 'evidence_reference':'fictional-roster-1',
        'intune_device_ids':list(ids), 'membership_complete':complete}}}


def model(bundle):
    return build_device_compliance(bundle, profile=bundle.get('assessment_profile'), evaluation_date=DAY)


def result(bundle):
    return build_assessment_result(expanded_findings(bundle, profile=bundle.get('assessment_profile'), evaluation_date=DAY),
        bundle, evaluation_date=DAY, expected_tenant_id=TENANT)


class DeviceEvidenceInterpretationTests(unittest.TestCase):
    def test_several_settings_one_device_reconcile_exact_evidence_ids(self):
        bundle = fixture(setting='BitLocker')
        records = bundle['assessment_sources']['compliance_setting_states'][0]['records']
        records.extend([dict(records[0],id='second',settingName='TPM'),dict(records[0],id='third',settingName='Firewall')])
        before = deepcopy(bundle)
        finding = expanded_findings(bundle,evaluation_date=DAY)[0]
        counts = finding['DeviceEvidenceCounts']
        self.assertEqual(counts['source_record_count'],3)
        self.assertEqual(counts['distinct_intune_devices'],1)
        self.assertEqual(counts['distinct_entra_devices'],1)
        self.assertEqual(counts['active_device_count'],1)
        self.assertEqual(counts['valid_timestamp_count'],3)
        self.assertEqual(counts['pilot_relevance_unknown_count'],1)
        self.assertEqual(counts['pilot_relevance_unknown_records'],3)
        self.assertEqual(finding['InvestigationEvidence']['records'],records)
        self.assertEqual(len(set(finding['DeviceEvidenceIds'])),3)
        self.assertTrue(all(identifier.startswith('EVD-') for identifier in finding['DeviceEvidenceIds']))
        self.assertEqual(bundle,before)

    def test_activity_security_and_reporting_have_disjoint_record_populations(self):
        bundle = fixture(setting='BitLocker')
        records = bundle['assessment_sources']['compliance_setting_states'][0]['records']
        records.extend(dict(records[0],id=str(index),settingName=name,state=state) for index,(name,state) in enumerate([
            ('Is active','noncompliant'),('Enrolled user exists','noncompliant'),('Firewall','error'),
            ('TPM','conflict'),('Encryption','unknown'),('Secure boot','notEvaluated')]))
        findings = expanded_findings(bundle,evaluation_date=DAY)
        self.assertEqual({row['DeviceFindingType'] for row in findings},
            {'UnconfirmedSettingState','ActivityOrEnrollmentFailure','ReportingErrorOrConflict'})
        self.assertEqual(sorted(row['DeviceEvidenceCounts']['source_record_count'] for row in findings),[1,2,4])
        self.assertCountEqual([item for row in findings for item in row['InvestigationEvidence']['records']],records)
        self.assertTrue(all(row['Disposition']!='Action' for row in findings if row['DeviceFindingType']!='ExplicitSecuritySettingFailure'))

    def test_compliant_inventory_does_not_override_explicit_failed_setting(self):
        bundle = fixture(setting='BitLocker')
        bundle['assessment_sources']['managed_devices'][0]['records'][0]['complianceState']='compliant'
        assessment = result(bundle)
        failed = next(row for row in assessment['recommendations'] if row.get('DeviceFindingType')=='ExplicitSecuritySettingFailure')
        inventory = next(row for row in assessment['recommendations'] if row.get('DeviceFindingType')=='InventoryComplianceState')
        self.assertEqual(failed['DeviceEvidenceCounts']['source_record_count'],1)
        self.assertIn('1 compliant',inventory['Observation'])
        self.assertEqual(failed['Priority'],'High')
        self.assertIn(failed['RecommendationId'],[row['RecommendationId'] for row in assessment['blockers']])

    def test_stale_missing_and_malformed_values_do_not_establish_current_risk(self):
        for raw, expected in [('2026-01-01','StaleDeviceRecord'),(None,'ActivityUnknown'),
                              ('0001-01-01','MalformedTimestamp'),(-1,'MalformedTimestamp'),
                              ('invalid','MalformedTimestamp'),('2999-01-01','MalformedTimestamp')]:
            with self.subTest(raw=raw):
                bundle=fixture(raw,'BitLocker')
                compliance=model(bundle)
                self.assertEqual(compliance['identity_model']['identities'][0]['relevance'],expected)
                finding=expanded_findings(bundle,evaluation_date=DAY)[0]
                self.assertNotEqual(finding['Priority'],'High')
                self.assertFalse(finding['DevicePilotBlockerEligible'])
                self.assertNotIn('739915',json.dumps(result(bundle),default=str))
                self.assertEqual(compliance['devices'][0]['device']['lastSyncDateTime'],raw)
                if expected=='MalformedTimestamp':
                    self.assertTrue(finding['DeviceEvidenceDiagnostics'])
                    self.assertEqual(finding['DeviceEvidenceCounts']['valid_timestamp_count'],0)
                    self.assertEqual(finding['DeviceEvidenceCounts']['malformed_timestamp_count'],1)
                    self.assertEqual(timestamp(raw,date.fromisoformat(DAY))['age_days'],None)
                    selected=next(row for row in result(bundle)['recommendations'] if row.get('DeviceFindingType')=='UnconfirmedSettingState')
                    self.assertEqual(selected['ObservationDate'],'')
                    self.assertEqual(selected['Freshness'],'unknown')

    def test_present_malformed_setting_date_does_not_fall_back_to_good_sync(self):
        bundle=fixture(setting='BitLocker',lastReportedDateTime='0001-01-01')
        setting=model(bundle)['settings'][0]
        self.assertEqual(setting['triage'],'historical_or_unknown_review')
        self.assertEqual(setting['evidence']['timestamp']['state'],'malformed')

    def test_stable_id_conflict_and_same_hostname_never_merge(self):
        bundle=fixture(setting='BitLocker')
        devices=bundle['assessment_sources']['managed_devices'][0]['records']
        devices.append(dict(devices[0],id='intune-2'))
        identities=model(bundle)['identity_model']['identities']
        self.assertEqual(len(identities),2)
        self.assertTrue(all(row['unresolved'] for row in identities))
        self.assertIsNone(model(bundle)['settings'][0]['device_id'])
        devices[1]['azureADDeviceId']='entra-2'
        identities=model(bundle)['identity_model']['identities']
        self.assertEqual(len(identities),2)
        self.assertTrue(all(not row['unresolved'] for row in identities))

    def test_conflicting_explicit_ids_or_summary_parent_cannot_resolve_identity(self):
        bundle=fixture(setting='BitLocker',managedDeviceId='other-intune')
        self.assertTrue(model(bundle)['settings'][0]['evidence']['unresolved'])
        bundle=fixture(setting='BitLocker')
        setting=bundle['assessment_sources']['compliance_setting_states'][0]['records'][0]
        setting.pop('deviceId');setting['ParentId']='intune-1'
        bundle['assessment_sources']['compliance_setting_summaries']=[dataset([{'id':'intune-1','settingName':'BitLocker'}])]
        bundle['assessment_sources']['compliance_setting_states'][0]['source']['source_api']='retained legacy summary read'
        self.assertTrue(model(bundle)['settings'][0]['evidence']['unresolved'])

    def test_identity_retains_namespaced_ids_and_inventory_metadata(self):
        bundle=fixture(setting='BitLocker')
        raw=bundle['assessment_sources']['managed_devices'][0]['records'][0]
        raw.update(serialNumber='fictional-serial',userId='fictional-user',userPrincipalName='fictional@example.invalid',
            managementState='managed',operatingSystem='Windows',deviceEnrollmentType='windowsAzureADJoin',managedDeviceOwnerType='company')
        bundle['assessment_sources']['machines']=[dataset([{'id':'defender-1','aadDeviceId':'entra-1',
            'lastSeen':DAY,'onboardingStatus':'Onboarded'}])]
        identity=model(bundle)['identity_model']['identities'][0]
        self.assertEqual((identity['intune_id'],identity['entra_id'],identity['defender_id']),('intune-1','entra-1','defender-1'))
        self.assertEqual(identity['serial_number'],'fictional-serial')
        self.assertEqual(identity['primary_user_id'],'fictional-user')
        self.assertEqual(identity['ownership'],'company')
        self.assertEqual(identity['platform'],'Windows')
        self.assertEqual(len(identity['sources']),2)
        self.assertTrue(all(entry['evidence_id'].startswith('EVD-') for entry in identity['sources']))

    def test_retired_and_duplicate_records_are_inventory_hygiene(self):
        for state in ('retired','duplicate'):
            bundle=fixture(setting='BitLocker')
            raw=bundle['assessment_sources']['managed_devices'][0]['records'][0]
            if state=='retired': raw['managementState']='retired'
            else: bundle['assessment_sources']['managed_devices'][0]['records'].append(dict(raw))
            finding=expanded_findings(bundle,evaluation_date=DAY)[0]
            self.assertEqual(finding['DevicePilotImpact'],'Inventory hygiene')
            self.assertFalse(finding['DevicePilotBlockerEligible'])

    def test_pilot_blocker_requires_current_confirmed_correlation(self):
        for profile, impact in [(None,'Confirm before pilot'),(roster('intune-1'),'Blocks pilot'),
                                (roster('other',complete=True),'Tenant hardening'),(roster('other'),'Confirm before pilot')]:
            bundle=fixture(setting='BitLocker')
            if profile: bundle['assessment_profile']=profile
            assessment=result(bundle)
            finding=next(row for row in assessment['recommendations'] if row.get('DeviceFindingType')=='ExplicitSecuritySettingFailure')
            self.assertEqual(finding['Priority'],'High')
            self.assertEqual(finding['DevicePilotImpact'],impact)
            self.assertEqual(finding['PilotImpact'],'Blocks readiness')
            self.assertIn(finding['RecommendationId'], [row['RecommendationId'] for row in assessment['blockers']])

    def test_old_roster_or_missing_scope_evidence_never_establishes_pilot(self):
        for field,value in [('reviewed_at','2025-01-01'),('evidence_reference',''),('intune_device_ids','intune-1')]:
            profile=roster('intune-1')
            profile['readiness_review']['pilot_scope'][field]=value
            bundle=fixture(setting='BitLocker');bundle['assessment_profile']=profile
            finding=expanded_findings(bundle,profile=profile,evaluation_date=DAY)[0]
            self.assertFalse(finding['DevicePilotBlockerEligible'])

    def test_unqualified_endpoint_control_failure_cannot_implicitly_block_pilot(self):
        self.assertFalse(pilot_blocker_eligible({'ControlId':'ENDPOINT.POSTURE','Priority':'High'}))
        self.assertTrue(pilot_blocker_eligible({'ControlId':'IDENTITY.AUTH','Priority':'High'}))

    def test_insufficient_source_support_does_not_publish_a_blocker_claim(self):
        bundle=fixture(setting='BitLocker')
        bundle['assessment_profile']=roster('intune-1')
        bundle['assessment_sources']['compliance_setting_states'][0]['source'].update(complete=False,availability_status='partial')
        assessed=result(bundle)
        row=next(row for row in assessed['recommendations'] if row.get('DeviceFindingType')=='ExplicitSecuritySettingFailure')
        self.assertEqual(row['PilotImpact'],'Evidence review')
        self.assertNotIn('Blocks pilot',row['Observation'])
        self.assertNotIn(row['RecommendationId'],[blocker['RecommendationId'] for blocker in assessed['blockers']])

    def test_rtp_population_units_freshness_and_invalid_values(self):
        bundle=fixture()
        bundle['assessment_sources']['windows_protection']=[dataset([
            {'id':'health-'+str(index),'ParentId':'intune-1','realTimeProtectionEnabled':False,
             'lastReportedDateTime':DAY if index<186 else '2026-01-01',
             'antiMalwareVersion':'0.0.0.0','lastQuickScanDateTime':'0001-01-01'} for index in range(252)])]
        findings=[row for row in expanded_findings(bundle,evaluation_date=DAY) if row.get('DeviceFindingType')=='DefenderRealTimeProtectionObservation']
        self.assertEqual([row['DeviceEvidenceCounts']['source_record_count'] for row in findings],[252])
        self.assertEqual(findings[0]['EndpointObservationSubsets']['current']['counts']['source_record_count'],186)
        self.assertEqual(findings[0]['EndpointObservationSubsets']['stale']['counts']['source_record_count'],66)
        self.assertTrue(all(row['DeviceEvidenceCounts']['distinct_intune_devices']==1 for row in findings))
        self.assertTrue(all(row['Priority']=='Low' and row['Disposition']=='Opportunity' for row in findings))
        self.assertTrue(all(not row['DevicePilotBlockerEligible'] for row in findings))
        self.assertTrue(all('effectiveness is not established' in row['Observation'] for row in findings))
        self.assertTrue(all(len(row['DeviceEvidenceDiagnostics'])==2*row['DeviceEvidenceCounts']['source_record_count'] for row in findings))
        self.assertTrue(all('0.0.0.0' not in row['Observation'] and '0001' not in row['Observation'] for row in findings))

    def test_approved_third_party_av_is_retained_without_false_assurance(self):
        bundle=fixture()
        bundle['assessment_sources']['windows_protection']=[dataset([{'id':'health','ParentId':'intune-1',
            'realTimeProtectionEnabled':False,'lastReportedDateTime':DAY,'thirdPartyAntivirusApproved':True,
            'thirdPartyAntivirus':'Fictional AV','avMode':'passive'}])]
        finding=next(row for row in expanded_findings(bundle,evaluation_date=DAY) if row.get('DeviceFindingType')=='DefenderRealTimeProtectionObservation')
        context=finding['DeviceEvidenceInterpretations'][0]['defender']
        self.assertEqual(context['approved_third_party_av'],True)
        self.assertEqual(context['defender_mode'],'passive')
        self.assertEqual(context['onboarding_state'],None)
        self.assertIn('effectiveness is not established',finding['Observation'])

    def test_five_machine_records_do_not_establish_fleet_coverage(self):
        from Core.operational_evidence import reconcile_devices
        bundle=fixture()
        bundle['assessment_sources']['machines']=[dataset([{'id':str(i),'aadDeviceId':'entra-'+str(i),
            'lastSeen':DAY,'healthStatus':'Active','onboardingStatus':'Onboarded'} for i in range(5)])]
        reconciliation=reconcile_devices(bundle,evaluation_date=DAY)
        self.assertIsNone(reconciliation['denominator'])
        self.assertIsNone(reconciliation['reporting_percent'])

    def test_fictional_reference_sized_populations_keep_units_and_rows(self):
        devices=[{'id':f'intune-{i}','azureADDeviceId':f'entra-{i}','lastSyncDateTime':DAY if i<157 else '2026-01-01',
                  'complianceState':'noncompliant' if i<277 else 'compliant' if i<492 else 'unknown'} for i in range(513)]
        settings=[{'id':f'current-{i}','deviceId':f'intune-{i%157}','settingName':'BitLocker', 'state':'noncompliant'} for i in range(394)]
        settings.extend({'id':f'older-{i}','deviceId':f'intune-{157+i%142}','settingName':'Is active' if i<172 else 'TPM',
                         'state':'noncompliant'} for i in range(239))
        settings.extend({'id':f'error-{i}','deviceId':f'intune-{i%244}','settingName':'Secure boot','state':'error'} for i in range(298))
        protection=[{'id':f'health-{i}','ParentId':f'intune-{i}','lastReportedDateTime':DAY if i<186 else '2026-01-01',
                     'realTimeProtectionEnabled':i>=252} for i in range(266)]
        bundle={'evaluation_date':DAY,'expected_tenant_id':TENANT,'assessment_sources':{
            'managed_devices':[dataset(devices)],'compliance_setting_states':[dataset(settings)],
            'windows_protection':[dataset(protection)],'machines':[dataset([
                {'id':f'defender-{i}','aadDeviceId':f'entra-{i}','lastSeen':DAY} for i in range(5)])]}}
        compliance=model(bundle)
        self.assertEqual(len(compliance['inventory']),513)
        findings=expanded_findings(bundle,evaluation_date=DAY)
        by_type={row['DeviceFindingType']:row for row in findings if row.get('DeviceFindingType')!='DefenderRealTimeProtectionObservation'}
        failed=by_type['ExplicitSecuritySettingFailure']['DeviceEvidenceCounts']
        self.assertEqual((failed['source_record_count'],failed['distinct_intune_devices']),(394,157))
        self.assertEqual(by_type['ActivityOrEnrollmentFailure']['DeviceEvidenceCounts']['source_record_count'],172)
        self.assertEqual(by_type['UnconfirmedSettingState']['DeviceEvidenceCounts']['source_record_count'],67)
        self.assertEqual(by_type['ReportingErrorOrConflict']['DeviceEvidenceCounts']['source_record_count'],298)
        self.assertCountEqual([raw for row in findings if row.get('DeviceFindingType')!='DefenderRealTimeProtectionObservation'
            for raw in row['InvestigationEvidence']['records']],settings)
        rtp=[row['DeviceEvidenceCounts']['source_record_count'] for row in findings if row.get('DeviceFindingType')=='DefenderRealTimeProtectionObservation']
        self.assertEqual(rtp,[252])
        self.assertEqual(len(protection),266)
        self.assertEqual(len(bundle['assessment_sources']['machines'][0]['records']),5)

    def test_legacy_conditional_access_claim_is_requalified(self):
        bundle=fixture()
        policies=[{'id':'fictional-ca','state':'enabled','grantControls':{'builtInControls':['compliantDevice']}}]
        bundle['assessment_sources']['ca_policies']=[dataset(policies)]
        old={'Feature':'Microsoft Intune (Plan A)','Service':'Entra','Priority':'High','Disposition':'Action',
            'ActionType':'Confirmation','OriginalFeature':'Microsoft Intune (Plan A)',
            'OriginalRecommendation':'Noncompliant devices bypass Intune security policies',
            'Observation':'Devices bypass Intune security policies and access Copilot',
            'Recommendation':'Noncompliant devices bypass Intune security policies'}
        assessment=build_assessment_result([old],bundle,evaluation_date=DAY,expected_tenant_id=TENANT)
        finding=next(row for row in assessment['recommendations'] if row.get('DeviceFindingType')=='CompliantDeviceAccessEnforcementNotEstablished')
        self.assertIn(CA_QUALIFICATION,finding['Observation'])
        self.assertNotIn('bypass Intune',finding['Observation']+finding['Recommendation'])
        self.assertEqual(finding['InvestigationEvidence']['records'],policies)
        self.assertEqual(finding['DeviceEvidenceCounts']['source_record_count'],1)
        self.assertFalse(finding['DevicePilotBlockerEligible'])

    def test_legacy_endpoint_date_cannot_generate_an_age_or_unprotected_claim(self):
        legacy={'Service':'Defender','FindingKey':'expanded.windows_protection.disabled.older',
            'Feature':'Restore real-time protection','Observation':'252 active unprotected devices',
            'Recommendation':'Fix immediately','Priority':'High','Disposition':'Action','ObservationDate':'0001-01-01'}
        assessed=build_assessment_result([legacy],{},evaluation_date=DAY,expected_tenant_id=TENANT)
        row=next(row for row in assessed['recommendations'] if row.get('DeviceFindingType')=='UnverifiedLegacyDeviceObservation')
        self.assertEqual(row['ObservationDate'],'')
        self.assertEqual(row['Freshness'],'unknown')
        self.assertNotIn('unprotected',row['Observation'])
        self.assertNotIn('739915',row['Qualification'])
        self.assertEqual(row['UnverifiedSourceObservation'],legacy['Observation'])
        self.assertEqual(row['Priority'],'Medium')
        self.assertFalse(row['DevicePilotBlockerEligible'])

    def test_invalid_retirement_date_cannot_establish_retirement(self):
        bundle=fixture(setting='BitLocker')
        bundle['assessment_sources']['managed_devices'][0]['records'][0]['retiredDateTime']='0001-01-01'
        identity=model(bundle)['identity_model']['identities'][0]
        self.assertEqual(identity['relevance'],'CurrentActiveDevice')
        self.assertTrue(model(bundle)['inventory'][0]['diagnostics'])


@fictional_publication
class DeviceEvidenceOutputTests(unittest.TestCase):
    def test_customer_summary_and_technical_raw_counts_replay(self):
        from openpyxl import load_workbook
        from Core.export_recommendations import export_to_excel
        from Core.workbook_layout import technical_workbook_path
        from Core.customer_report import render_customer_report
        from Core.assessment_serialization import write_assessment_result,read_assessment_result
        bundle=fixture(setting='BitLocker')
        device=bundle['assessment_sources']['managed_devices'][0]['records'][0]
        device.update(deviceName='Fictional protected laptop',userPrincipalName='fictional@example.invalid',osVersion='0.0.0.0')
        rows=expanded_findings(bundle,evaluation_date=DAY)
        bundle['assessment_result']=result(bundle)
        with TemporaryDirectory() as folder:
            path=Path(export_to_excel(rows,filename='Fictional.xlsx',evidence_bundle=bundle,output_dir=folder))
            assessment=bundle['assessment_result']
            customer=load_workbook(path);technical=load_workbook(technical_workbook_path(path))
            try:
                customer_text='\n'.join(str(cell.value or '') for sheet in customer for line in sheet for cell in line)
                self.assertNotIn('Fictional protected laptop',customer_text)
                self.assertNotIn('fictional@example.invalid',customer_text)
                self.assertNotIn('0.0.0.0',customer_text)
                self.assertIn('Device Detail',technical.sheetnames)
                self.assertIn('Device Finding Counts',technical.sheetnames)
                self.assertIn('Device Identity',technical.sheetnames)
                self.assertIn('Device Diagnostics',technical.sheetnames)
                raw=technical['Raw Managed Devices'];headers=[cell.value for cell in raw[1]]
                self.assertEqual(raw.cell(2,headers.index('osVersion')+1).value,'0.0.0.0')
                for finding in assessment['recommendations']:
                    if not finding.get('DeviceEvidenceVersion'):continue
                    count=finding['DeviceEvidenceCounts']['source_record_count']
                    self.assertEqual(count,finding['AssessmentEvidenceCount'])
                    self.assertEqual(count,len(finding['InvestigationEvidence']['records']))
                    self.assertEqual(count,len(finding['DeviceEvidenceIds']))
                    if count:
                        self.assertTrue(finding['TechnicalEvidenceRanges'])
                        raw_ids={cell.value for sheet in technical for line in sheet for cell in line if isinstance(cell.value,str) and cell.value.startswith('EVD-')}
                        self.assertTrue(set(finding['DeviceEvidenceIds'])<=raw_ids)
                summary=list(customer['Devices'].values)
                self.assertLessEqual(len(summary)-1,4)
                html=render_customer_report(assessment,bundle,'Fictional',path,technical_workbook_path=technical_workbook_path(path))
                from tools.audit_report_signal import public_report_source
                self.assertNotIn('739915',public_report_source(html))
                self.assertNotIn('fictional@example.invalid',public_report_source(html))
            finally:
                customer.close();technical.close()
            snapshot=Path(folder)/'result.json'
            write_assessment_result(snapshot,assessment)
            restored=read_assessment_result(snapshot)
            self.assertEqual(restored['device_compliance'],json.loads(json.dumps(assessment['device_compliance'],default=str)))
            self.assertEqual(html,render_customer_report(restored,bundle,'Fictional',path,technical_workbook_path=technical_workbook_path(path)))


if __name__=='__main__':
    unittest.main()
