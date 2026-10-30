"""Failed requirements and unresolved diagnostics must not become positive claims."""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from provenance_fixture import fictional_publication
from test_device_compliance_reasons import fixture
from test_device_example_details import render
from test_evidence_reconciliation import DAY
from Core.device_compliance import build_device_compliance
from Core.device_compliance_text import requirement_result, group_settings, group_state


def translated_fixture():
    bundle = fixture(setting='DefaultDeviceCompliancePolicy.RequireRemainContact')
    raw = bundle['assessment_sources']['compliance_setting_states'][0]['records']
    raw[0]['state'] = 'nonCompliant'
    raw.extend([dict(raw[0], id='activity-duplicate'),
                dict(raw[0], id='user', settingName='DefaultDeviceCompliancePolicy.RequireUserExistence'),
                dict(raw[0], id='assignment', settingName='DefaultDeviceCompliancePolicy.RequireDeviceCompliancePolicyAssigned', state='ERROR'),
                dict(raw[0], id='assignment-duplicate', settingName='DefaultDeviceCompliancePolicy.RequireDeviceCompliancePolicyAssigned', state='ERROR')])
    return bundle


class ResultTranslationTests(unittest.TestCase):
    def test_failed_builtin_labels_are_unmet_requirements(self):
        for labels, expected in [
            (['Is active', 'RequireRemainContact', 'DefaultDeviceCompliancePolicy.RequireRemainContact'], 'Device may be inactive/not recently checked in'),
            (['Enrolled user exists', 'RequireUserExistence', 'DefaultDeviceCompliancePolicy.RequireUserExistence'], 'Enrolled user could not be confirmed')]:
            for label in labels:
                with self.subTest(label=label):
                    self.assertEqual(requirement_result(label, 'nonCompliant'), expected)

    def test_unresolved_assignment_is_never_a_confirmed_missing_policy(self):
        for state in ['ERROR', 'unknown', 'notEvaluated', 'conflict', None]:
            text = requirement_result('Has a compliance policy assigned', state)
            self.assertEqual(text, 'Compliance policy assignment check: result unresolved')
            self.assertNotIn('No applicable', text)

    def test_passing_check_does_not_receive_a_failure_description(self):
        self.assertEqual(requirement_result('Is active', 'compliant'), 'Device activity/check-in requirement')

    def test_missing_status_is_an_unresolved_diagnostic(self):
        bundle = fixture(setting='Has a compliance policy assigned')
        bundle['assessment_sources']['compliance_setting_states'][0]['records'][0].pop('state')
        result, html = render(bundle)
        self.assertTrue(any(row.get('DeviceFindingType') == 'ReportingErrorOrConflict' for row in result['recommendations']))
        self.assertIn('Compliance policy assignment check: result unresolved', html)
        self.assertNotIn('No applicable compliance policy reported', html)

    def test_diagnostics_do_not_enter_device_noncompliance_causes(self):
        model = build_device_compliance(translated_fixture(), evaluation_date=DAY)
        device = model['devices'][0]
        self.assertIn('Device may be inactive/not recently checked in', device['reason'])
        self.assertIn('Enrolled user could not be confirmed', device['reason'])
        self.assertNotIn('assignment', device['reason'])
        self.assertIn('assignment check: result unresolved', device['diagnostics'])
        self.assertEqual(device['policy_assignment'], 'Unknown')

    def test_duplicate_checks_are_grouped_but_source_occurrences_survive(self):
        bundle = translated_fixture(); before = deepcopy(bundle)
        model = build_device_compliance(bundle, evaluation_date=DAY)
        self.assertEqual(len(model['settings']), 5)
        self.assertEqual(len(group_settings(model['settings'])), 3)
        self.assertEqual(bundle, before)
        result, html = render(bundle)
        row = next(row for row in result['recommendations'] if row.get('DeviceFindingType') == 'ReportingErrorOrConflict')
        self.assertIn('Distinct device/setting checks: Error: 1', row['DeviceSummary']['text'])
        self.assertIn('Compliance policy assignment check: result unresolved', html)
        self.assertNotIn('>Has a compliance policy assigned<', html)
        self.assertNotIn('>Is active<', html)
        self.assertNotIn('>Enrolled user exists<', html)
        self.assertIn('No error code or diagnostic message retained', html)
        activity = html.split('<summary>Review device compliance activity and unconfirmed failures</summary>')[1].split('</details>')[0]
        self.assertIn('Device may be inactive/not recently checked in', activity)
        self.assertIn('Enrolled user could not be confirmed', activity)
        self.assertEqual(activity.count('Device may be inactive/not recently checked in'), 1)

    def test_different_statuses_remain_conflicting_and_do_not_confirm_a_cause(self):
        bundle = fixture(setting='RequireRemainContact')
        raw = bundle['assessment_sources']['compliance_setting_states'][0]['records']
        raw.append(dict(raw[0], id='other', state='ERROR'))
        model = build_device_compliance(bundle, evaluation_date=DAY)
        group = group_settings(model['settings'])[0]
        self.assertEqual(group_state(group), 'conflictingResults')
        self.assertEqual(model['devices'][0]['reason'], 'Reason not retained')
        self.assertIn('result unresolved', model['devices'][0]['diagnostics'])
        self.assertEqual(model['devices'][0]['compliance_activity'], 'unknown')

    def test_unknown_setting_or_identity_is_not_collapsed_into_another_check(self):
        item = {'device_id': None, 'setting': 'Is active', 'record': {'state': 'ERROR'}}
        self.assertEqual(len(group_settings([item, deepcopy(item)])), 2)
        item.update(device_id='known', setting='Setting name not retained')
        self.assertEqual(len(group_settings([item, deepcopy(item)])), 2)

    def test_assignment_failure_and_error_preserve_uncertainty(self):
        bundle = fixture(setting='Has a compliance policy assigned')
        raw = bundle['assessment_sources']['compliance_setting_states'][0]['records']
        raw.append(dict(raw[0], id='unresolved', state='ERROR'))
        device = build_device_compliance(bundle, evaluation_date=DAY)['devices'][0]
        self.assertEqual(device['policy_assignment'], 'Conflicting assignment evidence')
        self.assertEqual(device['reason'], 'Reason not retained')
        self.assertIn('result unresolved', device['diagnostics'])


@fictional_publication
class ResultTranslationWorkbookTests(unittest.TestCase):
    def test_workbook_separates_diagnostics_and_groups_checks_with_all_lineage(self):
        from Core.export_recommendations import export_to_excel
        from Core.workbook_layout import technical_workbook_path
        from openpyxl import load_workbook
        bundle = translated_fixture()
        result, _ = render(bundle)
        bundle['assessment_result'] = result
        with TemporaryDirectory() as folder:
            path = Path(export_to_excel(result['recommendations'], filename='Fictional.xlsx', evidence_bundle=bundle, output_dir=folder))
            book = load_workbook(technical_workbook_path(path))
            try:
                values = list(book['Device Detail'].values)
                rows = [dict(zip(values[0], row)) for row in values[1:]]
                settings = [row for row in rows if row['Source dataset'] == 'compliance_setting_states']
                keys = [(row['ID'], row['Managed device ID'], row['Setting key']) for row in settings]
                self.assertEqual(len(keys), len(set(keys)))
                error = next(row for row in settings if row['Setting status'] == 'Error')
                self.assertEqual(error['Source occurrences'], 2)
                self.assertEqual(error['Result type'], 'Unresolved diagnostic')
                self.assertFalse(error['Compliance reason'])
                self.assertIn('result unresolved', error['Issue'])
                self.assertIn('No error code or diagnostic message retained', error['Diagnostic detail'])
                self.assertEqual(len(str(error['Evidence ID']).splitlines()), 2)
                raw = book['Raw Compliance Setting States']
                self.assertEqual(raw.max_row-1, 5)
                raw_text = '\n'.join(str(cell or '') for row in raw.values for cell in row)
                self.assertIn('RequireRemainContact', raw_text)
                self.assertIn('nonCompliant', raw_text)
                self.assertIn('ERROR', raw_text)
            finally:
                book.close()


if __name__ == '__main__':
    unittest.main()
