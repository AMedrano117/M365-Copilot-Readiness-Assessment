import os
import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from html import escape
from datetime import date
import re


from openpyxl import load_workbook

from Core.export_recommendations import (
    export_to_excel,
    export_to_html,
    print_recommendations_summary,
)


class ExportRecommendationsTests(unittest.TestCase):
    def setUp(self):
        self.original_cwd = os.getcwd()
        self.tempdir = tempfile.TemporaryDirectory()
        os.chdir(self.tempdir.name)

    def tearDown(self):
        os.chdir(self.original_cwd)
        self.tempdir.cleanup()

    def _sample_recommendations(self):
        return [
            {
                "RecommendationId": "DEF-001",
                "TenantId": "11111111-1111-1111-1111-111111111111",
                "Service": "Defender",
                "Feature": "Copilot Security Posture",
                "Status": "Warning",
                "Priority": "High",
                "Observation": "14 apps with high-level access were identified.",
                "Recommendation": "Review the flagged apps and security signals before rollout.",
                "LinkText": "Copilot Security",
                "LinkUrl": "https://example.com",
                "EvidenceAvailable": "Yes",
                "EvidenceSheet": "App Access Detail",
                "EvidenceSummary": "See App Access Detail for the flagged applications, reasons, and available activity.",
            },
            {
                "RecommendationId": "ENT-001",
                "TenantId": "11111111-1111-1111-1111-111111111111",
                "Service": "Entra",
                "Feature": "Conditional Access",
                "Status": "Success",
                "Priority": "",
                "Observation": "Conditional Access requires MFA for administrative access.",
                "Recommendation": "",
                "EvidenceAvailable": "Yes",
                "EvidenceBasis": "Tenant evidence",
                "Disposition": "Assurance",
                "ImpactArea": "Identity & access",
            }
        ]

    def _sample_evidence_bundle(self):
        return {
            "evaluation_date": "2026-09-15",
            "collection_context": {"collected_at": "2026-09-14T12:00:00Z", "tenant_id": "example-tenant", "mode": "offline"},
            "evidence_index": [
                {
                    "RecommendationId": "DEF-001",
                    "Service": "Defender",
                    "Feature": "Copilot Security Posture",
                    "Evidence Available": "Yes",
                    "Workbook Tab": "App Access Detail",
                    "Engineer Follow-Up": "See App Access Detail for the flagged applications, reasons, and available activity.",
                }
            ],
            "sheets": {
                "app_access_detail": {
                    "title": "App Access Detail",
                    "rows": [
                        {
                            "RecommendationId": "DEF-001",
                            "Flagged By": "Copilot Security Posture",
                            "App Display Name": "Risky App",
                            "Flagged Because": "High-privilege delegated permissions",
                            "Activity Band": "Unavailable",
                            "Last Activity": "",
                        }
                    ],
                }
            },
            "appendix_sections": [
                {
                    "key": "app_access_detail",
                    "title": "Appendix: Application Access Detail",
                    "workbook_tab": "App Access Detail",
                    "summary": "1 unique application was flagged.",
                    "details": [
                        "The workbook identifies which applications were reviewed.",
                        "Activity context was unavailable from the reviewed source data.",
                    ],
                    "preview_columns": ["App Display Name", "Flagged Because", "Activity Band"],
                    "preview_rows": [
                        {
                            "RecommendationId": "DEF-001",
                            "Flagged By": "Copilot Security Posture",
                            "App Display Name": "Risky App",
                            "Flagged Because": "High-privilege delegated permissions",
                            "Activity Band": "Unavailable",
                        }
                    ],
                }
            ],
        }

    def test_export_to_excel_creates_multi_sheet_evidence_workbook(self):
        recommendations = self._sample_recommendations()
        bundle = self._sample_evidence_bundle()
        excel_path = export_to_excel(recommendations, filename="tenant_report.xlsx", tenant_name="Contoso", evidence_bundle=bundle)
        workbook = load_workbook(excel_path)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames[:3], ["Action Plan", "Evidence Index", "Collection Coverage"])
        for title in ("Recommendations", "Evidence Index", "App Access Detail", "Assessment Summary", "Domain Results", "Evidence Observations"):
            self.assertIn(title, workbook.sheetnames)
            self.assertEqual(len(workbook[title].tables), 1)
        for worksheet in workbook.worksheets:
            if worksheet.tables:
                self.assertIsNone(worksheet.auto_filter.ref)
        rows = list(workbook["Action Plan"].values)
        actions = [dict(zip(rows[0], row)) for row in rows[1:]]
        self.assertEqual([row["RecommendationId"] for row in actions], [row["RecommendationId"] for row in bundle["assessment_result"]["actions"]])
        self.assertTrue(all(row["Responsible Role"] and row["Rollout Stage"] and row["Completion Evidence"] for row in actions))
        self.assertTrue(all(row["Target Date"] is None for row in actions))
        evidence_column = rows[0].index("Evidence") + 1
        self.assertTrue(all(workbook["Action Plan"].cell(i, evidence_column).hyperlink for i in range(2, len(rows)+1)))
        headers = [cell.value for cell in workbook["App Access Detail"][1]]
        self.assertIn("RecommendationId", headers)

    def test_export_to_html_renders_engineer_appendix_and_follow_up_reference(self):
        recommendations = self._sample_recommendations()
        evidence_bundle = self._sample_evidence_bundle()
        html_path = export_to_html(recommendations, filename="tenant_report.html", tenant_name="Contoso",
            evidence_bundle=evidence_bundle, excel_path=str(Path("Reports") / "tenant_report.xlsx"))
        html = Path(html_path).read_text(encoding="utf-8")
        action_plan = html.split('id="action-plan"', 1)[1].split("</section>", 1)[0]
        self.assertIn("What we found", action_plan)
        self.assertIn("What to do", action_plan)
        self.assertNotIn("Control ID", action_plan)
        self.assertNotIn("DEF-001", action_plan)
        self.assertIn("Observed strengths", html)
        self.assertIn("Why it matters", html)
        self.assertIn("Technical appendix and evidence workbook", html)
        self.assertIn('<details class="appendix-panel" id="engineer-appendix">', html)
        self.assertIn('href="tenant_report.xlsx"', html)
        appendix = html.split('id="engineer-appendix"', 1)[1]
        self.assertIn("App Access Detail", appendix)
        self.assertIn("DEF-001", appendix)
        self.assertIn("Review the flagged apps and security signals before rollout.", html)
        for anchor in re.findall(r'href="#(evidence-[^"]+)"', html):
            self.assertIn('id="' + anchor + '"', appendix)
        self.assertNotIn("Run manifest and collection outcomes", html)

    def test_export_to_html_explains_opportunity_and_external_ai_decision_effects(self):
        recommendations = self._sample_recommendations() + [
            {
                "RecommendationId": "M365-001",
                "Service": "M365",
                "Feature": "Teams pilot opportunity",
                "FindingKey": "adoption.reviewed_teams_pilot_opportunity",
                "Disposition": "Opportunity",
                "EvidenceBasis": "Reviewed pilot opportunity",
                "EvidenceKey": "m365_activity_detail",
                "ObservationDate": date.today().isoformat(),
                "EvidenceScope": "Reviewed 20-user Teams pilot cohort",
                "TenantId": "11111111-1111-1111-1111-111111111111",
                "EvidenceComplete": True,
                "SourceType": "reviewed_workshop",
                "SourceFile": "pilot-review.json",
                "Status": "Insight",
                "Priority": "Low",
                "Observation": "Teams activity suggests a focused pilot population.",
                "Recommendation": "Run a measured pilot with a named business owner.",
            }
        ]

        html_path = export_to_html(
            recommendations,
            filename="decision_explanation.html",
            tenant_name="Contoso",
        )

        html = Path(html_path).read_text(encoding="utf-8")
        self.assertIn("Adoption opportunities guide the value of a pilot", html)
        self.assertIn("They do not establish that security controls are effective", html)
        self.assertIn("Teams activity suggests a focused pilot population", html)
        self.assertIn("How to decide whether to expand", html)
        self.assertNotIn("Three readiness conclusions", html)
        self.assertNotIn("Provider and tier approval", html)
        self.assertNotIn("complete once per product and subscription tier", html)

    def test_html_surfaces_entra_license_data_scans_and_dlp_reference(self):
        evidence_bundle = self._sample_evidence_bundle()
        evidence_bundle.update({
            "entra_license_context": {
                "detected_tier": "Microsoft Entra ID P1",
                "rows": [{
                    "Capability": "Full Identity Protection risk data and risk-based Conditional Access",
                    "P1": "Not included", "P2": "Included", "Tenant": "Not included with detected P1",
                }],
            },
            "data_exposure": {"sources": {
                "sam": {"files_loaded": 0},
                "dspm": {"files_loaded": 1, "freshness": "fresh", "reports": [{
                    "source_file": "dspm-assessment.csv", "report_type": "dspm_assessment",
                    "report_date": "2026-09-14", "status": "selected", "records_read": 1,
                    "freshness": "fresh",
                }]},
            }},
            "purview_policy_summary": {
                "available": True, "total": 1, "enabled": 1,
                "rules_available": True, "total_rules": 1, "enabled_rules": 1,
                "rows": [{
                    "Policy": "Protect financial data", "Enabled": "Yes", "Mode": "Enable",
                    "Locations": "Exchange, SharePoint", "Rules": 1,
                    "Protection behavior": "Block access, Notify users",
                }],
                "rule_rows": [{
                    "Policy": "Protect financial data", "Rule": "Block financial records",
                    "Enabled": "Yes", "Conditions": "Sensitive information",
                    "Actions": "Block access, Notify users", "Severity": "High",
                }],
            },
        })

        evidence_bundle["sheets"]["purview_policy_detail"] = {
            "title": "Purview Policy Detail", "rows": [
                {"Detail Type": "DLP policy", **evidence_bundle["purview_policy_summary"]["rows"][0]},
                {"Detail Type": "DLP rule", **evidence_bundle["purview_policy_summary"]["rule_rows"][0]},
            ],
        }
        html_path = export_to_html(
            self._sample_recommendations(), filename="license_and_data.html",
            tenant_name="Contoso", evidence_bundle=evidence_bundle,
        )
        html = Path(html_path).read_text(encoding="utf-8")

        self.assertIn("Microsoft Entra ID P1", html)
        self.assertIn("Full Identity Protection risk data", html)
        self.assertIn("Confirm tenant sharing defaults", html)
        self.assertIn("dspm-assessment.csv", html)
        self.assertIn("2026-09-14", html)
        self.assertIn("selected", html)
        self.assertNotIn("Current report assessed", html)
        self.assertIn("Confirm data loss prevention coverage and enforcement", html)
        path = export_to_excel(self._sample_recommendations(), filename="license_and_data.xlsx", evidence_bundle=evidence_bundle)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        cells = [str(cell.value or "") for row in workbook["Purview Policy Detail"] for cell in row]
        self.assertIn("Protect financial data", cells)
        self.assertIn("Block financial records", cells)
        self.assertIn("Block access, Notify users", cells)

    def test_html_uses_curated_configured_strengths_and_excludes_license_entitlements(self):
        recommendations = self._sample_recommendations() + [{
            "Service": "Purview", "Feature": "Power Automate Free", "Status": "Success",
            "Priority": "", "Observation": "Power Automate Free is active in Power Automate Free.",
            "Recommendation": "", "Disposition": "Assurance", "EvidenceAvailable": "Yes",
            "EvidenceBasis": "Tenant evidence",
        }]
        evidence_bundle = self._sample_evidence_bundle()
        evidence_bundle["verified_strengths"] = [{
            "Area": "Data loss prevention", "TenantId": "11111111-1111-1111-1111-111111111111",
            "Strength": "DLP policies and their protection rules are enabled.",
            "Evidence": "2 enabled policies and 4 enabled rules were returned.",
            "Benefit": "Configured restrictions can be applied when sensitive content is shared.",
        }]

        html_path = export_to_html(
            recommendations, filename="curated_strengths.html", tenant_name="Contoso",
            evidence_bundle=evidence_bundle,
        )
        html = Path(html_path).read_text(encoding="utf-8")

        self.assertIn("DLP policies and their protection rules are enabled", html)
        path = export_to_excel(recommendations, filename="curated_strengths.xlsx", evidence_bundle=evidence_bundle)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        cells = [str(cell.value or "") for sheet in workbook for row in sheet for cell in row]
        self.assertTrue(any("2 enabled policies and 4 enabled rules" in value for value in cells))
        self.assertNotIn("Power Automate Free is active", html)

    def test_html_leads_with_tenant_specific_readiness_and_hides_empty_scope(self):
        evidence_bundle = self._sample_evidence_bundle()
        evidence_bundle.update({
            "verified_strengths": [{
                "Area": "Data loss prevention", "TenantId": "11111111-1111-1111-1111-111111111111",
                "Strength": "DLP rules are enabled.",
                "Evidence": "Four enabled rules were returned.",
                "Benefit": "Configured restrictions can be applied.",
            }],
            "data_exposure": {"sources": {
                "sam": {"files_loaded": 0},
                "dspm": {"files_loaded": 0},
            }},
            "ai_usage": {"copilot_usage": {
                "available": True, "refresh_date": "2026-09-14", "availability_status": "available",
                "selected_period": "D28",
                "periods": {"D28": {
                    "enabled_users": 6,
                    "active_users": 1,
                    "total_prompts": 85,
                }},
            }},
            "conclusions": {
                "provider_conclusion": "Not assessed",
                "use_case_conclusion": "Not assessed",
            },
        })

        html_path = export_to_html(
            self._sample_recommendations(), filename="customer_readout.html",
            tenant_name="Contoso", evidence_bundle=evidence_bundle,
        )
        html = Path(html_path).read_text(encoding="utf-8")

        result = evidence_bundle["assessment_result"]
        executive = html.split('id="executive"', 1)[1].split("</section>", 1)[0]
        self.assertIn(escape(result["decision"]), executive)
        self.assertIn("First actions", executive)
        self.assertIn("14 apps with high-level access were identified", html)
        self.assertNotIn("Built offline", executive)
        self.assertNotIn("source_file", executive)
        self.assertNotIn("Not established to Not established", html)
        for key, label in (("remediation", "Remediation actions"), ("confirmation", "Findings to confirm"), ("evidence_gaps", "Evidence checks"), ("strengths", "Verified strengths")):
            self.assertIn(f'<div class="value">{result["counts"][key]}</div><div class="stat-label">{label}</div>', executive)
        self.assertIn(f'>{len(result["actions"])} actions</span>', html)
        self.assertEqual(len(re.findall(r'<article class="action" ', html)), len(result["actions"]))
        self.assertIn("Active users", html)
        self.assertIn("85", html)
        self.assertNotIn('id="domain-external_ai"', html)
        self.assertNotIn("Three readiness conclusions", html)

    def test_html_shows_provider_and_use_case_review_only_when_records_are_supplied(self):
        evidence_bundle = self._sample_evidence_bundle()
        evidence_bundle["assessment_profile"] = {"scope": {"external_ai": True}}
        evidence_bundle["conclusions"] = {
            "provider_conclusion": "Approved",
            "use_case_conclusion": "Ready",
            "providers": [{
                "Provider": "Contoso AI", "Product": "Enterprise Chat", "Tier": "Business",
                "Approval Status": "Approved", "Reason": "Current review supplied.",
            }],
            "use_cases": [{
                "Use Case": "Proposal review", "Business Owner": "Sales Operations",
                "Provider": "Contoso AI", "Readiness": "Ready for controlled pilot",
                "Reason": "Required evidence is complete.",
            }],
        }

        html_path = export_to_html(
            self._sample_recommendations(), filename="provider_scope.html",
            tenant_name="Contoso", evidence_bundle=evidence_bundle,
        )
        html = Path(html_path).read_text(encoding="utf-8")

        self.assertIn('id="domain-external_ai"', html)
        self.assertIn("Contoso AI", html)
        self.assertIn("Proposal review", html)

    def test_html_translates_sharepoint_settings_into_customer_language(self):
        evidence_bundle = self._sample_evidence_bundle()
        evidence_bundle["sharepoint_governance"] = {
            "available": True,
            "site_count": 12,
            "anyone_site_count": 3,
            "settings": {
                "SharingCapability": "ExternalUserAndGuestSharing",
                "OneDriveSharingCapability": "ExternalUserSharingOnly",
                "DefaultSharingLinkType": "AnonymousAccess",
                "RequireAnonymousLinksExpireInDays": 0,
                "ExternalUserExpirationRequired": False,
                "LegacyAuthProtocolsEnabled": True,
            },
        }

        html_path = export_to_html(
            self._sample_recommendations(), filename="friendly_sharepoint.html",
            tenant_name="Contoso", evidence_bundle=evidence_bundle,
        )
        html = Path(html_path).read_text(encoding="utf-8")

        self.assertIn("Anyone links, new guests, and existing guests", html)
        self.assertIn("New and existing guests; no Anyone links", html)
        self.assertIn("<td>Anyone link</td>", html)
        self.assertIn("<td>Disabled</td>", html)
        self.assertNotIn("ExternalUserAndGuestSharing", html)
        self.assertNotIn("AnonymousAccess", html)

    def test_console_summary_uses_same_curated_strength_count_as_report(self):
        output = io.StringIO()
        bundle = self._sample_evidence_bundle()
        bundle["verified_strengths"] = [{"Area": "Data loss prevention", "TenantId": "11111111-1111-1111-1111-111111111111", "Strength": "DLP rules are enabled.", "Evidence": "Four rules were returned."}]
        recommendations = self._sample_recommendations()
        html = Path(export_to_html(recommendations, evidence_bundle=bundle)).read_text(encoding="utf-8")
        expected = bundle["assessment_result"]["counts"]["strengths"]
        self.assertGreater(expected, 0)
        with redirect_stdout(output):
            print_recommendations_summary(recommendations, tenant_name="Contoso", evidence_bundle=bundle)
        self.assertIn(f"Verified strengths: {expected}", output.getvalue())
        self.assertIn(f'<div class="value">{expected}</div><div class="stat-label">Verified strengths</div>', html)


if __name__ == "__main__":
    unittest.main()
