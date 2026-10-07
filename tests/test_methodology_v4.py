"""Methodology 4 collection, operational confirmation and lossless outputs."""

import asyncio
import copy
import json
import unittest
from datetime import datetime,timezone
from types import SimpleNamespace
from unittest.mock import AsyncMock,patch

from Core.assessment_catalog import CHECKS,DOMAIN_SPECS,collect_assessment_sources,context_validation
from Core.assessment_result import build_assessment_result
from Core.collector_registry import selected_collector_ids,profile_permission_resources
from Core.expanded_collection import collect_entra_details
from Core.expanded_findings import expanded_findings
from Core.operational_evidence import reconcile_devices,operational_results,signin_record
from Core.raw_evidence import prepare_raw_details,split_evidence_sheets,safe_record
from Core.investigation_details import prepare_investigation_details
from Core.investigation_contract import validate_investigation_coverage
from Core.tenant_baseline import conditional_access_facts

TENANT='11111111-1111-1111-1111-111111111111'
DAY='2026-09-15'


def dataset(records,**metadata):
    return [{'records':records,'source':{'available':True,'availability_status':'available','complete':True,
            'collected_at':'2026-09-14T12:00:00Z','scope':'Returned tenant records',**metadata}}]


def operational_profile(tenant=TENANT,day='2026-09-14'):
    from Core.operational_evidence import OPERATION_CHECKS
    return {'version':'1.0','readiness_review':{'version':'1.0','tenant_id':tenant,
        'tenant_scope':{'id':'tenant','description':'Assessed tenant, all active users and devices',
            'reviewed_at':day,'reviewer_role':'Assessment owner','evidence_reference':'Synthetic tenant-wide scope'},
        'check_reviews':[{'check_id':check,'scope_id':'tenant','reviewed_at':day,
            'reviewer_role':'Accountable safeguard owner','evidence_reference':'Synthetic tested behavior and scope '+check,
            'rationale':'Tested effective behavior across the stated tenant-wide scope; reviewed exclusions and reporting gaps.',
            'tested_behavior':'Tested the safeguard and retained the expected and actual outcome.',
            'tested_scope':'All active tenant users, devices and planned content; exclusions reviewed.',
            'evidence_level':'observed_operation','result':'pass'} for check in OPERATION_CHECKS.values()]}}


def device_bundle():
    return {'assessment_sources':{
        'directory_devices':dataset([{'id':'object','deviceId':'stable','approximateLastSignInDateTime':'2026-09-14T12:00:00Z'}]),
        'managed_devices':dataset([{'id':'intune','azureADDeviceId':'stable','lastSyncDateTime':'2026-09-14T12:00:00Z',
                                   'operatingSystem':'Windows','complianceState':'compliant'}]),
        'machines':dataset([{'id':'mde','aadDeviceId':'stable','lastSeen':'2026-09-14T12:00:00Z','onboardingStatus':'Onboarded','healthStatus':'Active'}])}}


class ExpandedCollectionTests(unittest.IsolatedAsyncioTestCase):
    async def test_shadow_ai_paging_and_entity_failure_preserve_returned_catalog(self):
        from Core.ai_usage import collect_shadow_ai_usage
        async def fetch(client,path):
            if path.endswith('/uploadedStreams'):
                return {'value':[{'id':'first'}],'@odata.nextLink':'https://graph.microsoft.com/beta/security/dataDiscovery/cloudAppDiscovery/uploadedStreams?$skiptoken=two'},''
            if '$skiptoken=two' in path:return {'value':[{'id':'second'}]},''
            if path.endswith('/users'):return None,'HTTP 403 permission missing for entity detail'
            if path.endswith('/name'):return {'value':[{'name':'unjoined device'}]},''
            if path.endswith('/ipAddress'):return {'value':[{'ipAddress':'192.0.2.1'}]},''
            return {'value':[{'id':'app','displayName':'ChatGPT','userCount':2,'nested':{'risk':'returned'}}]},''
        with patch('Core.ai_usage._get_json',side_effect=fetch):
            result=await collect_shadow_ai_usage(object())
        self.assertEqual(len(result['streams']),2)
        self.assertEqual(len(result['discovered_app_records']),2)
        self.assertEqual(result['application_count'],1)
        self.assertEqual(result['availability_status'],'partial')
        self.assertEqual(len(result['entity_records']),4)
        self.assertIn('403',str(result['partial_errors']))
        self.assertEqual(result['applications'][0]['nested'],{'risk':'returned'})

    async def test_graph_paging_failure_retains_original_request_metadata(self):
        from Core.get_graph_client import GraphRestClient,GraphRequestError
        client=GraphRestClient(object())
        self.addAsyncCleanup(client.aclose)
        client.get_json=AsyncMock(side_effect=[{'value':[{'id':'retained'}],
            '@odata.nextLink':'https://graph.microsoft.com/v1.0/devices?$skiptoken=next'},
            GraphRequestError(403,'HTTP 403 missing permission')])
        result=await client.get_collection('/v1.0/devices',params={'$top':'999'})
        self.assertEqual(result['availability_status'],'partial')
        self.assertEqual(result['value'],[{'id':'retained'}])
        self.assertEqual(result['request_params'],{'$top':'999'})
        self.assertEqual(len(result['page_requests']),2)
        self.assertTrue(result['collection_started_at'])
        self.assertTrue(result['collection_completed_at'])
    async def test_independent_failures_preserve_parent_ids_nested_rows_and_pages(self):
        client=SimpleNamespace(collection_status={},compliance_policies=[{'id':'policy'}],
            managed_devices=[{'id':'intune','operatingSystem':'Windows'}],service_principals=[{'id':'app'}],ca_policies=[])
        async def fetch(path):
            if path.endswith('/owners'): return {'available':False,'availability_status':'unavailable','value':[], 'reason':'HTTP 403 permission missing'}
            if path.endswith('/appRoleAssignments'): return {'available':True,'availability_status':'partial','truncated':True,'pages_collected':2,'value':[{'id':'grant','nested':{'roles':['one','two']}}]}
            if 'deviceCompliancePolicySettingStateSummaries' in path: return {'available':False,'availability_status':'unavailable','value':[], 'reason':'Active Intune license unavailable'}
            return {'available':True,'availability_status':'available','pages_collected':1,'value':[{'id':'returned'}]}
        await collect_entra_details(client,fetch)
        grant=client.assessment_datasets['application_permissions'][0]
        self.assertEqual(grant['records'][0]['ParentId'],'app')
        self.assertEqual(grant['records'][0]['nested']['roles'],['one','two'])
        self.assertEqual(grant['source']['pages_collected'],2)
        self.assertFalse(grant['source']['complete'])
        self.assertIn('403',client.collection_status['application_owners:app']['reason'])
        self.assertIn('license',client.collection_status['compliance_setting_summaries']['reason'])
        self.assertEqual(client.collection_status['entra_recommendations']['evidence_quality'],'preview')

    async def test_restricted_auto_never_requests_excluded_datasets(self):
        client=SimpleNamespace(collection_status={},service_principals=[{'id':'app'}],ca_policies=[],compliance_policies=[],managed_devices=[])
        fetch=AsyncMock(return_value={'available':True,'availability_status':'available','value':[]})
        await collect_entra_details(client,fetch,permission_profile='restricted')
        paths=[call.args[0] for call in fetch.await_args_list]
        self.assertFalse(any('/directory/recommendations' in path or '/servicePrincipals/' in path or path=='/v1.0/devices' for path in paths))
        self.assertEqual(client.collection_status['directory_devices']['availability_status'],'not_requested')

    async def test_standard_default_collects_microsoft_recommendations_and_impacted_resources(self):
        client = SimpleNamespace(collection_status={})
        recommendations = [
            {'id': 'recommendation-one', 'displayName': 'Use modern authentication', 'status': 'active',
             'actionSteps': [{'stepNumber': 1, 'text': 'Review the impacted applications'}]},
            {'id': 'recommendation-two', 'displayName': 'Review application credentials', 'status': 'completed'},
        ]
        async def fetch(path):
            if path == '/beta/directory/recommendations':
                return {'available': True, 'availability_status': 'available', 'pages_collected': 2,
                        'value': copy.deepcopy(recommendations)}
            if path.endswith('/recommendation-one/impactedResources'):
                return {'available': True, 'availability_status': 'available', 'pages_collected': 2,
                        'value': [{'id': 'app-one', 'resourceType': 'application', 'status': 'active'},
                                  {'id': 'app-two', 'additionalDetails': [{'key': 'Client', 'value': 'Legacy'}]}]}
            if path.endswith('/recommendation-two/impactedResources'):
                return {'available': True, 'availability_status': 'available', 'value': []}
            return {'available': True, 'availability_status': 'available', 'value': []}
        reader = AsyncMock(side_effect=fetch)
        await collect_entra_details(client, reader)
        paths = {call.args[0] for call in reader.await_args_list}
        self.assertTrue({'/beta/directory/recommendations',
                         '/beta/directory/recommendations/recommendation-one/impactedResources',
                         '/beta/directory/recommendations/recommendation-two/impactedResources'} <= paths)
        parent = client.assessment_datasets['entra_recommendations'][0]
        self.assertEqual(parent['records'], recommendations)
        self.assertEqual(parent['source']['pages_collected'], 2)
        self.assertEqual(parent['source']['evidence_quality'], 'preview')
        child = next(item for item in client.assessment_datasets['entra_impacted_resources'] if item['records'])
        self.assertTrue(child['source']['complete'])
        self.assertEqual(child['source']['pages_collected'], 2)
        self.assertEqual([row['ParentId'] for row in child['records']], ['recommendation-one'] * 2)
        self.assertEqual(child['records'][1]['additionalDetails'], [{'key': 'Client', 'value': 'Legacy'}])

    async def test_recommendation_child_failure_keeps_parent_and_successful_children(self):
        client = SimpleNamespace(collection_status={})
        async def fetch(path):
            if path == '/beta/directory/recommendations':
                return {'available': True, 'availability_status': 'available',
                        'value': [{'id': 'allowed'}, {'id': 'denied'}]}
            if path.endswith('/denied/impactedResources'):
                return {'available': False, 'availability_status': 'unavailable', 'status_code': 403,
                        'reason': 'Permission missing', 'value': []}
            if path.endswith('/allowed/impactedResources'):
                return {'available': True, 'availability_status': 'available', 'value': [{'id': 'retained'}]}
            return {'available': True, 'availability_status': 'available', 'value': []}
        await collect_entra_details(client, fetch)
        self.assertTrue(client.collection_status['entra_recommendations']['complete'])
        self.assertEqual(client.collection_status['entra_impacted_resources:denied']['status_code'], 403)
        self.assertFalse(client.collection_status['entra_impacted_resources:denied']['complete'])
        rows = [row for item in client.assessment_datasets['entra_impacted_resources'] for row in item['records']]
        self.assertEqual(rows, [{'id': 'retained', 'ParentId': 'allowed'}])

    async def test_explicit_none_does_not_request_directory_recommendations(self):
        client = SimpleNamespace(collection_status={})
        reader = AsyncMock(return_value={'available': True, 'availability_status': 'available', 'value': []})
        await collect_entra_details(client, reader, preview_collectors='none')
        self.assertFalse(any('/directory/recommendations' in call.args[0] for call in reader.await_args_list))

    async def test_unexpected_child_failure_is_retained_without_discarding_other_reads(self):
        client=SimpleNamespace(collection_status={},service_principals=[],ca_policies=[],compliance_policies=[],managed_devices=[])
        async def fetch(path):
            if path.endswith('/devices'): raise RuntimeError('Unavailable directory inventory')
            return {'available':True,'availability_status':'available','value':[]}
        await collect_entra_details(client,fetch,preview_collectors='none')
        self.assertEqual(client.collection_status['directory_devices']['availability_status'],'unavailable')
        self.assertTrue(client.collection_status['authentication_strengths']['complete'])


class PolicyAndOperationTests(unittest.TestCase):
    def test_standard_auto_selections_and_permissions(self):
        services={'run_m365':True,'run_entra':True}
        self.assertTrue({'copilot_audit','shadow_ai','entra_recommendations'} <= set(selected_collector_ids(services)))
        self.assertFalse({'copilot_audit','shadow_ai','entra_recommendations'} & set(selected_collector_ids(services,permission_profile='restricted')))
        resources=profile_permission_resources('standard')
        self.assertTrue({'DirectoryRecommendations.Read.All','CloudApp-Discovery.Read.All','AuditLogsQuery.Read.All'} <= set(resources['00000003-0000-0000-c000-000000000000']))
        self.assertFalse({'DirectoryRecommendations.Read.All','CloudApp-Discovery.Read.All','AuditLogsQuery.Read.All'} & set(profile_permission_resources('restricted')['00000003-0000-0000-c000-000000000000']))

    def test_all_nine_domains_and_every_catalog_check_are_reported_without_sources(self):
        result=build_assessment_result([],{},evaluation_date=DAY,expected_tenant_id=TENANT)
        self.assertEqual(len(result['assessment_domains']),9)
        self.assertEqual(len(result['domain_coverage']),len(CHECKS))
        self.assertEqual(len({row['Check ID'] for row in result['domain_coverage']}),len(CHECKS))
        self.assertEqual(result['methodology_version'],'4.0.0')
        self.assertEqual(result['evidence_schema_version'],'1.1.0')

    def test_signin_failures_report_only_prior_mfa_and_unknown_are_separate(self):
        row=signin_record({'status':{'errorCode':0},'authenticationRequirement':'multiFactorAuthentication',
            'authenticationDetails':[{'succeeded':True,'authenticationStepRequirement':'MultiFactorAuthentication',
                                      'authenticationStepResultDetail':'MFA requirement satisfied by claim in the token'}],
            'appliedConditionalAccessPolicies':[{'id':'p','result':'reportOnlySuccess'}]})
        self.assertEqual(row['MFASatisfaction'],'previously_satisfied')
        self.assertEqual(len(row['ReportOnlyResults']),1)
        self.assertEqual(signin_record({})['SignInOutcome'],'unknown')
        self.assertEqual(signin_record({'status':{'errorCode':50076}})['SignInOutcome'],'failure')
        self.assertEqual(signin_record({'status':{'errorCode':0},'authenticationRequirement':'singleFactorAuthentication'})['MFARequirement'],'not_required')

    def test_configured_only_review_and_pilot_scope_cannot_establish_operation(self):
        profile=operational_profile()
        profile['readiness_review']['check_reviews'][0]['evidence_level']='configuration'
        operations,_=operational_results({},profile,evaluation_date=DAY,tenant_id=TENANT)
        self.assertEqual(operations['IDENTITY.AUTH']['result'],'unknown')
        profile=operational_profile();profile['readiness_review']['tenant_scope']['reviewed_at']='2026-01-01'
        self.assertEqual(operational_results({},profile,evaluation_date=DAY,tenant_id=TENANT)[0]['DATA.DLP']['result'],'unknown')

    def test_operational_review_requires_matching_tenant_and_conflicts_remain_unknown(self):
        profile=operational_profile()
        operations,_=operational_results({},profile,evaluation_date=DAY,tenant_id=TENANT)
        self.assertEqual(operations['DATA.DLP']['result'],'pass')
        self.assertEqual(operational_results({},profile,evaluation_date=DAY,tenant_id='different')[0]['DATA.DLP']['result'],'unknown')
        duplicate=dict(profile['readiness_review']['check_reviews'][0],result='fail')
        profile['readiness_review']['check_reviews'].append(duplicate)
        self.assertEqual(operational_results({},profile,evaluation_date=DAY,tenant_id=TENANT)[0]['IDENTITY.AUTH']['result'],'conflict')

    def test_operation_label_without_tested_behavior_and_scope_is_unknown(self):
        profile=operational_profile()
        profile['readiness_review']['check_reviews'][0].pop('tested_behavior')
        self.assertEqual(operational_results({},profile,evaluation_date=DAY,tenant_id=TENANT)[0]['IDENTITY.AUTH']['result'],'unknown')

    def test_audit_requires_usable_current_copilot_event_metadata(self):
        row={'id':'event','createdDateTime':'2026-09-14T12:00:00Z','operation':'CopilotInteraction'}
        bundle={'assessment_sources':{'copilot_audit':dataset([row])}}
        self.assertEqual(operational_results(bundle,{},evaluation_date=DAY)[0]['DATA.AUDIT']['result'],'pass')
        for changed in ({},dict(row,createdDateTime='2026-08-01'),dict(row,operation='Unrelated')):
            bundle['assessment_sources']['copilot_audit']=dataset([changed])
            self.assertEqual(operational_results(bundle,{},evaluation_date=DAY)[0]['DATA.AUDIT']['result'],'unknown')

    def test_dlp_and_label_exclusions_do_not_establish_tenant_wide_scope(self):
        from Core.tenant_baseline import dlp_facts,assess_data_protection_baseline
        from tests.test_tenant_baseline import purview
        policy={'Name':'Scoped','Mode':'Enable','Enabled':True,'ExchangeLocation':['All'],'ExchangeLocationException':['excluded'],
            'Locations':[{'Location':'Other','Inclusions':[{'Identity':'All'}]},
                         {'Location':'Copilot.M365','Inclusions':[{'Identity':'group'}]}]}
        self.assertFalse(dlp_facts([policy])['enforced_tenant'])
        row=assess_data_protection_baseline(purview(labels={'policies':[dict(policy,Mode='Enforce')]}))[0]
        self.assertEqual(row['Disposition'],'Coverage')

    def test_exclusions_or_grant_alternatives_cannot_establish_tenant_wide_mfa(self):
        from tests.test_tenant_baseline import policy
        base=policy('MFA');base['conditions']['users']['excludeUsers']=[]
        self.assertTrue(conditional_access_facts([base])['enforced_mfa_all'])
        for changed in ('excluded','or','location','unknown_strength','apps'):
            row=copy.deepcopy(base)
            if changed=='excluded':row['conditions']['users']['excludeGroups']=['group']
            elif changed=='or':row['grantControls'].update(operator='OR',builtInControls=['mfa','compliantDevice'])
            elif changed=='location':row['conditions']['locations']={'includeLocations':['trusted']}
            elif changed=='unknown_strength':row['grantControls']={'authenticationStrength':{'id':'unknown'}}
            elif changed=='apps':row['conditions']['applications']['includeApplications']=['Office365']
            with self.subTest(changed=changed):self.assertFalse(conditional_access_facts([row])['enforced_mfa_all'])


class DeviceAndSourceTests(unittest.TestCase):
    def test_endpoint_string_booleans_conflicts_and_original_report_dates(self):
        bundle=device_bundle()
        bundle['assessment_sources'].update(windows_protection=dataset([{'ParentId':'intune','realTimeProtectionEnabled':True,
            'signatureUpdateOverdue':False,'lastReportedDateTime':'2026-09-14'}]),
            antivirus_health=dataset([{'machineId':'mde','avIsSignatureUpToDate':'True','dataRefreshTimestamp':'2026-09-14'}]))
        self.assertEqual(operational_results(bundle,{},evaluation_date=DAY)[0]['ENDPOINT.POSTURE']['result'],'pass')
        bundle['assessment_sources']['antivirus_health'][0]['records'].append({'machineId':'mde','avIsSignatureUpToDate':'False','dataRefreshTimestamp':'2026-08-01'})
        self.assertEqual(operational_results(bundle,{},evaluation_date=DAY)[0]['ENDPOINT.POSTURE']['result'],'unknown')
        finding=expanded_findings(bundle,evaluation_date=DAY)[0]
        self.assertEqual(finding['ObservationDate'],'2026-08-01')
        self.assertEqual(finding['InvestigationEvidence']['records'][0]['avIsSignatureUpToDate'],'False')
        result=build_assessment_result([dict(finding,TenantId=TENANT)],bundle,evaluation_date=DAY,expected_tenant_id=TENANT)
        row=next(row for row in result['actions'] if row.get('FindingKey')==finding['FindingKey'])
        self.assertEqual(row['ActionType'],'Confirmation')

    def test_percentages_require_all_population_fields_and_nonzero_denominator(self):
        from Core.evidence_contract import normalize_observation
        base={'unit':'%','value':50,'numerator':1,'denominator':2,'availability':'available','complete':True,
            'scope':'Returned records','population':'Known users','window':'D7','observed_at':DAY,'tenant_id':TENANT}
        self.assertEqual(normalize_observation(base,evaluation_date=DAY)['value'],50)
        for field in ('numerator','denominator','scope','population','window'):
            row=dict(base);row.pop(field)
            self.assertIsNone(normalize_observation(row,evaluation_date=DAY)['value'])
        self.assertIsNone(normalize_observation(dict(base,denominator=0),evaluation_date=DAY)['value'])
    def test_stable_device_ids_reconcile_and_names_do_not_join(self):
        bundle=device_bundle();result=reconcile_devices(bundle,evaluation_date=DAY)
        self.assertEqual(result['denominator'],1)
        self.assertEqual(result['reporting_percent'],100)
        bundle['assessment_sources']['machines'][0]['records'][0].pop('aadDeviceId')
        result=reconcile_devices(bundle,evaluation_date=DAY)
        self.assertEqual(len(result['records']),2)
        self.assertIsNone(result['reporting_percent'])

    def test_partial_unknown_zero_stale_and_conflicting_populations_have_no_percentage(self):
        for change in ('partial','empty','conflicting','unknown','future'):
            bundle=device_bundle()
            if change=='partial':bundle['assessment_sources']['directory_devices'][0]['source']['complete']=False
            elif change=='empty':bundle={'assessment_sources':{}}
            elif change=='conflicting':bundle['assessment_sources']['machines'][0]['records'].append(dict(bundle['assessment_sources']['machines'][0]['records'][0],onboardingStatus='CanBeOnboarded'))
            elif change=='unknown':bundle['assessment_sources']['machines'][0]['records'][0]['lastSeen']=''
            elif change=='future':bundle['assessment_sources']['machines'][0]['records'][0]['lastSeen']='2027-01-01'
            result=reconcile_devices(bundle,evaluation_date=DAY)
            if change=='unknown':self.assertFalse(result['records'][0]['Actively reporting'])
            else:self.assertIsNone(result['onboarding_percent'])
        bundle=device_bundle()
        for datasets in bundle['assessment_sources'].values():
            for record in datasets[0]['records']:
                for key in ('lastSeen','lastSyncDateTime','approximateLastSignInDateTime'):
                    if key in record:record[key]='2026-01-01'
        self.assertEqual(reconcile_devices(bundle,evaluation_date=DAY)['records'][0]['Activity'],'stale')

    def test_onboarding_reporting_and_healthy_protection_are_separate(self):
        bundle=device_bundle()
        self.assertEqual(operational_results(bundle,{},evaluation_date=DAY,tenant_id=TENANT)[0]['ENDPOINT.POSTURE']['result'],'unknown')
        bundle['assessment_sources'].update(windows_protection=dataset([{'ParentId':'intune','realTimeProtectionEnabled':True,'signatureUpdateOverdue':False,'lastReportedDateTime':'2026-09-14T12:00:00Z'}]),
            antivirus_health=dataset([{'machineId':'mde','avIsSignatureUpToDate':True,'dataRefreshTimestamp':'2026-09-14T12:00:00Z'}]))
        self.assertEqual(operational_results(bundle,{},evaluation_date=DAY,tenant_id=TENANT)[0]['ENDPOINT.POSTURE']['result'],'pass')
        bundle['assessment_sources']['windows_protection'][0]['records'][0]['realTimeProtectionEnabled']=False
        findings=expanded_findings(bundle,evaluation_date=DAY)
        self.assertEqual(findings[0]['Disposition'],'Action')
        self.assertEqual(findings[0]['InvestigationEvidence']['records'][0]['ParentId'],'intune')

    def test_per_device_protection_reads_produce_one_action_per_condition_and_freshness(self):
        dates=['2026-09-14T08:00:00Z','2026-09-13T09:00:00Z','2026-09-10T10:00:00Z','2026-07-06T00:00:00Z','2025-11-01T00:00:00Z']
        reads=[dataset([{'ParentId':f'device-{n}','realTimeProtectionEnabled':False,'signatureUpdateOverdue':False,
                         'lastReportedDateTime':date}],source_api=f'https://graph.microsoft.com/v1.0/deviceManagement/managedDevices/device-{n}/windowsProtectionState')[0]
               for n,date in enumerate(dates)]
        findings=expanded_findings({'assessment_sources':{'windows_protection':reads}},evaluation_date=DAY)
        self.assertEqual([row['FindingKey'] for row in findings],
                         ['expanded.windows_protection.disabled.current','expanded.windows_protection.disabled.older'])
        current,older=findings
        self.assertEqual(len(current['InvestigationEvidence']['records']),3)
        self.assertEqual(current['ObservationDate'],'2026-09-10')
        self.assertTrue(current['Observation'].startswith('3 Windows device'))
        self.assertIn('between 2026-09-10 and 2026-09-14',current['Observation'])
        self.assertIn('Endpoint security > Antivirus',current['Recommendation'])
        self.assertEqual(older['ObservationDate'],'2026-07-06')
        self.assertIn('{id}',current['InvestigationEvidence']['source']['source_api'])
        result=build_assessment_result([dict(row,TenantId=TENANT) for row in findings],{'assessment_sources':{'windows_protection':reads}},
                                       evaluation_date=DAY,expected_tenant_id=TENANT)
        titles=[row['Feature'] for row in result['actions'] if 'windows_protection' in str(row.get('FindingKey'))]
        self.assertEqual(sorted(titles),['Confirm the device protection baseline','Restore real-time antivirus protection'])

    def test_device_thresholds_exceptions_and_validation(self):
        profile={'assessment_context':{'device_activity_days':12,'endpoint_reporting_days':3,'exceptions':[
            {'id':'exception','device_id':'stable','owner':'Device owner','reason':'Approved retirement','approved_by':'Accountable owner',
             'approved_at':'2026-09-01','expires_at':'2026-09-20'}]}}
        result=reconcile_devices(device_bundle(),profile,evaluation_date=DAY)
        self.assertTrue(result['records'][0]['Excluded']);self.assertIsNone(result['denominator'])
        profile['assessment_context']['device_activity_days']=False
        self.assertTrue(context_validation(profile))

    def test_raw_sources_preserve_historical_rows_and_do_not_invent_completeness(self):
        sources=collect_assessment_sources([SimpleNamespace(managed_devices=[{'id':'saved','nested':{'values':['a','b']}}])],
            collected_at='2026-01-01',data_exposure={'sources':{'sam':{'files_loaded':1,'retained_records':[{'Item ID':'file','_evidence_date':'2025-12-01'}]}}})
        self.assertFalse(sources['managed_devices'][0]['source']['complete'])
        self.assertEqual(sources['content_exposure'][0]['records'][0]['_evidence_date'],'2025-12-01')


class FindingLinksTests(unittest.TestCase):
    def test_oversized_register_worksheets_preserve_every_row(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill,Font,Alignment
        from Core.export_recommendations import _append_dict_rows_to_sheet
        workbook=Workbook();self.addCleanup(workbook.close)
        sheet=workbook.active;sheet.title='Findings Register'
        with patch('Core.export_recommendations.EXCEL_MAX_DATA_ROWS',2):
            _append_dict_rows_to_sheet(sheet,[{'id':index,'value':'record'} for index in range(5)],PatternFill(),Font(),Alignment())
        self.assertEqual(workbook.sheetnames,['Findings Register','Findings Register 2','Findings Register 3'])
        self.assertEqual([row[0] for sheet in workbook for row in list(sheet.values)[1:]],[0,1,2,3,4])
    def test_all_findings_have_support_or_precise_gap_and_split_links_reconcile(self):
        rows=[{'id':str(index),'nested':{'value':'x'*40000}} for index in range(5)]
        rec={'RecommendationId':'REF','EvidenceKey':'conditional_access_detail','Disposition':'Reference',
             'EvidenceScope':'Returned configuration','ObservationDate':DAY,'EvidenceComplete':True,'EvidenceLevel':'configuration'}
        bundle={'sheets':{},'assessment_sources':{'ca_policies':dataset(rows)},'worksheet_row_capacity':2}
        result={'recommendations':[rec,{'RecommendationId':'GAP','Feature':'Missing source','EvidenceKey':'missing','Disposition':'Coverage'}], 'actions':[]}
        prepare_investigation_details(bundle,result)
        self.assertEqual(rec['InvestigationCount'],5)
        self.assertEqual(len(rec['InvestigationRanges']),3)
        self.assertFalse(validate_investigation_coverage(result,bundle))
        before=copy.deepcopy(rec['InvestigationRanges']);prepare_investigation_details(bundle,result)
        self.assertEqual(rec['InvestigationRanges'],before)
        self.assertTrue(result['recommendations'][1]['InvestigationQualification'])
        raw=[sheet for key,sheet in bundle['sheets'].items() if key.startswith('raw_source.')]
        self.assertEqual(sum(len(sheet['rows']) for sheet in raw),5)
        self.assertTrue(any('continued' in key for key in raw[0]['rows'][0]))

    def test_secrets_and_prompt_content_are_removed_while_usage_counts_survive(self):
        value=safe_record({'id':'row','passwordProfile':{'password':'secret'},'nested':{'PromptText':'private',
             'response':'private','promptsSubmitted':2,'roles':['read']}})
        self.assertNotIn('passwordProfile',value)
        self.assertEqual(value['nested'],{'promptsSubmitted':2,'roles':['read']})


if __name__=='__main__':unittest.main()
