import contextlib
import csv
import io
import json
import os
from pathlib import Path
import runpy
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from unittest.mock import patch

from Core.offline_collection import (
    empty_service_results, load_collection, save_collection, refresh_saved_freshness,
)
from Core.portal_report_import import route_portal_reports, validate_report_tenants

ROOT = Path(__file__).resolve().parents[1]


class OfflineReportTests(unittest.TestCase):
    def csv(self, path, rows):
        with path.open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.DictWriter(handle, fieldnames=rows[0])
            writer.writeheader()
            writer.writerows(rows)
        return path

    def saved(self, path, collected_at=None):
        results = empty_service_results()
        results['m365_result'] = [{'_client': SimpleNamespace(
            available=True, users_summary={'total': 3},
            external_connections=[], collection_status={},
            client_secret='DO_NOT_SAVE_SECRET', access_token='DO_NOT_SAVE_TOKEN',
        ), 'licenses': []}, [{
            'Service': 'M365', 'Feature': 'Saved control finding',
            'Observation': 'Evidence from the original collection.', 'Recommendation': 'Review this control.',
            'Status': 'Attention Required', 'Priority': 'Medium', 'Disposition': 'Action',
        }]]
        save_collection(path, tenant_id='11111111-1111-1111-1111-111111111111',
                        tenant_name='Example tenant', service_results=results, collected_at=collected_at)
        return path

    def test_collection_round_trip_excludes_secrets_and_preserves_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.saved(Path(directory) / 'collection.json')
            raw = path.read_text(encoding='utf-8')
            self.assertNotIn('DO_NOT_SAVE', raw)
            loaded = load_collection(path)
            self.assertEqual(loaded['service_results']['m365_result'][0]['_client'].users_summary['total'], 3)
            self.assertEqual(loaded['service_results']['m365_result'][1][0]['Feature'], 'Saved control finding')

    def test_assessment_snapshot_is_not_mistaken_for_collection(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'snapshot.json'
            path.write_text('{"recommendations": []}', encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'save-collection'):
                load_collection(path)

    def test_power_platform_inventory_and_legacy_transport_save_evidence_only(self):
        import asyncio
        import httpx
        from Core.power_platform_inventory import PowerPlatformInventoryData
        with tempfile.TemporaryDirectory() as directory:
            results = empty_service_results()
            inventory = PowerPlatformInventoryData()
            inventory.agents = [{'id': 'invented-agent', 'name': 'Example agent'}]
            results['power_platform_info']['_client'] = inventory
            transport = httpx.AsyncClient(headers={'Authorization': 'Bearer DO_NOT_SAVE'})
            try:
                transport.environments = [{'name': 'Example environment'}]
                transport.environment_summary = {'total': 1}
                results['copilot_studio_info']['_client'] = transport
                path = Path(directory) / 'collection.json'
                save_collection(path, tenant_id='example', tenant_name='Example', service_results=results)
                loaded = load_collection(path)['service_results']
                self.assertEqual(loaded['power_platform_info']['_client'].agents, inventory.agents)
                self.assertEqual(loaded['copilot_studio_info']['_client'].environment_summary, {'total': 1})
                self.assertNotIn('DO_NOT_SAVE', path.read_text(encoding='utf-8'))
                self.assertFalse(hasattr(loaded['copilot_studio_info']['_client'], 'headers'))
            finally:
                asyncio.run(transport.aclose())

    def test_no_executable_object_serialization(self):
        with tempfile.TemporaryDirectory() as directory:
            results = empty_service_results()
            results['m365_result'][0]['_client'] = object()
            with self.assertRaisesRegex(ValueError, 'non-data object'):
                save_collection(Path(directory) / 'bad.json', tenant_id='example',
                                tenant_name='Example', service_results=results)

    def test_old_usage_freshness_is_recomputed(self):
        results = empty_service_results()
        usage = {'available': True, 'freshness': 'Fresh',
                 'refresh_date': (datetime.now(timezone.utc) - timedelta(days=20)).date().isoformat()}
        results['m365_result'][0]['_client'] = SimpleNamespace(copilot_usage=usage)
        refresh_saved_freshness(results)
        self.assertEqual(usage['freshness'], 'Stale')

    def test_portal_routing_includes_header_only_label_report(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'unhelpful-filename.csv'
            path.write_text('Site ID,URL,Primary admin,Labeled files,SiteSensitivity,External sharing\n', encoding='utf-8')
            routed = route_portal_reports([directory])
            self.assertEqual(routed['sam'], [str(path)])
            self.assertFalse(routed['dspm'])

    def test_cross_tenant_export_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.csv(Path(directory) / 'report.csv', [{'Tenant ID': '22222222-2222-2222-2222-222222222222', 'Anyone link count': '1'}])
            with self.assertRaisesRegex(ValueError, 'does not match'):
                validate_report_tenants([str(path)], '11111111-1111-1111-1111-111111111111')

    def run_cli(self, args, directory):
        old_cwd = Path.cwd()
        try:
            os.chdir(directory)
            with patch('sys.argv', ['main.py', *args]), \
                 patch('Core.console_setup.setup_console_encoding'), \
                 patch('Core.credentials_check.load_env_file', side_effect=AssertionError('loaded credentials')), \
                 patch('Core.credentials_check.validate_credentials_or_exit', side_effect=AssertionError('validated credentials')), \
                 patch('socket.socket.connect', side_effect=AssertionError('network attempted')), \
                 patch('socket.create_connection', side_effect=AssertionError('network attempted')), \
                 patch('subprocess.Popen', side_effect=AssertionError('external process attempted')), \
                 patch('builtins.input', side_effect=AssertionError('interactive prompt attempted')), \
                 patch.dict(os.environ, {'SAM_DAG_REPORT_PATHS': str(Path(directory) / 'must-not-open.csv')}), \
                 contextlib.redirect_stdout(io.StringIO()), self.assertRaises(SystemExit) as exit_context:
                runpy.run_path(str(ROOT / 'main.py'), run_name='__main__')
            self.assertEqual(exit_context.exception.code, 0)
            return next((Path(directory) / 'Reports').glob('*.html'))
        finally:
            os.chdir(old_cwd)

    def test_cli_replay_builds_reports_without_auth_network_or_powershell(self):
        with tempfile.TemporaryDirectory() as directory:
            timestamp = (datetime.now(timezone.utc) - timedelta(days=40)).isoformat()
            path = self.saved(Path(directory) / 'collection.json', collected_at=timestamp)
            html = self.run_cli(['--collection-input', str(path)], directory)
            text = html.read_text(encoding='utf-8')
            self.assertIn('saved control finding', text.lower())
            self.assertIn('40 days old', text)
            self.assertTrue('Readiness unconfirmed' in text, 'Stale replay must not establish readiness.')
            self.assertTrue(list((Path(directory) / 'Reports').glob('*.xlsx')))
            self.assertNotIn('must-not-open', text)

    def test_repeating_reports_directory_reuses_packaged_readiness_snapshot_once(self):
        from Core.copilot_readiness_import import FLAG_COLUMNS
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            reports = folder / 'admin exports'
            reports.mkdir()
            row = {name: 'False' for name in FLAG_COLUMNS.values()}
            row.update({'Report Refresh Date': '2026-09-12', 'User Principal Name': 'invented@example.test', 'Report Period': '30'})
            self.csv(reports / 'readiness.csv', [row])
            collection = self.saved(folder / 'collection.json', collected_at='2026-09-15T12:00:00Z')
            original = collection.read_bytes()
            snapshots = []
            for index in range(2):
                snapshot = folder / f'snapshot-{index}.json'
                self.run_cli(['--mode', 'offline', '--collection-input', str(collection),
                              '--reports-dir', str(reports), '--evaluation-date', '2026-09-15',
                              '--snapshot-json', str(snapshot)], directory)
                snapshots.append(json.loads(snapshot.read_text(encoding='utf-8'))['assessment_result'])
            self.assertEqual(snapshots[0], snapshots[1])
            self.assertEqual(collection.read_bytes(), original)
            receipt = json.loads((folder / 'collection_package' / 'operator-log.jsonl').read_text(encoding='utf-8').splitlines()[-1])
            self.assertTrue(any(row['Status'] == 'duplicate' for row in receipt['import_receipt']))

    def test_distinct_readiness_snapshots_require_an_explicit_selection(self):
        from Core.portal_report_import import select_readiness_report
        with tempfile.TemporaryDirectory() as directory:
            first, second = Path(directory) / 'first.csv', Path(directory) / 'second.csv'
            first.write_text('first snapshot', encoding='utf-8')
            second.write_text('different snapshot', encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'More than one') as error:
                select_readiness_report([str(first), str(second)])
            self.assertIn('first.csv', str(error.exception))
            self.assertIn('second.csv', str(error.exception))
            self.assertIn('--reports-dir adds files', str(error.exception))
            self.assertIn('same tenant', str(error.exception))
            selected, receipt = select_readiness_report([str(first), str(second)], str(second))
            self.assertEqual(selected, [str(second)])
            self.assertEqual(receipt[0]['Status'], 'not_selected')

    def test_replay_user_usage_details_require_opt_in_again(self):
        from Core.cli_parser import parse_arguments
        from Core.offline_report import run_offline_report
        with tempfile.TemporaryDirectory() as directory:
            results = empty_service_results()
            results['m365_result'][0]['_client'] = SimpleNamespace(copilot_usage={
                'available': False, 'user_detail': [{'userPrincipalName': 'private@example.invalid'}],
            })
            path = Path(directory) / 'collection.json'
            save_collection(path, tenant_id='example', tenant_name='Example', service_results=results)
            for opted_in in (False, True):
                arguments = ['main.py', '--collection-input', str(path)]
                if opted_in:
                    arguments.append('--include-user-usage-detail')
                with patch('sys.argv', arguments):
                    args = parse_arguments(None, [])
                with patch('Core.processor.process_and_print_all_information', return_value={'html_path': 'test.html'}) as process:
                    self.assertEqual(run_offline_report(args), 0)
                usage = process.call_args.kwargs['m365_result'][0]['_client'].copilot_usage
                self.assertEqual(bool(usage.get('user_detail')), opted_in)

    def test_portal_only_cli_does_not_claim_complete_tenant_collection(self):
        with tempfile.TemporaryDirectory() as directory:
            report = self.csv(Path(directory) / 'CMA.csv', [{
                'Site name': 'Old site', 'URL': 'https://example.sharepoint.com/sites/old',
                'Is inactive': 'True', 'Is ownerless': 'True',
                'Email address of site owners': '', 'Site creation date (UTC)': '2020-01-01',
            }])
            html = self.run_cli(['--mode', 'offline', '--sam-report', str(report)], directory)
            text = html.read_text(encoding='utf-8')
            self.assertIn('Confirm sign-in policy coverage for the pilot', text)
            self.assertTrue('Readiness unconfirmed' in text, 'Portal exports alone must not establish readiness.')
            self.assertNotIn('must-not-open', text)

    def test_readiness_cli_import_is_separate_from_copilot_usage(self):
        with tempfile.TemporaryDirectory() as directory:
            from Core.copilot_readiness_import import FLAG_COLUMNS
            row = {'Report Refresh Date': '2026-09-12', 'Report Period': '30',
                   'User Principal Name': 'private-user@example.invalid',
                   **{header: 'Yes' for header in FLAG_COLUMNS.values()}}
            report = self.csv(Path(directory) / 'readiness.csv', [row])
            html = self.run_cli(['--offline', '--copilot-readiness-export', str(report)], directory)
            text = html.read_text(encoding='utf-8')
            self.assertIn('Copilot readiness portal export', text)
            self.assertNotIn('private-user@example.invalid', text)

    def test_prior_workbook_and_portal_export_build_together_offline(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as directory:
            prior = Path(directory) / 'prior.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'Recommendations'
            ws.append(['Service', 'Feature', 'Observation', 'Recommendation', 'Disposition', 'Priority'])
            ws.append(['Entra', 'Historical MFA finding', 'Prior MFA evidence.', 'Review sign-in protection.', 'Action', 'High'])
            manifest = wb.create_sheet('Run Manifest')
            manifest.append(['Item', 'Value'])
            manifest.append(['Tenant', 'Example tenant'])
            manifest.append(['Generation Time (UTC)', '2026-09-09T18:30:00+00:00'])
            detail = wb.create_sheet('Authentication Coverage')
            detail.append(['Metric', 'Value'])
            detail.append(['Prior MFA registration count', 12])
            wb.save(prior)
            wb.close()
            report = self.csv(Path(directory) / 'CMA.csv', [{
                'Site name': 'Old site', 'URL': 'https://example.sharepoint.com/sites/old',
                'Is inactive': 'True', 'Is ownerless': 'True',
                'Email address of site owners': '', 'Site creation date (UTC)': '2020-01-01',
            }])
            cache = Path(directory) / 'purview-cache.json'
            cache.write_text(json.dumps({
                'schema_version': 2, 'tenant_id': '11111111-1111-1111-1111-111111111111',
                'cached_at_epoch': datetime(2026, 9, 9, tzinfo=timezone.utc).timestamp(),
                'purview_data_json': json.dumps({'dlp_policies': {
                    'available': True, 'count': 1,
                    'policies': [{'Name': 'Example cached DLP policy', 'Mode': 'Enable'}],
                }}),
            }), encoding='utf-8')
            html = self.run_cli(['--mode', 'offline', '--prior-report', str(prior),
                                 '--purview-cache', str(cache), '--sam-report', str(report)], directory)
            body = html.read_text(encoding='utf-8')
            self.assertTrue('historical mfa finding' in body.lower(), 'Prior findings should appear with imported reports.')
            self.assertTrue('2026-09-09T18:30:00+00:00' in body, 'Prior generation date must be retained.')
            self.assertTrue('Content ownership and lifecycle' in body)
            self.assertTrue('Readiness unconfirmed' in body)
            self.assertTrue('purview-cache.json' in body, 'The Purview source must be identified.')
            combined = load_workbook(next((Path(directory) / 'Reports').glob('*.xlsx')), read_only=True)
            try:
                history = [s for s in combined if s.title.startswith('Prior')]
                self.assertTrue(history, 'Historical evidence must remain in the combined workbook.')
                values = [str(value) for sheet in history for row in sheet.values for value in row]
                self.assertIn('Prior MFA registration count', values)
                self.assertIn('Purview Policy Detail', combined.sheetnames)
            finally:
                combined.close()


if __name__ == '__main__':
    unittest.main()
