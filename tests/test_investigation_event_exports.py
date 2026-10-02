"""Customer callouts must resolve to the actual records in the saved collection."""

import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from urllib.parse import unquote

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from Core.evidence_layer import build_evidence_bundle
from Core.export_recommendations import export_to_excel, export_to_html
from Core.saved_control_checks import qualify_saved_recommendations


DAY = '2026-09-29T20:00:00Z'


def event(number, client='Authenticated SMTP', error=0):
    return {'id': f'event-{number}', 'userId': 'account-1', 'userPrincipalName': 'private-user@example.test',
            'userDisplayName': 'Private user', 'appId': 'app-1', 'appDisplayName': 'Private application',
            'resourceId': 'resource-1', 'resourceDisplayName': 'Mail', 'clientAppUsed': client,
            'createdDateTime': f'2026-09-29T10:00:{number:02d}Z', 'ipAddress': '192.0.2.42',
            'correlationId': f'correlation-{number}', 'isInteractive': False,
            'conditionalAccessStatus': 'failure' if error else 'notApplied',
            'status': {'errorCode': error, 'failureReason': 'Blocked request' if error else ''},
            'deviceDetail': {'deviceId': 'device-1', 'displayName': 'Private device'}}


def state():
    return {'available': True, 'availability_status': 'available', 'pages_collected': 2,
            'truncated': False, 'collected_at': DAY, 'source_api': '/v1.0/auditLogs/signIns',
            'window_start': '2026-09-22T20:00:00Z', 'filter': 'createdDateTime ge 2026-09-22T20:00:00Z'}


class InvestigationEventExportTests(unittest.TestCase):
    def setUp(self):
        self.cwd = os.getcwd()
        self.directory = tempfile.TemporaryDirectory()
        os.chdir(self.directory.name)

    def tearDown(self):
        os.chdir(self.cwd)
        self.directory.cleanup()

    def build(self, client, observation):
        original = {'Service': 'Entra', 'Feature': 'Microsoft Entra ID P1',
                    'Observation': observation, 'Recommendation': 'Review affected records.',
                    'Disposition': 'Action', 'Status': 'Action Required', 'Priority': 'High',
                    'ObservationDate': DAY, 'TenantId': '11111111-1111-1111-1111-111111111111'}
        recommendations = qualify_saved_recommendations([original], 'Entra', client)
        context = {'collected_at': DAY, 'source_file': 'saved-collection.json', 'mode': 'offline',
                   'tenant_id': original['TenantId']}
        bundle = build_evidence_bundle(recommendations, ({}, []), {'_client': client}, {}, {}, {}, {},
                                       collection_context=context)
        bundle['collection_context'] = context
        bundle['source_statuses'] = {'entra_' + key: value for key, value in client.collection_status.items()}
        return bundle

    def export(self, bundle):
        path = export_to_excel(bundle['recommendations'], filename='investigation.xlsx', evidence_bundle=bundle)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        return path, workbook

    def row(self, workbook, title, identifier):
        sheet = workbook[title]
        headers = {cell.value: cell.column for cell in sheet[1]}
        number = next(i for i in range(2, sheet.max_row + 1)
                      if sheet.cell(i, headers['RecommendationId']).value == identifier)
        return {key: sheet.cell(number, column) for key, column in headers.items()}

    def test_eight_events_export_individually_and_all_callouts_link_to_logs(self):
        logs = [event(i, error=50053 if i == 8 else 0) for i in range(1, 9)] + [event(9, 'Browser')]
        client = SimpleNamespace(signin_logs=logs, signin_summary={'legacy_auth_attempts': 8},
                                 collection_status={'signin_logs': state()})
        bundle = self.build(client, '8 legacy authentication sign-ins detected in the past 30 days, bypassing MFA and CA protections')
        path, workbook = self.export(bundle)
        sheet = workbook['Legacy Sign-In Detail']
        values = list(sheet.values)
        records = [dict(zip(values[0], row)) for row in values[1:]]
        self.assertEqual([row['Sign-In ID'] for row in records], [f'event-{i}' for i in range(1, 9)])
        self.assertEqual([row['Outcome'] for row in records].count('Success'), 7)
        self.assertEqual(records[-1]['Outcome'], 'Failure')
        for row in records:
            for field in ('User ID', 'User Principal Name', 'Application ID', 'Created UTC',
                          'IP Address', 'Correlation ID', 'Client App / Protocol'):
                self.assertTrue(row[field], field)
        rec = next(row for row in bundle['assessment_result']['actions'] if row.get('FindingKey') == 'entra.signins.legacy_auth')
        expected = f"'Legacy Sign-In Detail'!A2:{get_column_letter(sheet.max_column)}9"
        self.assertEqual(rec['InvestigationRange'], expected)
        for title, field in (('Action Plan', 'Investigation Details'), ('Evidence Index', 'Workbook Tab'),
                             ('Recommendations', 'Evidence Sheet')):
            row = self.row(workbook, title, rec['RecommendationId'])
            self.assertEqual(row[field].hyperlink.target, '#' + expected)
        action = self.row(workbook, 'Action Plan', rec['RecommendationId'])
        self.assertEqual(action['Investigation Details'].value, '8 sign-in records')
        self.assertNotIn('30 days', action['What We Found'].value)
        self.assertNotIn('bypassing', action['What We Found'].value)
        self.assertEqual(len([row for row in bundle['sheets']['investigation_items']['rows']
                              if row['Item Type'] == 'Sign-in event']), 8)
        html_path = export_to_html(bundle['recommendations'], filename='investigation.html', evidence_bundle=bundle, excel_path=path)
        html = Path(html_path).read_text(encoding='utf-8')
        self.assertIn('investigation.xlsx#' + expected, unquote(html))
        for private in ('private-user@example.test', '192.0.2.42', 'Private device', 'correlation-1'):
            self.assertNotIn(private, html)

    def test_mismatch_is_visible_and_does_not_pad_records_to_reach_the_summary(self):
        client = SimpleNamespace(signin_logs=[event(1), event(2), event(3)],
                                 signin_summary={'legacy_auth_attempts': 8}, collection_status={'signin_logs': state()})
        bundle = self.build(client, '8 legacy authentication sign-ins detected in the returned sign-in records.')
        _, workbook = self.export(bundle)
        self.assertEqual(workbook['Legacy Sign-In Detail'].max_row, 4)
        rec = next(row for row in bundle['assessment_result']['actions'] if row.get('FindingKey') == 'entra.signins.legacy_auth')
        qualification = self.row(workbook, 'Action Plan', rec['RecommendationId'])['Qualification'].value
        self.assertIn('8', qualification)
        self.assertIn('3', qualification)

    def test_aggregate_without_saved_logs_has_a_visible_missing_detail_explanation(self):
        client = SimpleNamespace(signin_summary={'legacy_auth_attempts': 8}, collection_status={'signin_logs': state()})
        bundle = self.build(client, '8 legacy authentication sign-ins detected in the returned sign-in records.')
        _, workbook = self.export(bundle)
        self.assertNotIn('Legacy Sign-In Detail', workbook.sheetnames)
        rec = next(row for row in bundle['assessment_result']['actions'] if row.get('FindingKey') == 'entra.signins.legacy_auth')
        action = self.row(workbook, 'Action Plan', rec['RecommendationId'])
        self.assertIn('Evidence unavailable', action['Investigation Details'].value)
        self.assertTrue(rec['InvestigationQualification'])
        self.assertIn(rec['InvestigationQualification'], action['Qualification'].value)
        html_path = export_to_html(bundle['recommendations'], filename='missing.html', evidence_bundle=bundle)
        html = Path(html_path).read_text(encoding='utf-8')
        self.assertIn('retained', html.lower())

    def test_mfa_review_lists_false_registrations_without_treating_unknown_as_false(self):
        registrations = [{'id': 'unregistered', 'userPrincipalName': 'unregistered@example.test',
                          'isMfaRegistered': False, 'userType': 'guest', 'methodsRegistered': ['email']},
                         {'id': 'registered', 'userPrincipalName': 'registered@example.test', 'isMfaRegistered': True},
                         {'id': 'unknown', 'userPrincipalName': 'unknown@example.test'}]
        client = SimpleNamespace(auth_methods_registration=registrations,
                                 collection_status={'auth_methods': state()})
        bundle = self.build(client, 'Only 1 of 3 users (33.3%) were enrolled in MFA.')
        _, workbook = self.export(bundle)
        values = list(workbook['MFA Registration Review'].values)
        self.assertEqual(len(values), 2)
        self.assertEqual(dict(zip(values[0], values[1]))['User ID'], 'unregistered')
        rec = next(row for row in bundle['assessment_result']['actions'] if row.get('FindingKey') == 'entra.authentication.mfa_registration')
        self.assertEqual(rec['InvestigationCount'], 1)
        self.assertIn('unknown', rec['InvestigationQualification'].lower())


if __name__ == '__main__':
    unittest.main()
